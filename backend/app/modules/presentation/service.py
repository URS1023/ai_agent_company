"""Built-in presentation template catalog for the V1 PPT plugin."""
from __future__ import annotations

from copy import deepcopy
import io
import json
from pathlib import Path
from collections import Counter

from sqlalchemy import select

from ...config import PROJECT_ROOT
from ...common.storage import get_storage
from ...errors import NotFoundError, ValidationError
from ...models import File as FileRow
from ..files.service import save_upload
from .preview import (
    BUILTIN_PREVIEW_DIR,
    builtin_preview_url,
    cache_custom_preview,
    custom_preview_key,
)


from .premium_catalog import PREMIUM_TEMPLATES

TEMPLATES = PREMIUM_TEMPLATES


_CUSTOM_PREVIEWS = [
    {"background": "#F2F5F8", "accent": "#40566F", "text": "#17212B"},
    {"background": "#F7F2EC", "accent": "#8B4A32", "text": "#2B211D"},
    {"background": "#EEF6F2", "accent": "#2E6B55", "text": "#172B24"},
    {"background": "#F3F0F7", "accent": "#6B4D86", "text": "#251D2D"},
]


def _builtin_public(item: dict) -> dict:
    """把内置模板目录条目转成对外结构（去掉 theme 等内部字段）。

    只有在磁盘上确实存在预览图/示例 PPTX 时才补上 ``image_url``/``example_url``，
    避免前端拿到 404 的地址；``example_page_count=14`` 与示例文件的实际页数保持一致，
    属于示例文档约定，不是用户在生成时能选的页数上限。
    """
    from .premium_demo import PROFILES

    out = {k: deepcopy(v) for k, v in item.items() if k != "theme"}
    image_url = builtin_preview_url(item["id"])
    if image_url:
        out.setdefault("preview", {})["image_url"] = image_url
        out["preview"]["kind"] = "rendered-deck-sheet"
    example_path = (
        Path(PROJECT_ROOT) / "examples" / "presentation-templates"
        / f"{item['id']}.pptx"
    )
    if example_path.exists():
        out["example_url"] = f"/api/presentation-templates/{item['id']}/example"
        out["example_page_count"] = 14
        out["demo_title"] = PROFILES[item["id"]]["title"]
    out["source"] = "builtin"
    return out


def _load_custom_analysis(row: FileRow) -> dict:
    """读取上传时旁存的模板画像（``<storage_key>.template.json``）。

    读取失败或缺失一律返回空字典：模板仍可列出，只是没有母版/配色分析，不能因为一张
    历史附件损坏就整份目录不可用。
    """

    storage = get_storage()
    sidecar = row.storage_key + ".template.json"
    if not storage.exists(sidecar):
        return {}
    try:
        return json.loads(storage.read_text(sidecar))
    except Exception:
        return {}


def _custom_public(row: FileRow) -> dict:
    """把上传的客户 PPTX 转成模板目录条目；ID 前缀 ``custom:`` 是区分内置/客户的唯一约定。

    没有渲染出真实缩略图时用按文件 ID 稳定选取的配色占位（同一模板每次配色一致），
    而不是伪造一张"预览图"；``image_url`` 只在预览真实存在时给出。
    """

    analysis = _load_custom_analysis(row)
    storage = get_storage()
    index = sum(ord(ch) for ch in row.id[-6:]) % len(_CUSTOM_PREVIEWS)
    display_name = Path(row.name).stem or "客户模板"
    preview = deepcopy(analysis.get("preview") or _CUSTOM_PREVIEWS[index])
    if storage.exists(custom_preview_key(row.storage_key)):
        preview["image_url"] = f"/api/presentation-templates/custom:{row.id}/preview"
        preview["kind"] = "rendered-slide"
    return {
        "id": f"custom:{row.id}",
        "name": display_name,
        "category": "客户模板",
        "description": "保留原始母版、布局、字体、配色与品牌元素",
        "recommended_for": ["品牌汇报", "客户专属模板"],
        "layout_family": analysis.get("layout_family") or "native",
        "density": analysis.get("density") or "balanced",
        "layouts": analysis.get("roles") or ["cover", "content"],
        "preview": preview,
        "analysis": analysis,
        "source": "custom",
        "file_id": row.id,
        "download_url": f"/api/files/{row.id}/download",
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def list_templates(db=None, workspace_id: str | None = None) -> list[dict]:
    """内置模板 + 当前工作区可用客户模板；不传 db/workspace 时只返回内置目录。

    客户模板只取 ``source_type="presentation_template"``、未软删除且 ``status="ready"`` 的行，
    因此上传解析未完成的文件不会出现在可选列表里。
    """

    items = [_builtin_public(item) for item in TEMPLATES]
    if db is not None and workspace_id:
        rows = db.execute(select(FileRow).where(
            FileRow.workspace_id == workspace_id,
            FileRow.source_type == "presentation_template",
            FileRow.deleted_at.is_(None),
            FileRow.status == "ready",
        ).order_by(FileRow.created_at.desc())).scalars().all()
        items.extend(_custom_public(row) for row in rows)
    return items


def get_template(template_id: str | None) -> dict:
    """按 ID 取内置模板；**未知或为空时回退到目录第一套**。

    注意与 ``reports`` 模块"必须显式选择模板、禁止默认回退"的规则相反：本函数的回退是
    PPT 插件的既有行为（仅在内置目录内生效，不涉及客户数据），因此调用方若需要"用户必须
    明确选模板"的语义，应在自己那一层先校验 template_id 再调用，不能依赖这里报错。
    返回深拷贝，调用方修改不会污染目录。
    """

    selected = next((item for item in TEMPLATES if item["id"] == template_id), TEMPLATES[0])
    return deepcopy(selected)


def save_custom_template(db, *, workspace_id: str, owner_user_id: str,
                         filename: str, mime_type: str, data: bytes,
                         display_name: str = "") -> dict:
    """校验并保存上传的客户 PPTX 模板，返回目录条目；副作用：写文件表、落盘附件与画像。

    校验顺序：后缀必须是 .pptx、非空且不超过 50MB、ZIP 魔数 ``PK`` 开头，最后用 python-pptx
    真正打开并确认至少有一个幻灯片布局——只信任解析成功，不靠后缀判断。备注（sidecar）
    ``.template.json`` 保存模板画像，供后续渲染按真实母版/配色出图，解析失败时该模板仍可用
    但退化为占位配色。抛 ``ValidationError`` 表示上传内容不可用。
    """

    if Path(filename or "").suffix.lower() != ".pptx":
        raise ValidationError("仅支持上传 .pptx 模板")
    if not data or len(data) > 50 * 1024 * 1024:
        raise ValidationError("PPT 模板不能为空且不得超过 50MB")
    if data[:2] != b"PK":
        raise ValidationError("文件不是有效的 PPTX")
    try:
        from pptx import Presentation
        deck = Presentation(io.BytesIO(data))
        if len(deck.slide_layouts) < 1:
            raise ValueError("缺少幻灯片布局")
    except Exception as exc:
        raise ValidationError(f"PPT 模板解析失败: {str(exc)[:240]}")
    safe_display = (display_name or Path(filename).stem or "客户模板").strip()[:100]
    stored_name = safe_display + ".pptx"
    row = save_upload(
        db, workspace_id=workspace_id, owner_user_id=owner_user_id,
        filename=stored_name,
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        data=data,
    )
    row.source_type = "presentation_template"
    db.flush()
    analysis = _analyze_template_deck(deck)
    storage = get_storage()
    analysis["preview_rendered"] = cache_custom_preview(
        storage, row.storage_key, data
    )
    storage.put_bytes(
        row.storage_key + ".template.json",
        json.dumps(analysis, ensure_ascii=False, indent=2).encode("utf-8"),
    )
    return _custom_public(row)


def read_template_preview(
    db, workspace_id: str, template_id: str
) -> bytes:
    """读取模板预览图字节：内置走随包 webp，``custom:` 前缀走客户模板渲染缓存。

    内置模板会校验 ID 确实在目录内，防止用任意文件名探测预览目录；客户模板必须能通过
    ``resolve_template`` 的工作区校验。无预览时抛 ``NotFoundError``（不是 500），
    前端可回退到占位配色。
    """

    if not template_id.startswith("custom:"):
        path = BUILTIN_PREVIEW_DIR / f"{template_id}.webp"
        if not path.exists() or not any(item["id"] == template_id for item in TEMPLATES):
            raise NotFoundError("模板预览不存在")
        return path.read_bytes()
    resolved = resolve_template(db, workspace_id, template_id)
    key = custom_preview_key(resolved["storage_key"])
    storage = get_storage()
    if not storage.exists(key):
        raise NotFoundError("客户模板尚无可用预览")
    return storage.read_bytes(key)


def read_template_example(template_id: str) -> tuple[bytes, str]:
    """读取内置模板的 14 页示例 PPTX，返回 ``(字节, 下载文件名)``。

    仅内置模板有示例文件，未生成时抛 ``NotFoundError``；文件名带模板中文名，
    由路由层做 RFC 5987 编码。
    """

    if not any(item["id"] == template_id for item in TEMPLATES):
        raise NotFoundError("模板示例不存在")
    path = (
        Path(PROJECT_ROOT) / "examples" / "presentation-templates"
        / f"{template_id}.pptx"
    )
    if not path.exists():
        raise NotFoundError("模板示例尚未生成")
    template = get_template(template_id)
    return path.read_bytes(), f"{template['name']}-14页完整示例.pptx"


def _analyze_template_deck(deck) -> dict:
    """Build a compact PPTX→template profile without copying source content.

    只统计"形态"不搬运内容：布局角色、母版/布局数量、比例、字体与配色频次、图片/表格/
    图表个数、文本密度。最多扫描前 12 页且每个取色步骤都吞异常——客户模板可能存在损坏的
    形状或不受支持的填充，不能因此让上传失败。``preview`` 用出现频次最高的颜色推导
    背景/正文/强调色，属于近似配色提示而非精确取色。
    """
    from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL

    colors: Counter[str] = Counter()
    fonts: Counter[str] = Counter()
    roles: Counter[str] = Counter()
    picture_count = table_count = chart_count = text_count = 0

    def rgb_value(color) -> str | None:
        """安全取色：非 RGB 颜色或读取失败返回 None（客户模板可能用主题色引用）。"""
        try:
            if color.type == MSO_COLOR_TYPE.RGB and color.rgb:
                return str(color.rgb)
        except Exception:
            return None
        return None

    def classify_layout(name: str, placeholders) -> str:
        """按母版版式名与占位符数量推断版式角色（cover/section/two_column/image/blank/
        multi_content/content），用于目录里展示该客户模板"能做哪些页"。"""
        lowered = (name or "").lower()
        if any(token in lowered for token in ("title slide", "cover", "封面")):
            return "cover"
        if any(token in lowered for token in ("section", "chapter", "章节")):
            return "section"
        if any(token in lowered for token in ("two", "comparison", "双", "比较")):
            return "two_column"
        if any(token in lowered for token in ("picture", "image", "图片")):
            return "image"
        if any(token in lowered for token in ("blank", "空白")):
            return "blank"
        if len(placeholders) >= 3:
            return "multi_content"
        return "content"

    layout_inventory = []
    for layout in deck.slide_layouts:
        placeholders = list(layout.placeholders)
        role = classify_layout(layout.name or "", placeholders)
        roles[role] += 1
        layout_inventory.append({
            "name": layout.name or "未命名布局",
            "role": role,
            "placeholder_count": len(placeholders),
            "placeholder_types": [str(item.placeholder_format.type) for item in placeholders],
        })

    for slide in list(deck.slides)[:12]:
        try:
            if slide.background.fill.type == MSO_FILL.SOLID:
                value = rgb_value(slide.background.fill.fore_color)
                if value:
                    colors[value] += 3
        except Exception:
            pass
        for shape in slide.shapes:
            if getattr(shape, "shape_type", None) == 13:
                picture_count += 1
            if getattr(shape, "has_table", False):
                table_count += 1
            if getattr(shape, "has_chart", False):
                chart_count += 1
            try:
                if shape.fill.type == MSO_FILL.SOLID:
                    value = rgb_value(shape.fill.fore_color)
                    if value:
                        colors[value] += 1
            except Exception:
                pass
            if not getattr(shape, "has_text_frame", False):
                continue
            text_count += 1
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts[run.font.name] += 1
                    try:
                        value = rgb_value(run.font.color)
                        if value:
                            colors[value] += 1
                    except Exception:
                        pass

    common_colors = [value for value, _ in colors.most_common(8)]
    background = common_colors[0] if common_colors else "F5F5F5"
    text = next((value for value in common_colors if value not in {background, "FFFFFF"}), "202020")
    accent = next((value for value in common_colors if value not in {background, text, "FFFFFF", "000000"}), "315A7D")
    density_score = text_count / max(1, len(deck.slides))
    density = "dense" if density_score >= 12 else "airy" if density_score <= 5 else "balanced"
    aspect = round(deck.slide_width / max(1, deck.slide_height), 3)
    return {
        "schema": "presentation_template_profile.v1",
        "slide_count": len(deck.slides),
        "master_count": len(deck.slide_masters),
        "layout_count": len(deck.slide_layouts),
        "aspect_ratio": "16:9" if abs(aspect - 1.778) < 0.12 else "4:3" if abs(aspect - 1.333) < 0.12 else str(aspect),
        "layout_family": "native",
        "density": density,
        "roles": list(roles.keys()) or ["cover", "content"],
        "layout_inventory": layout_inventory,
        "fonts": [name for name, _ in fonts.most_common(5)],
        "colors": common_colors,
        "content_features": {
            "pictures": picture_count,
            "tables": table_count,
            "charts": chart_count,
        },
        "preview": {
            "background": "#" + background,
            "accent": "#" + accent,
            "text": "#" + text,
            "layout": "native",
        },
    }


def resolve_template(db, workspace_id: str, template_id: str | None) -> dict:
    """把模板 ID 解析成渲染所需的完整描述，返回 ``source`` 与 ``template_path``。

    ``custom:<file_id>`` 必须同时满足：文件存在、属于当前工作区、``source_type`` 是
    ``presentation_template``、未软删除且 ``status="ready"``，否则一律抛 ``NotFoundError``
    ——跨工作区猜测 file_id 不会泄露模板。内置 ID（含空值/未知）走 ``get_template``，
    因此会回退到第一套内置模板（详见 ``get_template`` 的说明）。
    """

    if not template_id or not template_id.startswith("custom:"):
        template = get_template(template_id)
        return {**template, "source": "builtin", "template_path": None}
    file_id = template_id.split(":", 1)[1]
    row = db.get(FileRow, file_id)
    if (not row or row.workspace_id != workspace_id or
            row.source_type != "presentation_template" or
            row.deleted_at is not None or row.status != "ready"):
        raise NotFoundError("客户 PPT 模板不存在")
    path = get_storage().get_path(row.storage_key)
    return {
        **_custom_public(row),
        "template_path": path,
        "storage_key": row.storage_key,
    }


def delete_custom_template(db, workspace_id: str, template_id: str) -> None:
    """软删除工作区自定义模板：置 ``status="deleted"`` 并记录 ``deleted_at`` 后提交。

    只允许删除客户模板（内置目录不可删）；用 ``resolve_template`` 先做归属校验，
    因此删除他人工作区的模板只会得到 NotFoundError。已生成的文档不受影响。
    """

    resolved = resolve_template(db, workspace_id, template_id)
    if resolved.get("source") != "custom":
        raise ValidationError("内置模板不能删除")
    row = db.get(FileRow, resolved["file_id"])
    if row:
        row.status = "deleted"
        from datetime import datetime
        row.deleted_at = datetime.utcnow()
        db.commit()
