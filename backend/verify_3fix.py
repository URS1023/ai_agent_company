# -*- coding: utf-8 -*-
import io, sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from app.db import new_session
from app.models import Workspace, User
from app.providers.tool.gateway import ToolGateway, build_context

db = new_session(); ws = db.query(Workspace).first(); user = db.query(User).first()
gw = ToolGateway(db)
names = [s.name for s in gw.list_available(ws.id)]
print("list has python_execute:", "python_execute" in names)

code1 = """
import pandas as pd
df = pd.DataFrame({'a':[1,2,3],'b':[10,20,30]})
result = {'sum_a': int(df.a.sum()), 'mean_b': float(df.b.mean()), 'rows': len(df)}
"""
ctx = build_context(db=db, workspace_id=ws.id, user_id=user.id, run_id="verify_py")
r = gw.execute(ws.id, "python_execute", {"code": code1}, ctx)
print("python_execute:", r.status, "| data:", (r.data or {}).get("result") if r.data else (r.error_summary or "")[:120])

from app.modules.files.service import save_upload
from app.models import Artifact, Conversation, Run
from app.ids import new_id
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
wb = Workbook(); wsh = wb.active
wsh["A1"] = "债务重组分析表"; wsh["A1"].font = Font(name="微软雅黑", size=14, bold=True)
wsh.merge_cells("A2:C2"); wsh["A2"] = "客户姓名"; wsh["B3"] = "身份证号"; wsh["A4"] = "婚姻状况"
wsh.merge_cells("B4:C4")
buf = io.BytesIO(); wb.save(buf)
row = save_upload(db, workspace_id=ws.id, owner_user_id=user.id, filename="合并模板.xlsx",
                  mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", data=buf.getvalue())
db.commit()
conv = Conversation(id=new_id(), workspace_id=ws.id, owner_user_id=user.id, title="verify",
                    created_at=datetime.utcnow(), updated_at=datetime.utcnow())
db.add(conv); db.flush()
run = Run(id=new_id(), workspace_id=ws.id, conversation_id=conv.id, started_by_user_id=user.id,
          status="completed", created_at=datetime.utcnow(), model_profile="reasoning")
db.add(run); db.flush(); db.commit()
ctx2 = build_context(db=db, workspace_id=ws.id, user_id=user.id, run_id=run.id)
res = gw.execute(ws.id, "fill_template", {"file_id": row.id,
    "by_label":[{"label":"客户姓名","value":"高君华","direction":"below"}],
    "font":{"name":"宋体","size":11,"bold":True}}, ctx2)
print("fill merged:", res.status, (res.error_summary or "")[:120] if res.status!="success" else "")
if res.status == "success":
    art = db.get(Artifact, res.data["artifact_id"])
    from app.common.storage import get_storage
    wb2 = load_workbook(io.BytesIO(get_storage().read_bytes(art.storage_key)))
    print("A3 (below merged header) =", wb2.active["A3"].value, "| font:", wb2.active["A3"].font.name, wb2.active["A3"].font.size)
    print("A1 header font preserved:", wb2.active["A1"].font.name, wb2.active["A1"].font.size)
db.close();
