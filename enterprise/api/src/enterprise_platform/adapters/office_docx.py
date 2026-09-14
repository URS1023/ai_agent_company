"""Native DOCX paragraphs/tables adapted from the legacy renderer, with stable unit identity.

This is a pure byte renderer, not an authorization endpoint or artifact publisher.
No filesystem paths, network images or generated layout instructions are accepted.
"""

from io import BytesIO
from xml.etree import ElementTree as ET

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement

from enterprise_platform.application.contracts import JsonObject, canonical_json
from enterprise_platform.application.errors import InvalidInput
from enterprise_platform.application.office_edits import OfficeFileRecord
from enterprise_platform.domain.office_content import TableData
from enterprise_platform.domain.office_revision import OfficeText


def _xml_text(text: str) -> str:
    if any(
        not (
            ord(char) in {9, 10, 13}
            or 0x20 <= ord(char) <= 0xD7FF
            or 0xE000 <= ord(char) <= 0xFFFD
            or 0x10000 <= ord(char) <= 0x10FFFF
        )
        for char in text
    ):
        raise InvalidInput("office_document_text_invalid")
    return text


def _tag_unit(element: BaseOxmlElement, unit_id: str) -> None:
    parent = element.getparent()
    if parent is None:
        raise InvalidInput("office_document_structure_invalid")
    index = parent.index(element)
    control, properties, tag, content = (OxmlElement(name) for name in ("w:sdt", "w:sdtPr", "w:tag", "w:sdtContent"))
    tag.set(qn("w:val"), unit_id)
    properties.append(tag)
    control.append(properties)
    content.append(element)
    control.append(content)
    parent.insert(index, control)


def render_office_docx(record: OfficeFileRecord) -> bytes:
    if record.content.kind != "document":
        raise InvalidInput("office_document_required")
    metadata_payload: JsonObject = {
        "workspace_id": record.workspace_id,
        "template_id": record.template_id,
        "template_revision": record.template_revision,
        "source_snapshot_ids": list(record.source_snapshot_ids),
        "content": record.content.model_dump(mode="json"),
    }
    # Validate metadata before fingerprinting, which encodes the same strings as UTF-8.
    _xml_text(canonical_json(metadata_payload))
    metadata_payload["file_fingerprint"] = record.fingerprint()
    document = Document()
    document.core_properties.identifier = f"office:{record.content.file_id}:r{record.content.revision}"
    for unit in record.content.units:
        if unit.kind == "paragraph":
            paragraph = document.add_paragraph()
            for item in unit.content:
                if not isinstance(item, OfficeText):
                    raise InvalidInput("office_paragraph_content_invalid")
                paragraph.add_run(_xml_text(item.text))
            _tag_unit(paragraph._p, str(unit.unit_id))
        elif unit.kind == "table":
            data = unit.content[0]
            if not isinstance(data, TableData):
                raise InvalidInput("office_table_content_invalid")
            table = document.add_table(rows=len(data.rows) + 1, cols=len(data.headers))
            table.style = "Light Grid Accent 1"
            for row, values in zip(table.rows, (data.headers, *data.rows), strict=True):
                for cell, value in zip(row.cells, values, strict=True):
                    cell.text = "" if value is None else _xml_text(str(value))
            _tag_unit(table._tbl, str(unit.unit_id))
        else:
            raise InvalidInput("office_document_unit_invalid")
    metadata = ET.Element("{urn:enterprise:office:revision:v1}record")
    metadata.text = canonical_json(metadata_payload)
    part = Part(
        PackURI("/customXml/enterprise-office.xml"),
        "application/xml",
        ET.tostring(metadata, encoding="utf-8", xml_declaration=True),
        document.part.package,
    )
    document.part.relate_to(part, RT.CUSTOM_XML)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
