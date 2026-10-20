"""受限上传、授权下载与来源可追溯的文档生成接口。"""

from io import BytesIO
from pathlib import Path
import hashlib
import hmac
import json
import math
import time

from fastapi import APIRouter, Depends, Request, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from ...core.context import Actor, get_actor, require_permission
from ...core.database import get_session
from ...core.errors import DomainError
from ...core.files import FileService, UPLOAD_SUFFIXES, file_dict
from ...core.models import FileEntry, ResourceReference
from ...core.records import RecordStore
from ...core.settings import get_settings
from ...core.tasks import TaskService

router = APIRouter(prefix='/api/v1', tags=['文件与成果'])


@router.get('/files')
def files(request: Request, q: str = '', limit: int = 50, offset: int = 0,
          actor: Actor = Depends(require_permission('m17.read')), session: Session = Depends(get_session)):
    svc = FileService(session, actor, get_settings(request))
    rules = svc.conditions()
    if q:
        rules.append(FileEntry.name.contains(q, autoescape=True))
    total = session.scalar(select(func.count()).select_from(FileEntry).where(*rules)) or 0
    rows = session.scalars(select(FileEntry).where(*rules).order_by(FileEntry.created_at.desc())
                            .limit(min(max(limit, 1), 200)).offset(max(offset, 0)))
    return {'items': [file_dict(r) for r in rows], 'total': total}


@router.post('/files', status_code=201)
async def upload_file(file: UploadFile, request: Request,
                      actor: Actor = Depends(require_permission('m17.write')), session: Session = Depends(get_session)):
    if Path(file.filename or '').suffix.lower() not in UPLOAD_SUFFIXES:
        raise DomainError(422, 'unsupported_file_type', '当前文件类型不支持上传')
    content = await file.read(get_settings(request).max_upload_bytes + 1)
    return FileService(session, actor, get_settings(request)).store_bytes(
        file.filename, content, file.content_type or 'application/octet-stream')


@router.get('/files/{id}')
def file_metadata(id: str, request: Request, actor: Actor = Depends(require_permission('m17.read')),
                  session: Session = Depends(get_session)):
    return FileService(session, actor, get_settings(request)).get(id)


@router.post('/file-upload-sessions', status_code=201)
def upload_session(body: dict, request: Request, actor: Actor = Depends(require_permission('m17.write')),
                   session: Session = Depends(get_session)):
    name, size = str(body.get('name', '')), body.get('size', 0)
    if Path(name).suffix.lower() not in UPLOAD_SUFFIXES or not isinstance(size, int) or not 0 < size <= get_settings(request).max_upload_bytes:
        raise DomainError(422, 'invalid_file', '文件类型或大小不符合当前上传要求')
    row = RecordStore(session, actor).create('file-upload-session', {**body, 'status': 'pending',
        'expires_at': int(time.time()) + 900}, visibility='private')
    return {**row, 'upload_url': f"/api/v1/file-upload-sessions/{row['id']}/content"}


@router.put('/file-upload-sessions/{id}/content')
async def upload_content(id: str, request: Request, actor: Actor = Depends(require_permission('m17.write')),
                         session: Session = Depends(get_session)):
    store = RecordStore(session, actor)
    row = store.get('file-upload-session', id)
    if row['status'] != 'pending' or row['expires_at'] < time.time():
        raise DomainError(409, 'upload_session_expired', '上传会话已使用或过期，请重新上传')
    chunks, size = [], 0
    async for chunk in request.stream():
        size += len(chunk)
        if size > row['size'] or size > get_settings(request).max_upload_bytes:
            raise DomainError(413, 'file_too_large', '文件大小超过本次上传声明')
        chunks.append(chunk)
    if size != row['size']:
        raise DomainError(422, 'file_size_mismatch', '上传内容不完整，请重新上传')
    file = FileService(session, actor, get_settings(request)).store_bytes(
        row['name'], b''.join(chunks), row.get('content_type', 'application/octet-stream'),
        visibility=row.get('visibility', 'private'))
    store.update('file-upload-session', id, {'status': 'uploaded', 'file_id': file['id']}, row['revision'])
    return file


@router.post('/file-upload-sessions/{id}/completions')
def complete_upload(id: str, body: dict, request: Request,
                     actor: Actor = Depends(require_permission('m17.write')), session: Session = Depends(get_session)):
    store = RecordStore(session, actor)
    row = store.get('file-upload-session', id)
    if row.get('status') not in {'uploaded', 'completed'}:
        raise DomainError(409, 'file_not_uploaded', '请先上传文件内容')
    file = FileService(session, actor, get_settings(request)).get(row['file_id'])
    if body.get('sha256') and not hmac.compare_digest(str(body['sha256']), file['sha256']):
        raise DomainError(422, 'file_hash_mismatch', '文件校验不一致，请重新上传')
    if row['status'] != 'completed':
        store.update('file-upload-session', id, {'status': 'completed'}, row['revision'])
    return file


@router.post('/files/{id}/download-links')
def download_link(id: str, request: Request, actor: Actor = Depends(require_permission('m17.read')),
                   session: Session = Depends(get_session)):
    FileService(session, actor, get_settings(request)).get(id)
    expires = int(time.time()) + 300
    message = f'{id}:{actor.user_id}:{actor.workspace_id}:{expires}'
    signature = request.app.state.secret_vault.sign(message)
    return {'url': f'/api/v1/files/{id}/content?expires={expires}&signature={signature}',
            'expires_at': expires}


@router.get('/files/{id}/content')
def download(id: str, request: Request, expires: int, signature: str,
              actor: Actor = Depends(require_permission('m17.read')), session: Session = Depends(get_session)):
    if expires < time.time() or expires > time.time() + 3600:
        raise DomainError(403, 'download_expired', '下载链接已过期，请重新获取')
    expected = request.app.state.secret_vault.sign(f'{id}:{actor.user_id}:{actor.workspace_id}:{expires}')
    if not hmac.compare_digest(expected, signature):
        raise DomainError(403, 'download_signature_invalid', '下载链接无效')
    svc = FileService(session, actor, get_settings(request))
    file = svc.get(id)
    return FileResponse(svc.path(id), filename=file['name'], media_type=file['content_type'],
                        headers={'Cache-Control': 'private, no-store', 'X-Content-Type-Options': 'nosniff',
                                 'Content-Security-Policy': "default-src 'none'; sandbox"})


@router.post('/files/{id}/deletion-previews')
def deletion_preview(id: str, request: Request, actor: Actor = Depends(require_permission('m17.delete')),
                       session: Session = Depends(get_session)):
    references = FileService(session, actor, get_settings(request)).references(id)
    return {'id': id, 'can_delete': not references, 'references': references}


@router.delete('/files/{id}')
def delete_file(id: str, request: Request, revision: int,
                actor: Actor = Depends(require_permission('m17.delete')), session: Session = Depends(get_session)):
    svc = FileService(session, actor, get_settings(request))
    row = svc._get(id)
    if row.owner_id != actor.user_id and not actor.is_admin:
        raise DomainError(403, 'file_owner_required', '只有文件所有者或管理员可删除')
    if row.revision != revision:
        raise DomainError(409, 'revision_conflict', '文件状态已更新，请刷新后重试')
    if svc.references(id):
        raise DomainError(409, 'file_in_use', '文件仍被业务记录引用，请先查看引用影响')
    changed = session.execute(update(FileEntry).where(FileEntry.id == id, FileEntry.revision == revision,
        FileEntry.workspace_id == actor.workspace_id, FileEntry.deleted.is_(False))
        .values(deleted=True, revision=revision + 1)).rowcount
    if changed != 1:
        raise DomainError(409, 'revision_conflict', '文件状态已更新，请刷新后重试')
    RecordStore(session, actor).audit('file.delete', 'file', id)
    return {'id': id, 'deleted': True, 'revision': revision + 1}


@router.get('/artifacts')
def artifacts(q: str = '', actor: Actor = Depends(require_permission('m17.read')), session: Session = Depends(get_session)):
    return RecordStore(session, actor).list('artifact', q=q)


@router.get('/artifacts/{id}')
def artifact(id: str, actor: Actor = Depends(require_permission('m17.read')), session: Session = Depends(get_session)):
    return RecordStore(session, actor).get('artifact', id)


@router.post('/artifacts/{id}/archivals')
def archive_artifact(id: str, body: dict, actor: Actor = Depends(require_permission('m17.write')),
                       session: Session = Depends(get_session)):
    return RecordStore(session, actor).update('artifact', id, {'archived': True}, body.get('revision', 0))


def make_document(format: str, title: str, sections: list[dict]) -> tuple[bytes, str, str]:
    """生成真实办公格式；文本保持原文，表格按输入行列输出，不借用设备CSV替代。"""
    output = BytesIO()
    if format == 'docx':
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for section in sections:
            if section.get('title'):
                doc.add_heading(str(section['title']), 1)
            if section.get('text'):
                doc.add_paragraph(str(section['text']))
            rows = section.get('rows', [])
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=width)
                table.style = 'Table Grid'
                for row in rows:
                    cells = table.add_row().cells
                    for i, value in enumerate(row):
                        cells[i].text = str(value)
        doc.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font
        book = Workbook()
        sheet = book.active
        sheet.title = '报告'
        sheet.append([title])
        sheet.cell(1, 1).font = Font(size=18, bold=True)
        for section in sections:
            if section.get('title'):
                sheet.append([str(section['title'])])
            if section.get('text'):
                sheet.append([str(section['text'])])
            for row in section.get('rows', []):
                # 外部文本以单元格文本保存，防止公式注入。
                sheet.append([("'" + v if isinstance(v, str) and v[:1] in '=+-@' else v) for v in row])
        book.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif format == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        deck = Presentation()
        cover = deck.slides.add_slide(deck.slide_layouts[0])
        cover.shapes.title.text = title
        for section in sections:
            rows = section.get('rows', [])
            pages = [rows[i:i + 14] for i in range(0, len(rows), 14)] if rows else [[]]
            if len(deck.slides) + len(pages) > 100:
                raise DomainError(422, 'document_slide_budget', '文档超过100页，请拆分后生成')
            for page_index, page in enumerate(pages):
                slide = deck.slides.add_slide(deck.slide_layouts[5])
                slide.shapes.title.text = str(section.get('title', title))
                text = str(section.get('text', '')) if page_index == 0 else ''
                if text:
                    frame = slide.shapes.add_textbox(Inches(.5), Inches(1.2), Inches(9), Inches(.8)).text_frame
                    frame.text = text
                    for paragraph in frame.paragraphs:
                        paragraph.font.size = Pt(16)
                if page:
                    width = max(len(row) for row in page)
                    table = slide.shapes.add_table(len(page), width, Inches(.5),
                        Inches(2 if text else 1.3), Inches(9), Inches(4.5)).table
                    for row_index, row in enumerate(page):
                        for column, value in enumerate(row):
                            cell = table.cell(row_index, column)
                            cell.text = str(value)
                            for paragraph in cell.text_frame.paragraphs:
                                paragraph.font.size = Pt(12)
        deck.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    else:
        raise DomainError(422, 'document_format_invalid', '请选择Word、Excel或PowerPoint格式')
    return output.getvalue(), mime, format


def validate_sections(sections: list) -> None:
    """先限制内容结构和总工作量，再创建任务，避免格式错误产生半份文档。"""
    cells, characters = 0, 0
    for section in sections:
        if not isinstance(section, dict):
            raise DomainError(422, 'document_section_invalid', '文档章节格式不正确')
        for key in ('title', 'text'):
            value = section.get(key, '')
            if not isinstance(value, str) or len(value) > 20000:
                raise DomainError(422, 'document_text_invalid', '章节标题或内容过长')
            characters += len(value)
        rows = section.get('rows', [])
        if not isinstance(rows, list) or len(rows) > 2000:
            raise DomainError(422, 'document_rows_invalid', '单章节表格最多2000行')
        for row in rows:
            if not isinstance(row, list) or not 0 < len(row) <= 100:
                raise DomainError(422, 'document_columns_invalid', '表格每行需包含1到100个单元格')
            for value in row:
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    raise DomainError(422, 'document_cell_invalid', '单元格请填写文本、数字或布尔值')
                if isinstance(value, float) and not math.isfinite(value):
                    raise DomainError(422, 'document_cell_invalid', '表格数字必须是有限值')
                characters += len(str(value))
            cells += len(row)
    if cells > 20000 or characters > 2000000:
        raise DomainError(422, 'document_content_budget', '文档内容过大，请拆分后生成')


@router.post('/document-jobs', status_code=201)
def document_job(body: dict, request: Request, actor: Actor = Depends(require_permission('m17.write')),
                  session: Session = Depends(get_session)):
    title = str(body.get('title', '')).strip()
    sections = body.get('sections', [])
    if not title or len(title) > 200 or not isinstance(sections, list) or not 0 < len(sections) <= 100:
        raise DomainError(422, 'document_content_required', '请填写标题和文档内容')
    validate_sections(sections)
    source_ids = body.get('source_file_ids', [])
    if not isinstance(source_ids, list) or len(source_ids) > 50 or any(not isinstance(s, str) for s in source_ids):
        raise DomainError(422, 'document_sources_invalid', '请选择最多50份来源文件')
    source_ids = list(dict.fromkeys(source_ids))
    files = FileService(session, actor, get_settings(request))
    source_files = [files.get(source) for source in source_ids]
    source_versions = [{'id': f['id'], 'name': f['name'], 'revision': f['revision'], 'sha256': f['sha256']} for f in source_files]
    sections = [*sections]
    if source_files:
        sections.append({'title': '数据来源', 'rows': [['文件', '版本'],
            *[[f['name'], f['revision']] for f in source_files]]})
    content_hash = hashlib.sha256(json.dumps(sections, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    tasks = TaskService(session, actor)
    task = tasks.create('document', title, {'title': title, 'content_hash': content_hash,
                        'format': body.get('format', 'docx'), 'source_file_ids': source_ids, 'source_versions': source_versions},
                        request.headers.get('idempotency-key'))
    if task['status'] != 'queued':
        return task
    tasks.transition(task['id'], 'running')
    data, mime, ext = make_document(body.get('format', 'docx'), title, sections)
    file = files.store_bytes(title + '.' + ext, data, mime, {'type': 'document', 'id': task['id']})
    artifact = RecordStore(session, actor).create('artifact', {'name': title, 'type': 'document',
        'task_id': task['id'], 'file_id': file['id'], 'source_file_ids': source_ids,
        'source_versions': source_versions, 'content_hash': content_hash, 'format': ext}, visibility='private')
    for source in source_ids:
        session.add(ResourceReference(workspace_id=actor.workspace_id, source_type='artifact', source_id=artifact['id'],
                                     target_type='file', target_id=str(source)))
    return tasks.transition(task['id'], 'succeeded', outputs={'file': file, 'artifact': artifact})


@router.get('/document-jobs/{id}/previews')
def document_previews(id: str, actor: Actor = Depends(require_permission('m17.read')), session: Session = Depends(get_session)):
    task = TaskService(session, actor).get(id)
    if task['kind'] != 'document':
        raise DomainError(404, 'document_job_not_found', '文档任务不存在')
    return {'task_id': id, 'status': task['status'], 'file': task['outputs'].get('file'),
            'artifact': task['outputs'].get('artifact')}
