"""精提Test5 files/extractor.py的格式解析与来源块；不执行宏、公式、外链或OCR。

保留原实现的严格文本解码、逐页PDF、Office正文/表格和CSV解析方式；
增加格式白名单、解压预算和真实定位，Word不伪造页码，Excel公式按原文保留。
此模块可独立子进程运行，解析超时由父进程终止，不加载平台配置或凭据。
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
import sys
import zipfile

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_BLOCKS = 5000
MAX_ZIP_ENTRIES = 2000
MAX_ZIP_BYTES = 64 * 1024 * 1024
MAX_PDF_PAGES = 300
SUPPORTED_SUFFIXES = {'.txt', '.md', '.csv', '.json', '.log', '.pdf', '.docx', '.xlsx', '.pptx'}


class ExtractionError(ValueError):
    """可安全显示的解析失败，不携带本机路径或底层异常正文。"""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code, self.message = code, message


def _text(path: Path) -> str:
    data = path.read_bytes()
    encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
    try:
        text = data.decode(encoding, errors='strict')
    except UnicodeError as exc:
        raise ExtractionError('document_encoding_invalid', '文本编码无法识别，请保存为UTF-8后上传') from exc
    if any(ord(char) < 32 and char not in '\t\n\r\f' for char in text):
        raise ExtractionError('document_binary_text', '文件内容不是可解析的文本')
    return text


def _zip_budget(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES or sum(item.file_size for item in entries) > MAX_ZIP_BYTES:
                raise ExtractionError('document_expansion_limit', '文件展开后过大，请拆分后上传')
            if any(item.flag_bits & 1 for item in entries):
                raise ExtractionError('document_encrypted', '请先解除文件加密后上传')
    except zipfile.BadZipFile as exc:
        raise ExtractionError('document_corrupt', '文件已损坏或格式与扩展名不符') from exc


def extract_file(filename: str, path: str | Path) -> dict:
    """返回{text,location}块和解析元数据；达到预算时报错，绝不截断后称成功。"""
    path = Path(path)
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ExtractionError('document_format_unsupported', '请选择文本、PDF、DOCX、XLSX或PPTX文件')
    if not 0 < path.stat().st_size <= MAX_FILE_BYTES:
        raise ExtractionError('document_file_size', '请选择20MB以内的非空文件')
    if suffix in {'.docx', '.xlsx', '.pptx'}:
        _zip_budget(path)
    blocks, warnings, characters = [], [], 0

    def append(text: str, **location):
        nonlocal characters
        text = text.strip()
        if not text:
            return
        characters += len(text)
        if characters > MAX_TEXT_CHARS or len(blocks) >= MAX_BLOCKS:
            raise ExtractionError('document_text_limit', '文档正文过长，请拆分后上传')
        blocks.append({'text': text, 'location': location})

    try:
        if suffix == '.pdf':
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise ExtractionError('document_parser_unavailable', 'PDF解析组件尚未安装') from exc
            reader = PdfReader(path, strict=True)
            if reader.is_encrypted:
                raise ExtractionError('document_encrypted', '请先解除文件加密后上传')
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ExtractionError('document_page_limit', 'PDF页数超过300页，请拆分后上传')
            empty = []
            for index, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ''
                append(text, page=index)
                if not text.strip():
                    empty.append(index)
            if empty:
                warnings.append({'code': 'pages_without_text', 'pages': empty,
                                 'message': '部分页面没有可提取文字，未对图像执行识别'})
        elif suffix == '.docx':
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            document = Document(path)
            paragraph, table = 0, 0
            # 沿正文原顺序读取段落和表格，不把文末表格移到正文之后。
            for element in document.element.body.iterchildren():
                if element.tag.endswith('}p'):
                    paragraph += 1
                    append(Paragraph(element, document).text, paragraph=paragraph)
                elif element.tag.endswith('}tbl'):
                    table += 1
                    for number, row in enumerate(Table(element, document).rows, 1):
                        append(' | '.join(cell.text.strip() for cell in row.cells), table=table, row=number)
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            formula = False
            try:
                for sheet in workbook.worksheets:
                    # 空白工作表夸大的dimension也不能造成无界遍历。
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > 200_000:
                        raise ExtractionError('document_cell_limit', '工作表范围过大，请删除多余空白区域后上传')
                    for number, row in enumerate(sheet.iter_rows(), 1):
                        values = []
                        for cell in row:
                            formula = formula or cell.data_type == 'f'
                            values.append('' if cell.value is None else str(cell.value))
                        while values and not values[-1]:
                            values.pop()
                        append('\t'.join(values), sheet=sheet.title, row=number)
            finally:
                workbook.close()
            if formula:
                warnings.append({'code': 'formulas_as_text', 'message': '公式按原文保留，未计算结果'})
        elif suffix == '.pptx':
            from pptx import Presentation
            presentation = Presentation(path)
            for number, slide in enumerate(presentation.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                    if shape.has_table:
                        texts.extend(' | '.join(cell.text.strip() for cell in row.cells) for row in shape.table.rows)
                append('\n'.join(texts), slide=number)
        elif suffix == '.csv':
            # row为CSV逻辑记录序号，带引号的换行不会被当作两个业务记录。
            for number, row in enumerate(csv.reader(io.StringIO(_text(path), newline='')), 1):
                append(' | '.join(row), row=number)
        else:
            append(_text(path))
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError('document_corrupt', '文件解析失败，请确认文件完整且格式正确') from exc
    if not blocks:
        raise ExtractionError('document_no_text', '没有可提取的文字，请上传文字版文件')
    return {'blocks': blocks, 'metadata': {'name': 'platform-test5-extractor', 'version': 1,
        'blocks': len(blocks), 'warnings': warnings, 'format': suffix.removeprefix('.')}, 'char_count': characters}


def _main():
    """只接收父进程提供的受控文件路径；标准输出不含异常栈或本机配置。"""
    try:
        body = json.loads(sys.stdin.buffer.read(8192))
        result = extract_file(body['name'], body['path'])
        result = {'ok': True, **result}
    except ExtractionError as exc:
        result = {'ok': False, 'code': exc.code, 'message': exc.message}
    except Exception:
        result = {'ok': False, 'code': 'document_parse_failed', 'message': '文件解析未完成，请检查文件后重新上传'}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode('utf-8'))


if __name__ == '__main__':
    _main()
