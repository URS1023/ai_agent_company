"""将实际报表数据输出为办公表格与原生PPT图表，不用图片或CSV冒充。"""
from io import BytesIO
import math
from copy import deepcopy

from agent_platform.core.errors import DomainError


_STEP_CHART_FRAME = (.7, 1.35, 11.8, 5.0)
_STEP_PLOT_LAYOUT = (.10, .15, .84, .68)


def office_export(format, title, sections):
    total_cells = sum(sum(len(row) for row in section.get("rows", [])) for section in sections)
    if total_cells > 20000 or len(sections) > 200:
        raise DomainError(422, "office_export_budget_exceeded", "报表内容过大，请缩小范围再导出")
    if format != "pptx":
        from ..files.router import make_document
        return make_document(format, title, sections)
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = "数据报表"
    for section in sections:
        rows = section.get("rows", [])
        chart = section.get("chart")
        if chart and chart.get("values"):
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = str(section.get("title", title))
            if chart["type"] == "step":
                native = _add_step_chart(slide, chart)
            else:
                data = CategoryChartData()
                data.categories = chart["categories"]
                actual_series = chart.get("series") or [{"name": str(section.get("title", "数据")), "values": chart["values"]}]
                for item in actual_series:
                    data.add_series(item["name"], item["values"])
                kind = {"line": XL_CHART_TYPE.LINE, "area": XL_CHART_TYPE.AREA, "pie": XL_CHART_TYPE.PIE, "ring": XL_CHART_TYPE.DOUGHNUT, "radar": XL_CHART_TYPE.RADAR, "stacked": XL_CHART_TYPE.COLUMN_STACKED, "percent": XL_CHART_TYPE.COLUMN_STACKED_100, "horizontal_bar": XL_CHART_TYPE.BAR_CLUSTERED}.get(chart["type"], XL_CHART_TYPE.COLUMN_CLUSTERED)
                native = slide.shapes.add_chart(kind, Inches(0.7), Inches(1.35), Inches(11.8), Inches(5.4), data).chart
                if chart["type"] == "mixed":
                    _split_line_plots(native, actual_series)
                elif chart["type"] == "line":
                    for source, item in zip(actual_series, native.series):
                        item.smooth = bool(source.get("smooth"))
            _set_gap_policy(native)
        if rows:
            width = max(len(row) for row in rows)
            if width > 20:
                raise DomainError(422, "ppt_table_too_wide", "PPT表格超过20列，请减少导出字段")
            header, body = rows[0], rows[1:]
            chunks = [body[i:i + 18] for i in range(0, len(body), 18)] or [[]]
            if len(deck.slides) + len(chunks) > 100:
                raise DomainError(422, "ppt_slide_budget_exceeded", "PPT页数超过100页，请缩小数据范围")
            for index, chunk in enumerate(chunks):
                slide = deck.slides.add_slide(deck.slide_layouts[5])
                slide.shapes.title.text = str(section.get("title", title)) + (f"（{index + 1}/{len(chunks)}）" if len(chunks) > 1 else "")
                table = slide.shapes.add_table(len(chunk) + 1, width, Inches(0.6), Inches(1.35), Inches(12.1), Inches(min(5.5, 0.3 * (len(chunk) + 1)))).table
                for r, row in enumerate([header] + chunk):
                    for c in range(width):
                        cell = table.cell(r, c)
                        cell.text = str(row[c]) if c < len(row) else ""
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size = Pt(12 if width <= 8 else 9)
                            paragraph.font.bold = r == 0
                            paragraph.font.color.rgb = RGBColor.from_string("153454")
        elif not chart:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = str(section.get("title", title))
            slide.placeholders[1].text = str(section.get("text", ""))
    output = BytesIO()
    deck.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"


def _set_gap_policy(chart):
    """覆盖Office默认的面积图补零策略，所有原生图统一保留缺值断点。"""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    element = chart._chartSpace.find(qn("c:chart"))
    policy = element.find(qn("c:dispBlanksAs"))
    if policy is None:
        policy = OxmlElement("c:dispBlanksAs")
        following = element.find(qn("c:showDLblsOverMax"))
        following = following if following is not None else element.find(qn("c:extLst"))
        if following is not None:
            element.insert(list(element).index(following), policy)
        else:
            element.append(policy)
    policy.set("val", "gap")


def numeric_chart(widget, rows):
    """根据明确绑定或name/value识别图表，不能把无法确认的字段硬编成数据。"""
    if not rows:
        return None
    from .chart_data import CARTESIAN_COMPONENTS, series_data
    component = widget.get("component", widget.get("type", ""))
    component = {"line": "JLine", "bar": "JBar"}.get(component, component)
    if component in CARTESIAN_COMPONENTS:
        model = series_data({**widget, "component": component}, rows)
        if model["status"] != "ready" or len(model["categories"]) > 200:
            return None
        series = [{"name": item["name"], "values": [None if value is None else float(value) for value in item["values"]], "axis": item["axis"], "kind": item["kind"], "smooth": item["smooth"], "step": item["step"]} for item in model["series"]]
        kind = {"JStackBar": "stacked", "JPercentBar": "percent", "JMixLineBar": "mixed", "JMultipleLine": "line", "JSmoothLine": "line", "JStepLine": "step", "JLine": "line", "JArea": "area", "JHorizontalBar": "horizontal_bar"}.get(component, "bar")
        return {"type": kind, "categories": model["categories"], "values": series[0]["values"], "series": series}
    binding = widget.get("binding") or {}
    value_field = binding.get("value_field", "value")
    label_field = binding.get("label_field", "name")
    if value_field not in rows[0]:
        return None
    categories, values = [], []
    for i, row in enumerate(rows):
        try:
            value = float(row[value_field])
            if not math.isfinite(value):
                return None
        except (ValueError, TypeError):
            return None
        categories.append(str(row.get(label_field, i + 1)))
        values.append(value)
    if len(values) > 200:
        return None
    component = widget.get("component", widget.get("type", ""))
    type = {"JLine": "line", "JArea": "area", "JPie": "pie", "JRing": "ring", "JRadar": "radar", "line": "line"}.get(component, "bar")
    return {"type": type, "categories": categories, "values": values}


def _split_line_plots(chart, series):
    """把组合图的真实系列分入原生柱/折线图，显式右轴才创建第二数值轴。"""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    bars = area.find(qn("c:barChart"))
    axis_ids = [item.get("val") for item in bars.findall(qn("c:axId"))]
    line_groups = {}
    for source, element in zip(series, list(bars.findall(qn("c:ser")))):
        if source.get("kind") != "line":
            continue
        axis = source.get("axis", 0)
        if axis not in line_groups:
            line = OxmlElement("c:lineChart")
            grouping = OxmlElement("c:grouping")
            grouping.set("val", "standard")
            line.append(grouping)
            line_groups[axis] = line
        bars.remove(element)
        inverse = element.find(qn("c:invertIfNegative"))
        if inverse is not None:
            element.remove(inverse)
        smooth = OxmlElement("c:smooth")
        smooth.set("val", "1" if source.get("smooth") else "0")
        element.append(smooth)
        line_groups[axis].append(element)
    primary_axis = area.find(qn("c:valAx"))
    secondary_id = str(int(axis_ids[1]) + 113)
    if 1 in line_groups:
        secondary = deepcopy(primary_axis)
        secondary.find(qn("c:axId")).set("val", secondary_id)
        secondary.find(qn("c:axPos")).set("val", "r")
        crosses = secondary.find(qn("c:crosses"))
        if crosses is not None:
            crosses.set("val", "max")
        area.append(secondary)
    for axis, line in line_groups.items():
        for value in (axis_ids[0], secondary_id if axis else axis_ids[1]):
            element = OxmlElement("c:axId")
            element.set("val", value)
            line.append(element)
        area.insert(list(area).index(bars) + 1, line)


def _add_step_chart(slide, chart):
    """单XY系列用空y点断线，保留同字段颜色/图例；标签沿固定绘图区映射。"""
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    data = XyChartData()
    real_points = []
    for source in chart["series"]:
        series = data.add_series(source["name"])
        actual_indexes = []
        count = 0
        previous = None
        for index, value in enumerate(source["values"]):
            if value is None:
                # python-pptx保留缓存ptCount并省略这个索引的y值，工作簿同一格为空。
                series.add_data_point(index, None)
                count += 1
                previous = None
                continue
            if previous is not None:
                old_x, old_y = previous
                step = source.get("step", "end")
                if step == "start":
                    series.add_data_point(old_x, value)
                    count += 1
                elif step == "middle":
                    middle = (old_x + index) / 2
                    series.add_data_point(middle, old_y)
                    series.add_data_point(middle, value)
                    count += 2
                elif step:
                    series.add_data_point(index, old_y)
                    count += 1
            series.add_data_point(index, value)
            actual_indexes.append(count)
            count += 1
            previous = (index, value)
        real_points.append(actual_indexes)
    left, top, width, height = _STEP_CHART_FRAME
    native = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, *(Inches(value) for value in _STEP_CHART_FRAME), data).chart
    single = len(chart["categories"]) == 1
    native.category_axis.minimum_scale = -.5 if single else 0
    native.category_axis.maximum_scale = .5 if single else max(1, len(chart["categories"]) - 1)
    native.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    native.legend.position = XL_LEGEND_POSITION.TOP
    native.legend.include_in_layout = False
    _set_step_plot_layout(native)
    # 只标实际采样点，转折辅助点不添加标记；断点两侧的孤立样本也仍可见。
    for series, indexes in zip(native.series, real_points):
        series.marker.style = XL_MARKER_STYLE.NONE
        for index in indexes:
            series.points[index].marker.style = XL_MARKER_STYLE.CIRCLE
            series.points[index].marker.size = 4
    if len(chart["categories"]) <= 20:
        plot_x, plot_y, plot_w, plot_h = _STEP_PLOT_LAYOUT
        plot_left, plot_width = left + width * plot_x, width * plot_w
        label_width = min(1.25, plot_width / max(1, len(chart["categories"])))
        for index, label in enumerate(chart["categories"]):
            ratio = .5 if single else index / max(1, len(chart["categories"]) - 1)
            center = plot_left + plot_width * ratio
            box = slide.shapes.add_textbox(Inches(center - label_width / 2), Inches(top + height * (plot_y + plot_h) + .08), Inches(label_width), Inches(.42))
            box.text_frame.margin_left = box.text_frame.margin_right = 0
            box.text_frame.word_wrap = True
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = str(label)
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.size = Pt(10 if len(chart["categories"]) <= 10 else 8)
    return native


def _set_step_plot_layout(chart):
    """固定原生XY内部绘图区，使外置业务分类标签与x轴索引使用同一几何口径。"""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    layout = area.find(qn("c:layout"))
    if layout is None:
        layout = OxmlElement("c:layout")
        area.insert(0, layout)
    manual = OxmlElement("c:manualLayout")
    for name, value in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("wMode", "factor"), ("hMode", "factor"), *zip(("x", "y", "w", "h"), _STEP_PLOT_LAYOUT)):
        item = OxmlElement("c:" + name)
        item.set("val", str(value))
        manual.append(item)
    for item in list(layout):
        layout.remove(item)
    layout.append(manual)
