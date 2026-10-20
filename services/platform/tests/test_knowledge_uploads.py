"""原件摄取闭环与撤权边界，仅使用新临时SQLite/文件及向量替身。"""

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import sys
import zipfile

import pytest
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from fastapi.testclient import TestClient
from sqlalchemy import event, func, select

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from agent_platform.core.context import Actor, get_actor
from agent_platform.core.database import Database
from agent_platform.core.embeddings import EmbeddingClient
from agent_platform.core.errors import DomainError
from agent_platform.core.files import FileService
from agent_platform.core.models import (AuditEvent, FileEntry, Organization, ResourceGrant,
    ResourceReference, User, Workspace)
from agent_platform.core.settings import Settings
from agent_platform.modules.files.router import router as file_router
from agent_platform.modules.knowledge import extractor
from agent_platform.modules.knowledge.evidence import authorize_evidence
from agent_platform.modules.knowledge.ingestion import KnowledgeIngestion, parse_in_worker
from agent_platform.modules.knowledge.models import KnowledgeChunk, KnowledgeDocument
from agent_platform.modules.knowledge.router import router
from agent_platform.modules.knowledge.service import KnowledgeService


@pytest.fixture
def site(tmp_path, monkeypatch):
    settings = Settings(data_dir=tmp_path / 'new-kb', environment='test')
    database = Database(settings)
    database.initialize()
    actor = [Actor('u1', 'w1', frozenset({'m08.*'}))]
    calls = []
    def embed(self, texts, **kwargs):
        calls.append({'texts': list(texts), **kwargs})
        return [[1.0, 0.0] for _ in texts]
    monkeypatch.setattr(EmbeddingClient, 'embed', embed)
    # 大部分业务用例用纯解析器替身；另有独立真实子进程用例核超时边界/协议。
    monkeypatch.setattr('agent_platform.modules.knowledge.ingestion.parse_in_worker',
        lambda path, name, timeout: extractor.extract_file(name, path))
    with database.sessions.begin() as session:
        session.add(Organization(id='o1', name='本轮临时企业'))
        session.flush()
        session.add_all([Workspace(id='w1', organization_id='o1', name='甲'),
            Workspace(id='w2', organization_id='o1', name='乙'),
            User(id='u1', organization_id='o1', username='u1', display_name='甲用户', password_hash='unused'),
            User(id='u2', organization_id='o1', username='u2', display_name='乙用户', password_hash='unused')])
        session.flush()
        space = KnowledgeService(session, actor[0]).create_space('设备规程')
    app = FastAPI()
    app.state.settings, app.state.database = settings, database
    app.state.secret_vault = type('Vault', (), {'sign': lambda self, value: 'test-signature'})()
    app.dependency_overrides[get_actor] = lambda: actor[0]
    app.include_router(router)
    app.include_router(file_router)
    @app.exception_handler(DomainError)
    def error(request, exc):
        return JSONResponse(status_code=exc.status, content={'error': {'code': exc.code, 'message': exc.message}})
    with TestClient(app, raise_server_exceptions=True) as client:
        yield type('Site', (), {'client': client, 'database': database, 'settings': settings,
            'actor': actor, 'space': space, 'calls': calls})()
    database.dispose()


def upload(site, text='停机后检查轴承，温度上限75℃。', filename='设备规程.txt', data=None):
    content = data if data is not None else text.encode('utf-8')
    response = site.client.post('/api/v1/kb/spaces/' + site.space['id'] + '/uploads',
        files={'file': (filename, content, 'application/octet-stream')})
    assert response.status_code == 201, response.text
    return response.json()


def test_knowledge_only_user_can_upload_search_locate_and_download(site):
    site.actor[0] = Actor('u1', 'w1', frozenset({'m08.write'}))
    document = upload(site, '样本,温度\nA,75\nB,80', filename='温度.csv')
    assert document['status'] == 'ready' and document['source_file']['size'] > 0
    assert document['created_at'].endswith('+00:00')
    assert len(site.calls) == 1 and site.calls[0]['timeout_seconds'] <= 60
    site.actor[0] = Actor('u2', 'w1', frozenset({'m08.read'}))
    detail = site.client.get('/api/v1/kb/documents/' + document['id']).json()
    assert detail['source_file']['id'] == document['source_file']['id']
    chunks = site.client.get('/api/v1/kb/documents/' + document['id'] + '/chunks?limit=1&offset=1').json()
    assert chunks['total'] == 3 and chunks['items'][0]['source_location'] == {'row': 2}
    found = site.client.post('/api/v1/kb/search', json={'space_ids': [site.space['id']], 'query': '温度'}).json()
    assert found['items'][0]['source_file']['sha256'] == document['source_file']['sha256']
    original = site.client.get('/api/v1/kb/documents/' + document['id'] + '/source')
    assert original.status_code == 200 and original.content.decode() == '样本,温度\nA,75\nB,80'
    assert 'attachment' in original.headers['content-disposition']
    assert original.headers['cache-control'] == 'private, no-store'


def test_general_file_permission_cannot_bypass_knowledge_permission_even_owner(site):
    document = upload(site)
    file_id = document['source_file']['id']
    site.actor[0] = Actor('u1', 'w1', frozenset({'m17.read'}))
    listing = site.client.get('/api/v1/files').json()
    assert listing == {'items': [], 'total': 0}
    for path, method in [(f'/files/{file_id}', 'get'), (f'/files/{file_id}/download-links', 'post')]:
        assert getattr(site.client, method)('/api/v1' + path).status_code == 404
    assert site.client.get(f'/api/v1/files/{file_id}/content?expires=9999999999&signature=x').status_code == 403
    assert site.client.get('/api/v1/kb/documents/' + document['id'] + '/source').status_code == 403


def test_private_existing_file_is_not_shared_by_import_and_history_rechecks_access(site):
    with site.database.sessions.begin() as session:
        private = FileService(session, site.actor[0], site.settings).store_bytes('个人规程.txt', b'private rule', 'text/plain')
    response = site.client.post('/api/v1/kb/spaces/' + site.space['id'] + '/file-imports', json={'file_id': private['id']})
    assert response.status_code == 201 and response.json()['status'] == 'ready'
    document = response.json()
    query = {'space_ids': [site.space['id']], 'query': 'rule'}
    evidence = site.client.post('/api/v1/kb/search', json=query).json()['knowledge_sources']
    assert evidence['chunks']
    site.actor[0] = Actor('u2', 'w1', frozenset({'m08.read'}))
    assert site.client.get('/api/v1/kb/spaces/' + site.space['id'] + '/documents').json() == {'items': [], 'total': 0}
    assert site.client.get('/api/v1/kb/spaces').json()['items'][0]['document_count'] == 0
    before = len(site.calls)
    assert site.client.post('/api/v1/kb/search', json=query).json()['items'] == []
    assert len(site.calls) == before
    assert site.client.get('/api/v1/kb/documents/' + document['id'] + '/source').status_code == 404
    with site.database.sessions() as session:
        with pytest.raises(DomainError) as denied:
            authorize_evidence(session, site.actor[0], evidence)
        assert denied.value.status == 403
    # 明确授权原件后知识可见，撤销该授权后旧结果立即不可再读。
    with site.database.sessions.begin() as session:
        grant = ResourceGrant(workspace_id='w1', user_id='u2', resource_type='file',
            resource_id=private['id'], actions=['read'])
        session.add(grant)
    assert site.client.get('/api/v1/kb/spaces/' + site.space['id'] + '/documents').json()['total'] == 1
    with site.database.sessions.begin() as session:
        session.delete(session.get(ResourceGrant, grant.id))
    assert site.client.get('/api/v1/kb/documents/' + document['id']).status_code == 404


def test_cross_workspace_and_unbound_original_are_denied(site):
    document = upload(site)
    site.actor[0] = Actor('u2', 'w2', frozenset({'*'}), is_admin=True)
    assert site.client.get('/api/v1/kb/documents/' + document['id'] + '/source').status_code == 404
    site.actor[0] = Actor('u1', 'w1', frozenset({'m08.*'}))
    with site.database.sessions.begin() as session:
        reference = session.scalar(select(ResourceReference).where(ResourceReference.source_id == document['id']))
        session.delete(reference)
    assert site.client.get('/api/v1/kb/documents/' + document['id']).status_code == 403
    assert site.client.get('/api/v1/kb/spaces/' + site.space['id'] + '/documents').json()['total'] == 0


@pytest.mark.parametrize('filename,data,code', [
    ('broken.pdf', b'%PDF-1.4\nbroken', 'document_corrupt'),
    ('empty.txt', b'  \n ', 'document_no_text'),
    ('binary.txt', b'abc\x00def', 'document_binary_text'),
    ('legacy.docx', b'not a zip', 'document_corrupt')])
def test_parse_failure_keeps_original_and_failed_document_without_embedding(site, filename, data, code):
    document = upload(site, filename=filename, data=data)
    assert document['status'] == 'failed' and document['error'] and document['error_code'] == code
    assert document['chunk_count'] == 0 and site.calls == []
    assert site.client.get('/api/v1/kb/documents/' + document['id'] + '/source').content == data
    assert site.client.post('/api/v1/kb/documents/' + document['id'] + '/index-jobs').status_code == 422


def test_embedding_failure_persists_chunks_and_explicit_index_is_only_retry(site, monkeypatch):
    count = []
    def fail(self, texts, **kwargs):
        count.append(texts)
        raise DomainError(503, 'embedding_unavailable', '无法连接向量模型')
    monkeypatch.setattr(EmbeddingClient, 'embed', fail)
    document = upload(site)
    assert document['status'] == 'failed' and document['error_code'] == 'embedding_unavailable'
    assert document['chunk_count'] == 1 and len(count) == 1
    assert site.client.get('/api/v1/kb/documents/' + document['id'] + '/chunks').json()['total'] == 1
    assert site.client.post('/api/v1/kb/search', json={'space_ids': [site.space['id']], 'query': '温度'}).json()['items'] == []
    assert len(count) == 1
    monkeypatch.setattr(EmbeddingClient, 'embed', lambda self, texts, **kw: [[1.0, 0.0] for _ in texts])
    retried = site.client.post('/api/v1/kb/documents/' + document['id'] + '/index-jobs').json()
    assert retried['status'] == 'ready' and retried['error'] is None


@pytest.mark.parametrize('vectors', [[], [[float('nan')]], [[True]], [[1], [2]]])
def test_invalid_embedding_is_never_marked_ready(site, monkeypatch, vectors):
    monkeypatch.setattr(EmbeddingClient, 'embed', lambda *args, **kw: vectors)
    document = upload(site)
    assert document['status'] == 'failed' and document['error_code'] == 'embedding_invalid_response'


def test_delete_last_original_reference_makes_file_private_and_preserves_other_refs(site):
    document = upload(site)
    file_id = document['source_file']['id']
    with site.database.sessions.begin() as session:
        session.add(ResourceReference(workspace_id='w1', source_type='unrelated_record', source_id='other',
            target_type='file', target_id=file_id))
    assert site.client.delete('/api/v1/kb/documents/' + document['id']).status_code == 200
    with site.database.sessions() as session:
        file = session.get(FileEntry, file_id)
        assert file.visibility == 'private' and file.revision == 2 and not file.deleted
        refs = list(session.scalars(select(ResourceReference).where(ResourceReference.target_id == file_id)))
        assert len(refs) == 1 and refs[0].source_type == 'unrelated_record'
        assert FileService(session, Actor('u1', 'w1', frozenset({'m17.read'})), site.settings).path(file_id).is_file()
    site.actor[0] = Actor('u2', 'w1', frozenset({'m17.read'}))
    assert site.client.get('/api/v1/files').json()['total'] == 0


def test_deleting_one_of_multiple_knowledge_links_does_not_change_file_revision(site):
    document = upload(site)
    with site.database.sessions.begin() as session:
        space = KnowledgeService(session, site.actor[0]).create_space('另一知识空间')
    second = site.client.post('/api/v1/kb/spaces/' + space['id'] + '/file-imports',
        json={'file_id': document['source_file']['id']}).json()
    assert second['status'] == 'ready'
    site.client.delete('/api/v1/kb/documents/' + document['id'])
    with site.database.sessions() as session:
        assert session.get(FileEntry, second['source_file']['id']).revision == 1
    assert site.client.get('/api/v1/kb/documents/' + second['id'] + '/source').status_code == 200


def test_duplicate_file_import_does_not_reembed(site):
    document = upload(site)
    count = len(site.calls)
    result = site.client.post('/api/v1/kb/spaces/' + site.space['id'] + '/file-imports',
        json={'file_id': document['source_file']['id']}).json()
    assert result['id'] == document['id'] and len(site.calls) == count


def test_index_batches_are_bounded_and_old_unknown_reservation_requires_explicit_action(site):
    document = upload(site, '\n'.join('样本' + str(index) for index in range(65)), filename='样本.csv')
    assert document['status'] == 'ready'
    assert [len(call['texts']) for call in site.calls] == [32, 32, 1]
    with site.database.sessions.begin() as session:
        row = session.get(KnowledgeDocument, document['id'])
        row.status = 'indexing'
        row.parser_metadata = {**row.parser_metadata, 'index_attempt_id': 'previous-unknown',
            'index_deadline_at': '2000-01-01T00:00:00+00:00'}
    before = len(site.calls)
    assert site.client.get('/api/v1/kb/documents/' + document['id']).json()['status'] == 'indexing'
    assert len(site.calls) == before
    result = site.client.post('/api/v1/kb/documents/' + document['id'] + '/index-jobs').json()
    assert result['status'] == 'ready' and len(site.calls) == before + 3
    assert result['parser']['previous_unknown_attempt_id'] == 'previous-unknown'


def test_active_index_reservation_never_dispatches_a_second_request(site):
    document = upload(site)
    with site.database.sessions.begin() as session:
        row = session.get(KnowledgeDocument, document['id'])
        row.status = 'indexing'
        row.parser_metadata = {**row.parser_metadata,
            'index_deadline_at': '2999-01-01T00:00:00+00:00'}
    before = len(site.calls)
    assert site.client.post('/api/v1/kb/documents/' + document['id'] + '/index-jobs').status_code == 409
    assert len(site.calls) == before


def test_upload_limits_reject_before_creating_files_and_source_version_is_frozen(site):
    for filename, content in [('unsupported.zip', b'zip'), ('empty.txt', b'')]:
        result = site.client.post('/api/v1/kb/spaces/' + site.space['id'] + '/uploads',
            files={'file': (filename, content)})
        assert result.status_code in {413, 422}
    with site.database.sessions() as session:
        assert session.scalar(select(func.count()).select_from(FileEntry)) == 0
        assert session.scalar(select(func.count()).select_from(KnowledgeDocument)) == 0
    document = upload(site)
    evidence = site.client.post('/api/v1/kb/search', json={
        'space_ids': [site.space['id']], 'query': '轴承'}).json()['knowledge_sources']
    assert evidence['chunks'][0]['source_sha256'] == document['source_file']['sha256']
    with site.database.sessions.begin() as session:
        session.get(FileEntry, document['source_file']['id']).revision += 1
    assert site.client.get('/api/v1/kb/documents/' + document['id'] + '/source').status_code == 403
    assert site.client.get('/api/v1/kb/spaces/' + site.space['id'] + '/documents').json()['total'] == 0
    with site.database.sessions() as session:
        with pytest.raises(DomainError) as denied:
            authorize_evidence(session, site.actor[0], evidence)
        assert denied.value.code == 'knowledge_evidence_access_denied'


def test_new_file_is_removed_if_audit_or_sql_flush_fails(site, monkeypatch):
    from agent_platform.core.records import RecordStore
    with site.database.sessions() as session:
        existing = FileService(session, site.actor[0], site.settings).store_bytes('保留.txt', b'keep', 'text/plain')
        session.commit()
    before = set(site.settings.files_dir.rglob('*'))
    def fail(*args, **kwargs):
        raise RuntimeError('simulated audit failure')
    monkeypatch.setattr(RecordStore, 'audit', fail)
    with site.database.sessions() as session:
        with pytest.raises(RuntimeError, match='audit failure'):
            FileService(session, site.actor[0], site.settings).store_bytes('失败.txt', b'failed', 'text/plain')
        session.rollback()
    assert set(site.settings.files_dir.rglob('*')) == before
    with site.database.sessions() as session:
        assert FileService(session, site.actor[0], site.settings).path(existing['id']).read_bytes() == b'keep'


def test_sql_insert_failure_removes_only_this_new_original(site):
    with site.database.sessions.begin() as session:
        existing = FileService(session, site.actor[0], site.settings).store_bytes('原件.txt', b'keep', 'text/plain')
    def reject(mapper, connection, target):
        if target.name == '失败.txt':
            raise RuntimeError('simulated insert failure')
    event.listen(FileEntry, 'before_insert', reject)
    try:
        with site.database.sessions() as session:
            with pytest.raises(RuntimeError, match='insert failure'):
                FileService(session, site.actor[0], site.settings).store_bytes('失败.txt', b'bad', 'text/plain')
            session.rollback()
    finally:
        event.remove(FileEntry, 'before_insert', reject)
    assert [path.name for path in site.settings.files_dir.rglob('*') if path.is_file()] == [existing['id']]


def test_migration10_preserves_existing_text_and_adds_only_nullable_sources(tmp_path):
    from alembic import command
    from sqlalchemy import inspect, text
    from tools.manage_database import config
    settings = Settings(data_dir=tmp_path / 'new-migration10', environment='test')
    database = Database(settings)
    try:
        with database.engine.begin() as connection:
            command.upgrade(config(connection), '20261021_09')
            connection.execute(text("INSERT INTO kb_documents (id,workspace_id,space_id,name,source,"
                "char_count,chunk_count,status,error,created_by,created_at) VALUES "
                "('kept','w','s','原规程','paste',3,1,'ready',NULL,'u','2026-10-21')"))
            command.upgrade(config(connection), '20261021_10')
            command.upgrade(config(connection), '20261021_10')
            row = connection.execute(text('SELECT name,file_id,parser_metadata FROM kb_documents WHERE id=:id'),
                {'id': 'kept'}).one()
            assert row == ('原规程', None, None)
            columns = {column['name']: column for column in inspect(connection).get_columns('kb_documents')}
            assert all(columns[name]['nullable'] for name in (
                'file_id', 'source_sha256', 'source_revision', 'parser_metadata'))
            assert 'source_location' in {column['name'] for column in inspect(connection).get_columns('kb_chunks')}
    finally:
        database.dispose()


def test_office_zip_expansion_budget_and_text_budget_are_explicit(tmp_path, monkeypatch):
    path = tmp_path / 'large.docx'
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('many.xml', 'x' * 10000)
    monkeypatch.setattr(extractor, 'MAX_ZIP_BYTES', 9000)
    with pytest.raises(extractor.ExtractionError) as error:
        extractor.extract_file(path.name, path)
    assert error.value.code == 'document_expansion_limit'
    text = tmp_path / 'too-long.txt'
    text.write_text('文' * 201, encoding='utf-8')
    monkeypatch.setattr(extractor, 'MAX_TEXT_CHARS', 200)
    with pytest.raises(extractor.ExtractionError) as error:
        extractor.extract_file(text.name, text)
    assert error.value.code == 'document_text_limit'


def test_real_office_parsers_keep_source_order_and_formula(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    docx = tmp_path / '规程.docx'
    document = Document()
    document.add_paragraph('表前段落')
    document.add_table(1, 1).cell(0, 0).text = '表格数据'
    document.add_paragraph('表后段落')
    document.save(docx)
    blocks = extractor.extract_file(docx.name, docx)['blocks']
    assert [block['text'] for block in blocks] == ['表前段落', '表格数据', '表后段落']
    assert blocks[1]['location'] == {'table': 1, 'row': 1} and 'page' not in blocks[0]['location']
    xlsx = tmp_path / '指标.xlsx'
    workbook = Workbook()
    workbook.active.title = '能耗'
    workbook.active.append(['产量', '=1+2'])
    workbook.save(xlsx)
    result = extractor.extract_file(xlsx.name, xlsx)
    assert result['blocks'][0]['text'] == '产量\t=1+2'
    assert result['blocks'][0]['location'] == {'sheet': '能耗', 'row': 1}
    assert result['metadata']['warnings'][0]['code'] == 'formulas_as_text'
    pptx = tmp_path / '培训.pptx'
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = '安全培训'
    presentation.save(pptx)
    assert extractor.extract_file(pptx.name, pptx)['blocks'][0]['location'] == {'slide': 1}


def test_real_pdf_text_and_scanned_page_are_not_faked(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    path = tmp_path / '规程.pdf'
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=400)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
        DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 20 300 Td (Pump temperature 75 C) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.add_blank_page(width=300, height=400)
    writer.write(path)
    result = extractor.extract_file(path.name, path)
    assert 'temperature 75' in result['blocks'][0]['text']
    assert result['blocks'][0]['location'] == {'page': 1}
    assert result['metadata']['warnings'][0]['pages'] == [2]
    blank = tmp_path / 'scanned.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.write(blank)
    with pytest.raises(extractor.ExtractionError) as error:
        extractor.extract_file(blank.name, blank)
    assert error.value.code == 'document_no_text'


def test_real_isolated_parser_protocol_and_timeout_cleanup(tmp_path, monkeypatch):
    path = tmp_path / 'text.txt'
    path.write_text('真实子进程解析', encoding='utf-8')
    result = parse_in_worker(path, path.name, 5)
    assert result['ok'] and result['blocks'][0]['text'] == '真实子进程解析'
    import subprocess
    def timeout(*args, **kwargs):
        assert kwargs['timeout'] <= 20
        assert not any('TOKEN' in key or 'PASSWORD' in key for key in kwargs['env'])
        raise subprocess.TimeoutExpired(args[0], kwargs['timeout'])
    monkeypatch.setattr(subprocess, 'run', timeout)
    with pytest.raises(DomainError) as error:
        parse_in_worker(path, path.name, 0.2)
    assert error.value.code == 'document_parse_timeout'


def test_worker_invalid_top_level_has_safe_failure(tmp_path, monkeypatch):
    import subprocess
    from types import SimpleNamespace
    monkeypatch.setattr(subprocess, 'run', lambda *a, **kw:
        SimpleNamespace(returncode=0, stdout=b'[]', stderr=b''))
    with pytest.raises(DomainError) as error:
        parse_in_worker(tmp_path / 'controlled.txt', 'controlled.txt', 1)
    assert error.value.code == 'document_parse_failed'


def test_database_error_during_parse_persistence_does_not_query_or_compensate(site, monkeypatch):
    from sqlalchemy.exc import OperationalError
    with site.database.sessions() as session:
        service = KnowledgeService(session, site.actor[0], settings=site.settings)
        counts = {'failed_record': 0, 'sql_after_error': 0}
        failed = [False]
        def before_cursor(connection, cursor, statement, parameters, context, executemany):
            if failed[0]:
                counts['sql_after_error'] += 1
        event.listen(site.database.engine, 'before_cursor_execute', before_cursor)
        def parser(path, name, timeout):
            result = extractor.extract_file(name, path)
            def database_failure():
                failed[0] = True
                raise OperationalError('simulated statement', {}, RuntimeError('new tmp only'))
            monkeypatch.setattr(session, 'commit', database_failure)
            return result
        ingestion = KnowledgeIngestion(service, site.settings, parser=parser)
        monkeypatch.setattr(ingestion, '_failure', lambda *a, **kw: counts.__setitem__(
            'failed_record', counts['failed_record'] + 1))
        try:
            with pytest.raises(OperationalError):
                ingestion.upload(site.space['id'], '数据库失败.txt', b'real retained original', 'text/plain')
        finally:
            event.remove(site.database.engine, 'before_cursor_execute', before_cursor)
        assert counts == {'failed_record': 0, 'sql_after_error': 0}
    assert len([path for path in site.settings.files_dir.rglob('*') if path.is_file()]) == 1


def test_encrypted_pdf_is_rejected_before_any_text_is_indexed(tmp_path):
    from pypdf import PdfWriter
    path = tmp_path / 'encrypted.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.encrypt('not-saved')
    writer.write(path)
    with pytest.raises(extractor.ExtractionError) as error:
        extractor.extract_file(path.name, path)
    assert error.value.code == 'document_encrypted'
