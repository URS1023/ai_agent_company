"""图表变体使用真实宽表字段，验证数值口径、缺失、几何与防伪模型。"""
from io import BytesIO
import re

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.chart import XL_MARKER_STYLE
from pptx.oxml.ns import qn
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from agent_platform.core.context import Actor
from agent_platform.core.database import Base
from agent_platform.core.settings import Settings
from agent_platform.modules.data.service import DataService
from agent_platform.modules.reports.chart_data import public_series_model, series_data
from agent_platform.modules.reports.exports import numeric_chart, office_export
from agent_platform.modules.reports.rendering import render_html
from agent_platform.modules.reports.service import ReportService


def widget(component, **binding):
    return {"id": "wide", "component": component, "w": 640, "h": 280, "binding": {"label_field": "shift", "value_fields": ["good", "bad"], **binding}}


def test_percentage_normalizes_each_category_and_preserves_raw_values():
    rows = [{"shift": "一班", "good": "1", "bad": "3"}, {"shift": "二班", "good": "6", "bad": "4"}, {"shift": "缺数据", "good": "7", "bad": None}, {"shift": "零产量", "good": "0", "bad": "0"}]
    model = public_series_model(widget("JPercentBar"), rows)
    assert model["status"] == "ready" and model["percentage"]
    assert model["categories"] == ["一班", "二班", "缺数据", "零产量"]
    assert model["series"][0]["values"] == ["25.00", "60.0", None, None]
    assert model["series"][1]["values"] == ["75.00", "40.0", None, None]
    assert model["series"][0]["raw_values"] == ["1", "6", "7", "0"]
    invalid = public_series_model(widget("JPercentBar"), [{"shift": "退货", "good": "-1", "bad": "2"}])
    assert invalid["status"] == "invalid_data"


def test_grouped_stacked_and_combination_keep_axis_and_gap_semantics():
    rows = [{"shift": "一班", "good": "9007199254740993", "bad": "-2"}, {"shift": "二班", "good": None, "bad": "3"}]
    grouped = public_series_model(widget("JMultipleBar"), rows)
    assert grouped["series"][0]["raw_values"][0] == "9007199254740993"
    assert all(series["axis"] == "primary" and series["type"] == "bar" and "stack" not in series for series in grouped["series"])
    stacked = public_series_model(widget("JStackBar"), rows)
    assert all(series["stack"] == "total" and series["axis"] == "primary" for series in stacked["series"])
    mixed = public_series_model(widget("JMixLineBar", secondary_value_fields=["bad"]), rows)
    assert [(series["type"], series["axis"]) for series in mixed["series"]] == [("bar", "primary"), ("line", "secondary")]
    assert mixed["series"][0]["values"][1] is None
    assert series_data(widget("JMultipleLine", secondary_value_fields=["bad"]), rows)["status"] == "needs_mapping"
    assert series_data(widget("JMixLineBar", secondary_value_fields=["good"]), rows)["status"] == "needs_mapping"
    assert series_data(widget("JStackBar", value_fields=["absent"]), rows)["status"] == "needs_mapping"
    assert series_data(widget("JMultipleBar"), [{"good": 3, "bad": 4}])["status"] == "needs_mapping"


def test_svg_distinguishes_stack_group_percent_mix_smooth_step_and_negative():
    rows = [{"shift": "一班", "good": "1", "bad": "3"}, {"shift": "二班", "good": "6", "bad": "4"}]
    def rendered(component, **binding):
        return render_html({"name": "真实变体", "layout": {"canvas": {}, "widgets": [{**widget(component, **binding), "data": rows}]}})
    grouped, stacked = rendered("JMultipleBar"), rendered("JStackBar")
    group_x = re.findall(r'<rect x="([^"]+)"[^>]*data-value=', grouped)
    stack_x = re.findall(r'<rect x="([^"]+)"[^>]*data-value=', stacked)
    assert group_x[0] != group_x[2] and stack_x[0] == stack_x[2]
    percent = rendered("JPercentBar")
    assert 'data-value="25.00"' in percent and 'data-raw-value="1"' in percent and "100%" in percent
    assert 'data-kind="line"' in rendered("JMixLineBar", secondary_value_fields=["bad"]) and 'data-axis="1"' in rendered("JMixLineBar", secondary_value_fields=["bad"])
    smooth, step = rendered("JSmoothLine"), rendered("JStepLine")
    assert re.search(r'<path d="M[^\"]+C', smooth)
    assert re.search(r'<path d="M[^\"]+H[^\"]+V', step)
    rows[0]["good"] = "-1"
    negative = rendered("JNegativeBar", value_fields=["good"])
    assert 'data-value="-1"' in negative and 'height="' in negative
    # 缺点形成分段而非连接断点，更不能用0补点。
    rows.insert(1, {"shift": "缺测", "good": None, "bad": "2"})
    gap = rendered("JMultipleLine", value_fields=["good"])
    assert gap.count('<polyline points=') == 2 and 'data-value="0"' not in gap


def test_only_verified_query_evidence_generates_series_model(tmp_path):
    engine = create_engine("sqlite:///" + (tmp_path / "series.sqlite3").as_posix())
    Base.metadata.create_all(engine)
    actor = Actor("series-owner", "workspace", frozenset({"*"}))
    settings = Settings(data_dir=tmp_path / "series-data")
    with Session(engine) as session:
        data, reports = DataService(session, actor, settings), ReportService(session, actor, settings)
        dataset = data.import_csv({"name": "班组产量", "csv_text": "name,value\n一班,17\n二班,23\n"})
        query = data.execute_query({"datasource_id": dataset["datasource_id"], "sql": f'SELECT * FROM "{dataset["table_name"]}"'})
        report = reports.create({"name": "不信任客户端模型", "layout": {"canvas": {}, "widgets": [{"id": "trend", "component": "JBar", "binding": {"query_id": query["id"]}, "series_model": {"status": "ready", "series": [{"values": [9999]}]}}]}})
        assert "series_model" not in report["layout"]["widgets"][0]
        refreshed = reports.refresh(report["id"], report["revision"])
        model = refreshed["layout"]["widgets"][0]["series_model"]
        assert model["series"][0]["values"] == ["17", "23"]
        refreshed["layout"]["widgets"][0]["series_model"]["series"][0]["values"] = ["9999"]
        patched = reports.patch(report["id"], {"revision": refreshed["revision"], "layout": refreshed["layout"]})
        assert patched["layout"]["widgets"][0]["series_model"]["series"][0]["values"] == ["17", "23"]
        assert reports.publish(report["id"], patched["revision"])["status"] == "published"
        session.rollback()
    engine.dispose()


def test_native_ppt_preserves_multiseries_stack_percent_mix_and_step():
    rows = [{"shift": "一班", "good": "1", "bad": "3"}, {"shift": "二班", "good": "6", "bad": "4"}]
    def native(component, **binding):
        chart = numeric_chart(widget(component, **binding), rows)
        content, _, _ = office_export("pptx", component, [{"title": component, "chart": chart, "rows": [["shift", "good", "bad"], ["一班", "1", "3"], ["二班", "6", "4"]]}])
        deck = Presentation(BytesIO(content))
        return next(shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart)
    grouped = native("JMultipleBar")
    assert len(grouped.series) == 2 and grouped.series[0].values == (1, 6) and grouped.series[1].values == (3, 4)
    assert native("JStackBar").chart_type == XL_CHART_TYPE.COLUMN_STACKED
    percent = native("JPercentBar")
    assert percent.chart_type == XL_CHART_TYPE.COLUMN_STACKED_100
    assert percent.series[0].values == (25, 60) and percent.series[1].values == (75, 40)
    mixed = native("JMixLineBar", secondary_value_fields=["bad"])
    assert len(mixed.plots) == 2 and len(mixed.series) == 2
    assert "<c:barChart>" in mixed._chartSpace.xml and "<c:lineChart>" in mixed._chartSpace.xml
    assert mixed._chartSpace.xml.count("<c:valAx>") == 2
    assert mixed.series[0].values == (1, 6) and mixed.series[1].values == (3, 4)
    smooth = native("JSmoothLine")
    assert all(series.smooth for series in smooth.series)
    step = native("JStepLine", value_fields=["good"])
    assert step.chart_type == XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS
    assert "<c:xVal>" in step._chartSpace.xml and "<c:yVal>" in step._chartSpace.xml
    assert step.series[0].values == (1, 1, 6)


def _native_deck(config, rows):
    """纯内存生成小型PPT，不运行Office、浏览器或数据库。"""
    chart = numeric_chart(config, rows)
    content, _, _ = office_export("pptx", "图表边界", [{"title": config["component"], "chart": chart}])
    deck = Presentation(BytesIO(content))
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                return deck, slide, shape
    raise AssertionError("没有产生实际原生图表")


def test_step_option_start_middle_end_is_preserved_in_model_svg_and_ppt():
    rows = [{"shift": "一班", "good": "1"}, {"shift": "二班", "good": "6"}]
    expected = {"start": ([0, 0, 1], (1, 6, 6)), "middle": ([0, .5, .5, 1], (1, 1, 6, 6)), "end": ([0, 1, 1], (1, 1, 6))}
    for step, (xs, ys) in expected.items():
        config = {**widget("JStepLine", value_fields=["good"]), "config": {"option": {"series": [{"step": step}]}}}
        assert public_series_model(config, rows)["series"][0]["step"] == step
        text = render_html({"name": "阶梯位置", "layout": {"canvas": {}, "widgets": [{**config, "data": rows}]}})
        path = re.search(r'<path d="(M[^"]+)" fill="none"', text).group(1)
        assert (path.index("V") < path.index("H")) == (step == "start")
        _, _, shape = _native_deck(config, rows)
        native = shape.chart
        assert native.series[0].values == ys
        actual = [float(value.text) for value in native.series[0]._ser.findall('.//' + qn('c:xVal') + '/' + qn('c:numRef') + '/' + qn('c:numCache') + '/' + qn('c:pt') + '/' + qn('c:v'))]
        assert actual == xs
    object_series = {**widget("JStepLine", value_fields=["good"]), "config": {"option": {"series": {"step": "start"}}}}
    assert public_series_model(object_series, rows)["series"][0]["step"] == "start"


def test_ppt_area_missing_point_is_explicit_gap():
    rows = [{"shift": "一班", "good": "1"}, {"shift": "缺测", "good": None}, {"shift": "二班", "good": "6"}]
    _, _, shape = _native_deck(widget("JArea", value_fields=["good"]), rows)
    native = shape.chart
    assert native.series[0].values == (1, None, 6)
    policy = native._chartSpace.find(qn("c:chart")).find(qn("c:dispBlanksAs"))
    assert policy.get("val") == "gap"
    cache = native.series[0]._ser.find(qn("c:val")).find(qn("c:numRef")).find(qn("c:numCache"))
    assert cache.find(qn("c:ptCount")).get("val") == "3"
    assert [point.get("idx") for point in cache.findall(qn("c:pt"))] == ["0", "2"]


def test_ppt_step_gap_keeps_one_series_real_markers_and_aligned_labels():
    rows = [{"shift": "一班", "good": "1"}, {"shift": "缺测", "good": None}, {"shift": "二班", "good": "6"}]
    _, slide, shape = _native_deck(widget("JStepLine", value_fields=["good"]), rows)
    native = shape.chart
    assert len(native.series) == 1 and native.series[0].values == (1, None, 6)
    assert native._chartSpace.find(qn("c:chart")).find(qn("c:dispBlanksAs")).get("val") == "gap"
    assert native.series[0].points[0].marker.style == XL_MARKER_STYLE.CIRCLE
    assert native.series[0].points[2].marker.style == XL_MARKER_STYLE.CIRCLE
    assert native.series[0].marker.style == XL_MARKER_STYLE.NONE
    manual = native._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")).find(qn("c:layout")).find(qn("c:manualLayout"))
    assert manual.find(qn("c:layoutTarget")).get("val") == "inner"
    left = shape.left + shape.width * float(manual.find(qn("c:x")).get("val"))
    width = shape.width * float(manual.find(qn("c:w")).get("val"))
    labels = {item.text: item for item in slide.shapes if item.has_text_frame and item.text in {"一班", "缺测", "二班"}}
    for name, fraction in (("一班", 0), ("缺测", .5), ("二班", 1)):
        actual = labels[name].left + labels[name].width / 2
        assert abs(actual - (left + width * fraction)) <= 2
    assert native.category_axis.minimum_scale == 0 and native.category_axis.maximum_scale == 2


def test_ppt_single_step_point_is_visible_and_centered():
    _, slide, shape = _native_deck(widget("JStepLine", value_fields=["good"]), [{"shift": "首件", "good": "2"}])
    native = shape.chart
    assert native.category_axis.minimum_scale == -.5 and native.category_axis.maximum_scale == .5
    assert native.series[0].points[0].marker.style == XL_MARKER_STYLE.CIRCLE
    manual = native._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")).find(qn("c:layout")).find(qn("c:manualLayout"))
    center = shape.left + shape.width * (float(manual.find(qn("c:x")).get("val")) + float(manual.find(qn("c:w")).get("val")) / 2)
    label = next(item for item in slide.shapes if item.has_text_frame and item.text == "首件")
    assert abs(label.left + label.width / 2 - center) <= 2


def test_svg_negative_bars_share_the_visible_zero_baseline():
    config = {**widget("JNegativeBar", value_fields=["good"]), "data": [{"shift": "一班", "good": "-2"}, {"shift": "二班", "good": "4"}]}
    text = render_html({"name": "正负值", "layout": {"canvas": {}, "widgets": [config]}})
    baseline = re.search(r'<line data-zero-axis="primary"[^>]+y1="([^"]+)"[^>]+y2="([^"]+)"', text)
    zero = float(baseline.group(1))
    assert baseline.group(1) == baseline.group(2) and zero != 215
    negative = re.search(r'<rect[^>]+y="([^"]+)"[^>]+height="([^"]+)"[^>]+data-value="-2"', text)
    positive = re.search(r'<rect[^>]+y="([^"]+)"[^>]+height="([^"]+)"[^>]+data-value="4"', text)
    assert abs(float(negative.group(1)) - zero) <= .002
    assert abs(float(positive.group(1)) + float(positive.group(2)) - zero) <= .002
