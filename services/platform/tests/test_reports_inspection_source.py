"""在新隔离库验证检验到正式报告，历史事实和当前授权必须同时成立。"""
from copy import deepcopy
from io import BytesIO, StringIO
import csv
import json

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import pytest
from pydantic import ValidationError as InputValidationError
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from agent_platform.core.context import Actor
from agent_platform.core.database import Base
from agent_platform.core.errors import DomainError
from agent_platform.core.models import Task
from agent_platform.core.settings import Settings
from agent_platform.modules.data.service import DataService
from agent_platform.modules.manufacturing import models as manufacturing_models
from agent_platform.modules.manufacturing.schemas import InspectionImport, Metric
from agent_platform.modules.manufacturing.service import ManufacturingService
from agent_platform.modules.reports.inspection_document import business_time, office_sections, render_html
from agent_platform.modules.reports.inspection_source import InspectionSource
from agent_platform.modules.reports.schemas import InspectionReportGenerationInput
from agent_platform.modules.reports.service import ReportService, validate_layout
from agent_platform.modules.reports import models as report_models


@pytest.fixture
def context(tmp_path):
    """仅创建本次临时SQLite和导入文件，不读取旧业务配置或数据库。"""
    engine = create_engine("sqlite:///" + (tmp_path / "inspection-reports.sqlite3").as_posix())
    Base.metadata.create_all(engine)
    actor = Actor("quality-owner", "quality-workspace", frozenset({"*"}), ("admin",))
    settings = Settings(data_dir=tmp_path / "inspection-report-data")
    with Session(engine) as session:
        yield session, actor, settings
        session.rollback()
    engine.dispose()


def quality_run(context, *, count=2, expected=None, reviewed=True, three_dimensions=True):
    """用一份显式新CSV执行真实受控取数、样本映射、判定和可选复核。"""
    session, actor, settings = context
    manufacturing = ManufacturingService(session, actor)
    data = DataService(session, actor, settings)
    product = manufacturing.create(manufacturing_models.Product, {"name": "高蛋白饮品"})
    batch = manufacturing.create(manufacturing_models.ProductBatch, {"name": "十月首批", "external_ref": "BATCH-20261020", "product_id": product["id"]})
    metrics = [Metric(id="weight", name="净含量", unit="g", standard_min="200", standard_max="205").model_dump()]
    if three_dimensions:
        metrics.extend([Metric(id="protein", name="蛋白含量", category="component", component_name="蛋白质", unit="g/100g", standard_min="3.0").model_dump(), Metric(id="surface", name="包装外观", category="appearance").model_dump()])
    standard = manufacturing.create(manufacturing_models.InspectionTemplate, {"name": "饮品出厂检验", "target_type": "product", "metrics": metrics})
    inspection = manufacturing.create_inspection({"name": "出厂质检", "target_type": "product", "product_id": product["id"], "batch_id": batch["id"], "template_id": standard["id"], "template_revision": 1, "expected_samples": expected if expected is not None else count})
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["sample_id", "weight", "weight_unit", "protein", "protein_unit", "image_ref", "defects"])
    for index in range(count):
        writer.writerow([f"S-{index + 1:03d}", "202.000000000000000001", "g", "3.2", "g/100g", f"db/image/{index + 1}", "[]"])
    dataset = data.import_csv({"name": "本轮质检来源", "csv_text": buffer.getvalue()})
    query = data.execute_query({"datasource_id": dataset["datasource_id"], "sql": f'SELECT * FROM "{dataset["table_name"]}"', "max_rows": count + 1})
    mappings = [{"metric_id": "weight", "value_field": "weight", "unit_field": "weight_unit"}]
    if three_dimensions:
        mappings.extend([{"metric_id": "protein", "value_field": "protein", "unit_field": "protein_unit"}, {"metric_id": "surface", "image_ref_field": "image_ref", "defects_field": "defects"}])
    imported = manufacturing.import_inspection(inspection["id"], InspectionImport(revision=inspection["revision"], query_id=query["id"], mappings=mappings), settings)
    result = manufacturing.run_inspection(inspection["id"], imported["revision"])
    if reviewed:
        result = manufacturing.review(inspection["id"], {"revision": result["revision"], "outcome": result["outcome"], "note": "已核查样本及原始记录"})["inspection"]
    return manufacturing, data, ReportService(session, actor, settings), result, dataset, query


def generate(reports, inspection, *, name="饮品批次检验报告"):
    return reports.generate_inspection({"name": name, "inspection_id": inspection["id"], "run_id": inspection["result"]["run_id"], "inspection_revision": inspection["revision"]})


def test_trusted_three_dimension_report_and_readable_historical_document(context):
    manufacturing, _, reports, inspection, _, _ = quality_run(context)
    report = generate(reports, inspection)
    source = report["layout"]["inspection_source"]
    assert report["layout"]["kind"] == "inspection_report"
    assert source["metadata"]["batch_number"] == "BATCH-20261020"
    assert source["standard_snapshot"]["name"] == "饮品出厂检验"
    assert source["review_snapshot"][0]["note"] == "已核查样本及原始记录"
    assert len(next(widget["data"] for widget in report["layout"]["widgets"] if widget["binding"]["view"] == "metrics")) == 6
    body, media, _ = reports.export(report["id"], "html")
    assert "成分检验明细" in body and "蛋白质" in body and "数值检验明细" in body and "外观检验明细" in body
    assert "202.000000000000000001" in body and "BATCH-20261020" in body and "零部件" not in body
    assert media.startswith("text/html") and "table-header-group" in body and "overflow:hidden" not in body
    published = reports.publish(report["id"], report["revision"])
    assert published["status"] == "published"
    # 当前标准和产品改名不改写任务的标准快照与报告生成时的人类可读名称。
    template = manufacturing.get(manufacturing_models.InspectionTemplate, inspection["template_id"])
    manufacturing.change(manufacturing_models.InspectionTemplate, template.id, {"name": "新版标准"}, template.revision)
    assert "饮品出厂检验" in reports.export(report["id"], "html", revision=1)[0]
    assert business_time("2026-10-20T00:01:59+00:00") == "2026-10-20 08:01"


def test_source_spoofing_mixed_binding_and_formal_body_edit_are_rejected(context):
    _, _, reports, inspection, _, _ = quality_run(context)
    report = generate(reports, inspection)
    tampered = deepcopy(report["layout"])
    tampered["inspection_source"]["review_snapshot"][0]["note"] = "伪造复核"
    tampered["widgets"][3]["data"] = [{"value": "999999"}]
    with pytest.raises(DomainError, match="正文") as error:
        reports.patch(report["id"], {"revision": report["revision"], "layout": tampered})
    assert error.value.code == "inspection_report_layout_read_only"
    with pytest.raises(DomainError) as error:
        reports.create({"name": "伪造正式报告", "layout": tampered})
    assert error.value.code == "inspection_generation_required"
    binding = deepcopy(report["layout"]["widgets"][3]["binding"])
    binding["query_id"] = "query-cannot-mix"
    with pytest.raises(DomainError) as error:
        validate_layout({"widgets": [{"id": "mixed", "component": "JScrollBoard", "binding": binding}]})
    assert error.value.code == "inspection_binding_invalid"
    with pytest.raises(InputValidationError):
        InspectionReportGenerationInput(name="伪造事实", inspection_id=inspection["id"], run_id=inspection["result"]["run_id"], inspection_revision=inspection["revision"], data=[{"outcome": "passed"}])
    # 普通自定义报表也由服务端重建检验投影，不接受客户端data和来源上下文。
    tampered.pop("kind")
    custom = reports.create({"name": "检验数据看板", "layout": tampered})
    assert custom["layout"]["inspection_source"]["review_snapshot"][0]["note"] == "已核查样本及原始记录"
    assert custom["layout"]["widgets"][3]["data"][0]["value"] == "202.000000000000000001"
    changed_source = deepcopy(custom["layout"])
    for widget in changed_source["widgets"]:
        widget["binding"]["run_id"] = "different-run"
    with pytest.raises(DomainError) as error:
        reports.patch(custom["id"], {"revision": custom["revision"], "layout": changed_source})
    assert error.value.code == "inspection_source_change_requires_generation"
    changed_source["widgets"][0].pop("binding")
    with pytest.raises(DomainError) as error:
        reports.patch(custom["id"], {"revision": custom["revision"], "layout": changed_source})
    assert error.value.code == "inspection_bindings_required"


def test_copy_and_template_cannot_launder_inspection_dependency(context):
    session, _, settings = context
    _, _, reports, inspection, _, _ = quality_run(context)
    report = generate(reports, inspection)
    copied = deepcopy(report["layout"])
    copied.pop("kind")
    for widget in copied["widgets"]:
        widget.pop("binding")
    revoked = ReportService(session, Actor("reader", "quality-workspace", frozenset({"m07.*"})), settings)
    for model in (report_models.Report, report_models.ReportTemplate):
        with pytest.raises(DomainError) as error:
            revoked.create({"name": "不能洗掉检验来源", "layout": copied}, model=model)
        assert error.value.code == "inspection_bindings_required"
    copied.pop("inspection_source")
    with pytest.raises(DomainError) as error:
        revoked.create({"name": "残留来源不能冒充静态数据", "layout": copied})
    assert error.value.code == "inspection_bindings_required"


def test_review_append_is_frozen_until_explicit_refresh_and_publication_rechecks(context):
    manufacturing, _, reports, inspection, _, _ = quality_run(context)
    report = generate(reports, inspection)
    revised = manufacturing.review(inspection["id"], {"revision": inspection["revision"], "outcome": "passed", "note": "复核后补充追溯记录"}, kind="deviation")["inspection"]
    assert "复核后补充追溯记录" not in reports.export(report["id"], "html")[0]
    with pytest.raises(DomainError) as error:
        reports.publish(report["id"], report["revision"])
    assert error.value.code == "report_source_changed"
    refreshed = reports.refresh(report["id"], report["revision"])
    assert refreshed["layout"]["inspection_source"]["inspection_revision"] == revised["revision"]
    assert "复核后补充追溯记录" in reports.export(report["id"], "html")[0]
    assert "复核后补充追溯记录" not in reports.export(report["id"], "html", revision=1)[0]
    published = reports.publish(report["id"], refreshed["revision"])
    assert published["status"] == "published"
    restored = reports.restore(report["id"], published["revision"], 1)
    assert restored["status"] == "draft"
    assert "复核后补充追溯记录" not in reports.export(report["id"], "html")[0]


def test_unreviewed_run_can_draft_but_unrelated_review_cannot_publish(context):
    manufacturing, _, reports, inspection, _, _ = quality_run(context, reviewed=False)
    first_run = inspection["result"]["run_id"]
    first = generate(reports, inspection)
    with pytest.raises(DomainError) as error:
        reports.publish(first["id"], first["revision"])
    assert error.value.code == "inspection_review_required"
    # 后续再次计算并复核只属于第二次任务，不得给第一次运行补上合格复核。
    second = manufacturing.run_inspection(inspection["id"], inspection["revision"])
    reviewed = manufacturing.review(second["id"], {"revision": second["revision"], "outcome": "passed", "note": "仅复核第二次运行"})["inspection"]
    old = reports.generate_inspection({"name": "第一次运行历史草稿", "inspection_id": inspection["id"], "run_id": first_run, "inspection_revision": reviewed["revision"]})
    assert old["layout"]["inspection_source"]["review_snapshot"] == []
    with pytest.raises(DomainError) as error:
        reports.publish(old["id"], old["revision"])
    assert error.value.code == "report_source_changed"


def test_revoked_quality_permission_hides_only_sensitive_reports_and_counts(context):
    session, _, settings = context
    _, _, reports, inspection, _, _ = quality_run(context)
    report = generate(reports, inspection)
    plain = reports.create({"name": "公开班报", "layout": {"widgets": []}})
    revoked = ReportService(session, Actor("reader", "quality-workspace", frozenset({"m07.*"})), settings)
    listing = revoked.list(limit=1)
    assert listing["total"] == 1 and listing["items"][0]["id"] == plain["id"]
    assert revoked.list(offset=1)["items"] == []
    for operation in (lambda: revoked.get(report["id"]), lambda: revoked.revisions(report["id"]), lambda: revoked.export(report["id"], "json", revision=1)):
        with pytest.raises(DomainError) as error:
            operation()
        assert error.value.status == 403
    # 检验报告使用源字段授权，不要求给质检读者整个数据分析工作台权限。
    reader = ReportService(session, Actor("quality-reader", "quality-workspace", frozenset({"m07.*", "m11.read"}), ("operator",)), settings)
    assert "S-001" in reader.export(report["id"], "html")[0]
    assert reader.refresh(report["id"], report["revision"])["revision"] == report["revision"] + 1


def test_historical_run_rechecks_its_own_query_after_new_run_uses_other_source(context):
    session, _, settings = context
    manufacturing, data, reports, inspection, dataset, _ = quality_run(context, reviewed=False, three_dimensions=False)
    report = generate(reports, inspection)
    fresh = data.import_csv({"name": "第二来源", "csv_text": "sample_id,weight,weight_unit\nS-001,203,g\nS-002,203,g\n"})
    query = data.execute_query({"datasource_id": fresh["datasource_id"], "sql": f'SELECT * FROM "{fresh["table_name"]}"'})
    imported = manufacturing.import_inspection(inspection["id"], InspectionImport(revision=inspection["revision"], query_id=query["id"], mappings=[{"metric_id": "weight", "value_field": "weight", "unit_field": "weight_unit"}]), settings)
    manufacturing.run_inspection(inspection["id"], imported["revision"])
    source = data.get_source(dataset["datasource_id"])
    data.patch_source(source.id, {"revision": source.revision, "field_roles": {dataset["table_name"]: {"weight": ["manager"]}}})
    reader = ReportService(session, Actor("quality-reader", "quality-workspace", frozenset({"m07.*", "m11.read"}), ("operator",)), settings)
    with pytest.raises(DomainError) as error:
        reader.export(report["id"], "html", revision=1)
    assert error.value.status == 403 and error.value.code == "inspection_evidence_access_denied"
    assert reader.list()["total"] == 0


def test_revoked_query_report_does_not_block_authorized_static_report(context):
    session, _, settings = context
    _, data, reports, _, dataset, query = quality_run(context, three_dimensions=False)
    reports.create({"name": "受限取数报表", "layout": {"widgets": [{"id": "numbers", "component": "JBar", "binding": {"query_id": query["id"], "label_field": "sample_id", "value_field": "weight"}, "data_provenance": {"type": "query", "query_id": query["id"], "truncated": False}, "data": query["rows"]}]}})
    plain = reports.create({"name": "普通静态报表", "layout": {"widgets": []}})
    source = data.get_source(dataset["datasource_id"])
    data.patch_source(source.id, {"revision": source.revision, "field_roles": {dataset["table_name"]: {"weight": ["manager"]}}})
    reader = ReportService(session, Actor("reader", "quality-workspace", frozenset({"m07.*"}), ("operator",)), settings)
    listing = reader.list(limit=1)
    assert listing["total"] == 1 and listing["items"][0]["id"] == plain["id"]
    assert reader.list(offset=1)["items"] == []


def test_wrong_workspace_run_standard_or_task_cannot_be_substituted(context):
    session, actor, settings = context
    _, _, reports, inspection, _, _ = quality_run(context)
    other = InspectionSource(session, Actor(actor.user_id, "other-workspace", frozenset({"*"})), settings)
    with pytest.raises(DomainError) as error:
        other.capture(inspection["id"], inspection["result"]["run_id"])
    assert error.value.status == 404
    with pytest.raises(DomainError) as error:
        reports.generate_inspection({"name": "过期版本", "inspection_id": inspection["id"], "run_id": inspection["result"]["run_id"], "inspection_revision": 1})
    assert error.value.code == "inspection_revision_conflict"
    run = session.get(manufacturing_models.InspectionRun, inspection["result"]["run_id"])
    task = session.get(Task, run.task_id)
    task.inputs = {**task.inputs, "template_revision": 99}
    session.flush()
    with pytest.raises(DomainError) as error:
        generate(reports, inspection)
    assert error.value.code == "inspection_standard_evidence_mismatch"


def test_full_rows_reach_real_office_files_and_missing_samples_stay_explicit(context):
    _, _, reports, inspection, _, _ = quality_run(context, count=25, expected=26, three_dimensions=False)
    report = generate(reports, inspection)
    summary = next(widget["data"] for widget in report["layout"]["widgets"] if widget["binding"]["view"] == "summary")
    assert next(row["value"] for row in summary if row["name"] == "missing_samples") == 1
    assert report["layout"]["inspection_source"]["metadata"]["system_outcome"] == "incomplete"
    body, _, _ = reports.export(report["id"], "html")
    assert "S-025" in body and "缺少样本" in body and "数据不完整" in body
    docx, _, _ = reports.export(report["id"], "docx")
    doc = Document(BytesIO(docx))
    assert any("S-025" in cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    xlsx, _, _ = reports.export(report["id"], "xlsx")
    book = load_workbook(BytesIO(xlsx), read_only=True, data_only=True)
    assert any(value == "S-025" for row in book.active.iter_rows(values_only=True) for value in row)
    book.close()
    pptx, _, _ = reports.export(report["id"], "pptx")
    deck = Presentation(BytesIO(pptx))
    tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
    assert any("S-025" == cell.text for table in tables for row in table.rows for cell in row.cells)
    assert any(shape.has_chart for slide in deck.slides for shape in slide.shapes)
    # 多页正文不以单张PNG冒充完整文档；PDF由root串行实际渲染验首末页。
    with pytest.raises(DomainError) as error:
        reports.export(report["id"], "png")
    assert error.value.code == "inspection_document_png_unsupported"


def _display_report():
    """纯内存正文夹具覆盖同名不同外观标准，不执行数据库或浏览器。"""
    standard = [
        {"name": "净含量", "category": "numeric", "component_name": "", "standard_min": "200", "standard_max": "205", "unit": "g", "minimum_confidence": "0.8", "blocking_severities": "moderate, severe"},
        {"name": "包装外观", "category": "appearance", "component_name": "", "standard_min": None, "standard_max": None, "unit": "", "minimum_confidence": "0.80", "blocking_severities": "moderate, severe"},
        {"name": "包装外观", "category": "appearance", "component_name": "", "standard_min": None, "standard_max": None, "unit": "", "minimum_confidence": "0.95", "blocking_severities": "severe"},
    ]
    metrics = []
    for sample in range(1, 26):
        for item in standard:
            metrics.append({**item, "metric": item["name"], "standard_unit": item["unit"],
                            "sample": f"S-{sample:03d}", "value": "202.000000000000000001" if item["category"] == "numeric" else None,
                            "image_ref": "", "outcome": "passed", "reason": "在标准内" if item["category"] == "numeric" else "检测结果明确返回无缺陷",
                            "defects": "", "observation_note": ""})
    return {"name": "25样本排版", "revision": 1, "status": "draft", "layout": {
        "kind": "inspection_report", "inspection_source": {"reviewer_names": {}},
        "widgets": [{"binding": {"view": view}, "data": rows} for view, rows in (
            ("meta", [{"name": "target_type", "value": "product"}]), ("summary", []), ("outcomes", []),
            ("metrics", metrics), ("standard", standard), ("reviews", []))]}}


def test_display_omits_empty_columns_and_keeps_precise_values_with_readable_width():
    report = _display_report()
    before = deepcopy(report)
    sections = office_sections(report)
    numeric = next(section for section in sections if section["title"] == "数值检验明细")
    assert "图像来源" not in numeric["rows"][0]
    assert len(numeric["rows"]) == 26 and all(len(row) == len(numeric["rows"][0]) for row in numeric["rows"])
    value_index = numeric["rows"][0].index("实测")
    assert numeric["column_widths"][value_index] >= 30
    assert numeric["rows"][-1][0] == "S-025" and numeric["rows"][-1][value_index] == "202.000000000000000001 g"
    content = render_html(report)
    assert '<col style="width:32.0%">' in content and "font-size:12px" in content
    assert "S-025" in content and "202.000000000000000001 g" in content
    assert report == before
    metric_rows = next(widget["data"] for widget in report["layout"]["widgets"] if widget["binding"]["view"] == "metrics")
    metric_rows[0]["image_ref"] = "db/image/first"
    numeric = next(section for section in office_sections(report) if section["title"] == "数值检验明细")
    assert numeric["rows"][0][-1] == "图像来源" and numeric["rows"][1][-1] == "db/image/first"
    assert numeric["rows"][-1][-1] == "—"


def test_appearance_rows_reference_distinct_standards_without_repeating_constraints():
    report = _display_report()
    sections = office_sections(report)
    appearance = next(section for section in sections if section["title"] == "外观检验明细")
    assert appearance["rows"][0] == ["样本", "检验项", "标准", "判定", "依据"]
    assert len(appearance["rows"]) == 51 and appearance["rows"][-1][0] == "S-025"
    assert [row[2] for row in appearance["rows"][1:3]] == ["标准 2", "标准 3"]
    assert not any("判定置信度" in str(value) for row in appearance["rows"] for value in row)
    standard = next(section for section in sections if section["title"] == "标准快照")
    assert len(standard["rows"]) == 4
    assert "0.80" in standard["rows"][2][-1] and "中等、严重" in standard["rows"][2][-1]
    assert "0.95" in standard["rows"][3][-1] and "严重" in standard["rows"][3][-1]
    content = render_html(report)
    assert content.count("判定置信度") == 2 and "moderate" not in content and "severe" not in content
    metrics = next(widget["data"] for widget in report["layout"]["widgets"] if widget["binding"]["view"] == "metrics")
    metrics[1]["defects"] = json.dumps([{"defect_type": "划痕", "confidence": "0.91", "severity": "severe", "bbox": [0.1, 0.2, 0.3, 0.4]}], ensure_ascii=False)
    content = render_html(report)
    assert "划痕（置信度 0.91，严重，位置 [0.1, 0.2, 0.3, 0.4]）" in content
