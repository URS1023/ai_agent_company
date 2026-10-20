"""验证真实办公文件、原生PPT图表/表格与离线浏览器打印文件。"""
from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from agent_platform.core.context import Actor
from agent_platform.core.database import Base
from agent_platform.core.settings import Settings
from agent_platform.modules.data.service import DataService
from agent_platform.modules.reports.service import ReportService
from agent_platform.modules.reports.rendering import render_html


def test_office_pdf_and_png_have_real_data_and_native_shapes(tmp_path):
    engine = create_engine("sqlite:///" + (tmp_path / "office.sqlite3").as_posix())
    Base.metadata.create_all(engine)
    actor = Actor("office-test", "workspace", frozenset({"*"}), ("admin",))
    settings = Settings(data_dir=tmp_path / "office-data")
    with Session(engine) as session:
        data, reports = DataService(session, actor, settings), ReportService(session, actor, settings)
        dataset = data.import_csv({"name": "班组统计", "csv_text": "name,value\n一班,17\n二班,23\n"})
        query = data.execute_query({"datasource_id": dataset["datasource_id"], "sql": f'SELECT * FROM "{dataset["table_name"]}"'})
        report = reports.create({"name": "设备统计报表", "layout": {"canvas": {}, "widgets": [{"id": "chart", "component": "JBar", "componentName": "班组设备数", "binding": {"query_id": query["id"]}}]}})
        report = reports.refresh(report["id"], report["revision"])
        word, _, _ = reports.export(report["id"], "docx")
        doc = Document(BytesIO(word))
        assert doc.tables[0].cell(1, 0).text == "一班"
        excel, _, _ = reports.export(report["id"], "xlsx")
        workbook = load_workbook(BytesIO(excel))
        assert any("23" in str(cell.value) for row in workbook.active for cell in row)
        slides, _, _ = reports.export(report["id"], "pptx")
        deck = Presentation(BytesIO(slides))
        charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
        tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
        assert charts and tables
        assert charts[0].series[0].values == (17.0, 23.0)
        assert tables[0].cell(1, 0).text == "一班"
        pdf, _, _ = reports.export(report["id"], "pdf")
        assert pdf.startswith(b"%PDF-") and b"/Type /Page" in pdf and pdf.rstrip().endswith(b"%%EOF")
        image, _, _ = reports.export(report["id"], "png")
        assert image.startswith(b"\x89PNG\r\n\x1a\n")
        (tmp_path / "manufacturing-report.pdf").write_bytes(pdf)
        (tmp_path / "manufacturing-report.png").write_bytes(image)
        session.rollback()
    engine.dispose()


def test_offline_canvas_keeps_position_and_chart_type():
    layout = {"canvas": {"width": 1440, "height": 900}, "widgets": [{"id": "bar", "component": "JBar", "x": 123, "y": 87, "w": 420, "h": 280, "data": [{"name": "一班", "value": "17"}, {"name": "二班", "value": "23"}]}]}
    content = render_html({"name": "柱状图", "layout": layout})
    assert "left:123.0px;top:87.0px;width:420.0px;height:280.0px" in content
    assert "<rect" in content and "<polyline" not in content
    layout["widgets"][0]["component"] = "JLine"
    assert "<polyline" in render_html({"name": "折线图", "layout": layout})
