"""展示组件使用真实字段与精确数值；仅纯内存及本次新建临时SQLite。"""
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from agent_platform.core.context import Actor
from agent_platform.core.database import Base
from agent_platform.core.errors import DomainError
from agent_platform.core.settings import Settings
from agent_platform.modules.data.service import DataService
from agent_platform.modules.reports.factory import component_instance
from agent_platform.modules.reports.models import ReportTemplate
from agent_platform.modules.reports.presentation import (
    NAMES, PROGRESS_TYPES, SPECIAL_CHART_TYPES, clock_text,
    office_section, presentation_model, render_presentation, sanitize_html,
)
from agent_platform.modules.reports.schemas import ComponentInstanceInput
from agent_platform.modules.reports.schemas import GenerationInput
from agent_platform.modules.reports.service import ReportService, SUPPORTED_COMPONENTS, component_catalog, validate_layout


def widget(component, **binding):
    return {"id": component, "component": component, "w": 640, "h": 280, **({"binding": binding} if binding else {}), "config": {}}


def test_source_instances_keep_original_config_without_sample_values():
    items = [item for item in component_catalog()["items"] if item["compType"] in NAMES]
    assert {item["compType"] for item in items} == set(NAMES)
    for item in items:
        actual = component_instance(ComponentInstanceInput(source_id=item["id"]), SUPPORTED_COMPONENTS)
        assert actual["source_component_id"] == item["id"] and actual["component"] == item["compType"]
        assert actual["config"]["chartDataParsed"] == [] and actual["config"]["chartData"] == "[]"
        validate_layout({"canvas": {}, "widgets": [actual]})


def test_rank_order_uses_exact_decimal_and_keeps_missing_last():
    rows = [{"name": "较小", "value": "900719925474099312345678901234567890"}, {"name": "较大", "value": "900719925474099312345678901234567891"}, {"name": "缺测", "value": None}]
    model = presentation_model(widget("JRankingList"), rows)
    assert model["status"] == "ready"
    assert [row["name"] for row in model["rows"]] == ["较大", "较小", "缺测"]
    assert model["rows"][0]["value"] == rows[1]["value"] and model["rows"][-1]["value"] is None
    assert presentation_model(widget("JNumber"), rows)["status"] == "needs_mapping"
    assert presentation_model(widget("JNumber", value_field="absent"), rows[:1])["status"] == "needs_mapping"


def test_table_explicit_columns_labels_and_office_follow_same_projection():
    config = widget("JCommonTable", fields=["value", "name"])
    config["config"]["column_labels"] = {"value": "实际产量"}
    rows = [{"name": "一班", "value": "1.2300", "private_note": "不导出"}, {"name": "二班", "value": None, "private_note": "不导出"}]
    model = presentation_model(config, rows)
    assert [column["label"] for column in model["columns"]] == ["实际产量", "name"]
    assert list(model["rows"][0]) == ["value", "name"]
    section = office_section(config, rows)
    assert section["rows"] == [["实际产量", "name"], ["1.2300", "一班"], ["—", "二班"]]
    content = render_presentation(config, rows, "#fff", ["#38bdf8"])
    assert "1.2300" in content and "private_note" not in content and "不导出" not in content
    assert presentation_model(widget("JCommonTable", fields=["absent"]), rows)["status"] == "needs_mapping"


@pytest.mark.parametrize("component", sorted(PROGRESS_TYPES))
def test_progress_has_actual_target_bounds_and_explicit_missing(component):
    config = widget(component, total_field="target")
    assert presentation_model(config, [{"name": "完成", "value": "1", "target": "4"}])["rows"][0]["percentage"] == "25.00"
    for value, target in (("-1", "4"), ("5", "4"), ("1", "0"), ("1", None)):
        assert presentation_model(config, [{"name": "完成", "value": value, "target": target}])["status"] == "invalid_data"
    missing = presentation_model(config, [{"name": "未采集", "value": None, "target": "4"}])
    assert missing["status"] == "empty" and missing["rows"][0]["percentage"] is None


def test_list_preserves_zero_text_and_only_selected_optional_fields():
    config = widget("JList", label_field="code", description_field="detail", time_field="observed")
    model = presentation_model(config, [{"code": 0, "detail": "待处理", "observed": "2026-10-21 12:00"}])
    assert model["status"] == "ready" and model["rows"][0]["name"] == "0"
    assert model["rows"][0]["description"] == "待处理" and model["rows"][0]["time"] == "2026-10-21 12:00"
    assert "value" not in office_section(config, [{"code": 0, "detail": "待处理", "observed": "2026-10-21 12:00"}])["rows"][0]


def test_rich_text_strips_execution_attributes_and_timezone_is_explicit():
    value = '<p onclick="alert(1)">说明<strong style="color:red">合格</strong><script>bad()</script><img src="https://outside.invalid/trace"><iframe>隐藏</iframe></p>'
    assert sanitize_html(value) == "<p>说明<strong>合格</strong></p>"
    config = widget("JDragEditor")
    config["config"] = {"option": {"body": {"text": value}}}
    saved = validate_layout({"canvas": {}, "widgets": [config]})
    assert saved["widgets"][0]["config"]["option"]["body"]["text"] == "<p>说明<strong>合格</strong></p>"
    assert office_section(config, [])["text"] == "说明合格"
    now = datetime(2026, 10, 20, 17, 2, 3, tzinfo=timezone.utc)
    assert clock_text({}, now) == "2026-10-21 01:02:03"
    assert clock_text({"timezone": "UTC", "date_only": True}, now) == "2026-10-20"
    invalid = widget("JCurrentTime")
    invalid["config"] = {"option": {"timezone": "Mars"}}
    with pytest.raises(DomainError, match="请选择北京时间"):
        validate_layout({"canvas": {}, "widgets": [invalid]})


@pytest.mark.parametrize("component", sorted(SPECIAL_CHART_TYPES))
def test_special_shapes_are_not_bar_fallbacks_and_reject_negative(component):
    rows = [{"name": "单项", "value": "25"}] if component == "JSemiGauge" else [{"name": "一班", "value": "25"}, {"name": "二班", "value": "10"}]
    config = widget(component)
    assert presentation_model(config, rows)["status"] == "ready"
    text = render_presentation(config, rows, "#fff", ["#38bdf8", "#22d3a6"])
    assert "<svg" in text and ('<polygon' in text or '<rect' in text or '<path' in text)
    assert presentation_model(config, [{"name": "异常", "value": "-1"}])["status"] == "invalid_data"
    assert office_section(config, rows).get("chart") is None
    if component == "JRose":
        assert '<circle data-value="25"' in render_presentation(config, rows[:1], "#fff", ["#38bdf8"])


def test_zero_gauge_and_radar_missing_do_not_fabricate_values():
    gauge = render_presentation(widget("JSemiGauge"), [{"name": "电流", "value": "0"}], "#fff", ["#38bdf8"])
    assert "<svg" in gauge and ">0</text>" in gauge
    assert presentation_model(widget("JCircleRadar"), [{"name": "温度", "value": "1"}, {"name": "电流", "value": None}])["status"] == "invalid_data"
    assert "暂无可绘制数值" in render_presentation(widget("JRose"), [{"name": "未生产", "value": "0"}], "#fff", ["#38bdf8"])


def test_query_evidence_projection_publish_export_and_forged_model(tmp_path):
    """预算2次取数：1条初始查询、1次共享刷新；不启动浏览器、服务或访问现有库。"""
    engine = create_engine("sqlite:///" + (tmp_path / "presentation.sqlite3").as_posix())
    Base.metadata.create_all(engine)
    settings = Settings(data_dir=tmp_path / "presentation-data")
    actor = Actor("presentation-owner", "workspace", frozenset({"*"}))
    with Session(engine) as session:
        data, reports = DataService(session, actor, settings), ReportService(session, actor, settings)
        dataset = data.import_csv({"name": "展示临时样本", "csv_text": "name,value,target,note\n一班,17.00,100,保留字段\n二班,23.00,100,保留字段\n"})
        query = data.execute_query({"datasource_id": dataset["datasource_id"], "sql": f'SELECT * FROM "{dataset["table_name"]}"', "max_rows": 10})
        first = widget("JRankingList", query_id=query["id"])
        first["presentation_model"] = {"status": "ready", "rows": [{"name": "伪造", "value": "9999"}]}
        table = widget("JCommonTable", query_id=query["id"], fields=["name", "value"])
        report = reports.create({"name": "可信投影", "layout": {"canvas": {}, "widgets": [first, table]}})
        assert "presentation_model" not in report["layout"]["widgets"][0]
        refreshed = reports.refresh(report["id"], report["revision"])
        actual = refreshed["layout"]["widgets"][0]["presentation_model"]
        assert actual["rows"][0]["name"] == "二班" and actual["rows"][0]["value"] == "23.00"
        forged = deepcopy(refreshed["layout"])
        forged["widgets"][0]["presentation_model"]["rows"][0]["value"] = "9999"
        patched = reports.patch(report["id"], {"revision": refreshed["revision"], "layout": forged})
        assert patched["layout"]["widgets"][0]["presentation_model"]["rows"][0]["value"] == "23.00"
        published = reports.publish(report["id"], patched["revision"])
        assert published["status"] == "published"
        doc, _, _ = reports.export(report["id"], "docx")
        document = Document(BytesIO(doc))
        assert len(document.tables) == 2
        assert [[cell.text for cell in row.cells] for row in document.tables[1].rows] == [["name", "value"], ["一班", "17.00"], ["二班", "23.00"]]
        csv, _, _ = reports.export(report["id"], "csv")
        assert "17.00" in csv and "保留字段" not in csv and "target" not in csv
        changed = deepcopy(patched["layout"])
        changed["widgets"][0]["binding"]["value_field"] = "missing"
        modified = reports.patch(report["id"], {"revision": published["revision"], "layout": changed})
        with pytest.raises(DomainError) as error:
            reports.publish(report["id"], modified["revision"])
        assert error.value.code == "presentation_data_not_ready"
        session.rollback()
    engine.dispose()


def test_generation_accepts_legacy_query_ids_and_explicit_mapping(tmp_path):
    """新临时库预算2次取数：1初始+1模板刷新；同源多组件共享查询。"""
    engine = create_engine("sqlite:///" + (tmp_path / "generation.sqlite3").as_posix())
    Base.metadata.create_all(engine)
    settings = Settings(data_dir=tmp_path / "generation-data")
    actor = Actor("generation-owner", "workspace", frozenset({"*"}))
    with Session(engine) as session:
        data, reports = DataService(session, actor, settings), ReportService(session, actor, settings)
        dataset = data.import_csv({"name": "模板临时样本", "csv_text": "team,output\n一班,17\n二班,23\n"})
        query = data.execute_query({"datasource_id": dataset["datasource_id"], "sql": f'SELECT * FROM "{dataset["table_name"]}"', "max_rows": 10})
        template = reports.create({"name": "列名映射模板", "layout": {"canvas": {}, "widgets": [widget("JBar"), widget("JRankingList")]}}, model=ReportTemplate)
        payload = GenerationInput(name="映射生成", template_id=template["id"], bindings={key: {"query_id": query["id"], "label_field": "team", "value_field": "output"} for key in ("JBar", "JRankingList")})
        report = reports.generate(payload.model_dump(exclude_none=True))
        assert reports.publish(report["id"], report["revision"])["status"] == "published"
        assert report["layout"]["widgets"][0]["series_model"]["categories"] == ["一班", "二班"]
        assert report["layout"]["widgets"][1]["presentation_model"]["rows"][0]["name"] == "二班"
        assert GenerationInput(name="旧客户端", template_id=template["id"], bindings={"JBar": query["id"]}).bindings["JBar"] == query["id"]
        session.rollback()
    engine.dispose()
