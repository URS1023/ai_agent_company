"""12个精确来源的报表展示回归；只用纯函数与本次新建的临时SQLite。"""
from copy import deepcopy
from io import BytesIO

from docx import Document
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from agent_platform.core.context import Actor
from agent_platform.core.database import Base
from agent_platform.core.errors import DomainError
from agent_platform.core.settings import Settings
from agent_platform.modules.data.service import DataService
from agent_platform.modules.reports.capabilities import SOURCE_COMPONENTS, SOURCE_TYPES, source_supported
from agent_platform.modules.reports.factory import component_instance
from agent_platform.modules.reports.presentation import office_section, presentation_model, render_presentation
from agent_platform.modules.reports.schemas import ComponentInstanceInput
from agent_platform.modules.reports.service import ReportService, SUPPORTED_COMPONENTS, component_catalog, validate_layout


def widget(component, **binding):
    source = next(source for source, kind in SOURCE_COMPONENTS.items() if kind == component)
    result = component_instance(ComponentInstanceInput(source_id=source), SUPPORTED_COMPONENTS)
    if binding:
        result["binding"] = binding
    return result


def test_exact_sources_preserve_254_existing_entries_and_enable_only_the_reviewed_twelve():
    catalog = component_catalog()["items"]
    enabled = {item["id"] for item in catalog if item["renderer_supported"]}
    old = {item["id"] for item in catalog if item["compType"] in set(SUPPORTED_COMPONENTS) - SOURCE_TYPES}
    assert len(old) == 254 and len(enabled) == 266
    assert enabled == old | set(SOURCE_COMPONENTS)
    assert len(SOURCE_COMPONENTS) == 12
    assert all(not item["renderer_supported"] for item in catalog if item["compType"] in {"JCommon", "JWeatherForecast", "JCustomEchart"})
    for source, component in SOURCE_COMPONENTS.items():
        actual = component_instance(ComponentInstanceInput(source_id=source), SUPPORTED_COMPONENTS)
        assert actual["source_component_id"] == source and actual["component"] == component
        assert actual["config"]["chartDataParsed"] == [] and actual["config"]["chartData"] == "[]"
        validate_layout({"widgets": [actual]})
        if component.endswith("Calendar"):
            assert actual["config"]["option"]["calendar"] == {"timezone": "Asia/Shanghai"}
    from agent_platform.modules.reports.router import components
    contracts = {row["id"]: row for row in components(actor=Actor("owner", "workspace", frozenset({"*"})))["items"]}
    assert all(contracts[component]["source_required"] for component in SOURCE_TYPES)
    assert contracts["JBar"]["source_required"] is False


def test_source_identity_cannot_be_missing_forged_or_changed_to_a_shared_type():
    actual = widget("JCardScroll")
    for source, component in ((None, "JCardScroll"), ("invented", "JCardScroll"), (actual["source_component_id"], "JBar")):
        changed = {**actual, "component": component}
        if source is None:
            changed.pop("source_component_id")
        else:
            changed["source_component_id"] = source
        with pytest.raises(DomainError) as error:
            validate_layout({"widgets": [changed]})
        assert error.value.code == "component_unsupported"
    assert source_supported(None, "JBar", SUPPORTED_COMPONENTS)


@pytest.mark.parametrize("option", [
    {"min": 10, "max": 10}, {"min": True}, {"max": "NaN"},
    {"gauge_stops": []}, {"gauge_stops": [{"value": "99", "color": "#123456"}]},
    {"gauge_stops": [{"value": "80", "color": "#123456"}, {"value": "70", "color": "#123456"}]},
    {"gauge_stops": [{"value": "100", "color": "url(https://invalid)"}]},
])
def test_gauge_option_errors_are_field_values_not_silently_clipped(option):
    actual = widget("JColorGauge")
    actual["config"]["option"] = option
    with pytest.raises(DomainError) as error:
        validate_layout({"widgets": [actual]})
    assert error.value.code == "component_options_invalid"


@pytest.mark.parametrize("component", ["JColorGauge", "JAntvGauge"])
def test_gauges_keep_exact_value_explicit_range_and_distinct_static_rendering(component):
    actual = widget(component, unit_field="unit")
    actual["config"]["option"] = {"min": "-20", "max": "80", "gauge_stops": [{"value": "5.00", "color": "#36d7b7"}, {"value": "80", "color": "#f7c948"}]}
    rows = [{"name": "温度", "value": "-5.1200", "unit": "℃"}]
    model = presentation_model(actual, rows)
    assert model["status"] == "ready" and model["rows"][0]["value"] == "-5.1200"
    assert (model["minimum"], model["maximum"]) == ("-20", "80")
    assert model["gauge_stops"][0]["value"] == "5.00"
    output = render_presentation(actual, rows, "#ffffff", ["#36d7b7"])
    assert 'data-value="-5.1200"' in output and '<line' in output
    assert ('data-gauge="segment"' in output) == (component == "JColorGauge")
    assert ('<linearGradient' in output) == (component == "JAntvGauge")
    assert presentation_model(actual, [{**rows[0], "value": "80.0001"}])["status"] == "invalid_data"
    assert presentation_model(actual, [{**rows[0], "value": None}])["status"] == "empty"
    assert presentation_model(actual, rows * 2)["status"] == "needs_mapping"


@pytest.mark.parametrize("component", ["JCapsuleChart", "JBackgroundBar"])
def test_bar_shapes_keep_negative_zero_and_missing_separate(component):
    actual = widget(component)
    rows = [{"name": name, "value": value} for name, value in (("负偏差", "-12.3400"), ("正偏差", "8"), ("零值", "0"), ("缺测", None))]
    model = presentation_model(actual, rows)
    assert model["status"] == "ready" and [row["value"] for row in model["rows"]] == ["-12.3400", "8", "0", None]
    output = render_presentation(actual, rows, "#ffffff", ["#36d7b7"])
    assert output.count('data-value=') == 3 and 'data-value="0"' in output and "—" in output
    assert ('data-background="capsule"' in output) == (component == "JCapsuleChart")
    assert ('data-background="column"' in output) == (component == "JBackgroundBar")


@pytest.mark.parametrize("component", ["JCalendar", "JPermanentCalendar", "JMultiViewCalendar", "JDevHistory"])
def test_events_use_actual_shanghai_day_and_require_explicit_timezone(component):
    actual = widget(component, time_field="observed", description_field="detail")
    rows = [{"name": "次日", "observed": "2026-10-01T16:30:00Z", "detail": "跨自然日"},
            {"name": "月底", "observed": "2026-09-30", "detail": "日期保留"}]
    model = presentation_model(actual, rows)
    assert model["status"] == "ready"
    assert [(row["name"], row["date"]) for row in model["rows"]] == [("月底", "2026-09-30"), ("次日", "2026-10-02")]
    assert model["rows"][1]["time"] == "2026-10-01T16:30:00Z"
    assert model["rows"][1]["time_display"] == "2026-10-02 00:30"
    assert not any("_sort_time" in row for row in model["rows"])
    if component.endswith("Calendar"):
        assert model["calendar"] == {"month": "2026-09", "timezone": "Asia/Shanghai"}
        actual["config"]["option"]["calendar"]["month"] = "2026-10"
        output = render_presentation(actual, rows, "#ffffff", ["#36d7b7"])
        assert 'data-date="2026-10-02"' in output
        # 月格仅作事项索引，完整说明由画布后的可分页日程承载，避免静态导出裁掉描述。
        from agent_platform.modules.reports.source_presentations import source_document_detail
        detail = source_document_detail(actual, rows)
        assert '跨自然日' in str(detail["rows"])
        assert "月底" not in output
    for invalid in ("2026-10-01T16:30:00", "2026-02-30", "随意描述", None):
        assert presentation_model(actual, [{**rows[0], "observed": invalid}])["status"] == "invalid_data"
    assert presentation_model(widget(component), rows)["status"] == "needs_mapping"


def test_cards_keep_three_source_variants_when_renamed_and_do_not_require_fake_numbers():
    for source in [source for source, component in SOURCE_COMPONENTS.items() if component == "JCardScroll"]:
        actual = component_instance(ComponentInstanceInput(source_id=source), SUPPORTED_COMPONENTS)
        actual.update(name="改名不会换类型", componentName="改名不会换类型", binding={"label_field": "label", "description_field": "detail"})
        rows = [{"label": "<script>待复核</script>", "detail": "等待现场确认"}]
        model = presentation_model(actual, rows)
        assert model["status"] == "ready" and model["rows"][0]["value"] is None
        assert model["variant"] == {"direction": "vertical" if "竖向" in source else "horizontal", "show_index": "竖向" in source, "highlight": "高亮" in source, "mode": "scroll"}
        output = render_presentation(actual, rows, "#ffffff", ["#36d7b7"])
        assert '<script>' not in output and '&lt;script&gt;' in output
        assert 'source-index' in output if "竖向" in source else 'source-index' not in output
    carousel = widget("JCardCarousel", label_field="label")
    assert presentation_model(carousel, [{"label": "设备01"}])["variant"]["mode"] == "carousel"
    carousel["binding"]["time_field"] = "observed"
    assert presentation_model(carousel, [{"label": "设备01", "observed": "2026-10-01T16:30:00Z"}])["rows"][0]["time_display"] == "2026-10-02 00:30"
    assert presentation_model(carousel, [{"label": "设备01", "observed": "等待现场确认"}])["rows"][0]["time_display"] == "等待现场确认"


def test_history_orders_instants_across_offsets_and_gradient_ids_do_not_collide_for_chinese_names():
    actual = widget("JDevHistory", time_field="observed")
    rows = [{"name": "后发生", "observed": "2026-10-02T00:40:00+08:00"}, {"name": "先发生", "observed": "2026-10-01T16:30:00Z"}]
    assert [row["name"] for row in presentation_model(actual, rows)["rows"]] == ["先发生", "后发生"]
    html = render_presentation(actual, rows, "#fff", ["#36d7b7"])
    assert "2026-10-02 00:30" in html and "T16:30:00Z" not in html
    first, second = widget("JAntvGauge"), widget("JAntvGauge")
    first["id"], second["id"] = "仪表甲", "仪表乙"
    one = render_presentation(first, [{"name": "温度", "value": "12"}], "#fff", ["#36d7b7"])
    two = render_presentation(second, [{"name": "温度", "value": "12"}], "#fff", ["#36d7b7"])
    assert one.split('linearGradient id="')[1].split('"')[0] != two.split('linearGradient id="')[1].split('"')[0]


def test_gauge_geometry_does_not_collapse_adjacent_decimal_values_after_float_conversion():
    actual = widget("JColorGauge")
    actual["config"]["option"] = {"min": "900719925474099312345678901234567890", "max": "900719925474099312345678901234567891"}
    rows = [{"name": "精密读数", "value": "900719925474099312345678901234567890.5"}]
    model = presentation_model(actual, rows)
    assert model["status"] == "ready" and model["gauge_position"] == "0.5"
    assert [stop["position"] for stop in model["gauge_stops"]] == ["0.33", "0.66", "1"]
    assert [tick["position"] for tick in model["gauge_ticks"]] == ["0", "0.2", "0.4", "0.6", "0.8", "1"]
    assert model["gauge_ticks"][1]["value"] == "900719925474099312345678901234567890.2"
    assert model["rows"][0]["value"] == rows[0]["value"]
    assert '<line data-value="900719925474099312345678901234567890.5"' in render_presentation(actual, rows, "#fff", ["#36d7b7"])


@pytest.mark.parametrize("component,option", [
    ("JCalendar", {"calendar": {"month": "2026-13"}}),
    ("JCalendar", {"calendar": {"timezone": "UTC"}}),
    ("JCardCarousel", {"carousel": {"page_size": True}}),
    ("JCardCarousel", {"carousel": {"interval_seconds": 2}}),
    ("JCardCarousel", {"carousel": {"autoplay": "yes"}}),
])
def test_calendar_and_carousel_controls_reject_invalid_values(component, option):
    actual = widget(component)
    actual["config"]["option"] = option
    with pytest.raises(DomainError):
        validate_layout({"widgets": [actual]})


def test_office_keeps_all_events_and_only_explicit_fields():
    actual = widget("JCalendar", time_field="observed", description_field="note")
    rows = [{"name": f"样本{index}", "observed": f"2026-10-{index:02}", "note": "1.2300", "private_field": "不要导出"} for index in range(1, 32)]
    section = office_section(actual, rows)
    assert len(section["rows"]) == 32 and "样本31" in section["rows"][-1]
    assert "不要导出" not in str(section) and section.get("chart") is None


def test_all_twelve_sources_share_one_fresh_query_and_keep_trusted_models(tmp_path):
    """唯一新CSV/SQLite，预算2条业务SQL：初始查询1次，12组件共享刷新1次。"""
    engine = create_engine("sqlite:///" + (tmp_path / "source-presentations.sqlite3").as_posix())
    Base.metadata.create_all(engine)
    settings = Settings(data_dir=tmp_path / "source-presentation-data")
    actor = Actor("source-owner", "workspace", frozenset({"*"}))
    with Session(engine) as session:
        data, reports = DataService(session, actor, settings), ReportService(session, actor, settings)
        dataset = data.import_csv({"name": "12展示临时数据", "csv_text": "name,value,observed,note,private\n设备01,12.3400,2026-10-01T16:30:00Z,检验完成,不要导出\n"})
        query = data.execute_query({"datasource_id": dataset["datasource_id"], "sql": f'SELECT * FROM "{dataset["table_name"]}"', "max_rows": 10})
        widgets = []
        for source in SOURCE_COMPONENTS:
            actual = component_instance(ComponentInstanceInput(source_id=source), SUPPORTED_COMPONENTS)
            actual["binding"] = {"query_id": query["id"], "time_field": "observed", "description_field": "note"}
            actual["presentation_model"] = {"status": "ready", "rows": [{"name": "伪造", "value": "9999"}]}
            widgets.append(actual)
        created = reports.create({"name": "12来源真实字段", "layout": {"widgets": widgets}})
        assert all("presentation_model" not in item for item in created["layout"]["widgets"])
        refreshed = reports.refresh(created["id"], created["revision"])
        assert all(item["presentation_model"]["status"] == "ready" for item in refreshed["layout"]["widgets"])
        assert len({item["data_provenance"]["query_id"] for item in refreshed["layout"]["widgets"]}) == 1
        assert all(item["presentation_model"]["rows"][0]["name"] == "设备01" for item in refreshed["layout"]["widgets"])
        forged = deepcopy(refreshed["layout"])
        forged["widgets"][0]["presentation_model"]["rows"][0]["value"] = "9999"
        modified = reports.patch(created["id"], {"revision": refreshed["revision"], "layout": forged})
        assert modified["layout"]["widgets"][0]["presentation_model"]["rows"][0]["value"] == "12.3400"
        published = reports.publish(created["id"], modified["revision"])
        assert published["status"] == "published"
        html, _, _ = reports.export(created["id"], "html")
        assert '<linearGradient' in html and 'data-date="2026-10-02"' in html and 'source-history' in html
        assert "不要导出" not in html and "9999" not in html
        office, _, _ = reports.export(created["id"], "docx")
        document = Document(BytesIO(office))
        assert len(document.tables) == 12
        assert "不要导出" not in str([[cell.text for row in table.rows for cell in row.cells] for table in document.tables])
        # 撤销真实字段授权后，不能借已保存的12种组件模型读回历史事实。
        source = data.get_source(dataset["datasource_id"])
        schema = deepcopy(source.schema_snapshot)
        next(column for column in schema[0]["columns"] if column["columnName"] == "value")["allowed_roles"] = ["quality-data-reader"]
        source.schema_snapshot = schema
        session.flush()
        lost_source_permission = ReportService(session, Actor("source-owner", "workspace", frozenset({"m07.read"})), settings)
        with pytest.raises(DomainError) as error:
            lost_source_permission.get(created["id"])
        assert error.value.status == 403
        session.rollback()
    engine.dispose()


def test_client_cannot_publish_truncated_evidence_by_changing_provenance_flag(tmp_path):
    """唯一新CSV两行，1条LIMIT预算查询；客户端篡改false也不能发布部分事实。"""
    engine = create_engine("sqlite:///" + (tmp_path / "truncated-source.sqlite3").as_posix())
    Base.metadata.create_all(engine)
    settings = Settings(data_dir=tmp_path / "truncated-source-data")
    actor = Actor("truncated-owner", "workspace", frozenset({"*"}))
    with Session(engine) as session:
        data, reports = DataService(session, actor, settings), ReportService(session, actor, settings)
        dataset = data.import_csv({"name": "截断临时样本", "csv_text": "name,value\n一班,1\n二班,2\n"})
        query = data.execute_query({"datasource_id": dataset["datasource_id"], "sql": f'SELECT * FROM "{dataset["table_name"]}"', "max_rows": 1})
        assert query["truncated"] is True and len(query["rows"]) == 1
        actual = widget("JCapsuleChart", query_id=query["id"])
        actual["data"] = deepcopy(query["rows"])
        actual["data_provenance"] = {"type": "query", "query_id": query["id"], "datasource_id": query["datasource_id"], "truncated": False}
        created = reports.create({"name": "来源标记不能伪造完整性", "layout": {"widgets": [actual]}})
        assert created["layout"]["widgets"][0]["presentation_model"]["status"] == "ready"
        with pytest.raises(DomainError) as error:
            reports.publish(created["id"], created["revision"])
        assert error.value.code == "report_data_not_ready"
        assert reports.get(created["id"]).status == "draft"
        session.rollback()
    engine.dispose()
