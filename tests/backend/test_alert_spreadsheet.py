"""预警读取直接交付 Excel，保留权限过滤结果、分页边界与原始数值。"""
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from app.modules.chat_artifacts.alert_spreadsheet import alert_workbook_spec, export_alert_spreadsheet


def example_view():
    """用两份报告与一条问题模拟分别分页的结果，不含任何外部来源。"""
    return {"resultType": "alert_problem_list", "title": "设备预警与问题",
        "summary": {"returnedReports": 2, "returnedProblems": 1,
                    "reportsMayHaveMore": True, "problemsMayHaveMore": False,
                    "offset": 50, "scopeNote": "仅本页，不是全量"},
        "reports": [
            {"id": "report-1", "title": "运行报告", "deviceName": "设备一", "rows": 72,
             "issues": 5, "outcomeLabel": "发现问题", "executionStatusLabel": "执行完成",
             "createTime": "2026-09-25T08:00:00", "href": "/alerts/reports?report=report-1"},
            {"id": "report-2", "title": "第二份报告", "deviceName": "设备二", "rows": 10,
             "issues": 1}],
        "problems": [{"id": "problem-1", "reportId": "report-2", "message": "液压 29.4 MPa",
                      "severity": "critical", "statusLabel": "待处理", "note": "=未执行公式"}],
        "processingHistory": None,
        "provenance": {"retrievedAt": "2026-09-26T02:00:00", "workspaceId": "workspace-1",
                       "source": "持久化预警报告与问题；未重新请求设备"}}


def test_direct_alert_export_preserves_saved_page_and_never_requeries(monkeypatch):
    from app.modules.artifact import service
    from app.modules.business_ai import service as business
    from app.providers.tool import builtin_provider
    captured, reads = {}, []
    view = example_view()
    original = deepcopy(view)
    run = SimpleNamespace(id="run-1", workspace_id="workspace-1", started_by_user_id="user-1")
    ctx = SimpleNamespace(db=SimpleNamespace(get=lambda model, identity: run if identity == run.id else None),
                          workspace_id="workspace-1", user_id="user-1")

    def overview(actor, payload):
        assert actor is ctx
        reads.append(payload)
        return SimpleNamespace(model_dump=lambda **kwargs: deepcopy(view))

    def create(db, **kwargs):
        captured.update(kwargs)
        return SimpleNamespace(id="artifact-1")

    monkeypatch.setattr(business, "alert_overview", overview)
    monkeypatch.setattr(service, "create_artifact", create)
    monkeypatch.setattr(service, "artifact_reference", lambda row: {"id": row.id, "name": captured["name"]})
    specs = {spec.name: spec for spec in builtin_provider.builtin_specs()}
    result = builtin_provider.invoke(specs["view_device_alerts"],
        {"limit": 50, "offset": 50, "export_name": "设备预警信息表"},
        {"auth_context": ctx, "run_id": run.id})

    assert len(reads) == 1 and reads[0].offset == 50
    assert view == original
    assert result.status == "success" and result.data["reports"] == view["reports"]
    assert result.artifacts == [{"id": "artifact-1", "name": "设备预警信息表.xlsx"}]
    workbook = load_workbook(BytesIO(captured["data"]))
    assert workbook.sheetnames == ["预警报告", "问题记录", "数据来源"]
    assert workbook["预警报告"].max_row == 3 and workbook["问题记录"].max_row == 2
    assert workbook["预警报告"]["F2"].value == 72
    assert workbook["预警报告"]["G2"].value == 5
    assert workbook["问题记录"]["E2"].value == "液压 29.4 MPa"
    assert workbook["问题记录"]["H2"].data_type == "s"
    source = dict(workbook["数据来源"].values)
    assert source["报告还有后续页"] is True and source["问题还有后续页"] is False
    assert source["分页起点"] == 50 and source["统计范围"] == "仅本页，不是全量"
    assert captured["metadata"]["delivery_evidence"]["scope"]["returnedReports"] == 2
    assert "name" in specs["export_query_spreadsheet"].input_schema["properties"]
    assert "export_name" in specs["view_device_alerts"].input_schema["properties"]


def test_alert_export_keeps_empty_evidence_and_processing_history_scope():
    view = example_view()
    view["reports"], view["problems"] = [], []
    view["processingHistory"] = {"problemId": "problem-1", "offset": 20, "hasMore": True,
        "events": [{"id": "event-1", "actorId": "actor-1", "createTime": "2026-09-26T02:00:00",
                    "previousStatus": "open", "status": "acknowledged", "note": "现场已确认"}]}
    spec = alert_workbook_spec(view)
    assert spec["sheets"][0]["rows"] == spec["sheets"][1]["rows"] == []
    assert spec["sheets"][2]["rows"][0][-1] == "现场已确认"
    source = dict(spec["sheets"][3]["rows"])
    assert source["处置日志分页起点"] == 20 and source["处置日志还有后续页"] is True


def test_alert_export_rejects_run_from_another_actor():
    from app.errors import PermissionErrorEx
    run = SimpleNamespace(id="run-1", workspace_id="workspace-1", started_by_user_id="another-user")
    ctx = SimpleNamespace(db=SimpleNamespace(get=lambda *args: run), workspace_id="workspace-1", user_id="user-1")
    with pytest.raises(PermissionErrorEx):
        export_alert_spreadsheet(ctx, {"run_id": run.id}, example_view(), "预警.xlsx")
