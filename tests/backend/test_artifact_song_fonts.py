"""Regression for Chinese chart glyphs and generated-document font policy.

中文说明：中文（宋体）字形与生成文档字体策略的回归：Word 标题/正文/表格/页眉页脚字体覆盖、
文档与表格渲染器设置宋体、PDF 内嵌字体而不是依赖阅读器的 CID 字体、PPT 字段字体策略与 schema 顺序不变。
隔离方式：monkeypatch + tmp_path 生成真实文档文件后解包检查，不连数据库、不调用模型。
边界：沙箱内中文字体覆盖被修复，且内部临时文件不得发布为产物。
"""
import io
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import pytest
from app.common.artifact_fonts import normalize_office_fonts, song_font_path, W
from app.modules.artifact.renderers import render_document, render_spreadsheet, render_pdf


def test_word_font_covers_titles_body_tables_headers_and_footers():
    from docx import Document
    doc = Document()
    doc.add_heading('巡检报告', 0)
    doc.add_paragraph('设备运行状态')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '温度 80 ℃'
    doc.sections[0].header.paragraphs[0].text = '设备分析'
    doc.sections[0].footer.paragraphs[0].text = '仅为演示数据'
    output = io.BytesIO(); doc.save(output)
    data = normalize_office_fonts(output.getvalue(), '.docx')
    Document(io.BytesIO(data))
    with ZipFile(io.BytesIO(data)) as archive:
        for filename in ['word/document.xml', 'word/styles.xml', 'word/header1.xml', 'word/footer1.xml']:
            xml = etree.fromstring(archive.read(filename))
            fonts = list(xml.iter(f'{{{W}}}rFonts'))
            assert fonts
            for font in fonts:
                for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                    assert font.get(f'{{{W}}}{name}') == '宋体'
                assert not any(k.endswith('Theme') for k in font.attrib)
    assert normalize_office_fonts(data, '.docx') == data


def test_document_and_spreadsheet_renderers_set_song_font():
    from docx import Document
    from openpyxl import load_workbook
    document = Document(io.BytesIO(render_document({'title': '报告', 'sections': [{'blocks': [{'type': 'p', 'text': '设备状态'}]}]})))
    assert all(run.font.name == '宋体' for p in document.paragraphs for run in p.runs)
    wb = load_workbook(io.BytesIO(render_spreadsheet({'columns': ['设备'], 'rows': [['演示设备']]})))
    assert wb.active['A1'].font.name == wb.active['A2'].font.name == '宋体'
    assert wb.active['A1'].font.bold
    assert wb.active['A1'].font.scheme is None
    assert wb.active['A2'].font.scheme is None


def test_pdf_embeds_song_font_instead_of_viewer_dependent_cid_font():
    try: song_font_path()
    except ValueError: pytest.skip('Song font not installed')
    data = render_pdf({'title': '设备温度分析', 'sections': [{'blocks': [{'type': 'p', 'text': '开机时间与转速'}]}]})
    assert b'/FontFile2' in data
    assert b'SimSun' in data


def test_sandbox_chinese_font_override_is_repaired_and_internal_files_not_published(tmp_path, monkeypatch):
    try: song_font_path()
    except ValueError: pytest.skip('Song font not installed')
    from app.config import settings
    from app.providers.tool.sandbox_provider import invoke
    from app.providers.tool.base import ToolSpec
    monkeypatch.setattr(settings, 'STORAGE_ROOT', str(tmp_path))
    code = '''import warnings
import matplotlib.pyplot as plt
warnings.filterwarnings('error', message='Glyph .* missing from font')
plt.rcParams['font.sans-serif'] = ['DefinitelyMissingFont', 'DejaVu Sans']
fig, ax = plt.subplots(figsize=(8, 4))
ax.bar(['电石炉供电变压器', '设备温度'], [6, 8])
ax.set_title('设备巡检问题统计')
fig.tight_layout()
fig.savefig('chinese.png')
result = {'font': ax.title.get_fontproperties().get_file()}
'''
    result = invoke(ToolSpec(name='python_execute', provider='sandbox', description='font test', timeout_seconds=30), {'code': code}, {'run_id': 'chinese-font', 'allow_network': False})
    assert result.status == 'success'
    assert result.data['files'] == ['chinese.png']
    assert result.data['result']['font'].endswith('simsun.ttc')
    assert Path(result.artifacts[0]['path']).stat().st_size > 5000


def test_office_font_policy_covers_ppt_fields_and_preserves_schema_order():
    from pptx import Presentation
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from app.common.artifact_fonts import A
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = '运行分析'
    run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    run.hyperlink.address = 'https://example.invalid/report'
    out = io.BytesIO(); deck.save(out)
    data = normalize_office_fonts(out.getvalue(), '.PPTX')
    parsed = Presentation(io.BytesIO(data))
    assert parsed.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.name == '宋体'
    with ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if name.startswith('ppt/') and name.endswith('.xml'):
                xml = etree.fromstring(archive.read(name))
                for prop in xml.iter(f'{{{A}}}rPr'):
                    fonts = [prop.find(f'{{{A}}}{kind}') for kind in ('latin', 'ea', 'cs')]
                    assert all(font is not None and font.get('typeface') == '宋体' for font in fonts)
                    link = prop.find(f'{{{A}}}hlinkClick')
                    if link is not None:
                        assert all(list(prop).index(font) < list(prop).index(link) for font in fonts)
    doc = Document()
    style = doc.styles.add_style('NoFontYet', WD_STYLE_TYPE.PARAGRAPH)
    style.paragraph_format.keep_with_next = True
    out = io.BytesIO(); doc.save(out)
    data = normalize_office_fonts(out.getvalue(), '.DOCX')
    with ZipFile(io.BytesIO(data)) as archive:
        xml = etree.fromstring(archive.read('word/styles.xml'))
        elem = next(s for s in xml.iter(f'{{{W}}}style') if s.get(f'{{{W}}}styleId') == 'NoFontYet')
        children = [etree.QName(child).localname for child in elem]
        assert children.index('rPr') > children.index('pPr') > children.index('name')
