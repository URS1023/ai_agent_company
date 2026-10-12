"""Unified documents/PPT center: real files, explicit templates and traceability.

中文说明：文档/PPT 中心：真实生成 docx/pdf/xlsx/pptx；模板必须显式选择（内置 20 套）；
来源需授权并快照、绝不带连接密钥；客户模板可上传预览删除；分阶段 PPT 与 AI 失败如实反馈；
权限与完成后不可变。
隔离方式：monkeypatch + client 夹具（内存 SQLite + tmp_path），不连业务库、不调用外部模型。
"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
from sqlalchemy import select

from app.errors import ModelError
from app.ids import new_id
from app.models import Artifact, DataSource, WorkspaceMember
from app.modules.analytics.models import SourceMetadata
from app.modules.documents.models import DocumentGeneration, DocumentGenerationEvent
from test_system_foundation import activate_user, create_role, create_user


API = "/api/v1/documents"


def result(response, status=200):
    assert response.status_code == status, response.text
    return response.json()["data"]


def create_generation(client, headers, payload):
    return result(client.post(API + "/generations", headers=headers, json=payload), 201)


def execute(client, headers, generation):
    return result(client.post(
        API + f"/generations/{generation['id']}/execute",
        headers=headers,
        json={"version": generation["version"]},
    ))


def document_spec(title="设备运行月报"):
    return {
        "title": title,
        "metadata": {"author": "质量中心"},
        "sections": [{
            "title": "本月结论",
            "blocks": [
                {"type": "p", "text": "设备运行稳定，关键指标均已完成复核。"},
                {"type": "bullet", "text": "一号生产线保持正常运行"},
                {"type": "table", "rows": [["设备", "状态"], ["A-01", "正常"]]},
            ],
        }],
    }


def presentation_spec():
    return {
        "title": "财务与设备经营分析",
        "subtitle": "真实经营口径",
        "slides": [
            {
                "title": "经营指标趋势",
                "layout": "chart",
                "key_message": "收入增长同时设备停机率下降，经营与生产效率同步改善。",
                "items": [{
                    "type": "chart",
                    "chart_type": "column",
                    "categories": ["一月", "二月", "三月"],
                    "series": [{"name": "收入", "values": [120, 138, 156]}],
                    "source": "财务生产库",
                }],
            },
            {
                "title": "设备健康判断",
                "layout": "cards",
                "key_message": "核心设备总体稳定，重点关注温度波动和保养计划执行。",
                "items": [
                    {"title": "运行", "description": "主设备连续运行时间符合计划，未发生非计划停机。"},
                    {"title": "温度", "description": "二号设备存在短时波动，已纳入下一班次复核。"},
                    {"title": "维护", "description": "本月保养任务按期完成，备件库存能够覆盖近期需求。"},
                ],
            },
            {
                "title": "下月行动",
                "layout": "content",
                "key_message": "围绕效率、可靠性和成本三个方向执行并逐周复盘。",
                "items": [
                    {"title": "效率", "description": "优化换线节拍并跟踪单位工时产出变化。"},
                    {"title": "可靠性", "description": "增加温度异常设备的巡检频率与复核责任人。"},
                    {"title": "成本", "description": "对能耗和备件消耗进行分设备归集与趋势分析。"},
                ],
            },
        ],
    }


def test_templates_require_explicit_selection_and_have_twenty_builtins(client, administrator):
    templates = result(client.get(API + "/templates", headers=administrator))
    builtins = [item for item in templates["items"] if item["source"] == "builtin"]
    assert templates["builtinCount"] == 20
    assert len(builtins) == 20
    assert templates["selectionRequired"] is True
    assert "defaultId" not in templates and "default_id" not in templates
    assert all(item.get("preview", {}).get("image_url") for item in builtins)

    rejected = client.post(API + "/generations", headers=administrator, json={
        "name": "没有选模板", "outputType": "presentation", "mode": "manual",
        "pageCount": 5, "spec": presentation_spec(),
    })
    assert rejected.status_code == 422
    assert "不会套用默认模板" in rejected.text
    invalid = client.post(API + "/generations", headers=administrator, json={
        "name": "不存在模板", "outputType": "presentation", "mode": "manual",
        "templateId": "missing-template", "pageCount": 5, "spec": presentation_spec(),
    })
    assert invalid.status_code == 404


def test_manual_generation_creates_real_docx_pdf_xlsx_and_pptx(client, administrator, db, monkeypatch):
    doc = execute(client, administrator, create_generation(client, administrator, {
        "name": "设备运行月报", "outputType": "document", "mode": "manual", "spec": document_spec(),
    }))
    response = client.get(API + f"/generations/{doc['id']}/download", headers=administrator)
    assert response.status_code == 200 and response.content[:2] == b"PK"
    opened_doc = Document(io.BytesIO(response.content))
    assert "设备运行月报" in "\n".join(paragraph.text for paragraph in opened_doc.paragraphs)

    pdf = execute(client, administrator, create_generation(client, administrator, {
        "name": "设备质量报告", "outputType": "pdf", "mode": "manual",
        "spec": document_spec("设备质量报告 <已复核>"),
    }))
    response = client.get(API + f"/generations/{pdf['id']}/download", headers=administrator)
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("attachment;")
    assert response.status_code == 200 and response.content.startswith(b"%PDF")
    assert len(PdfReader(io.BytesIO(response.content)).pages) >= 1
    assert pdf["actualPageCount"] >= 1
    # fetch 字节模式不携带附件头，仍须交付同一份完整 PDF，不能用空响应冒充成功。
    fetched = client.get(API + f"/generations/{pdf['id']}/download?attachment=false", headers=administrator)
    assert fetched.status_code == 200 and fetched.content == response.content
    assert fetched.headers["content-type"] == "application/octet-stream"
    assert fetched.headers["x-content-type-options"] == "nosniff"
    assert "content-disposition" not in fetched.headers

    workbook_generation = execute(client, administrator, create_generation(client, administrator, {
        "name": "经营台账", "outputType": "spreadsheet", "mode": "manual",
        "spec": {"sheets": [
            {"name": "财务", "columns": ["月份", "收入"], "rows": [["一月", 120], ["二月", 138]]},
            {"name": "设备", "columns": ["设备", "备注"], "rows": [["A-01", "=HYPERLINK(\"bad\")"]]},
        ]},
    }))
    response = client.get(API + f"/generations/{workbook_generation['id']}/download", headers=administrator)
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["财务", "设备"]
    assert workbook["设备"]["B2"].value.startswith("'")
    assert workbook_generation["actualPageCount"] == 2

    ppt = execute(client, administrator, create_generation(client, administrator, {
        "name": "财务设备大屏汇报", "outputType": "presentation", "mode": "manual",
        "templateId": "executive-ivory", "pageCount": 5, "spec": presentation_spec(),
    }))
    response = client.get(API + f"/generations/{ppt['id']}/download", headers=administrator)
    deck = Presentation(io.BytesIO(response.content))
    assert len(deck.slides) == 5
    assert sum(1 for slide in deck.slides for shape in slide.shapes if getattr(shape, "has_chart", False)) == 1
    assert ppt["actualPageCount"] == 5
    assert ppt["artifact"]["metadata"]["nativeChartCount"] == 1
    assert ppt["artifact"]["metadata"]["templateId"] == "executive-ivory"

    monkeypatch.setattr("app.modules.documents.service.render_pptx_preview", lambda *_args, **_kwargs: None)
    preview = client.get(API + f"/generations/{ppt['id']}/preview", headers=administrator)
    assert preview.status_code == 200 and preview.headers["content-type"].startswith("image/svg+xml")
    events = result(client.get(API + f"/generations/{ppt['id']}/events", headers=administrator))
    assert [item["stage"] for item in events][-1] == "completed"
    assert {"validating", "authoring", "rendering", "finalizing", "completed"}.issubset({item["stage"] for item in events})

    db.expire_all()
    stored = db.get(DocumentGeneration, ppt["id"])
    artifact = db.get(Artifact, stored.artifact_id)
    assert stored.run_id and artifact.run_id == stored.run_id and artifact.sha256


def test_sources_are_authorized_snapshotted_and_never_include_connection_secrets(client, administrator, db):
    workspace_id = administrator["X-Workspace-Id"]
    owner_id = db.scalar(select(WorkspaceMember.user_id).where(WorkspaceMember.workspace_id == workspace_id))
    now = datetime.utcnow()
    source = DataSource(
        id=new_id(), workspace_id=workspace_id, owner_user_id=owner_id, name="财务生产库", provider="sqlite",
        connection_config={"version": 1, "databaseName": "finance", "password": "never-leak"},
        status="active", visibility="workspace", created_at=now, updated_at=now,
    )
    metadata = SourceMetadata(
        workspace_id=workspace_id, owner_user_id=owner_id, datasource_id=source.id, source_version=1,
        status="ready", error_message="", tables_json=[{
            "tableName": "orders", "columns": [
                {"columnName": "amount", "dataType": "DECIMAL", "comment": "含税金额", "nullable": False, "primaryKey": False},
            ],
        }],
    )
    db.add_all([source, metadata])
    db.commit()
    generation = execute(client, administrator, create_generation(client, administrator, {
        "name": "财务口径", "outputType": "document", "mode": "manual", "spec": document_spec("财务口径"),
        "sourceRefs": [{"type": "datasource", "id": source.id, "field": "orders.amount"}],
    }))
    assert generation["sources"][0]["name"] == "财务生产库"
    assert generation["sources"][0]["field"] == "orders.amount"
    assert "never-leak" not in str(generation)
    assert "never-leak" not in str(generation["artifact"]["metadata"])


def test_customer_template_upload_preview_and_delete(client, administrator, monkeypatch):
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "客户品牌模板"
    buffer = io.BytesIO()
    deck.save(buffer)
    monkeypatch.setattr("app.modules.presentation.service.cache_custom_preview", lambda *_args, **_kwargs: False)
    uploaded = result(client.post(
        API + "/templates",
        headers=administrator,
        data={"name": "蓝色客户模板"},
        files={"file": ("customer.pptx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    ), 201)
    assert uploaded["source"] == "custom" and uploaded["analysis"]["layout_count"] >= 1
    templates = result(client.get(API + "/templates", headers=administrator))
    assert uploaded["id"] in {item["id"] for item in templates["items"]}
    preview = client.get(API + f"/templates/{uploaded['id']}/preview", headers=administrator)
    assert preview.status_code == 200 and preview.headers["content-type"].startswith("image/svg+xml")
    deleted = result(client.delete(API + f"/templates/{uploaded['id']}", headers=administrator))
    assert deleted["deleted"] is True
    assert client.get(API + f"/templates/{uploaded['id']}/preview", headers=administrator).status_code == 404


def test_staged_ppt_and_ai_failure_are_truthful(client, administrator, monkeypatch):
    generation = create_generation(client, administrator, {
        "name": "分阶段经营汇报", "outputType": "presentation", "mode": "manual",
        "templateId": "executive-ivory", "pageCount": 5, "spec": {},
    })
    plan = result(client.post(API + f"/generations/{generation['id']}/presentation-plan", headers=administrator, json={
        "version": generation["version"], "title": "分阶段经营汇报", "slides": [
            {"title": "经营结论", "layout": "cards"},
            {"title": "设备状态", "layout": "cards"},
            {"title": "行动计划", "layout": "content"},
        ],
    }))
    assert plan["total_page_count"] == 5 and plan["missing_slide_numbers"] == [1, 2, 3]
    batch = result(client.post(API + f"/generations/{generation['id']}/presentation-slides", headers=administrator, json={
        "version": plan["version"], "slides": [
            {"slide_number": 1, **presentation_spec()["slides"][1], "title": "经营结论"},
            {"slide_number": 2, **presentation_spec()["slides"][1], "title": "设备状态"},
            {"slide_number": 3, **presentation_spec()["slides"][2], "title": "行动计划"},
        ],
    }))
    assert batch["missing_slide_numbers"] == []
    completed = result(client.post(API + f"/generations/{generation['id']}/execute", headers=administrator, json={
        "version": batch["version"],
    }))
    assert completed["status"] == "completed" and completed["actualPageCount"] == 5

    ai = create_generation(client, administrator, {
        "name": "AI 质量报告", "outputType": "pdf", "mode": "ai", "prompt": "根据质量数据生成报告", "spec": {},
    })
    monkeypatch.setattr(
        "app.modules.documents.service.model.chat_for_run",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ModelError("模型服务当前不可用")),
    )
    failed = client.post(API + f"/generations/{ai['id']}/execute", headers=administrator, json={"version": ai["version"]})
    assert failed.status_code == 502 and "模型服务当前不可用" in failed.text
    stored = result(client.get(API + f"/generations/{ai['id']}", headers=administrator))
    assert stored["status"] == "failed" and stored["runId"]
    assert stored["errorMessage"] == "模型服务当前不可用"


def test_document_permissions_and_completed_immutability(client, administrator):
    generation = execute(client, administrator, create_generation(client, administrator, {
        "name": "不可覆盖报告", "outputType": "document", "mode": "manual", "spec": document_spec(),
    }))
    overwrite = client.put(API + f"/generations/{generation['id']}", headers=administrator, json={
        "name": "覆盖", "outputType": "document", "mode": "manual", "spec": document_spec(),
        "version": generation["version"],
    })
    assert overwrite.status_code == 409
    copied = result(client.post(API + f"/generations/{generation['id']}/copy", headers=administrator, json={"name": "新草稿"}), 201)
    assert copied["status"] == "draft" and copied["artifact"] is None

    role = create_role(client, administrator, code="document_viewer", permissions=["home.view", "documents.view"], scope="workspace")
    create_user(client, administrator, [role["id"]], username="document-viewer")
    viewer = activate_user(client, username="document-viewer")
    listed = result(client.get(API + "/generations", headers=viewer))
    assert generation["id"] in {item["id"] for item in listed["items"]}
    assert client.post(API + "/generations", headers=viewer, json={
        "name": "越权", "outputType": "document", "mode": "manual", "spec": document_spec(),
    }).status_code == 403
    assert client.get(API + f"/generations/{generation['id']}/download", headers=viewer).status_code == 403
    assert client.get(API + f"/generations/{generation['id']}/download?attachment=false", headers=viewer).status_code == 403
