"""B20 逐份成果修订：真实 Worker、临时 SQLite 和最终文件，禁止源查询与模型调用。"""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import pytest

from app.business_worker import chat, native_deliveries, tools
from app.common.storage import get_storage
from app.models import Artifact, Run
from app.modules.analytics import execution, metadata, queries
from app.providers import model
from test_business_worker import prepare, worker_context  # noqa: F401
from test_dsh_compound_contract import deck_bytes, local_outputs, present_bytes  # noqa: F401


OLD_TITLE = "隔离原始标题"
NEW_TITLE = "隔离修订标题"
WORD_TEXT = "本说明保持原文，PPT 标题修订不要求重新生成本文件。"
SHEET_ROWS = [("设备", "数量"), ("设备甲", 12), ("设备乙", 18)]


@pytest.fixture(autouse=True)
def no_business_io(monkeypatch, record_property):
    """模型配置仍用 worker_context；任何真实模型/源访问一经尝试立即失败。"""
    counts = {"source_queries": 0, "source_connections": 0, "source_metadata": 0, "model_calls": 0}

    def forbidden(kind):
        def reject(*args, **kwargs):
            counts[kind] += 1
            pytest.fail(f"修订合同验收不得访问业务源或模型：{kind}")
        return reject

    monkeypatch.setattr(queries, "execute_read", forbidden("source_queries"))
    monkeypatch.setattr(execution, "readonly_connection", forbidden("source_connections"))
    monkeypatch.setattr(metadata, "inspect_sql_source", forbidden("source_metadata"))
    monkeypatch.setattr(model, "chat", forbidden("model_calls"))
    monkeypatch.setattr(model, "chat_for_run", forbidden("model_calls"))
    yield counts
    # 即使主断言失败，也把零外部访问计数保留到该单例的独立 XML。
    record_property("source_model_counts", json.dumps(counts, ensure_ascii=False, sort_keys=True))
    assert counts == {"source_queries": 0, "source_connections": 0, "source_metadata": 0, "model_calls": 0}


def three_page_deck(title):
    """沿用已有真实非空三页夹具，仅替换第一张幻灯片标题。"""
    deck = Presentation(BytesIO(deck_bytes(3)))
    deck.slides[0].shapes.title.text = title
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def stored_file(db, artifact_id):
    """读取标准存储最终字节，独立计算 SHA；不能只信工具返回的文件名。"""
    db.expire_all()
    row = db.get(Artifact, artifact_id)
    assert row is not None and row.storage_key, artifact_id
    data = get_storage().read_bytes(row.storage_key)
    digest = sha256(data).hexdigest()
    assert digest == row.sha256
    return data, {"id": row.id, "run_id": row.run_id, "type": row.type,
                  "sha256": digest, "source_sha256": (row.metadata_json or {}).get("source_sha256"),
                  "parent_artifact_id": row.parent_artifact_id, "status": row.status}


def slide_titles(data):
    """python-pptx 按正文实际页序读取，不使用 metadata 页数或文件名。"""
    deck = Presentation(BytesIO(data))
    assert len(deck.slides) == 3
    assert all(slide.shapes.title is not None for slide in deck.slides)
    return [slide.shapes.title.text for slide in deck.slides]


def assert_word_and_sheet(word, sheet):
    """读取实际 DOCX 和 XLSX 正文，验证旧成果仍包含初始交付内容。"""
    paragraphs = [paragraph.text for paragraph in Document(BytesIO(word)).paragraphs]
    assert WORD_TEXT in paragraphs
    workbook = load_workbook(BytesIO(sheet), data_only=False)
    try:
        assert workbook.sheetnames == ["固定明细"]
        assert list(workbook["固定明细"].values) == SHEET_ROWS
    finally:
        workbook.close()
    return {"word_paragraphs": paragraphs, "sheet_name": "固定明细", "sheet_rows": SHEET_ROWS}


def test_unchanged_prior_ppt_represented_in_new_run_cannot_complete_title_revision(
    local_outputs, db, record_property, no_business_io,
):
    """真实旧文件跨 Run 原样 present，不能满足明确要求的新标题。"""
    first = prepare(local_outputs, "制作一份3页PPT")
    source = three_page_deck(OLD_TITLE)
    original = present_bytes(first, "原始报告.pptx", source)
    original_bytes, original_proof = stored_file(db, original["id"])
    assert slide_titles(original_bytes)[0] == OLD_TITLE
    completed = chat.settle(first["run_id"], "completed", "已交付原始三页文件")
    assert completed["status"] == "completed", completed
    db.expire_all()
    original_snapshot = deepcopy(db.get(Run, first["run_id"]).context_snapshot)

    revised = prepare(local_outputs, f"修改刚才的PPT，把第1页标题改成{NEW_TITLE}",
                      task_id=first["task"]["task_id"])
    assert revised["task"]["task_id"] == first["task"]["task_id"]
    assert revised["run_id"] != first["run_id"]
    assert revised["task"]["requires_new_delivery"] is True
    assert revised["task"]["required_deliveries"] == first["task"]["required_deliveries"]
    source_path = Path(first["cwd"]) / "原始报告.pptx"
    assert Path(revised["cwd"]).resolve() == Path(first["cwd"]).resolve()
    assert source_path.read_bytes() == source
    # 不重写源文件；同一稳定会话目录中的旧字节经真实原生交付入口再次登记。
    receipt = native_deliveries.present(revised["run_id"], [{"path": source_path.name}])
    assert receipt["count"] == 1
    represented = receipt["artifacts"][0]
    represented_bytes, represented_proof = stored_file(db, represented["id"])
    assert represented_proof["source_sha256"] == original_proof["source_sha256"] == sha256(source).hexdigest()
    actual_titles = slide_titles(represented_bytes)
    assert actual_titles == slide_titles(original_bytes)
    assert actual_titles[0] == OLD_TITLE and NEW_TITLE not in actual_titles

    settled = chat.settle(revised["run_id"], "completed", "声称已完成标题修改")
    preserved_bytes, preserved_proof = stored_file(db, original["id"])
    assert preserved_bytes == original_bytes and preserved_proof["sha256"] == original_proof["sha256"]
    assert db.get(Run, first["run_id"]).context_snapshot == original_snapshot
    assert not any(no_business_io.values())
    record_property("revision_evidence", json.dumps({
        "task_id": first["task"]["task_id"], "old": original_proof, "represented": represented_proof,
        "requested_title": NEW_TITLE, "actual_titles": actual_titles,
        "settled_status": settled["status"], "missing": settled["task"].get("missing_deliveries"),
        "constraint_failures": settled["task"].get("constraint_failures"),
    }, ensure_ascii=False, sort_keys=True))
    # 行为红例：即使新建 Artifact 和当前 Run 的文件计数成立，旧内容仍未完成新要求。
    assert settled["status"] == "partial", settled


def test_only_ppt_title_revision_keeps_original_word_and_excel_fulfilled(
    local_outputs, db, record_property, no_business_io,
):
    """只修订 PPT，真实且未要求修改的 Word/Excel 应继续满足原交付义务。"""
    first = prepare(local_outputs, "制作一份3页PPT，再生成一份Word和一份Excel")
    assert {item["type"]: item["count"] for item in first["task"]["required_deliveries"]} == {
        "presentation": 1, "document": 1, "spreadsheet": 1}
    original = present_bytes(first, "原始报告.pptx", three_page_deck(OLD_TITLE))
    word = tools.execute(first["run_id"], "original-word", 1, "render_document", {
        "name": "保留说明", "document_spec": {"title": "保留说明", "sections": [
            {"title": "原始交付", "blocks": [{"type": "p", "text": WORD_TEXT}]}]}})
    assert word["status"] == "success" and len(word["artifacts"]) == 1, word
    sheet = tools.execute(first["run_id"], "original-sheet", 2, "render_spreadsheet", {
        "name": "保留明细", "spec": {"sheet_name": "固定明细", "columns": list(SHEET_ROWS[0]),
                                      "rows": [list(row) for row in SHEET_ROWS[1:]]}})
    assert sheet["status"] == "success" and len(sheet["artifacts"]) == 1, sheet
    original_bytes, original_proof = stored_file(db, original["id"])
    word_bytes, word_proof = stored_file(db, word["artifacts"][0]["id"])
    sheet_bytes, sheet_proof = stored_file(db, sheet["artifacts"][0]["id"])
    original_titles = slide_titles(original_bytes)
    assert original_titles[0] == OLD_TITLE
    original_content = assert_word_and_sheet(word_bytes, sheet_bytes)
    completed = chat.settle(first["run_id"], "completed", "三个实际文件均已交付")
    assert completed["status"] == "completed", completed
    db.expire_all()
    original_snapshot = deepcopy(db.get(Run, first["run_id"]).context_snapshot)

    revised = prepare(local_outputs,
                      f"只修改刚才的PPT第1页标题为{NEW_TITLE}，其他已交付文件保留",
                      task_id=first["task"]["task_id"])
    assert revised["task"]["task_id"] == first["task"]["task_id"]
    assert revised["run_id"] != first["run_id"]
    assert revised["task"]["requires_new_delivery"] is True
    # 不传显式替换清单；保留三种原义务，不能靠删掉 Word/Excel 要求把测试变绿。
    assert revised["task"]["required_deliveries"] == first["task"]["required_deliveries"]
    changed = present_bytes(revised, "修订报告.pptx", three_page_deck(NEW_TITLE))
    changed_bytes, changed_proof = stored_file(db, changed["id"])
    assert changed_proof["id"] != original_proof["id"]
    assert changed_proof["run_id"] == revised["run_id"]
    changed_titles = slide_titles(changed_bytes)
    assert changed_titles == [NEW_TITLE, *original_titles[1:]]
    assert changed_proof["sha256"] != original_proof["sha256"]

    settled = chat.settle(revised["run_id"], "completed", "已修订PPT，Word和Excel保留原文件")
    old_ppt_after, _ = stored_file(db, original["id"])
    word_after, word_after_proof = stored_file(db, word_proof["id"])
    sheet_after, sheet_after_proof = stored_file(db, sheet_proof["id"])
    assert old_ppt_after == original_bytes
    assert word_after == word_bytes and sheet_after == sheet_bytes
    assert word_after_proof["sha256"] == word_proof["sha256"]
    assert sheet_after_proof["sha256"] == sheet_proof["sha256"]
    assert assert_word_and_sheet(word_after, sheet_after) == original_content
    assert db.get(Run, first["run_id"]).context_snapshot == original_snapshot
    for proof in (word_proof, sheet_proof):
        retained = [item for item in settled["deliveries"] if item.get("artifact_id") == proof["id"]]
        assert len(retained) == 1 and retained[0]["status"] == "ready"
        assert retained[0]["source_run_id"] == first["run_id"]
        assert not any(item["type"] == proof["type"] and item["source_run_id"] == revised["run_id"]
                       for item in settled["deliveries"])
    assert not any(no_business_io.values())
    record_property("revision_evidence", json.dumps({
        "task_id": first["task"]["task_id"], "old_ppt": original_proof, "new_ppt": changed_proof,
        "word": word_after_proof, "sheet": sheet_after_proof, "old_content": original_content,
        "actual_new_titles": changed_titles, "settled_status": settled["status"],
        "missing": settled["task"].get("missing_deliveries"),
        "constraint_failures": settled["task"].get("constraint_failures"),
    }, ensure_ascii=False, sort_keys=True))
    # 行为红例：新 PPT 已真实修改，旧 Word/Excel 可核验且未失效，不应要求重复生成。
    assert settled["status"] == "completed", settled
    assert settled["task"]["missing_deliveries"] == []


def completed_decks(context, titles):
    """实际完成具名 PPT，后续修订仍按原任务身份发起。"""
    first = prepare(context, "，再".join(f"生成{title}3页PPT" for title in titles))
    files = {title: present_bytes(first, title + ".pptx", three_page_deck(title)) for title in titles}
    result = chat.settle(first["run_id"], "completed", "具名文件已生成")
    assert result["status"] == "completed", result
    return first, files


def test_failed_title_revision_can_continue_with_real_changed_file(local_outputs, db):
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    present_bytes(revised, "原样.pptx", three_page_deck("ERP介绍"))
    partial = chat.settle(revised["run_id"], "completed", "尚未修改")
    assert partial["status"] == "partial"
    continued = prepare(local_outputs, "继续", task_id=first["task"]["task_id"])
    assert continued["task"]["delivery_objects"] == partial["task"]["delivery_objects"]
    fresh = present_bytes(continued, "真正修订.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(continued["run_id"], "completed", "已修改指定标题")
    assert result["status"] == "completed", result
    obj = result["task"]["delivery_objects"][0]
    assert obj["revision"] == 2 and obj["fulfillment"]["id"] == fresh["id"]
    assert obj["base"]["id"] == files["ERP介绍"]["id"]
    assert slide_titles(stored_file(db, fresh["id"])[0])[0] == NEW_TITLE


@pytest.mark.parametrize("change", ["zip_properties", "other_page", "wrong_base"])
def test_new_byte_digest_alone_cannot_prove_requested_title_revision(local_outputs, db, change):
    from zipfile import ZipFile, ZIP_STORED
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    original, _ = stored_file(db, files["ERP介绍"]["id"])
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    if change == "zip_properties":
        output = BytesIO()
        with ZipFile(BytesIO(original)) as source, ZipFile(output, "w", compression=ZIP_STORED) as target:
            for item in source.infolist():
                raw = source.read(item.filename)
                if item.filename == "docProps/core.xml":
                    raw = raw.replace(b'</cp:coreProperties>', b'<cp:keywords>changed package only</cp:keywords></cp:coreProperties>')
                target.writestr(item.filename, raw)
        modified = output.getvalue()
    else:
        deck = Presentation(BytesIO(original))
        deck.slides[1].shapes.title.text = "无关页被替换"
        if change == "wrong_base":
            deck.slides[0].shapes.title.text = NEW_TITLE
        output = BytesIO()
        deck.save(output)
        modified = output.getvalue()
    assert sha256(modified).hexdigest() != sha256(original).hexdigest()
    present_bytes(revised, "换名修订.pptx", modified)
    result = chat.settle(revised["run_id"], "completed", "仅字节改变不足以证明完成")
    assert result["status"] == "partial", result
    assert any(row["code"] == "delivery_revision_unverified" for row in result["task"]["constraint_failures"])


def test_named_revision_keeps_other_named_deck_and_bound_page_contract(local_outputs, db):
    first, files = completed_decks(local_outputs, ["ERP介绍", "MES介绍"])
    untouched, proof = stored_file(db, files["MES介绍"]["id"])
    revised = prepare(local_outputs, f"把ERP介绍PPT第1页标题改成{NEW_TITLE}", task_id=first["task"]["task_id"])
    current = present_bytes(revised, "ERP修订.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(revised["run_id"], "completed", "仅修订指定文件")
    assert result["status"] == "completed", result
    assert {row["revision"] for row in result["task"]["delivery_objects"]} == {1, 2}
    assert {row["artifact_id"] for row in result["task"]["constraint_receipts"]} == {current["id"], files["MES介绍"]["id"]}
    assert stored_file(db, proof["id"])[0] == untouched


def test_ambiguous_same_type_revision_does_not_guess_which_old_file(local_outputs):
    first, _ = completed_decks(local_outputs, ["ERP介绍", "MES介绍"])
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    present_bytes(revised, "不明对象.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(revised["run_id"], "completed", "未明确哪一份")
    assert result["status"] == "partial", result
    assert any(row["code"] == "delivery_revision_ambiguous" for row in result["task"]["constraint_failures"])


def test_same_revision_can_be_completed_across_two_new_runs(local_outputs):
    first, _ = completed_decks(local_outputs, ["ERP介绍", "MES介绍"])
    revised = prepare(local_outputs, f"修改两份PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    first_new = present_bytes(revised, "第一份修订.pptx", three_page_deck(NEW_TITLE))
    partial = chat.settle(revised["run_id"], "completed", "只完成一份修订")
    assert partial["status"] == "partial", partial
    assert len(partial["task"]["revision_receipts"]) == 1
    continued = prepare(local_outputs, "继续", task_id=first["task"]["task_id"])
    second_new = present_bytes(continued, "第二份修订.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(continued["run_id"], "completed", "两份修订已完成")
    assert result["status"] == "completed", result
    assert {row["revision"] for row in result["task"]["delivery_objects"]} == {2}
    assert {row["id"] for row in result["task"]["revision_receipts"]} == {first_new["id"], second_new["id"]}


def test_queue_edit_cannot_clear_an_unfulfilled_revision(local_outputs):
    first, _ = completed_decks(local_outputs, ["ERP介绍"])
    content = f"修改刚才的PPT第1页标题为{NEW_TITLE}"
    revised = prepare(local_outputs, content, task_id=first["task"]["task_id"])
    before = deepcopy(revised["task"]["delivery_objects"])
    edited = chat.queue_update(revised["run_id"], "edit", content)
    assert edited["task"]["delivery_objects"] == before
    result = chat.settle(revised["run_id"], "completed", "尚未交付新内容")
    assert result["status"] == "partial", result


def test_revised_task_can_add_an_independent_delivery_without_losing_slots(local_outputs):
    first, _ = completed_decks(local_outputs, ["ERP介绍"])
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    present_bytes(revised, "修订.pptx", three_page_deck(NEW_TITLE))
    assert chat.settle(revised["run_id"], "completed", "已修订")["status"] == "completed"
    added = prepare(local_outputs, "再生成一份Word", task_id=first["task"]["task_id"])
    word = tools.execute(added["run_id"], "added-word", 1, "render_document", {
        "name": "新增说明", "document_spec": {"title": "新增说明", "sections": []}})
    assert word["status"] == "success", word
    result = chat.settle(added["run_id"], "completed", "新增文件已交付")
    assert result["status"] == "completed", result
    assert {row["type"] for row in result["task"]["revision_receipts"]} == {"presentation", "document"}


def test_more_historical_files_than_slots_does_not_erase_revision_baseline(local_outputs):
    first = prepare(local_outputs, "生成一份3页PPT")
    present_bytes(first, "第一稿.pptx", three_page_deck("ERP介绍"))
    present_bytes(first, "备选稿.pptx", three_page_deck("MES介绍"))
    assert chat.settle(first["run_id"], "completed", "存在两稿")["status"] == "completed"
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    present_bytes(revised, "未确定基线.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(revised["run_id"], "completed", "未明确哪稿")
    assert result["status"] == "partial", result
    assert any(row["code"] == "delivery_revision_ambiguous" for row in result["task"]["constraint_failures"])


def test_late_previous_run_file_cannot_fulfill_new_revision(local_outputs):
    first, _ = completed_decks(local_outputs, ["ERP介绍"])
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    late = present_bytes(first, "旧轮迟到.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(revised["run_id"], "completed", "旧轮次迟到")
    assert any(row["id"] == late["id"] for row in result["deliveries"])
    assert result["status"] == "partial", result


def test_wrong_page_count_does_not_pin_revision_to_unrepairable_file(local_outputs):
    first, _ = completed_decks(local_outputs, ["ERP介绍"])
    revised = prepare(local_outputs, "重新生成刚才的PPT", task_id=first["task"]["task_id"])
    wrong = present_bytes(revised, "错误两页.pptx", deck_bytes(2))
    partial = chat.settle(revised["run_id"], "completed", "页数不足")
    assert partial["status"] == "partial", partial
    assert partial["task"]["delivery_objects"][0]["fulfillment"] is None
    continued = prepare(local_outputs, "继续", task_id=first["task"]["task_id"])
    correct = present_bytes(continued, "正确三页.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(continued["run_id"], "completed", "已补正")
    assert result["status"] == "completed", result
    assert result["task"]["delivery_objects"][0]["fulfillment"]["id"] == correct["id"] != wrong["id"]


@pytest.mark.parametrize("fault", ["stored_bytes", "files_scope"])
def test_inherited_file_is_rechecked_and_can_recover_without_recreation(local_outputs, db, client, fault):
    from test_dsh_compound_boundaries import client_context
    context, service = client_context(client, db, local_outputs)
    first = prepare(context, "制作一份3页PPT和Word")
    present_bytes(first, "原PPT.pptx", three_page_deck(OLD_TITLE))
    word = tools.execute(first["run_id"], "retained-word", 1, "render_document", {
        "name": "原Word", "document_spec": {"title": "原Word", "sections": []}})
    assert word["status"] == "success", word
    word_id = word["artifacts"][0]["id"]
    old_bytes, _ = stored_file(db, word_id)
    assert chat.settle(first["run_id"], "completed", "原交付")["status"] == "completed"
    revised = prepare(context, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    present_bytes(revised, "新PPT.pptx", three_page_deck(NEW_TITLE))
    path = Path(get_storage().get_path(db.get(Artifact, word_id).storage_key))
    scopes = list(service.scopes)
    if fault == "stored_bytes":
        path.write_bytes("已损坏的文件".encode())
    else:
        service.scopes = [scope for scope in scopes if scope != "files.view"]
        db.commit()
    partial = chat.settle(revised["run_id"], "completed", "当前证据不可用")
    assert partial["status"] == "partial", partial
    inherited = next(row for row in partial["task"]["delivery_objects"] if row["type"] == "document")
    assert inherited["fulfillment"]["id"] == word_id
    path.write_bytes(old_bytes)
    service.scopes = scopes
    db.commit()
    continued = prepare(context, "继续", task_id=first["task"]["task_id"])
    result = chat.settle(continued["run_id"], "completed", "原证据恢复")
    assert result["status"] == "completed", result
    assert next(row for row in result["task"]["revision_receipts"] if row["type"] == "document")["id"] == word_id


@pytest.mark.parametrize("kind", ["Word", "Excel"])
@pytest.mark.parametrize("changed", [False, True], ids=["unchanged", "changed"])
def test_word_and_excel_regeneration_requires_real_content_change(local_outputs, db, kind, changed):
    def payload(value):
        if kind == "Word":
            document = Document()
            document.add_paragraph(value)
            output = BytesIO()
            document.save(output)
            return "docx", output.getvalue()
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["内容", value])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return "xlsx", output.getvalue()

    first = prepare(local_outputs, f"生成一份{kind}")
    extension, original = payload("原始正文")
    present_bytes(first, "原始." + extension, original)
    assert chat.settle(first["run_id"], "completed", "原交付")["status"] == "completed"
    revised = prepare(local_outputs, f"重新生成刚才的{kind}", task_id=first["task"]["task_id"])
    current = present_bytes(revised, "重新交付." + extension, payload("新的正文")[1] if changed else original)
    actual, _ = stored_file(db, current["id"])
    if kind == "Word":
        assert Document(BytesIO(actual)).paragraphs[0].text == ("新的正文" if changed else "原始正文")
    else:
        workbook = load_workbook(BytesIO(actual))
        assert workbook.active["B1"].value == ("新的正文" if changed else "原始正文")
        workbook.close()
    result = chat.settle(revised["run_id"], "completed", "再次交付")
    assert result["status"] == ("completed" if changed else "partial"), result


def test_ppt_revision_keeps_saved_dashboard_version_and_content(local_outputs, db):
    from test_dsh_dashboard_delivery import saved_dashboard
    first = prepare(local_outputs, "生成一份3页PPT和应收大屏")
    present_bytes(first, "原PPT.pptx", three_page_deck(OLD_TITLE))
    _, report = saved_dashboard(db, first["run_id"])
    version, content = report.version, deepcopy(report.content)
    assert chat.settle(first["run_id"], "completed", "原交付")["status"] == "completed"
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    present_bytes(revised, "新PPT.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(revised["run_id"], "completed", "仅修订PPT")
    assert result["status"] == "completed", result
    proof = next(row for row in result["task"]["revision_receipts"] if row["type"] == "dashboard")
    assert proof["report_id"] == report.id and proof["report_version"] == version
    db.refresh(report)
    assert report.version == version and report.content == content


@pytest.mark.parametrize("changed", [False, True], ids=["same-content", "changed-content"])
def test_new_dashboard_identity_requires_content_change_for_regeneration(local_outputs, db, changed):
    from test_dsh_dashboard_delivery import saved_dashboard
    from app.ids import new_id
    from app.modules.reports.models import Report
    first = prepare(local_outputs, "生成应收大屏")
    _, report = saved_dashboard(db, first["run_id"])
    assert chat.settle(first["run_id"], "completed", "原交付")["status"] == "completed"
    revised = prepare(local_outputs, "重新生成刚才的大屏", task_id=first["task"]["task_id"])
    # 夹具保存真实 Report 及既有合成 Q；不伪造工具回执，不触发源查询。
    new_report = Report(id=new_id(), workspace_id=report.workspace_id, owner_user_id=report.owner_user_id,
        name="新的应收大屏" if changed else report.name, datasource_ids=list(report.datasource_ids), status=report.status,
        content=deepcopy(report.content), lineage={**report.lineage, "invocationRunId": revised["run_id"]})
    db.add(new_report)
    db.commit()
    result = chat.settle(revised["run_id"], "completed", "保存新对象")
    assert result["status"] == ("completed" if changed else "partial"), result
    assert new_report.content == report.content
    if changed:
        proof = result["task"]["revision_receipts"][0]
        assert proof["report_id"] == new_report.id and proof["label"] == "新的应收大屏"


def test_first_revision_cannot_lose_old_baseline_when_file_scope_is_revoked(local_outputs, db, client):
    from test_dsh_compound_boundaries import client_context
    from app.errors import AppError
    from sqlalchemy import select
    context, service = client_context(client, db, local_outputs)
    first = prepare(context, "生成一份Word")
    document = Document()
    document.add_paragraph("原始内容不能变成无基线")
    stream = BytesIO()
    document.save(stream)
    original = stream.getvalue()
    present_bytes(first, "原文.docx", original)
    assert chat.settle(first["run_id"], "completed", "原交付")["status"] == "completed"
    scopes = list(service.scopes)
    service.scopes = [scope for scope in scopes if scope != "files.view"]
    db.commit()
    before = set(db.scalars(select(Run.id)).all())
    # 当前无法鉴权读取既有基线时拒绝开始修订，不能保存一个虚假的“无旧成果”基线。
    with pytest.raises(AppError):
        prepare(context, "重新生成刚才的Word", task_id=first["task"]["task_id"])
    assert set(db.scalars(select(Run.id)).all()) == before
    service.scopes = scopes
    db.commit()
    revised = prepare(context, "重新生成刚才的Word", task_id=first["task"]["task_id"])
    present_bytes(revised, "原样再交付.docx", original)
    assert chat.settle(revised["run_id"], "completed", "原样不是修改")["status"] == "partial"


def test_first_revision_requires_readable_existing_file_baseline(local_outputs, db):
    from app.errors import ValidationError
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    original, _ = stored_file(db, files["ERP介绍"]["id"])
    path = Path(get_storage().get_path(db.get(Artifact, files["ERP介绍"]["id"]).storage_key))
    path.write_bytes(b"broken")
    with pytest.raises(ValidationError, match="原交付"):
        prepare(local_outputs, "重新生成刚才的PPT", task_id=first["task"]["task_id"])
    path.write_bytes(original)
    revised = prepare(local_outputs, "重新生成刚才的PPT", task_id=first["task"]["task_id"])
    assert revised["task"]["delivery_objects"][0]["base"]["id"] == files["ERP介绍"]["id"]


def test_ppt_revision_preserves_unmodified_csv_final_bytes(local_outputs, db):
    first = prepare(local_outputs, "生成一份3页PPT，再生成 CSV")
    assert {row["type"] for row in first["task"]["required_deliveries"]} == {"presentation", "csv"}
    present_bytes(first, "原PPT.pptx", three_page_deck(OLD_TITLE))
    csv = present_bytes(first, "原明细.csv", "设备,数量\n设备甲,12\n".encode())
    original, proof = stored_file(db, csv["id"])
    assert chat.settle(first["run_id"], "completed", "原交付")["status"] == "completed"
    revised = prepare(local_outputs, f"修改刚才的PPT第1页标题为{NEW_TITLE}", task_id=first["task"]["task_id"])
    present_bytes(revised, "新PPT.pptx", three_page_deck(NEW_TITLE))
    result = chat.settle(revised["run_id"], "completed", "原CSV保留")
    assert result["status"] == "completed", result
    assert stored_file(db, csv["id"])[0] == original
    retained = next(row for row in result["task"]["revision_receipts"] if row["type"] == "csv")
    assert retained["id"] == csv["id"] and retained["sha256"] == proof["sha256"]


def test_invalid_additional_relationship_part_is_unverified_not_worker_crash(local_outputs):
    from zipfile import ZipFile
    first, _ = completed_decks(local_outputs, ["ERP介绍"])
    revised = prepare(local_outputs, "重新生成刚才的PPT", task_id=first["task"]["task_id"])
    output = BytesIO(three_page_deck("ERP介绍"))
    with ZipFile(output, "a") as archive:
        archive.writestr("ppt/_rels/unused.xml.rels", b"<broken")
    present_bytes(revised, "关系损坏.pptx", output.getvalue())
    result = chat.settle(revised["run_id"], "completed", "附加关系分片损坏")
    assert result["status"] == "partial", result


def changed_page_titles(data, changes):
    deck = Presentation(BytesIO(data))
    for page, title in changes.items():
        deck.slides[page - 1].shapes.title.text = title
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def test_two_title_changes_in_one_request_cannot_drop_second_obligation(local_outputs, db, record_property):
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    original, _ = stored_file(db, files["ERP介绍"]["id"])
    revised = prepare(local_outputs, "修改刚才的PPT第1页标题为第一修订，修改第2页标题为第二修订",
                      task_id=first["task"]["task_id"])
    current = present_bytes(revised, "只完成第一页.pptx", changed_page_titles(original, {1: "第一修订"}))
    actual = slide_titles(stored_file(db, current["id"])[0])
    assert actual[0] == "第一修订" and actual[1] != "第二修订"
    result = chat.settle(revised["run_id"], "completed", "第二项尚未执行")
    record_property("revision_loss", json.dumps({"actual_titles": actual, "status": result["status"]}, ensure_ascii=False))
    assert result["status"] == "partial", result


def test_next_run_keeps_unfinished_title_change_when_adding_another(local_outputs, db, record_property):
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    original, _ = stored_file(db, files["ERP介绍"]["id"])
    pending = prepare(local_outputs, "修改刚才的PPT第1页标题为第一修订", task_id=first["task"]["task_id"])
    assert chat.settle(pending["run_id"], "completed", "第一项尚未执行")["status"] == "partial"
    next_run = prepare(local_outputs, "修改刚才的PPT第2页标题为第二修订，同时保留上次第1页修改要求",
                       task_id=first["task"]["task_id"])
    current = present_bytes(next_run, "只完成第二页.pptx", changed_page_titles(original, {2: "第二修订"}))
    actual = slide_titles(stored_file(db, current["id"])[0])
    assert actual[0] != "第一修订" and actual[1] == "第二修订"
    result = chat.settle(next_run["run_id"], "completed", "仍缺第一项")
    record_property("revision_loss", json.dumps({"actual_titles": actual, "status": result["status"]}, ensure_ascii=False))
    assert result["status"] == "partial", result


@pytest.mark.parametrize("unused_relation", [False, True], ids=["zip-only", "unused-relationship"])
def test_unreferenced_word_xml_cannot_prove_regenerated_body(local_outputs, db, record_property, unused_relation):
    from zipfile import ZipFile
    first = prepare(local_outputs, "生成一份Word")
    document = Document()
    document.add_paragraph("只有这段真实正文")
    stream = BytesIO()
    document.save(stream)
    original = present_bytes(first, "原正文.docx", stream.getvalue())
    old, old_proof = stored_file(db, original["id"])
    assert chat.settle(first["run_id"], "completed", "原交付")["status"] == "completed"
    revised = prepare(local_outputs, "重新生成刚才的Word", task_id=first["task"]["task_id"])
    changed = BytesIO()
    with ZipFile(BytesIO(old)) as source, ZipFile(changed, "w") as archive:
        for info in source.infolist():
            raw = source.read(info.filename)
            if unused_relation and info.filename == "word/_rels/document.xml.rels":
                raw = raw.replace(b"</Relationships>", b'<Relationship Id="rIdRevisionMarker" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml" Target="revision-marker.xml"/></Relationships>')
            archive.writestr(info, raw)
        archive.writestr("word/revision-marker.xml", "<revisionMarker>孤立标记</revisionMarker>".encode())
    current = present_bytes(revised, "仅加孤立分片.docx", changed.getvalue())
    actual, proof = stored_file(db, current["id"])
    assert proof["sha256"] != old_proof["sha256"]
    paragraphs = [row.text for row in Document(BytesIO(actual)).paragraphs]
    assert paragraphs == [row.text for row in Document(BytesIO(old)).paragraphs] == ["只有这段真实正文"]
    with ZipFile(BytesIO(actual)) as archive:
        assert "word/revision-marker.xml" in archive.namelist()
        assert (b"revision-marker" in archive.read("word/_rels/document.xml.rels")) is unused_relation
        assert b"rIdRevisionMarker" not in archive.read("word/document.xml")
    result = chat.settle(revised["run_id"], "completed", "正文仍未改变")
    record_property("orphan_content", json.dumps({"paragraphs": paragraphs, "old_sha": old_proof["sha256"],
                    "new_sha": proof["sha256"], "status": result["status"]}, ensure_ascii=False))
    assert result["status"] == "partial", result


@pytest.mark.parametrize("separate_runs", [False, True], ids=["same-request", "cumulative"])
def test_all_requested_titles_complete_only_after_both_actual_changes(local_outputs, db, separate_runs):
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    original, _ = stored_file(db, files["ERP介绍"]["id"])
    if separate_runs:
        pending = prepare(local_outputs, "修改刚才的PPT第1页标题为第一修订", task_id=first["task"]["task_id"])
        assert chat.settle(pending["run_id"], "completed", "未完成第一项")["status"] == "partial"
        content = "修改刚才的PPT第2页标题为第二修订，同时保留上次第1页修改要求"
    else:
        content = "修改刚才的PPT第1页标题为第一修订，修改第2页标题为第二修订"
    revised = prepare(local_outputs, content, task_id=first["task"]["task_id"])
    current = present_bytes(revised, "两页真实修改.pptx", changed_page_titles(original, {1: "第一修订", 2: "第二修订"}))
    actual = slide_titles(stored_file(db, current["id"])[0])
    assert actual[:2] == ["第一修订", "第二修订"]
    result = chat.settle(revised["run_id"], "completed", "两项实际修改均完成")
    assert result["status"] == "completed", result


def test_unparsed_additional_change_is_not_silently_dropped(local_outputs, db):
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    original, _ = stored_file(db, files["ERP介绍"]["id"])
    revised = prepare(local_outputs, "修改刚才的PPT第1页标题为第一修订，修改第2页字体颜色为红色",
                      task_id=first["task"]["task_id"])
    present_bytes(revised, "仅标题已改.pptx", changed_page_titles(original, {1: "第一修订"}))
    result = chat.settle(revised["run_id"], "completed", "第二项未验证")
    assert result["status"] == "partial", result
    assert any(row["kind"] == "unsupported" for row in result["task"]["delivery_objects"][0]["change"]["checks"])


def test_explicit_new_title_replaces_only_same_pending_field(local_outputs, db):
    first, files = completed_decks(local_outputs, ["ERP介绍"])
    original, _ = stored_file(db, files["ERP介绍"]["id"])
    pending = prepare(local_outputs, "修改刚才的PPT第1页标题为第一修订", task_id=first["task"]["task_id"])
    assert chat.settle(pending["run_id"], "completed", "尚未完成")["status"] == "partial"
    revised = prepare(local_outputs, "把刚才的PPT第1页标题改成第二修订", task_id=first["task"]["task_id"])
    present_bytes(revised, "最新指定标题.pptx", changed_page_titles(original, {1: "第二修订"}))
    result = chat.settle(revised["run_id"], "completed", "按最新明确要求完成")
    assert result["status"] == "completed", result
    assert result["task"]["delivery_objects"][0]["change"]["checks"] == [
        {"kind": "slide_title", "page": 1, "value": "第二修订"}]
