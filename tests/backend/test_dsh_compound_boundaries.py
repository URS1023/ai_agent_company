"""复合交付的实时授权、原生字节及异常收尾边界；只运行隔离平台 SQLite。"""
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from sqlalchemy import select

from test_business_worker import prepare, worker_context  # noqa: F401
from test_dsh_compound_contract import (local_outputs, deck_bytes, present_bytes,
    dashboard_with_complete_query, render_image, render_word)
from app.business_worker import chat, tasks, native_deliveries
from app.common.storage import get_storage
from app.errors import ValidationError
from app.models import Artifact, Run, ToolCall
from app.modules.system_tools.models import IntegrationClient


def client_context(client, db, original):
    """经真实 HTTP 签发有限权限令牌；只替换模型配置，不跳过任何业务网关。"""
    scopes = ["chat.create", "chat.view", "chat.execute", "runs.view", "files.view",
              "analytics.view", "analytics.export", "datasources.view", "dashboards.view"]
    response = client.post("/api/v1/system-tools/integration-clients", headers=original[0],
        json={"name": "交付授权验收", "redirectUris": [], "scopes": scopes})
    assert response.status_code == 201, response.text
    registered = response.json()
    response = client.post("/api/auth/client-token", json={"clientId": registered["clientId"],
        "clientSecret": registered["clientSecret"], "workspaceId": original[0]["X-Workspace-Id"]})
    assert response.status_code == 200, response.text
    headers = {"Authorization": "Bearer " + response.json()["accessToken"],
               "X-Workspace-Id": original[0]["X-Workspace-Id"]}
    response = client.post("/api/chat/conversations", headers=headers, json={"title": "受限范围"})
    assert response.status_code == 201, response.text
    return (headers, response.json()["id"]), db.get(IntegrationClient, registered["id"])


@pytest.mark.parametrize("revoked", [None, "files.view", "analytics.export", "datasources.view", "dashboards.view", "runs.view"])
def test_current_client_scope_controls_evidence_without_reissuing_tools(local_outputs, db, client, revoked):
    context, service = client_context(client, db, local_outputs)
    run = prepare(context, "生成应收大屏和Word，将应收大屏放到Word里面")
    report = dashboard_with_complete_query(db, run["run_id"])
    image_id = render_image(run["run_id"], report.id)
    render_word(run["run_id"], image_id)
    before = len(db.scalars(select(ToolCall).where(ToolCall.run_id == run["run_id"])).all())
    if revoked:
        service.scopes = [scope for scope in service.scopes if scope != revoked]
        db.commit()
    result = chat.settle(run["run_id"], "completed", "已交付")
    assert result["status"] == ("partial" if revoked not in {None, "runs.view"} else "completed"), result
    assert len(db.scalars(select(ToolCall).where(ToolCall.run_id == run["run_id"])).all()) == before


@pytest.mark.parametrize("invalid_source", [False, True])
def test_scope_restored_after_success_and_invalid_saved_source(local_outputs, db, invalid_source):
    run = prepare(local_outputs, "生成应收大屏和Word，将应收大屏嵌入Word")
    report = dashboard_with_complete_query(db, run["run_id"])
    image_id = render_image(run["run_id"], report.id)
    render_word(run["run_id"], image_id)
    if invalid_source:
        from app.modules.analytics.models import QueryHistory
        query = db.get(QueryHistory, report.lineage["queryIds"][0])
        query.result_json = {"data": [{"应收金额": "125.00"}], "columns": [{"columnName": "应收金额"}], "truncated": False}
        db.commit()
    keys = ("resource_access_criteria", "actor_user_id", "actor_workspace_id")
    previous = {key: db.info[key] for key in keys if key in db.info}
    try:
        db.info.update(resource_access_criteria=(), actor_user_id="原调用方", actor_workspace_id="原工作区")
        expected = {key: db.info[key] for key in keys}
        result = tasks.verify(db, db.get(Run, run["run_id"]), "completed")
        assert result[0] == ("partial" if invalid_source else "completed")
        assert {key: db.info[key] for key in keys} == expected
    finally:
        for key in keys:
            if key in previous:
                db.info[key] = previous[key]
            else:
                db.info.pop(key, None)


def test_native_word_with_actual_dashboard_picture_can_complete(local_outputs, db):
    run = prepare(local_outputs, "生成应收大屏和Word，将应收大屏嵌入Word")
    report = dashboard_with_complete_query(db, run["run_id"])
    image_id = render_image(run["run_id"], report.id)
    word_id = render_word(run["run_id"], image_id)
    word = db.get(Artifact, word_id)
    data = get_storage().read_bytes(word.storage_key)
    word.status = "superseded"
    db.commit()
    native = present_bytes(run, "本机制作.docx", data)
    result = chat.settle(run["run_id"], "completed", "已完成")
    assert result["status"] == "completed", result
    assert result["task"]["constraint_receipts"][0]["artifact_id"] == native["id"]


def test_explicit_existing_dashboard_can_be_embedded_without_recreating_it(local_outputs, db):
    old = prepare(local_outputs, "生成应收大屏")
    report = dashboard_with_complete_query(db, old["run_id"])
    chat.settle(old["run_id"], "completed", "已保存")
    run = prepare(local_outputs, "生成Word，把应收大屏放入Word")
    assert run["task"]["task_id"] != old["task"]["task_id"]
    image_id = render_image(run["run_id"], report.id)
    render_word(run["run_id"], image_id)
    result = chat.settle(run["run_id"], "completed", "已嵌入既有大屏")
    assert result["status"] == "completed", result
    assert result["task"]["constraint_receipts"][0]["embedded_image"]["report_id"] == report.id


def test_archive_limits_are_checked_before_font_normalization(local_outputs, monkeypatch):
    run = prepare(local_outputs, "生成15页PPT")
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("ppt/presentation.xml", b"x" * (12 * 1024 * 1024 + 1))
    def forbidden(*args, **kwargs):
        raise AssertionError("超限输入不得到达会解压字体规范化的工件创建函数")
    monkeypatch.setattr(native_deliveries, "create_artifact", forbidden)
    with pytest.raises(ValidationError, match="安全解析范围"):
        present_bytes(run, "超限.pptx", stream.getvalue())


@pytest.mark.parametrize("fault", ["blank_pages", "metadata_only", "stored_bytes_changed"])
def test_empty_pages_metadata_and_changed_final_bytes_do_not_prove_completion(local_outputs, db, fault):
    from pptx import Presentation
    run = prepare(local_outputs, "制作15页PPT")
    data = deck_bytes(15 if fault == "stored_bytes_changed" else 1)
    if fault == "blank_pages":
        deck = Presentation()
        for _ in range(15):
            deck.slides.add_slide(deck.slide_layouts[6])
        stream = BytesIO()
        deck.save(stream)
        data = stream.getvalue()
    artifact = present_bytes(run, "15页完整文件.pptx", data)
    row = db.get(Artifact, artifact["id"])
    row.metadata_json = {**row.metadata_json, "actual_page_count": 15, "requested_page_count": 15}
    if fault == "stored_bytes_changed":
        Path(get_storage().get_path(row.storage_key)).write_bytes(deck_bytes(1))
    db.commit()
    result = chat.settle(run["run_id"], "completed", "已完成15页")
    assert result["status"] == "partial" and result["task"]["constraint_failures"]


def test_continuation_cannot_switch_bound_dashboard_to_another_with_same_name(local_outputs, db):
    run = prepare(local_outputs, "生成应收大屏和Word，将应收大屏嵌入Word")
    original = dashboard_with_complete_query(db, run["run_id"])
    incomplete = chat.settle(run["run_id"], "completed", "只生成大屏")
    assert incomplete["task"]["delivery_constraints"][0]["bound_report_id"] == original.id
    original.status = "deleted"
    db.commit()
    continued = prepare(local_outputs, "继续原任务，先完成Word")
    replacement = dashboard_with_complete_query(db, continued["run_id"])
    image_id = render_image(continued["run_id"], replacement.id)
    render_word(continued["run_id"], image_id)
    result = chat.settle(continued["run_id"], "completed", "已完成")
    assert result["status"] == "partial"
    assert result["task"]["delivery_constraints"][0]["bound_report_id"] == original.id


@pytest.mark.parametrize("text", ["生成PPT，15页", "制作PPT，总共十五页", "生成15页PPT和Word", "生成PPT15页和Word"])
def test_adjacent_or_mixed_format_page_request_cannot_be_lost(local_outputs, text):
    run = prepare(local_outputs, text)
    assert run["task"]["delivery_constraints"][0]["minimum"] == 15


def test_queue_new_request_edit_rebuilds_contract_but_steer_page_edit_is_explicit(local_outputs):
    from test_dsh_steering import submit
    run = prepare(local_outputs, "制作15页PPT")
    edited = chat.queue_update(run["run_id"], "edit", "生成Word")
    assert not edited["task"]["delivery_constraints"]
    assert edited["task"]["required_deliveries"] == [{"type": "document", "count": 1}]
    another = prepare(local_outputs, "制作15页PPT")
    chat.append_event(another["run_id"], "run.started")
    submit(local_outputs, another["run_id"], "改为3页")
    present_bytes(another, "三页.pptx", deck_bytes(3))
    result = chat.settle(another["run_id"], "completed", "交付三页")
    assert result["status"] == "completed", result
    assert result["task"]["delivery_constraints"][0]["minimum"] == 3


def test_original_compound_request_completes_with_actual_deliveries(local_outputs, db):
    from openpyxl import Workbook
    import json
    text = ("帮我分析一下设备巡检，近期有多少设备需要巡检，但是未巡检的需要图饼并导出为word和excel，"
            "再帮我做一份15页的ppt然后做一份数字大屏放到word里面")
    run = prepare(local_outputs, text)
    present_bytes(run, "巡检分析.pptx", deck_bytes(15))
    book = Workbook()
    book.active.append(["范围", "说明"])
    book.active.append(["隔离验收", "这是合成平台快照，不代表真实巡检结论"])
    stream = BytesIO()
    book.save(stream)
    present_bytes(run, "巡检明细.xlsx", stream.getvalue())
    report = dashboard_with_complete_query(db, run["run_id"], "巡检数字大屏")
    image_id = render_image(run["run_id"], report.id)
    render_word(run["run_id"], image_id)
    result = chat.settle(run["run_id"], "completed", "已按保存快照交付")
    assert result["status"] == "completed", json.dumps(result["task"], ensure_ascii=False)
    assert {row["type"]: row["count"] for row in result["task"]["required_deliveries"]} == {
        "presentation": 1, "document": 1, "spreadsheet": 1, "dashboard": 1}
    assert {row["kind"] for row in result["task"]["delivery_constraints"]} == {"page_count", "embedded_dashboard"}
    assert next(row for row in result["task"]["constraint_receipts"] if row.get("kind") == "pptx")["pages"] == 15


def test_explicit_list_replacement_removes_obligations_of_removed_types(local_outputs):
    run = prepare(local_outputs, "生成15页PPT及Word")
    chat.settle(run["run_id"], "completed", "未完成")
    continued = prepare(local_outputs, "本次只生成Word", task_id=run["task"]["task_id"],
                        required_deliveries=[{"type": "document", "count": 1}])
    assert not continued["task"]["delivery_constraints"]
    assert continued["task"]["required_deliveries"] == [{"type": "document", "count": 1}]


@pytest.mark.parametrize("prefix", ["帮我做", "分析并做", "请帮我制作"])
def test_request_verbs_are_not_presentation_title(local_outputs, prefix):
    run = prepare(local_outputs, prefix + "一份15页的PPT")
    present_bytes(run, "巡检汇报.pptx", deck_bytes(15))
    result = chat.settle(run["run_id"], "completed", "已生成")
    assert result["status"] == "completed", result


def test_named_page_scope_in_later_sentence_reuses_existing_parser(local_outputs):
    from pptx import Presentation
    run = prepare(local_outputs, "生成库存数据PPT，再做一份ERP系统介绍PPT。ERP介绍为通用产品介绍，六到八页即可。")
    constraints = run["task"]["delivery_constraints"]
    assert len(constraints) == 1 and constraints[0]["target"] == "erp系统介绍"
    for title, count in [("库存数据", 4), ("ERP系统介绍", 7)]:
        deck = Presentation(BytesIO(deck_bytes(count)))
        deck.slides[0].shapes.title.text = title
        stream = BytesIO()
        deck.save(stream)
        present_bytes(run, title + ".pptx", stream.getvalue())
    assert chat.settle(run["run_id"], "completed", "已交付")["status"] == "completed"


def test_generic_and_named_deck_constraints_do_not_double_count_or_steal_file(local_outputs):
    from pptx import Presentation
    run = prepare(local_outputs, "生成15页PPT，再生成ERP介绍15页PPT")
    # 故意先登记具名文件，不能让前面的泛指要求先占走它。
    deck = Presentation(BytesIO(deck_bytes(15)))
    deck.slides[0].shapes.title.text = "ERP介绍"
    stream = BytesIO()
    deck.save(stream)
    named = present_bytes(run, "ERP介绍.pptx", stream.getvalue())
    generic = present_bytes(run, "另一份.pptx", deck_bytes(15))
    result = chat.settle(run["run_id"], "completed", "两份均已生成")
    assert result["status"] == "completed", result
    receipts = result["task"]["constraint_receipts"]
    assert {row["artifact_id"] for row in receipts} == {named["id"], generic["id"]}
    assert len(receipts) == 2


def test_continue_with_no_more_queries_keeps_original_task(local_outputs):
    run = prepare(local_outputs, "生成15页PPT")
    chat.settle(run["run_id"], "completed", "未完成")
    continued = prepare(local_outputs, "请继续完成上一条请求里尚未完成的产物，不要继续查询数据。")
    assert continued["task"]["task_id"] == run["task"]["task_id"]
    assert continued["task"]["delivery_constraints"] == run["task"]["delivery_constraints"]


def test_continuation_clause_can_explicitly_change_page_requirement(local_outputs):
    run = prepare(local_outputs, "生成6页PPT")
    chat.settle(run["run_id"], "completed", "未完成")
    continued = prepare(local_outputs, "继续原任务，把PPT改成15页")
    present_bytes(continued, "新十五页.pptx", deck_bytes(15))
    result = chat.settle(continued["run_id"], "completed", "完成新要求")
    assert result["status"] == "completed", result
    assert result["task"]["delivery_constraints"][0]["minimum"] == 15


@pytest.mark.parametrize("version", [1, 2])
def test_explicit_dashboard_version_cannot_be_replaced_with_current_version(local_outputs, db, version):
    run = prepare(local_outputs, "生成应收大屏和Word，把应收大屏的版本1嵌入Word")
    report = dashboard_with_complete_query(db, run["run_id"])
    report.version = version
    db.commit()
    image_id = render_image(run["run_id"], report.id)
    render_word(run["run_id"], image_id)
    result = chat.settle(run["run_id"], "completed", "已嵌入大屏")
    assert result["status"] == ("partial" if version == 2 else "completed"), result


@pytest.mark.parametrize("embedded", [False, True])
def test_each_word_requires_its_own_actual_embedding(local_outputs, db, embedded):
    run = prepare(local_outputs, "生成两份Word和应收大屏，将应收大屏嵌入每份Word。")
    report = dashboard_with_complete_query(db, run["run_id"])
    image_id = render_image(run["run_id"], report.id)
    first_word = render_word(run["run_id"], image_id)
    from app.business_worker import tools
    result = tools.execute(run["run_id"], "plain-second-word", 3, "render_document", {"name": "另一份报告",
        "document_spec": {"title": "说明", "sections": [{"title": "正文", "blocks": [
            {"type": "image", "artifact_id": image_id} if embedded else
            {"type": "p", "text": "本报告没有大屏图片，不应符合每份都要嵌图的要求。"}]}]}})
    assert result["status"] == "success", result
    settled = chat.settle(run["run_id"], "completed", "两份已全部完成")
    assert settled["status"] == ("completed" if embedded else "partial") and not settled["task"]["missing_deliveries"], settled
    if embedded:
        assert {item["artifact_id"] for item in settled["task"]["constraint_receipts"]} == {first_word, result["artifacts"][0]["id"]}


def test_version_word_in_dashboard_title_does_not_set_version_constraint(local_outputs, db):
    run = prepare(local_outputs, "生成版本1使用说明大屏和Word，把版本1使用说明大屏嵌入Word")
    assert "requested_report_version" not in run["task"]["delivery_constraints"][0]
    report = dashboard_with_complete_query(db, run["run_id"], "版本1使用说明大屏")
    report.version = 2
    db.commit()
    render_word(run["run_id"], render_image(run["run_id"], report.id))
    assert chat.settle(run["run_id"], "completed", "已交付")["status"] == "completed"
