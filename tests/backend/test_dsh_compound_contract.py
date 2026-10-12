"""复合交付验收：真实 Worker、临时 SQLite 和最终 Office 字节，不请求模型或业务库。"""
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from app.business_worker import chat, native_deliveries, tools
from app.config import settings
from app.common.storage import get_storage
from app.models import Artifact, Run
from test_business_worker import prepare, worker_context  # noqa: F401
from test_dsh_dashboard_delivery import saved_dashboard


@pytest.fixture
def local_outputs(worker_context, monkeypatch, tmp_path):
    """本用例的原生工作文件与标准产物均在独立目录内。"""
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path / "objects"))
    return worker_context


def deck_bytes(pages):
    """每页都有独立文字，不能靠空白占位页伪造通过。"""
    deck = Presentation()
    for number in range(1, pages + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = f"隔离验收主题 {number}"
        slide.placeholders[1].text = f"本页用于验证实际页序与总页数：{number}"
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def present_bytes(run, name, data):
    """走真实原生交付，不直接插入 ready 工件伪装成文件。"""
    path = Path(run["cwd"]) / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return native_deliveries.present(run["run_id"], [{"path": path.name}])["artifacts"][0]


def test_one_real_slide_cannot_complete_requested_fifteen_pages(local_outputs):
    run = prepare(local_outputs, "制作一份15页PPT")
    present_bytes(run, "15页完整报告.pptx", deck_bytes(1))
    result = chat.settle(run["run_id"], "completed", "全部完成")
    assert result["status"] == "partial"
    assert result["native_status"] == "completed"
    assert any(item["code"] == "presentation_page_count" for item in result["task"]["constraint_failures"])


def test_word_without_requested_dashboard_picture_cannot_complete(local_outputs, db):
    run = prepare(local_outputs, "生成一份Word报告，并生成应收大屏，然后把应收大屏放到Word里面")
    saved_dashboard(db, run["run_id"])
    result = tools.execute(run["run_id"], "word", 1, "render_document", {
        "name": "报告", "document_spec": {"title": "设备分析", "sections": [
            {"title": "说明", "blocks": [{"type": "p", "text": "这是合法正文，但没有嵌入指定大屏图片。"}]}]}})
    assert result["status"] == "success", result
    settled = chat.settle(run["run_id"], "completed", "大屏和Word全部完成")
    assert settled["status"] == "partial"
    assert any(item["code"] == "embedded_dashboard_missing" for item in settled["task"]["constraint_failures"])


def test_long_continuation_keeps_original_task_draft_and_all_requirements(local_outputs):
    first = prepare(local_outputs, "生成两份PPT，再生成Word和Excel")
    original = chat.settle(first["run_id"], "completed", "尚未交付")
    assert original["status"] == "partial"
    continued = prepare(local_outputs, "请继续完成上一条请求里尚未完成的产物，沿用原数据范围，不要重新查询。")
    assert continued["task"]["task_id"] == first["task"]["task_id"]
    assert continued["task"]["draft_id"] == first["task"]["draft_id"]
    assert continued["task"]["required_deliveries"] == first["task"]["required_deliveries"]
    assert chat.settle(continued["run_id"], "completed", "只回复完成")["status"] == "partial"


def test_explicit_continuation_cannot_replace_original_list_with_remaining_one(local_outputs):
    first = prepare(local_outputs, "生成两份PPT，再生成Word和Excel")
    chat.settle(first["run_id"], "completed", "尚未交付")
    continued = prepare(local_outputs, "继续原任务，先生成剩余的一份PPT", task_id=first["task"]["task_id"])
    assert continued["task"]["required_deliveries"] == first["task"]["required_deliveries"]


@pytest.mark.parametrize("text,pages", [("制作一份15页PPT", 15), ("制作一份十五页PPT", 15),
    ("制作一份6到8页PPT", 7), ("生成一份PPT", 1)])
def test_actual_page_count_can_complete_and_receipt_uses_final_digest(local_outputs, db, text, pages):
    run = prepare(local_outputs, text)
    delivered = present_bytes(run, "合法文件.pptx", deck_bytes(pages))
    result = chat.settle(run["run_id"], "completed", "交付附件")
    assert result["status"] == "completed", result
    row = db.get(Artifact, delivered["id"])
    if result["task"]["delivery_constraints"]:
        evidence = result["task"]["constraint_receipts"][0]
        assert evidence["sha256"] == row.sha256
        assert evidence["pages"] == len(Presentation(BytesIO(get_storage().read_bytes(row.storage_key))).slides) == pages


@pytest.mark.parametrize("text", ["怎么制作15页PPT", "不要生成15页PPT", "生成15页Word文档", "生成PPT，修改第15页标题"])
def test_inquiry_refusal_or_word_page_is_not_ppt_total(local_outputs, text):
    assert not prepare(local_outputs, text)["task"]["delivery_constraints"]


def dashboard_with_complete_query(db, run_id, name="应收大屏"):
    """只补齐平台留存结果协议，不访问配置中没有地址的合成数据源。"""
    from app.modules.analytics.schemas import QueryResult
    query, report = saved_dashboard(db, run_id)
    report.name = name
    query.result_json = QueryResult(queryId=query.id, runId=run_id, generatedSql=query.generated_sql,
        queryType="natural_language", executionTime=1, data=query.result_json["data"], totalCount=1,
        truncated=False, columns=[{"columnName": "应收金额", "dataType": "decimal"}],
        datasourceId=query.datasource_id, datasourceVersion=1).model_dump(mode="json")
    db.commit()
    return report


def render_image(run_id, report_id, call_id="image"):
    result = tools.execute(run_id, call_id, 1, "render_dashboard_image", {"report_id": report_id})
    assert result["status"] == "success", result
    return result["data"]["artifact_id"]


def render_word(run_id, image_id):
    result = tools.execute(run_id, "word-embedded", 2, "render_document", {"name": "嵌图报告",
        "document_spec": {"title": "已保存大屏", "sections": [{"title": "数据快照", "blocks": [
            {"type": "p", "text": "下图来自已保存的查询，不是实时刷新。"},
            {"type": "image", "artifact_id": image_id, "caption": "授权数据快照"}]}]}})
    assert result["status"] == "success", result
    return result["artifacts"][0]["id"]


@pytest.mark.parametrize("fault", [None, "wrong_dashboard", "old_version", "no_receipt", "changed_bytes", "orphan_media"])
def test_word_requires_real_referenced_picture_of_requested_dashboard(local_outputs, db, fault):
    from zipfile import ZipFile
    from docx import Document
    from app.models import ToolCall
    from sqlalchemy import delete

    run = prepare(local_outputs, "生成应收大屏和Word报告，并把应收大屏嵌入Word")
    target = dashboard_with_complete_query(db, run["run_id"])
    selected = dashboard_with_complete_query(db, run["run_id"], "库存大屏") if fault == "wrong_dashboard" else target
    image_id = render_image(run["run_id"], selected.id)
    image = db.get(Artifact, image_id)
    if fault == "orphan_media":
        doc = Document()
        doc.add_paragraph("只有文字的真实 Word，压缩包内的孤立图片不等于正文嵌图。")
        stream = BytesIO()
        doc.save(stream)
        with ZipFile(stream, "a") as archive:
            archive.writestr("word/media/orphan.png", get_storage().read_bytes(image.storage_key))
        present_bytes(run, "孤立图片.docx", stream.getvalue())
    else:
        render_word(run["run_id"], image_id)
    if fault == "old_version":
        target.version += 1
    elif fault == "no_receipt":
        db.execute(delete(ToolCall).where(ToolCall.run_id == run["run_id"], ToolCall.tool_name == "render_dashboard_image"))
    elif fault == "changed_bytes":
        Path(get_storage().get_path(image.storage_key)).write_bytes(b"changed")
    db.commit()
    result = chat.settle(run["run_id"], "completed", "已全部生成")
    assert result["status"] == ("partial" if fault else "completed"), result
    assert any(item["type"] == "document" and item["status"] == "ready" for item in result["deliveries"])
    if not fault:
        evidence = result["task"]["constraint_receipts"][0]
        assert evidence["embedded_image"]["report_id"] == target.id
        assert evidence["embedded_image"]["sha256"] in evidence["images"]


def test_long_continuation_preserves_real_presentation_draft(local_outputs):
    from test_presentation_request_pages import page
    run = prepare(local_outputs, "生成3页PPT及Word")
    plan = tools.execute(run["run_id"], "plan", 1, "create_presentation_plan", {"title": "隔离主题",
        "include_cover": False, "include_closing": False, "slides": [page(number) for number in range(1, 4)]})
    assert plan["status"] == "success", plan
    draft_id = plan["data"]["draft_id"]
    written = tools.execute(run["run_id"], "write", 2, "write_presentation_slides", {
        "draft_id": draft_id, "slides": [{**page(1), "slide_number": 1}]})
    assert written["status"] == "success", written
    chat.settle(run["run_id"], "completed", "还需要继续")
    continued = prepare(local_outputs, "请继续完成刚才未完成的内容，沿用原来的草稿。")
    restored = tools.execute(continued["run_id"], "restore", 1, "create_presentation_plan", {"title": "隔离主题",
        "include_cover": False, "include_closing": False, "slides": [page(number) for number in range(1, 4)]})
    assert restored["status"] == "success", restored
    assert restored["data"]["draft_id"] == draft_id
    assert restored["data"]["written_count"] == 1


def test_steer_and_queue_continuation_do_not_erase_original_constraints(local_outputs, db):
    from test_dsh_steering import submit
    run = prepare(local_outputs, "制作15页PPT，再生成Word及Excel")
    chat.append_event(run["run_id"], "run.started")
    submit(local_outputs, run["run_id"], "先生成Word，其他项目继续保留")
    result = chat.settle(run["run_id"], "completed", "尚未交付")
    assert result["task"]["delivery_constraints"][0]["minimum"] == 15
    continued = prepare(local_outputs, "继续原任务，先生成剩余的一份PPT")
    edited = chat.queue_update(continued["run_id"], "edit", "继续原任务，先生成Word")
    assert edited["task"]["delivery_constraints"] == result["task"]["delivery_constraints"]
    assert edited["task"]["required_deliveries"] == result["task"]["required_deliveries"]


def test_legacy_contract_recovers_first_request_after_steer(local_outputs, db):
    from app.business_worker import tasks
    run = prepare(local_outputs, "制作15页PPT")
    present_bytes(run, "一页.pptx", deck_bytes(1))
    saved = db.get(Run, run["run_id"])
    task = {key: value for key, value in saved.context_snapshot["business_task"].items()
            if key not in {"delivery_constraints", "delivery_contract_version"}}
    saved.context_snapshot = {**saved.context_snapshot, "business_task": task, "user_request": "改成蓝色",
                              "steering_messages": [{"accepted": True, "content": "改成蓝色"}]}
    db.commit()
    status, checked, _ = tasks.verify(db, saved, "completed")
    assert status == "partial" and checked["delivery_constraints"][0]["minimum"] == 15


@pytest.mark.parametrize("kind", ["alert", "inventory"])
def test_grounded_business_answer_does_not_override_missing_file_constraint(local_outputs, monkeypatch, kind):
    from app.modules.business_ai import alert_evidence, inventory_evidence
    run = prepare(local_outputs, "生成15页PPT")
    delivered = present_bytes(run, "一页.pptx", deck_bytes(1))
    monkeypatch.setattr(alert_evidence, "verified_alert_reply", lambda *args: None)
    target, name = (alert_evidence, "verified_alert_reply") if kind == "alert" else (inventory_evidence, "verified_inventory_reply")
    monkeypatch.setattr(target, name, lambda *args: {"content": "这是已核实的业务正文。", "status": "completed", "error": None})
    result = chat.settle(run["run_id"], "completed", "完成")
    assert result["status"] == "partial"
    assert "这是已核实的业务正文。" in result["content"] and "交付核验仍有缺项" in result["content"]
    assert any(item.get("artifact_id") == delivered["id"] and item["status"] == "ready" for item in result["deliveries"])


def test_two_fifteen_page_decks_require_two_distinct_valid_files(local_outputs):
    run = prepare(local_outputs, "生成两份15页PPT")
    present_bytes(run, "第一份.pptx", deck_bytes(15))
    present_bytes(run, "第二份.pptx", deck_bytes(1))
    result = chat.settle(run["run_id"], "completed", "已完成")
    assert result["status"] == "partial" and not result["task"]["missing_deliveries"]
    assert result["task"]["constraint_failures"][0]["expected"]["count"] == 2


@pytest.mark.parametrize("correct", [True, False])
def test_named_decks_bind_actual_cover_text_not_filename(local_outputs, correct):
    run = prepare(local_outputs, "生成ERP介绍PPT6页，再生成MES介绍PPT8页")
    for title, pages in [("ERP介绍", 6), ("MES介绍", 8)]:
        deck = Presentation(BytesIO(deck_bytes(pages)))
        deck.slides[0].shapes.title.text = title if correct else "文件名并不代表实际主题"
        stream = BytesIO()
        deck.save(stream)
        present_bytes(run, title + ".pptx", stream.getvalue())
    result = chat.settle(run["run_id"], "completed", "已完成")
    assert len(result["task"]["delivery_constraints"]) == 2
    assert result["status"] == ("completed" if correct else "partial"), result


def test_failed_ppt_is_corrected_by_new_final_bytes_on_continuation(local_outputs):
    run = prepare(local_outputs, "生成15页PPT及Word")
    present_bytes(run, "报告.pptx", deck_bytes(1))
    result = chat.settle(run["run_id"], "completed", "已生成部分")
    assert result["status"] == "partial" and result["task"]["constraint_failures"]
    continued = prepare(local_outputs, "继续完成原任务剩余的文件")
    present_bytes(continued, "报告.pptx", deck_bytes(15))
    word = tools.execute(continued["run_id"], "word", 1, "render_document", {"name": "正文", "document_spec": {
        "title": "说明", "sections": [{"title": "范围", "blocks": [{"type": "p", "text": "实际完成的说明。"}]}]}})
    assert word["status"] == "success", word
    result = chat.settle(continued["run_id"], "completed", "交付完成")
    assert result["status"] == "completed", result


@pytest.mark.parametrize("state", ["cancelled", "interrupted", "failed"])
def test_page_constraint_never_changes_non_completed_terminal_state(local_outputs, state):
    run = prepare(local_outputs, "制作15页PPT")
    present_bytes(run, "部分.pptx", deck_bytes(1))
    result = chat.settle(run["run_id"], state, "未完成")
    assert result["status"] == state and result["task"]["constraint_failures"]


@pytest.mark.parametrize("text", ["另外做一份新的PPT", "不要继续原任务", "如何继续上次的任务"])
def test_new_negative_or_inquiry_message_does_not_reuse_partial_task(local_outputs, text):
    initial = prepare(local_outputs, "制作15页PPT")
    chat.settle(initial["run_id"], "completed", "未完成")
    following = prepare(local_outputs, text)
    assert following["task"]["task_id"] != initial["task"]["task_id"]
