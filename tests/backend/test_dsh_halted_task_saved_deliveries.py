"""真实 Worker 停止新取数后交付留存结果；仅临时 SQLite、零模型调用。"""
from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
import sqlite3
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.business_worker import chat, tools
from app.common.storage import get_storage
from app.config import settings
from app.models import Artifact, Run, ToolCall, User, Workspace
from app.modules.analytics import execution, metadata, queries
from app.modules.analytics.connections import connection_failure, managed_database_path
from app.modules.analytics.models import QueryHistory
from app.modules.reports.models import Report
from app.providers import model
from test_business_worker import prepare, worker_context  # noqa: F401


SQL = "SELECT id, device, reading FROM readings ORDER BY id"
ROWS = [{"id": 1, "device": "设备甲", "reading": 12}, {"id": 2, "device": "设备乙", "reading": 18}]


@pytest.fixture
def saved_delivery_case(worker_context, client, db, monkeypatch, tmp_path, request):
    """两份专属 SQLite 源，A 真查两行，B 在连接边界注入一次故障后禁止任何新源访问。"""
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    headers, _ = worker_context
    workspace = headers["X-Workspace-Id"]
    sources = []
    for label in ("A", "B"):
        path = managed_database_path(workspace, f"halt-delivery-{label}.sqlite3", must_exist=False)
        path.parent.mkdir(parents=True, exist_ok=True)
        with sqlite3.connect(path) as connection:
            connection.execute("CREATE TABLE readings (id INTEGER PRIMARY KEY, device TEXT, reading INTEGER)")
            connection.executemany("INSERT INTO readings VALUES (?, ?, ?)", [tuple(row.values()) for row in ROWS])
        response = client.post("/api/v1/analytics/datasources", headers=headers, json={
            "name": f"隔离来源{label}", "type": "sqlite", "databaseName": path.name,
            "allowedTables": ["readings"], "timeoutSeconds": 2,
        })
        assert response.status_code == 201, response.text
        sources.append(response.json()["data"])
    case = SimpleNamespace(headers=headers, sources=sources, reads=[], connections=[], halted=False)
    if getattr(request, "param", None) == "member":
        from test_system_foundation import activate_user, create_role, create_user

        case.role = create_role(client, headers, "saved-delivery-member", permissions=[
            "chat.view", "chat.create", "chat.execute", "runs.view", "analytics.view", "analytics.execute", "analytics.export",
            "datasources.view", "datasources.execute", "dashboards.view", "dashboards.create", "dashboards.execute",
            "dashboards.edit", "templates.view", "files.view", "models.view"], scope="workspace")
        create_user(client, headers, [case.role["id"]])
        member_headers = {**activate_user(client), "X-Workspace-Id": workspace}
        conversation = client.post("/api/chat/conversations", headers=member_headers, json={"title": "普通成员留存交付"})
        assert conversation.status_code == 201, conversation.text
        worker_context = (member_headers, conversation.json()["id"])
    original_read = queries.execute_read
    original_connection = execution.readonly_connection

    def observe_read(*args, **kwargs):
        assert not case.halted, "halt 后不得再次尝试业务查询"
        case.reads.append(args[1].id)
        assert len(case.reads) <= 2, "每例最多两次有界源查询尝试"
        return original_read(*args, **kwargs)

    @contextmanager
    def observe_connection(db, source):
        assert not case.halted, "halt 后不得再次连接业务源"
        case.connections.append(source.id)
        if source.id == sources[1]["id"]:
            raise connection_failure(source, TimeoutError("隔离故障注入，未发送慢查询"))
        with original_connection(db, source) as connection:
            yield connection

    def forbidden_io(*args, **kwargs):
        pytest.fail("留存结果交付不得调用模型或刷新来源结构")

    monkeypatch.setattr(queries, "execute_read", observe_read)
    monkeypatch.setattr(execution, "readonly_connection", observe_connection)
    monkeypatch.setattr(metadata, "inspect_sql_source", forbidden_io)
    monkeypatch.setattr(model, "chat_for_run", forbidden_io)
    monkeypatch.setattr(model, "chat", forbidden_io)
    case.first = prepare(worker_context, "查询两个已选来源的设备读数，分别制作可核验的数字大屏和Excel",
                         datasource_ids=[item["id"] for item in sources])
    run_id = case.first["run_id"]
    first = tools.execute(run_id, "query-a", 1, "query_business_data", {
        "datasource_id": sources[0]["id"], "mode": "sql", "query": SQL, "task": "来源A设备读数明细"})
    assert first["status"] == "success", first
    case.query_id = first["data"]["queryId"]
    db.expire_all()
    case.original_query = deepcopy(db.get(QueryHistory, case.query_id).result_json)
    assert case.original_query["data"] == ROWS
    failed = tools.execute(run_id, "query-b", 2, "query_business_data", {
        "datasource_id": sources[1]["id"], "mode": "sql", "query": SQL, "task": "来源B设备读数明细"})
    assert failed["status"] == "failed", failed
    assert failed["metadata"]["failure_domain"] == "datasource", failed
    assert failed["metadata"]["recovery"] == "stop_datasource", failed
    case.halted = True
    db.expire_all()
    assert db.get(Run, run_id).context_snapshot["database_access_halted"] is True
    case.original_calls = {call.id: deepcopy(call.result_json) for call in db.scalars(
        select(ToolCall).where(ToolCall.run_id == run_id)).all()}
    assert len(case.original_calls) == 2
    assert case.reads == case.connections == [item["id"] for item in sources]
    return case


def assert_report(db, result, query_id):
    """核验实际静态报表与每一份显示数据，不能仅凭成功状态判断交付。"""
    assert result["status"] == "success", result
    db.expire_all()
    report = db.get(Report, result["data"]["reportId"])
    assert report.refresh_seconds == 0
    assert report.lineage["queryIds"] == report.lineage["sourceQueryIds"] == [query_id]
    widgets = [item for item in report.content["widgets"] if item.get("dataStatus") == "verified"]
    assert widgets
    for item in widgets:
        assert item["data"] == item["config"]["chartDataParsed"] == ROWS
        assert item["config"]["queryId"] == query_id
    return report


def test_halted_worker_delivers_real_query_dashboard_and_excel(saved_delivery_case, db, monkeypatch):
    from app.modules.analytics.snapshot_access import SnapshotAccess

    case = saved_delivery_case
    run_id = case.first["run_id"]
    checked_contents = []
    original_content = SnapshotAccess.content

    def observe_content(self, content, *args, **kwargs):
        checked_contents.append(deepcopy(content))
        return original_content(self, content, *args, **kwargs)

    monkeypatch.setattr(SnapshotAccess, "content", observe_content)
    dashboard = tools.execute(run_id, "saved-dashboard", 3, "create_dashboard_from_queries", {
        "query_ids": [case.query_id], "title": "来源A留存结果，来源B失败"})
    report = assert_report(db, dashboard, case.query_id)
    assert checked_contents, "真实保存路径必须执行 SnapshotAccess"
    created = tools.execute(run_id, "business-dashboard", 4, "create_business_dashboard", {
        "query_ids": [case.query_id], "title": "来源A静态交付", "requirement": "使用已存结果制作大屏"})
    assert_report(db, created, case.query_id)
    fetched = tools.execute(run_id, "get-dashboard", 5, "get_business_dashboard", {"report_id": report.id})
    assert fetched["status"] == "success" and fetched["data"]["queryIds"] == [case.query_id], fetched
    changed = tools.execute(run_id, "update-dashboard", 6, "update_business_dashboard", {
        "report_id": report.id, "theme": "dark-tech"})
    assert_report(db, changed, case.query_id)
    sheet = tools.execute(run_id, "saved-excel", 7, "export_query_spreadsheet", {
        "query_id": case.query_id, "name": "来源A留存读数.xlsx"})
    assert sheet["status"] == "success", sheet
    artifact = db.get(Artifact, sheet["artifacts"][0]["id"])
    workbook = load_workbook(BytesIO(get_storage().read_bytes(artifact.storage_key)))
    assert list(workbook["查询结果"].values) == [("id", "device", "reading"), (1, "设备甲", 12), (2, "设备乙", 18)]
    assert artifact.metadata_json["delivery_evidence"]["source_query_id"] == case.query_id
    assert artifact.metadata_json["task_id"] == case.first["task"]["task_id"]
    assert_original_receipts(db, case)
    assert case.reads == case.connections == [item["id"] for item in case.sources]


def assert_original_receipts(db, case):
    """交付成功或拒绝均不得重写原查询与失败回执，保留来源B未完成的事实。"""
    db.expire_all()
    assert db.get(QueryHistory, case.query_id).result_json == case.original_query
    for identity, original in case.original_calls.items():
        assert db.get(ToolCall, identity).result_json == original
    assert db.get(Run, case.first["run_id"]).context_snapshot["database_access_halted"] is True


@pytest.mark.parametrize("continuation", ["继续", "继续，使用已有结果制作大屏和Excel"],
                         ids=["short-continue", "saved-results-natural"])
def test_same_task_continue_delivers_q_but_keeps_sql_schema_and_profile_stopped(
    saved_delivery_case, worker_context, db, continuation,
):
    from app.business_worker.dispatch_state import task_datasource_halted
    from app.modules.chat_artifacts.resume import validated_resume_calls

    case = saved_delivery_case
    chat.settle(case.first["run_id"], "failed", "来源B异常已停止，只交付来源A已存结果")
    continued = prepare(worker_context, continuation)
    assert continued["task"]["task_id"] == case.first["task"]["task_id"]
    db.expire_all()
    run = db.get(Run, continued["run_id"])
    assert task_datasource_halted(db, run)
    calls = validated_resume_calls(db, run)
    assert len(calls) == 1 and calls[0].result_json["data"]["queryId"] == case.query_id
    dashboard = tools.execute(run.id, "continue-dashboard", 3, "create_dashboard_from_queries", {
        "query_ids": [case.query_id], "title": "继续交付来源A"})
    assert_report(db, dashboard, case.query_id)
    sheet = tools.execute(run.id, "continue-sheet", 4, "export_query_spreadsheet", {"query_id": case.query_id})
    assert sheet["status"] == "success", sheet
    artifact = db.get(Artifact, sheet["artifacts"][0]["id"])
    assert artifact.run_id == run.id
    workbook = load_workbook(BytesIO(get_storage().read_bytes(artifact.storage_key)))
    assert list(workbook["查询结果"].values)[1:] == [(1, "设备甲", 12), (2, "设备乙", 18)]
    for sequence, (source, mode) in enumerate(((case.sources[0], "sql"), (case.sources[1], "sql"),
                                              (case.sources[0], "schema"), (case.sources[0], "profile")), 5):
        result = tools.execute(run.id, f"blocked-{sequence}", sequence, "query_business_data", {
            "datasource_id": source["id"], "mode": mode, "query": SQL, "task": "重新取数"})
        assert result["error_code"] == "database_access_halted", result
    assert_original_receipts(db, case)
    assert case.reads == case.connections == [item["id"] for item in case.sources]


@pytest.mark.parametrize("restriction", ["disabled", "columns", "query_owner", "query_workspace", "conversation"])
def test_halted_deliveries_keep_current_source_and_identity_boundaries(
    saved_delivery_case, worker_context, client, db, monkeypatch, restriction,
):
    from app.modules.analytics import sources

    case = saved_delivery_case
    run_id = case.first["run_id"]
    if restriction in {"disabled", "columns"}:
        # 正式配置编辑只改平台授权；刷新业务结构属于 halt 禁止的新访问。
        monkeypatch.setattr(sources, "_automatic_metadata_sync", lambda *_: None)
        source = case.sources[0]
        edited = client.put("/api/v1/analytics/datasources/" + source["id"], headers=case.headers, json={
            "name": source["name"], "type": "sqlite", "databaseName": source["databaseName"],
            "version": source["version"], "enabled": restriction != "disabled", "allowedTables": ["readings"],
            "allowedColumns": {"readings": ["id"]} if restriction == "columns" else {},
        })
        assert edited.status_code == 200, edited.text
    elif restriction == "conversation":
        chat.settle(run_id, "failed", "来源B停止")
        other_id = client.post("/api/chat/conversations", headers=case.headers, json={"title": "另一会话"}).json()["id"]
        other = prepare((case.headers, other_id), "用已有查询制作大屏和Excel",
                        datasource_ids=[item["id"] for item in case.sources])
        run_id = other["run_id"]
        db.expire_all()
        run = db.get(Run, run_id)
        # 即使任务标识碰撞，另一会话仍不能继承原成功回执。
        run.context_snapshot = {**run.context_snapshot, "task_id": case.first["task"]["task_id"],
            "business_task": {**run.context_snapshot["business_task"], "task_id": case.first["task"]["task_id"]},
            "database_access_halted": True}
        db.commit()
    else:
        query = db.get(QueryHistory, case.query_id)
        if restriction == "query_owner":
            db.add(User(id="foreign-boundary", display_name="其他留存结果所有者", status="active",
                        created_at=query.created_at, updated_at=query.created_at))
        else:
            current_workspace = db.get(Workspace, query.workspace_id)
            db.add(Workspace(id="foreign-boundary", name="其他工作区",
                             organization_id=current_workspace.organization_id, created_at=query.created_at))
        db.flush()
        setattr(query, "owner_user_id" if restriction == "query_owner" else "workspace_id", "foreign-boundary")
        db.commit()
    before_reports = set(db.scalars(select(Report.id)).all())
    before_artifacts = set(db.scalars(select(Artifact.id)).all())
    for sequence, (name, args) in enumerate((
        ("create_dashboard_from_queries", {"query_ids": [case.query_id], "title": "当前权限校验"}),
        ("create_business_dashboard", {"query_ids": [case.query_id], "requirement": "制作静态大屏"}),
        ("export_query_spreadsheet", {"query_id": case.query_id}),
    ), 3):
        denied = tools.execute(run_id, f"denied-{sequence}", sequence, name, args)
        assert denied["status"] == "failed" and denied["error_code"] != "database_access_halted", denied
        assert not denied["artifacts"] and denied["data"] is None, denied
    db.expire_all()
    assert set(db.scalars(select(Report.id)).all()) == before_reports
    assert set(db.scalars(select(Artifact.id)).all()) == before_artifacts
    assert_original_receipts(db, case)
    assert case.reads == case.connections == [item["id"] for item in case.sources]


@pytest.mark.parametrize("saved_delivery_case", ["member"], indirect=True)
def test_halted_delivery_rechecks_role_revocation_from_official_api(saved_delivery_case, client, db):
    """普通成员先真查、真交付，再正式撤回权限，既有调用重放与新交付都不能泄露旧数据。"""
    case = saved_delivery_case
    run_id = case.first["run_id"]
    calls = [("create_dashboard_from_queries", {"query_ids": [case.query_id], "title": "成员原交付"}),
             ("export_query_spreadsheet", {"query_id": case.query_id})]
    receipts = []
    for sequence, (name, args) in enumerate(calls, 3):
        receipt = tools.execute(run_id, f"before-revoke-{sequence}", sequence, name, args)
        assert receipt["status"] == "success", receipt
        receipts.append(receipt)
    revoked = client.patch("/api/v1/system/roles/" + case.role["id"], headers=case.headers, json={
        "version": case.role["version"],
        "permissions": [permission for permission in case.role["permissions"]
                        if permission not in {"dashboards.create", "analytics.export"}],
    })
    assert revoked.status_code == 200, revoked.text
    before = (set(db.scalars(select(Report.id)).all()), set(db.scalars(select(Artifact.id)).all()))
    for sequence, (name, args) in enumerate(calls, 3):
        for call_id in (f"before-revoke-{sequence}", f"new-after-revoke-{sequence}"):
            result = tools.execute(run_id, call_id, sequence, name, args)
            assert result["status"] == "failed" and result["error_code"] == "tool_not_available", result
            assert not result["artifacts"] and result["data"] is None
    db.expire_all()
    assert (set(db.scalars(select(Report.id)).all()), set(db.scalars(select(Artifact.id)).all())) == before
    completed = db.scalars(select(ToolCall).where(ToolCall.run_id == run_id, ToolCall.tool_name.in_([item[0] for item in calls]))).all()
    assert [call.result_json for call in sorted(completed, key=lambda call: call.started_at)] == receipts
    assert_original_receipts(db, case)
    assert case.reads == case.connections == [item["id"] for item in case.sources]


@pytest.mark.parametrize("name", ["create_dashboard_from_queries", "create_business_dashboard",
                                  "get_business_dashboard", "update_business_dashboard", "view_device_alerts"])
def test_static_delivery_exception_requires_trusted_dsh_builtin_context(name):
    """同名非 DSH 入口和数据库 provider 仍可能新取数，不能因模型参数放行。"""
    assert tools.accesses_datasource(name, "builtin")
    assert tools.accesses_datasource(name, "builtin", execution_owner="openwebui")
    assert tools.accesses_datasource(name, "database", execution_owner="dsh")
    assert tools.accesses_datasource(name, "http", execution_owner="dsh")
    assert tools.accesses_datasource(name, "mcp", execution_owner="dsh")
    assert not tools.accesses_datasource(name, "builtin", execution_owner="dsh")


def test_real_report_save_rejects_old_star_snapshot_after_column_revocation(saved_delivery_case, client, db, monkeypatch):
    """独立进入真实保存链，证明旧星号含额外列由 SnapshotAccess 拒绝，而非仅靠 Q 版本比较。"""
    from app.errors import PermissionErrorEx
    from app.modules.analytics import sources
    from app.modules.reports import service
    from app.modules.reports.schemas import ReportInput

    case = saved_delivery_case
    source = case.sources[0]
    monkeypatch.setattr(sources, "_automatic_metadata_sync", lambda *_: None)
    response = client.put("/api/v1/analytics/datasources/" + source["id"], headers=case.headers, json={
        "name": source["name"], "type": "sqlite", "databaseName": source["databaseName"],
        "version": source["version"], "allowedTables": ["readings"], "allowedColumns": {"readings": ["id"]},
    })
    assert response.status_code == 200, response.text
    sql = "SELECT * FROM readings ORDER BY id"
    content = {"widgets": [{"id": "old-rows", "component": "JScrollBoard", "x": 0, "y": 0, "w": 800, "h": 500,
        "data": deepcopy(case.original_query["data"]), "query": {"datasourceId": source["id"], "sql": sql},
        "config": {"dataType": 1, "datasourceId": source["id"], "queryId": case.query_id, "sql": sql,
                   "chartDataParsed": deepcopy(case.original_query["data"]), "snapshotOnly": True}}]}
    before = set(db.scalars(select(Report.id)).all())
    actor = chat.run_actor(db, db.get(Run, case.first["run_id"]))
    with pytest.raises(PermissionErrorEx) as denied:
        service.save_report(actor, ReportInput(reportName="旧星号快照", datasourceIds=[source["id"]],
                                              content=content, refreshSeconds=0))
    assert denied.value.code == "snapshot_scope_revoked"
    db.rollback()
    assert set(db.scalars(select(Report.id)).all()) == before
    assert_original_receipts(db, case)
    assert case.reads == case.connections == [item["id"] for item in case.sources]
