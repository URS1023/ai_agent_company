"""同一 DSH 任务的多份内容型 PPT 分稿写入、交付与下一轮修改。"""
from io import BytesIO
from types import SimpleNamespace

from pptx import Presentation
import pytest

from test_openwebui_query_dashboard import dashboard_case
from test_presentation_request_pages import page

from app.common.storage import get_storage
from app.config import settings
from app.errors import ToolExecutionError
from app.models import Artifact
from app.modules.presentation import draft
from app.modules.runtime.task_control import current
from app.providers.tool import builtin_provider


@pytest.fixture
def decks(dashboard_case, monkeypatch, tmp_path):
    case = dashboard_case
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    control = SimpleNamespace(business_run_id=case.run.id, business_task_id="task", business_workspace_id=case.ctx.workspace_id)
    token = current.set(control)
    try:
        yield case, control
    finally:
        current.reset(token)


def _plan(case, title):
    return _invoke(case, "create_presentation_plan", {"title": title, "include_cover": False,
        "include_closing": False, "slides": [{**page(number), "title": f"{title}主题{number}"} for number in range(1, 4)]}).data


def _invoke(case, name, args):
    spec = next(item for item in builtin_provider.builtin_specs() if item.name == name)
    return builtin_provider.invoke(spec, args, case.tool)


def _write(case, plan, indices):
    return _invoke(case, "write_presentation_slides", {"draft_id": plan["draft_id"], "slides": [
        {**page(number), "title": plan["outline"][number - 1]["title"], "slide_number": number} for number in indices]})


def _deck_text(case, plan):
    result = _invoke(case, "finalize_presentation", {"draft_id": plan["draft_id"], "name": plan["title"]})
    assert result.status == "success", result.to_dict()
    artifact = case.ctx.db.get(Artifact, result.data["artifact_id"])
    presentation = Presentation(BytesIO(get_storage().read_bytes(artifact.storage_key)))
    return artifact, "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)


def test_two_decks_keep_distinct_outline_progress_and_actual_files(decks):
    case, _ = decks
    erp, mes = _plan(case, "ERP介绍"), _plan(case, "MES介绍")
    assert erp["draft_id"] != mes["draft_id"]
    assert mes["written_count"] == 0 and not mes.get("already_created")
    _write(case, erp, [1])
    _write(case, mes, [2])
    assert draft.load_draft(case.ctx.workspace_id, case.run.id, draft_id=erp["draft_id"])["slides"][0]["written"]
    assert not draft.load_draft(case.ctx.workspace_id, case.run.id, draft_id=mes["draft_id"])["slides"][0]["written"]
    _write(case, erp, [2, 3])
    _write(case, mes, [1, 3])
    first, erp_text = _deck_text(case, erp)
    second, mes_text = _deck_text(case, mes)
    assert first.id != second.id and first.name == "ERP介绍.pptx" and second.name == "MES介绍.pptx"
    assert "ERP介绍主题1" in erp_text and "MES介绍主题1" not in erp_text
    assert "MES介绍主题1" in mes_text and "ERP介绍主题1" not in mes_text
    assert len(_plan(case, "ERP介绍")["available_drafts"]) == 2
    assert _plan(case, "ERP介绍")["written_count"] == 3


def test_later_run_can_edit_first_named_draft_without_affecting_second(decks):
    case, control = decks
    erp, mes = _plan(case, "ERP介绍"), _plan(case, "MES介绍")
    _write(case, erp, [1, 2, 3])
    _write(case, mes, [1, 2, 3])
    control.business_run_id = "next-run"
    draft.write_draft_slide_batches(case.ctx.workspace_id, "next-run",
        [{**page(1), "title": "ERP第一处修订", "slide_number": 1}], draft_id=erp["draft_id"])
    assert draft.final_spec(case.ctx.workspace_id, "next-run", draft_id=erp["draft_id"])["slides"][0]["title"] == "ERP第一处修订"
    assert draft.final_spec(case.ctx.workspace_id, "next-run", draft_id=mes["draft_id"])["slides"][0]["title"] == "MES介绍主题1"


@pytest.mark.parametrize("identity", ["../another", "deck-" + "f" * 20])
def test_unknown_or_outside_draft_never_falls_back_to_current(decks, identity):
    case, _ = decks
    _plan(case, "ERP介绍")
    with pytest.raises(ToolExecutionError):
        draft.load_draft(case.ctx.workspace_id, case.run.id, draft_id=identity)


def test_draft_id_is_scoped_to_its_task(decks):
    case, control = decks
    plan = _plan(case, "ERP介绍")
    control.business_task_id = "other-task"
    with pytest.raises(ToolExecutionError):
        draft.load_draft(case.ctx.workspace_id, case.run.id, draft_id=plan["draft_id"])


def test_delete_current_draft_clears_named_copy_but_preserves_other_deck(decks):
    case, _ = decks
    erp, mes = _plan(case, "ERP介绍"), _plan(case, "MES介绍")
    _write(case, erp, [1])
    _write(case, mes, [1, 2])
    draft.delete_draft(case.ctx.workspace_id, case.run.id)
    with pytest.raises(ToolExecutionError):
        draft.load_draft(case.ctx.workspace_id, case.run.id, draft_id=mes["draft_id"])
    assert draft.draft_status(draft.load_draft(case.ctx.workspace_id, case.run.id), case.run.id)["written_count"] == 1
    assert _plan(case, "MES介绍")["written_count"] == 0
    assert draft.draft_status(draft.load_draft(case.ctx.workspace_id, case.run.id, draft_id=erp["draft_id"]), case.run.id)["written_count"] == 1


def test_old_single_native_draft_is_archived_before_creating_another(decks):
    import json

    case, _ = decks
    plan = _plan(case, "ERP介绍")
    _write(case, plan, [1])
    old = draft.load_draft(case.ctx.workspace_id, case.run.id)
    old.pop("draft_id")
    old.pop("available_drafts")
    storage = get_storage()
    storage.delete(draft._named_key(case.ctx.workspace_id, case.run.id, plan["draft_id"]))
    storage.put_bytes(draft._key(case.ctx.workspace_id, case.run.id), json.dumps(old, ensure_ascii=False).encode("utf-8"))
    second = _plan(case, "MES介绍")
    assert len(second["available_drafts"]) == 2
    restored = _plan(case, "ERP介绍")
    assert restored["written_count"] == 1 and restored["draft_id"] == plan["draft_id"]
