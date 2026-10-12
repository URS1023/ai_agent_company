"""原生 present 文件交付回归：临时工作目录、平台 SQLite 与真实下载字节，不访问业务库。"""
from io import BytesIO
from pathlib import Path

import pytest
from sqlalchemy import select

from app.business_worker import chat, native_deliveries, tasks
from app.config import settings
from app.errors import AuthError, ValidationError
from app.models import Artifact, Message, Run, User
from test_business_worker import prepare, worker_context


@pytest.fixture
def native_run(worker_context, monkeypatch, tmp_path):
    """路径由服务端 prepare 返回，测试文件与正式 DSH 工作空间规则一致。"""
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path / "objects"))
    prepared = prepare(worker_context, "生成两份PPT", required_deliveries=[{"type": "presentation", "count": 2}])
    Path(prepared["cwd"]).mkdir(parents=True)
    return prepared


def local_file(run, name="真实报告.csv", data=b"name,value\nA,12\n"):
    """仅创建本项测试自己的原生工作文件。"""
    path = Path(run["cwd"]) / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def presentation_bytes(title):
    """生成有效 PPTX，交付校验不使用伪扩展名文本代替办公文件。"""
    from pptx import Presentation
    value = Presentation()
    value.slides.add_slide(value.slide_layouts[0]).shapes.title.text = title
    output = BytesIO()
    value.save(output)
    return output.getvalue()


def test_native_files_become_real_downloads_and_history_cards(native_run, worker_context, client, db):
    path = local_file(native_run)
    receipt = native_deliveries.present(native_run["run_id"], [{"path": path.name, "description": "文件说明，不是展示文件名"}])
    assert receipt["run_id"] == native_run["run_id"] and receipt["count"] == 1
    reference = receipt["artifacts"][0]
    assert reference["name"] == path.name and reference["type"] == "csv" and reference["status"] == "ready"
    response = client.get(reference["download_url"], headers=worker_context[0])
    assert response.status_code == 200 and response.content == path.read_bytes()
    db.expire_all()
    artifact = db.get(Artifact, reference["id"])
    assert artifact.owner_user_id == native_run["user_id"] and artifact.run_id == native_run["run_id"]
    assert artifact.metadata_json["native_path"] == path.name
    assert artifact.metadata_json["description"] == "文件说明，不是展示文件名"
    message = db.get(Message, native_run["message_id"])
    assert [part["artifact"]["id"] for part in message.parts_json if part["type"] == "artifact"] == [artifact.id]


def test_saved_business_artifact_url_reuses_existing_card_without_filesystem_search(native_run, db):
    from app.business_worker import tools
    receipt = tools.execute(native_run["run_id"], "saved-ppt", 50, "render_document", {
        "name": "已生成说明", "document_spec": {"title": "说明", "sections": []}})
    assert receipt["status"] == "success", receipt
    assert receipt["metadata"]["delivery"]["already_presented"] is True
    artifact = receipt["artifacts"][0]
    first = native_deliveries.reference_saved(native_run["run_id"], [{"path": artifact["download_url"]}])
    second = native_deliveries.reference_saved(native_run["run_id"], [{"path": artifact["download_url"]}])
    assert first == second and first["already_presented"] is True
    db.expire_all()
    assert len(db.scalars(select(Artifact)).all()) == 1
    assert [part["artifact"]["id"] for part in db.get(Message, native_run["message_id"]).parts_json if part["type"] == "artifact"] == [artifact["id"]]


@pytest.mark.parametrize("boundary", ["task", "owner", "status", "missing_storage"])
def test_saved_artifact_reference_rejects_other_task_owner_or_unavailable_file(native_run, worker_context, db, boundary):
    path = local_file(native_run)
    reference = native_deliveries.present(native_run["run_id"], [{"path": path.name}])["artifacts"][0]
    db.expire_all()
    row = db.get(Artifact, reference["id"])
    source = db.get(Run, native_run["run_id"])
    if boundary == "task":
        chat.settle(native_run["run_id"], "completed", "完成")
        target = prepare(worker_context, "独立新请求")
    else:
        target = native_run
        if boundary == "owner":
            row.owner_user_id = "other-user"
        elif boundary == "status":
            row.status = "superseded"
        else:
            row.storage_key = "does-not-exist"
        db.commit()
    with pytest.raises((ValidationError, AuthError)):
        native_deliveries.reference_saved(target["run_id"], [{"path": reference["download_url"]}])


def test_saved_reference_validates_entire_batch_before_touching_message(native_run, db):
    path = local_file(native_run)
    reference = native_deliveries.present(native_run["run_id"], [{"path": path.name}])["artifacts"][0]
    db.expire_all()
    message = db.get(Message, native_run["message_id"])
    message.parts_json = []
    db.commit()
    with pytest.raises(ValidationError):
        native_deliveries.reference_saved(native_run["run_id"], [{"path": reference["download_url"]}, {"path": "local.txt"}])
    db.expire_all()
    assert db.get(Message, native_run["message_id"]).parts_json == []


def test_file_library_write_and_rename_keep_download_reference(worker_context):
    from app.business_worker import tools
    prepared = prepare(worker_context, "写一份说明文件")
    created = tools.execute(prepared["run_id"], "file-create", 60, "write_file", {"file_name": "说明.txt", "content": "真实内容"})
    assert created["status"] == "success", created
    renamed = tools.execute(prepared["run_id"], "file-rename", 61, "rename_file", {"file_id": created["data"]["file_id"], "new_name": "新说明.txt"})
    assert renamed["status"] == "success", renamed
    assert renamed["data"]["download_url"] == created["data"]["download_url"]
    assert renamed["metadata"]["delivery"]["channel"] == "file_library"


def test_same_run_path_and_hash_are_idempotent_across_absolute_and_relative_paths(native_run, db):
    path = local_file(native_run, "reports/明细.csv")
    first = native_deliveries.present(native_run["run_id"], [{"path": str(path)}, {"path": "reports/明细.csv"}])
    again = native_deliveries.present(native_run["run_id"], [{"path": "reports/明细.csv", "description": "再次展示"}])
    assert first["count"] == again["count"] == 1 and first["artifacts"][0]["id"] == again["artifacts"][0]["id"]
    assert len(db.scalars(select(Artifact)).all()) == 1
    db.expire_all()
    assert len([part for part in db.get(Message, native_run["message_id"]).parts_json if part["type"] == "artifact"]) == 1


def test_changed_content_is_versioned_without_counting_one_ppt_twice(native_run, db):
    original = presentation_bytes("第一版")
    path = local_file(native_run, "报告.pptx", original)
    first = native_deliveries.present(native_run["run_id"], [{"path": path.name}])["artifacts"][0]
    path.write_bytes(presentation_bytes("修改版"))
    second = native_deliveries.present(native_run["run_id"], [{"path": path.name}])["artifacts"][0]
    assert first["id"] != second["id"]
    db.expire_all()
    assert db.get(Artifact, first["id"]).status == "superseded"
    assert db.get(Artifact, second["id"]).parent_artifact_id == first["id"]
    run = db.get(Run, native_run["run_id"])
    status, task, delivered = tasks.verify(db, run, "completed")
    assert status == "partial" and task["missing_deliveries"] == [{"type": "presentation", "count": 2, "missing": 1}]
    assert sum(item["status"] == "ready" for item in delivered) == 1
    # 重新交付原内容复用原指纹工件；当前卡片与 ready 版本仍只有一份。
    db.rollback()
    path.write_bytes(original)
    restored = native_deliveries.present(native_run["run_id"], [{"path": path.name}])["artifacts"][0]
    assert restored["id"] == first["id"]
    db.expire_all()
    assert db.get(Artifact, second["id"]).status == "superseded"
    assert len(db.scalars(select(Artifact)).all()) == 2
    assert [part["artifact"]["id"] for part in db.get(Message, native_run["message_id"]).parts_json if part["type"] == "artifact"] == [first["id"]]


def test_identical_files_at_different_real_paths_are_separate_deliveries(native_run, db):
    a, b = local_file(native_run, "甲.csv"), local_file(native_run, "乙.csv")
    receipt = native_deliveries.present(native_run["run_id"], [{"path": str(a)}, {"path": str(b)}])
    assert receipt["count"] == 2 and len({item["id"] for item in receipt["artifacts"]}) == 2
    assert {item["name"] for item in receipt["artifacts"]} == {"甲.csv", "乙.csv"}


@pytest.mark.parametrize("invalid", [[], [{"path": "missing.txt"}], [{"path": "."}],
    [{"path": "exists.txt", "name": "伪造展示名.pptx"}], [{"artifact_id": "existing-id"}],
    [{"path": "exists.txt", "id": "existing-id"}], [{"path": "exists.txt", "description": 4}],
    [{"path": "exists.txt"}] * 17])
def test_invalid_present_arguments_create_no_artifact(native_run, db, invalid):
    local_file(native_run, "exists.txt", b"real")
    with pytest.raises(ValidationError):
        native_deliveries.present(native_run["run_id"], invalid)
    assert db.scalars(select(Artifact)).all() == []


def test_all_files_are_checked_before_any_artifact_is_written(native_run, db, tmp_path):
    path = local_file(native_run)
    outside = tmp_path / "outside.csv"
    outside.write_bytes(b"outside")
    with pytest.raises(ValidationError):
        native_deliveries.present(native_run["run_id"], [{"path": str(path)}, {"path": str(outside)}])
    assert db.scalars(select(Artifact)).all() == []
    assert path.is_file() and outside.read_bytes() == b"outside"


def test_other_conversation_workspace_cannot_be_presented(native_run, db):
    own = Path(native_run["cwd"])
    other = own.parent.parent / "other-session" / "workspace" / "other.csv"
    other.parent.mkdir(parents=True)
    other.write_bytes(b"other")
    with pytest.raises(ValidationError):
        native_deliveries.present(native_run["run_id"], [{"path": str(other)}])
    assert db.scalars(select(Artifact)).all() == []


def test_present_size_limits_are_checked_before_storage(native_run, db, monkeypatch):
    local_file(native_run, "a.csv", b"12345")
    monkeypatch.setattr(native_deliveries, "MAX_FILE_BYTES", 4)
    with pytest.raises(ValidationError):
        native_deliveries.present(native_run["run_id"], [{"path": "a.csv"}])
    monkeypatch.setattr(native_deliveries, "MAX_FILE_BYTES", 10)
    monkeypatch.setattr(native_deliveries, "MAX_PRESENT_TOTAL_BYTES", 8)
    local_file(native_run, "b.csv", b"12345")
    with pytest.raises(ValidationError):
        native_deliveries.present(native_run["run_id"], [{"path": "a.csv"}, {"path": "b.csv"}])
    assert db.scalars(select(Artifact)).all() == []


def test_disabled_run_actor_cannot_register_new_files(native_run, db):
    local_file(native_run)
    db.get(User, native_run["user_id"]).status = "disabled"
    db.commit()
    with pytest.raises(AuthError):
        native_deliveries.present(native_run["run_id"], [{"path": "真实报告.csv"}])
    assert db.scalars(select(Artifact)).all() == []


def test_failed_batch_rolls_back_artifacts_and_keeps_native_files(native_run, db, monkeypatch):
    paths = [local_file(native_run, name) for name in ("甲.csv", "乙.csv")]
    original, calls = native_deliveries.create_artifact, []
    def fail_second(*args, **kwargs):
        calls.append(kwargs["name"])
        if len(calls) == 2:
            raise RuntimeError("模拟第二个工件登记失败")
        return original(*args, **kwargs)
    monkeypatch.setattr(native_deliveries, "create_artifact", fail_second)
    with pytest.raises(RuntimeError, match="第二个"):
        native_deliveries.present(native_run["run_id"], [{"path": str(path)} for path in paths])
    assert db.scalars(select(Artifact)).all() == []
    assert all(path.is_file() for path in paths)
    assert not list((Path(settings.STORAGE_ROOT) / "artifacts").rglob("*.csv"))
