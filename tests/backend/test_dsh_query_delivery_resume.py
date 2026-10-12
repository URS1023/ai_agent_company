"""原生 DSH 同一任务跨轮复用查询快照，隔离验证真实交付文件。"""
from datetime import timedelta
from io import BytesIO

from openpyxl import load_workbook
from pptx import Presentation
import pytest
from sqlalchemy import select

from test_openwebui_query_dashboard import dashboard_case

from app.common.storage import get_storage
from app.config import settings
from app.errors import AppError
from app.models import Artifact, Run, ToolCall
from app.modules.chat_artifacts.chart_renderer import load_authorized_query, render_chart_image
from app.modules.chat_artifacts.query_dashboard import create_dashboard_from_queries
from app.modules.chat_artifacts.query_presentation import export_query_presentation
from app.modules.chat_artifacts.query_spreadsheet import export_query_spreadsheet


def _next_turn(case):
    """模拟 Worker 持久化的稳定任务；本轮不会再次生成查询回执。"""
    case.run.status = "completed"
    case.run.context_snapshot = {**case.run.context_snapshot, "chat_engine": "dsh", "task_id": "task",
        "business_task": {"task_id": "task"}, "user_request": "查询巡检设备数量"}
    case.query.result_json = {**case.query.result_json, "generatedSql": case.query.generated_sql,
        "queryType": "SELECT", "executionTime": 10, "datasourceId": case.source.id,
        "datasourceVersion": case.query.datasource_version,
        "columns": [{"columnName": "状态", "dataType": "text"}, {"columnName": "设备数", "dataType": "int"}]}
    run = Run(id="next-run", workspace_id=case.ctx.workspace_id, started_by_user_id=case.ctx.user_id,
        conversation_id=case.run.conversation_id, status="running", created_at=case.run.created_at + timedelta(minutes=1),
        context_snapshot={"chat_engine": "dsh", "task_id": "task", "business_task": {"task_id": "task"},
            "datasource_ids": [case.source.id], "template_id": "executive-ivory",
            "plugin_inputs": {"aspect_ratio": "16:9"}, "user_request": "用刚才的结果导出文件并换个标题"})
    case.ctx.db.add(run)
    case.ctx.db.commit()
    return run, {**case.tool, "run_id": run.id}


def test_same_dsh_task_exports_all_formats_without_requery(dashboard_case, monkeypatch, tmp_path):
    case = dashboard_case
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    run, tool = _next_turn(case)
    sheet = export_query_spreadsheet(case.ctx, tool, {"query_id": case.query.id, "name": "继续导出.xlsx"})
    deck = export_query_presentation(case.ctx, tool, {"query_ids": [case.query.id], "title": "继续导出"})
    chart = render_chart_image(case.ctx, tool, {"query_id": case.query.id, "chart_type": "bar",
        "label_column": "状态", "value_column": "设备数"})
    board = create_dashboard_from_queries(case.ctx, tool, {"query_ids": [case.query.id], "title": "继续导出"})
    files = [case.ctx.db.get(Artifact, result.artifacts[0]["id"]) for result in (sheet, deck, chart)]
    assert all(item.run_id == run.id for item in files)
    raw = [get_storage().read_bytes(item.storage_key) for item in files]
    workbook = load_workbook(BytesIO(raw[0]))
    assert list(workbook["查询结果"].values) == [("状态", "设备数"), ("已巡检", 15), ("未巡检", 5)]
    presentation = Presentation(BytesIO(raw[1]))
    assert len(presentation.slides) == deck.data["page_count"] == 4
    assert deck.data["presentation_mode"] == "report"
    series = [shape.chart.series[0] for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
    assert len(series) == 1 and list(series[0].values) == [15.0, 5.0]
    assert raw[2].startswith(b"\x89PNG")
    assert board.data["queryIds"] == [case.query.id]
    assert len(case.ctx.db.scalars(select(ToolCall)).all()) == 1


@pytest.mark.parametrize("change", ["task", "identity", "conversation", "sources", "future", "active", "failed", "task_contract"])
def test_dsh_reuse_cannot_pull_unrelated_or_unconfirmed_results(dashboard_case, change):
    case = dashboard_case
    run, tool = _next_turn(case)
    if change == "task":
        case.run.context_snapshot = {**case.run.context_snapshot, "task_id": "other-task"}
    elif change == "identity":
        case.run.started_by_user_id = "other"
    elif change == "conversation":
        case.run.conversation_id = "other"
    elif change == "sources":
        case.run.context_snapshot = {**case.run.context_snapshot, "datasource_ids": ["other"]}
    elif change == "future":
        case.run.created_at = run.created_at + timedelta(minutes=1)
    elif change == "active":
        case.run.status = "running"
    elif change == "failed":
        case.call.result_json = {"status": "failed", "data": {"queryId": case.query.id}}
    else:
        run.context_snapshot = {**run.context_snapshot, "business_task": {"task_id": "other"}}
    case.ctx.db.commit()
    with pytest.raises(AppError):
        load_authorized_query(case.ctx, tool, case.query.id)
