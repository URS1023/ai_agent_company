"""模板填充的字号按磅换算，中文字体与未覆盖样式保持明确。"""
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app.modules.files.template_fill import fill_docx


@pytest.mark.parametrize("size", [12, 12.5])
def test_docx_template_font_size_is_points(tmp_path, size):
    source = tmp_path / "template.docx"
    target = tmp_path / "filled.docx"
    document = Document()
    run = document.add_paragraph().add_run("设备名称：{device}")
    run.bold = True
    run.font.size = Pt(10)
    document.save(source)
    result = fill_docx(str(source), str(target), replace={"device": "一号设备"}, font={"name": "宋体", "size": size})
    assert result["written"] == 1
    run = Document(target).paragraphs[0].runs[0]
    assert run.text == "设备名称：一号设备"
    assert run.font.size.pt == size
    assert run.bold is True
    assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert Document(source).paragraphs[0].text == "设备名称：{device}"
