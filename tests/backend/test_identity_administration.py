"""Contracts for user XLSX exchange, self profile and permission change history.

中文说明：身份管理契约：用户 XLSX 模板预览与原子导入的往返、导出的更新数据不含口令且防公式注入并留审计、
个人资料有版本且不能借改资料改授权、权限变更历史包含角色前后值且有独立权限、OpenAPI 描述二进制/表单契约。
隔离方式：client 夹具（内存 SQLite + TestClient），不连业务库。
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.models import WorkspaceMember


SYSTEM = "/api/v1/system"
HEADERS = [
    "登录账号", "显示姓名", "初始临时密码", "账号状态", "部门编码",
    "岗位编码", "角色编码", "员工编号", "联系邮箱", "联系电话",
]


def workbook_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用户导入"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def preview(client, headers, data: bytes, *, update: bool = False):
    return client.post(
        SYSTEM + "/users/import-preview",
        headers=headers,
        data={"update_existing": str(update).lower()},
        files={"file": ("users.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def commit_import(client, headers, data: bytes, checksum: str, *, update: bool = False):
    return client.post(
        SYSTEM + "/users/import",
        headers=headers,
        data={"update_existing": str(update).lower(), "expected_checksum": checksum},
        files={"file": ("users.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_template_preview_and_atomic_import_round_trip(client, administrator):
    template = client.get(SYSTEM + "/users/import-template", headers=administrator)
    assert template.status_code == 200, template.text
    assert template.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    workbook = load_workbook(BytesIO(template.content), data_only=False)
    assert workbook.sheetnames == ["用户导入", "_选项", "填写说明"]
    assert [cell.value for cell in workbook["用户导入"][1]] == HEADERS
    assert workbook["_选项"].sheet_state == "hidden"

    invalid = workbook_bytes([
        ["batch-one", "批量成员一", "Temporary12345", "启用", "", "", "", "B001", "one@example.com", ""],
        ["bad account", "错误成员", "short", "未知", "", "", "", "", "not-an-email", ""],
    ])
    checked = preview(client, administrator, invalid)
    assert checked.status_code == 200, checked.text
    result = checked.json()
    assert result["valid"] is False
    assert {issue["field"] for issue in result["issues"]} >= {"登录账号", "账号状态", "联系邮箱"}
    assert client.get(SYSTEM + "/users?q=batch-one", headers=administrator).json()["total"] == 0
    rejected = commit_import(client, administrator, invalid, result["checksum"])
    assert rejected.status_code == 422
    assert client.get(SYSTEM + "/users?q=batch-one", headers=administrator).json()["total"] == 0

    valid = workbook_bytes([
        ["batch-one", "批量成员一", "Temporary12345", "启用", "", "", "", "B001", "one@example.com", "13800000001"],
        ["batch-two", "批量成员二", "Temporary12345", "停用", "", "", "", "B002", "two@example.com", ""],
    ])
    checked = preview(client, administrator, valid)
    assert checked.status_code == 200, checked.text
    result = checked.json()
    assert result["valid"] is True and result["create_count"] == 2
    assert all("password" not in str(row).lower() and "Temporary12345" not in str(row) for row in result["preview"])
    mismatch = commit_import(client, administrator, valid + b"changed", result["checksum"])
    assert mismatch.status_code == 409
    imported = commit_import(client, administrator, valid, result["checksum"])
    assert imported.status_code == 201, imported.text
    assert imported.json()["created_count"] == 2
    users = client.get(SYSTEM + "/users?q=batch-", headers=administrator).json()["items"]
    assert {item["username"]: item["status"] for item in users} == {"batch-one": "active", "batch-two": "disabled"}


def test_update_export_is_password_free_formula_safe_and_audited(client, administrator):
    created = client.post(SYSTEM + "/users", headers=administrator, json={
        "username": "sheet-user", "display_name": "=公式姓名", "password": "Temporary12345",
        "employee_code": "S001", "email": "before@example.com", "role_ids": [],
    })
    assert created.status_code == 201, created.text

    exported = client.get(SYSTEM + "/users/export?q=sheet-user", headers=administrator)
    assert exported.status_code == 200, exported.text
    book = load_workbook(BytesIO(exported.content), data_only=False)
    values = [cell.value for cell in book["用户导入"][2]]
    assert values[2] in {None, ""}
    assert values[1].startswith("'=")
    assert "Temporary12345" not in exported.content.decode("latin1", errors="ignore")

    update = workbook_bytes([
        ["sheet-user", "已更新姓名", "", "启用", "", "", "", "S001", "after@example.com", "13900000000"],
    ])
    checked = preview(client, administrator, update, update=True)
    assert checked.status_code == 200, checked.text
    assert checked.json()["update_count"] == 1
    applied = commit_import(client, administrator, update, checked.json()["checksum"], update=True)
    assert applied.status_code == 201, applied.text
    current = client.get(SYSTEM + "/users?q=sheet-user", headers=administrator).json()["items"][0]
    assert (current["display_name"], current["email"], current["phone"]) == (
        "已更新姓名", "after@example.com", "13900000000",
    )
    history = client.get(SYSTEM + "/permission-history?q=sheet-user&page=1&page_size=50", headers=administrator)
    assert history.status_code == 200, history.text
    assert any(item["action"] == "user.update" and item["source"] == "bulk_import" for item in history.json()["items"])
    assert any(item["action"] == "user.bulk_import" for item in client.get(
        SYSTEM + "/permission-history?action=user.bulk_import&page=1&page_size=50", headers=administrator,
    ).json()["items"])


def test_self_profile_is_versioned_and_cannot_change_authorization(client, administrator, db):
    profile = client.get("/api/auth/profile", headers=administrator)
    assert profile.status_code == 200, profile.text
    before = profile.json()
    response = client.patch("/api/auth/profile", headers=administrator, json={
        "version": before["version"],
        "display_name": "平台管理员",
        "email": "admin@example.com",
        "phone": "13800000000",
    })
    assert response.status_code == 200, response.text
    after = response.json()
    assert after["display_name"] == "平台管理员"
    assert after["organization_name"] and after["workspace_name"]
    stale = client.patch("/api/auth/profile", headers=administrator, json={
        "version": before["version"], "phone": "13900000000",
    })
    assert stale.status_code == 409
    forbidden_field = client.patch("/api/auth/profile", headers=administrator, json={
        "version": after["version"], "role_ids": [],
    })
    assert forbidden_field.status_code == 422
    member = db.query(WorkspaceMember).filter_by(user_id=after["id"], workspace_id=administrator["X-Workspace-Id"]).one()
    assert member.role == "admin"


def test_permission_history_contains_role_before_after_and_has_separate_permission(client, administrator):
    created = client.post(SYSTEM + "/roles", headers=administrator, json={
        "code": "history-role", "name": "历史角色", "permissions": ["home.view"], "data_scope": "self",
    })
    assert created.status_code == 201, created.text
    role = created.json()
    updated = client.patch(SYSTEM + "/roles/" + role["id"], headers=administrator, json={
        "version": role["version"], "permissions": ["home.view", "files.view"], "data_scope": "workspace",
    })
    assert updated.status_code == 200, updated.text
    history = client.get(
        SYSTEM + f"/permission-history?target_type=role&target_id={role['id']}&page=1&page_size=20",
        headers=administrator,
    )
    assert history.status_code == 200, history.text
    change = next(item for item in history.json()["items"] if item["action"] == "role.update")
    assert change["before"]["permissions"] == ["home.view"]
    assert change["after"]["permissions"] == ["files.view", "home.view"]
    assert "新增 1 项权限" in change["summary"]

    viewer_role = client.post(SYSTEM + "/roles", headers=administrator, json={
        "code": "catalog-only", "name": "仅看权限目录",
        "permissions": ["home.view", "system.permissions.view"], "data_scope": "workspace",
    }).json()
    client.post(SYSTEM + "/users", headers=administrator, json={
        "username": "catalog-viewer", "display_name": "目录查看者",
        "password": "Temporary12345", "role_ids": [viewer_role["id"]],
    })
    login = client.post("/api/auth/login", json={"username": "catalog-viewer", "password": "Temporary12345"}).json()
    limited = {"Authorization": "Bearer " + login["token"]}
    changed = client.post("/api/auth/password", headers=limited, json={
        "current_password": "Temporary12345", "new_password": "ChangedPassword12345",
    })
    assert changed.status_code == 200
    login = client.post("/api/auth/login", json={"username": "catalog-viewer", "password": "ChangedPassword12345"}).json()
    limited = {"Authorization": "Bearer " + login["token"]}
    assert client.get(SYSTEM + "/permissions", headers=limited).status_code == 200
    assert client.get(SYSTEM + "/permission-history", headers=limited).status_code == 403


def test_identity_openapi_describes_binary_and_multipart_contracts(client, administrator):
    spec = client.get("/api/v1/platform/openapi.json", headers=administrator).json()
    paths = spec["paths"]
    assert paths[SYSTEM + "/users/export"]["get"]["summary"]
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in paths[SYSTEM + "/users/export"]["get"]["responses"]["200"]["content"]
    preview_op = paths[SYSTEM + "/users/import-preview"]["post"]
    assert preview_op["requestBody"]["content"]["multipart/form-data"]
    assert preview_op["description"].startswith("只读校验")
    assert paths["/api/auth/profile"]["patch"]["responses"]["409"]["description"]
    collection = client.get("/api/v1/platform/api-collection", headers=administrator).json()
    system_group = next(group for group in collection["item"] if group["name"] == "系统管理")
    names = {item["name"] for item in system_group["item"]}
    assert {"下载当前工作区用户导入模板", "校验用户导入文件并生成脱敏预览", "原子提交已校验的用户导入文件"}.issubset(names)
