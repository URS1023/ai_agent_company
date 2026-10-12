"""Actual sandbox/Word/snapshot dashboard delivery with synthetic report evidence.

中文说明：真实的沙箱/Word/快照大屏交付：质检计数表示检查结果而不是完成率；
导出生成内嵌中文图表的 Word 并保留与快照大屏一致；报告工具做时区归一但 API 不变。
隔离方式：monkeypatch + client 夹具（内存 SQLite）+ 合成报告证据，不连业务库、不调用外部模型。
"""
import io,json
from zipfile import ZipFile
import pytest
from sqlalchemy import select
from app.models import AuthIdentity, Artifact
from app.modules.alerts.inspection_export import inspection_facts, export_report
from app.modules.runtime import service as runtime
from app.modules.business_ai.context import auth_context_from_tool
from app.modules.artifact.document_images import resolve_document_images
from app.modules.artifact.renderers import render_document
from app.errors import ValidationError,NotFoundError


def reports():
    return [{'executionId':'E'*26,'workflowId':'W'*26,'workflowName':'设备巡检','createTime':'2026-09-22T01:00:00+00:00',
             'reportUrl':'/workflows/results/test','data':[
        {'设备':'110kV主变压器','项目':'巡检到期','结果':'预警','说明':'已超过计划周期'},
        {'设备':'液压站','项目':'完成记录','结果':'待确认','说明':'缺少记录'},
        {'设备':'循环水泵','项目':'巡检到期','结果':'正常','说明':'未超过计划周期'}]}]


def test_inspection_counts_are_check_outcomes_not_completion_rates():
    facts=inspection_facts(reports()+reports())
    assert facts['checkCount']==3 and facts['deviceNameCount']==3
    assert facts['counts']=={'异常':1,'待确认':1,'正常':1}
    assert not any('Rate' in k or 'completion' in k for k in facts)
    assert '不推算' in facts['scopeNote']


def context_and_run(client,administrator,db):
    user=db.scalar(select(AuthIdentity).where(AuthIdentity.username_normalized=='workspace-admin')).user_id
    cid=client.post('/api/chat/conversations',headers=administrator,json={}).json()['id']
    run=runtime.create_run(db,workspace_id=administrator['X-Workspace-Id'],conversation_id=cid,user_id=user,
        context_snapshot={'user_request':'生成巡检 Word 报告和数字看板'})
    db.commit()
    return auth_context_from_tool({'db':db,'workspace_id':run.workspace_id,'user_id':user}),run


def test_inspection_export_creates_word_with_embedded_chinese_chart_and_dashboard(client,administrator,db,monkeypatch):
    from app.modules.workflows import business_results
    from app.common.storage import get_storage
    from docx import Document
    monkeypatch.setattr(business_results,'list_reports',lambda *a,**kw:{'items':reports(),'hasMore':False})
    ctx,run=context_and_run(client,administrator,db)
    result=export_report(ctx,run.id,{'include_dashboard':True})
    assert len(result['artifacts'])==2 and result['dashboard']['previewPath'].startswith('/report/view/')
    word=db.get(Artifact,result['artifacts'][1]['id'])
    data=get_storage().read_bytes(word.storage_key)
    doc=Document(io.BytesIO(data))
    assert len(doc.inline_shapes)==1
    assert all(r.font.name=='宋体' for p in doc.paragraphs for r in p.runs)
    assert '100%' not in '\n'.join(p.text for p in doc.paragraphs)
    source_url=result['artifacts'][0]['preview_url']
    resolved=resolve_document_images({'sections':[{'blocks':[{'type':'image','source_url':source_url}]}]},db,ctx.workspace_id)
    assert len(Document(io.BytesIO(render_document(resolved))).inline_shapes)==1
    with pytest.raises(NotFoundError):resolve_document_images({'sections':[{'blocks':[{'type':'image','source_url':source_url}]}]},db,'foreign-workspace')
    with pytest.raises(ValidationError):resolve_document_images({'sections':[{'blocks':[{'type':'image','path':'C:/Windows/win.ini'}]}]},db,ctx.workspace_id)
    with pytest.raises(ValueError):render_document({'sections':[{'blocks':[{'type':'image','source_url':'/missing.png'}]}]})


def test_report_tool_normalizes_timezone_and_uses_offset_without_changing_api(monkeypatch):
    from app.modules.workflows import business_results as br
    from app.providers.tool.builtin_provider import builtin_specs
    from app.providers.tool.validator import validate_arguments
    observed=[]
    monkeypatch.setattr(br,'list_reports',lambda *a,**kw:observed.append(kw) or {'items':[],'kind':'alert','hasMore':False})
    spec=next(s for s in builtin_specs() if s.name=='query_workflow_results')
    args=validate_arguments({'kind':'alert','history':True,'start':'2026-09-01T00:00:00','offset':'10','limit':20},spec.input_schema)
    br.tool_result(None,args)
    assert observed[0]['offset']==10 and observed[0]['limit']==20
    assert observed[0]['start']=='2026-09-01T00:00:00+08:00'
    assert args['start']=='2026-09-01T00:00:00'
    with pytest.raises(ValidationError): br._date(args['start'])
