"""交付核验只认真实文件与正文引用，使用合成 Office 包验证缺项边界。"""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.modules.openwebui.delivery import DeliveryIncomplete, DeliveryTracker, inspect_file, requested_outputs


def picture(color="blue"):
    """生成有可见内容的合成图片；图片本身不会被认作可信业务图表。"""
    buffer = BytesIO()
    image = Image.new("RGB", (400, 240), "white")
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((50, 40, 140, 210), fill=color)
    drawing.rectangle((170, 110, 260, 210), fill="green")
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def word(image=None):
    """生成真实可解析的 Word，按需插入一张正文图片。"""
    document = Document()
    document.add_paragraph("合成交付报告")
    if image is not None:
        document.add_picture(BytesIO(image))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def presentation(pages):
    """页数来自真实演示文稿关系，不依赖文件命名。"""
    document = Presentation()
    for index in range(pages):
        slide = document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2)).text = f"第 {index + 1} 页合成分析内容"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def append_member(data, name, payload):
    """向包内添加未被正文引用的文件，模拟遗留或无用附件。"""
    buffer = BytesIO(data)
    with ZipFile(buffer, "a") as archive:
        archive.writestr(name, payload)
    return buffer.getvalue()


def replace_member(data, name, transform):
    """精确改写一个 Office 分片，避免 ZIP 重复条目掩盖真正被读取的版本。"""
    result = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(result, "w") as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == name:
                payload = transform(payload)
                if payload is None:
                    continue
            target.writestr(entry.filename, payload)
    return result.getvalue()


def workbook(value=None):
    """创建真实工作簿，空白版本只用于验证空壳不能通过验收。"""
    document = Workbook()
    if value is not None:
        document.active.append(["设备数量", value])
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def observe_image(monkeypatch, tracker, *, name="python_execute", filename="设备大屏.png",
                  metadata=None, data=None, image=None, run_id="run", result_status="success",
                  trusted_source_run_id=None, snapshot=None):
    """只用内存工件与存储替身模拟已经返回的工具回执，不连接数据库。"""
    import app.modules.openwebui.delivery as delivery

    artifact = SimpleNamespace(id="image", name=filename, workspace_id="workspace", run_id=run_id,
                               storage_key="image-key", status="ready", metadata_json=metadata or {})
    monkeypatch.setattr(delivery, "get_storage", lambda: SimpleNamespace(read_bytes=lambda key: image or picture()))
    run = SimpleNamespace(id="run", workspace_id="workspace", context_snapshot=snapshot or {})
    result = SimpleNamespace(status=result_status, metadata={}, data=data or {}, artifacts=[{"id": "image"}],
                             error_code=None, error_summary="失败样例" if result_status != "success" else None)
    tracker.observe(SimpleNamespace(get=lambda cls, key: artifact), run, name, {}, result,
                    trusted_source_run_id=trusted_source_run_id)
    return artifact


@pytest.mark.parametrize("content", [
    "怎么生成 Word 和 Excel？", "如何做 15 页 PPT？", "不要生成 Word 和 Excel", "停止制作 15 页 PPT",
])
def test_inquiry_negation_or_stop_does_not_create_delivery_authorization(content):
    assert requested_outputs(content) == {}


def test_negated_clause_does_not_hide_positive_word_request():
    assert requested_outputs("不要生成 Excel，帮我制作 Word") == {"docx": True}


def test_composite_delivery_is_not_satisfied_by_word_only():
    tracker = DeliveryTracker("帮我制作饼图并导出 Word 和 Excel，再做 15 页 PPT 和数字大屏放入 Word")
    tracker.files["word"] = inspect_file("报告.docx", word())
    missing = tracker.status()["missing"]
    assert "Word 文件" not in missing
    assert {"Excel 文件", "15 页 PPT", "可打开的数字大屏", "图表图片"}.issubset(missing)
    assert any("Word 中嵌入" in item for item in missing)


def test_presentation_requires_exact_requested_page_count():
    tracker = DeliveryTracker("做 15 页 PPT")
    tracker.files["ppt"] = inspect_file("报告.pptx", presentation(14))
    assert tracker.status()["missing"] == ["15 页 PPT"]
    tracker.files["ppt"] = inspect_file("报告.pptx", presentation(15))
    assert tracker.status()["missing"] == []


def test_orphan_slide_cannot_satisfy_requested_page_count():
    data = presentation(14)
    with ZipFile(BytesIO(data)) as archive:
        payload = archive.read("ppt/slides/slide1.xml")
    data = append_member(data, "ppt/slides/slide15.xml", payload)
    assert inspect_file("报告.pptx", data)["pages"] == 14


def test_embedded_dashboard_requires_matching_picture_hash():
    dashboard = picture("blue")
    unrelated = picture("red")
    tracker = DeliveryTracker("制作大屏放入 Word")
    tracker.dashboard = True
    tracker.dashboard_images.add(sha256(dashboard).hexdigest())
    tracker.files["word"] = inspect_file("报告.docx", word(unrelated))
    assert any("Word 中嵌入" in item for item in tracker.status()["missing"])
    tracker.files["word"] = inspect_file("报告.docx", word(dashboard))
    assert tracker.status()["missing"] == []


def test_orphan_dashboard_media_does_not_count_as_embedded():
    dashboard = picture()
    data = append_member(word(), "word/media/dashboard.png", dashboard)
    tracker = DeliveryTracker("制作大屏放入 Word")
    tracker.dashboard = True
    tracker.dashboard_images.add(sha256(dashboard).hexdigest())
    tracker.files["word"] = inspect_file("报告.docx", data)
    assert tracker.files["word"]["images"] == []
    assert any("Word 中嵌入" in item for item in tracker.status()["missing"])


def test_invalid_office_xml_is_not_a_completed_file():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "not an office document")
    assert inspect_file("报告.docx", buffer.getvalue()) == {}


def test_repeated_no_progress_raises_incomplete_and_preserves_missing():
    tracker = DeliveryTracker("导出 Word 和 Excel")
    tracker.files["word"] = inspect_file("报告.docx", word())
    assert "Excel 文件" in tracker.continuation()
    assert "Excel 文件" in tracker.continuation()
    with pytest.raises(DeliveryIncomplete, match="Excel 文件"):
        tracker.continuation()
    assert tracker.status()["missing"] == ["Excel 文件"]


def test_blocked_database_dependency_does_not_auto_retry():
    tracker = DeliveryTracker("导出 Word")
    tracker.blocker = "数据库查询超时，需核对已发出的查询"
    with pytest.raises(DeliveryIncomplete, match="数据库查询超时"):
        tracker.continuation()
    assert tracker.continuations == 0


def test_polite_action_question_still_requires_files():
    """礼貌疑问语气的明确制作请求仍要核对附件。"""
    assert requested_outputs("能帮我生成 Word 和 Excel 吗？") == {"docx": True, "xlsx": True}


@pytest.mark.parametrize("size,color", [((2, 2), "white"), ((400, 240), "white"), ((400, 240), "blue")])
def test_tiny_or_solid_images_cannot_count_as_delivery(size, color):
    buffer = BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    assert inspect_file("实际饼图.png", buffer.getvalue()) == {}


def test_transparent_rgb_noise_does_not_count_as_visible_content():
    image = Image.new("RGBA", (400, 240), (255, 0, 0, 0))
    ImageDraw.Draw(image).rectangle((50, 50, 200, 200), fill=(0, 255, 0, 0))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    assert inspect_file("大屏.png", buffer.getvalue()) == {}


def test_nonempty_picture_structure_is_not_chart_evidence():
    evidence = inspect_file("分析饼图.png", picture())
    assert evidence["kind"] == "image"
    assert evidence["width"] == 400
    assert "chart_types" not in evidence


def test_empty_workbook_does_not_count_as_excel_delivery():
    assert inspect_file("数据.xlsx", workbook()) == {}


def test_zero_is_valid_spreadsheet_content():
    evidence = inspect_file("数据.xlsx", workbook(0))
    assert evidence["kind"] == "xlsx"
    assert evidence["nonempty_cells"] == 2


def test_formula_workbook_is_structural_evidence_not_computed_truth():
    evidence = inspect_file("数据.xlsx", workbook("=SUM(1,2)"))
    assert evidence["formula_cells"] == 1
    assert evidence["validation"] == "structure_and_nonempty_content"


def test_excel_error_cells_do_not_count_as_complete():
    assert inspect_file("数据.xlsx", workbook("#REF!")) == {}


def test_excel_broken_sheet_relationship_is_rejected():
    data = replace_member(workbook(12), "xl/worksheets/sheet1.xml", lambda _: None)
    assert inspect_file("数据.xlsx", data) == {}


def test_excel_orphan_sheet_cannot_replace_missing_declared_sheet():
    original = workbook(12)
    with ZipFile(BytesIO(original)) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml")
    data = append_member(original, "xl/worksheets/sheet2.xml", sheet)
    data = replace_member(data, "xl/worksheets/sheet1.xml", lambda _: None)
    assert inspect_file("数据.xlsx", data) == {}


def test_blank_slides_do_not_satisfy_page_count():
    document = Presentation()
    for _ in range(15):
        document.slides.add_slide(document.slide_layouts[6])
    buffer = BytesIO()
    document.save(buffer)
    assert inspect_file("15页.pptx", buffer.getvalue()) == {}


def test_slide_with_broken_picture_relationship_is_rejected():
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(picture()), Inches(1), Inches(1))
    buffer = BytesIO()
    document.save(buffer)
    assert inspect_file("报告.pptx", buffer.getvalue())["pages"] == 1
    broken = replace_member(buffer.getvalue(), "ppt/media/image1.png", lambda _: None)
    assert inspect_file("报告.pptx", broken) == {}


def test_empty_word_is_rejected():
    document = Document()
    buffer = BytesIO()
    document.save(buffer)
    assert inspect_file("报告.docx", buffer.getvalue()) == {}


def test_word_with_missing_referenced_picture_is_rejected():
    broken = replace_member(word(picture()), "word/media/image1.png", lambda _: None)
    assert inspect_file("报告.docx", broken) == {}


def test_office_without_package_entrypoint_is_rejected():
    broken = replace_member(word(), "_rels/.rels", lambda _: None)
    assert inspect_file("报告.docx", broken) == {}


def test_pdf_header_alone_does_not_count_as_a_document():
    assert inspect_file("报告.pdf", b"%PDF-1.7\nnot a valid file") == {}


def test_empty_pdf_is_rejected():
    from pypdf import PdfWriter

    document = PdfWriter()
    document.add_blank_page(width=595, height=842)
    buffer = BytesIO()
    document.write(buffer)
    assert inspect_file("报告.pdf", buffer.getvalue()) == {}


def test_nonempty_python_picture_is_retained_but_not_trusted_as_chart(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, metadata={"tool": "python_execute"},
                  data={"chart_types": ["pie"]})
    assert tracker.files["image"]["kind"] == "image"
    assert tracker.status()["missing"] == ["图表图片"]


def test_filename_cannot_claim_dashboard_origin(monkeypatch):
    tracker = DeliveryTracker("制作大屏放入 Word")
    dashboard = picture()
    observe_image(monkeypatch, tracker, image=dashboard, filename="真实大屏_dashboard.png")
    tracker.dashboard = True
    tracker.files["word"] = inspect_file("报告.docx", word(dashboard))
    assert not tracker.dashboard_images
    assert any("Word 中嵌入" in item for item in tracker.status()["missing"])


def chart_metadata(chart_type="pie"):
    """模拟专用服务端渲染工具写入的工件来源字段。"""
    return {"tool": "render_chart_image", "delivery_evidence": {
        "producer": "server_renderer.v1", "kind": "chart", "chart_types": [chart_type], "source_query_id": "query-1"}}


def test_server_rendered_chart_satisfies_matching_type(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, name="render_chart_image", metadata=chart_metadata())
    assert tracker.status()["missing"] == []
    assert tracker.files["image"]["source_query_id"] == "query-1"


def test_bar_chart_cannot_replace_requested_pie(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, name="render_chart_image", metadata=chart_metadata("bar"))
    assert tracker.status()["missing"] == ["图表图片"]


def test_failed_chart_tool_does_not_satisfy_delivery(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, name="render_chart_image", metadata=chart_metadata(), result_status="failed")
    assert tracker.status()["missing"] == ["图表图片"]


def test_unrelated_previous_run_cannot_provide_new_chart_proof(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, name="render_chart_image", metadata=chart_metadata(), run_id="other-run")
    assert tracker.status()["missing"] == ["图表图片"]


def test_multi_resume_accepts_ancestor_chart_after_server_validates_receipt(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, name="render_chart_image", metadata=chart_metadata(), run_id="original-run",
                  snapshot={"resumed_from_run_id": "previous-resume"}, trusted_source_run_id="original-run")
    assert tracker.status()["missing"] == []


def test_snapshot_source_run_list_cannot_authorize_ancestor_chart_by_itself(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, name="render_chart_image", metadata=chart_metadata(), run_id="original-run",
                  snapshot={"resumed_from_run_id": "previous-resume", "task_resume": {"source_run_ids": ["original-run"]}})
    assert tracker.status()["missing"] == ["图表图片"]


def test_python_cannot_forge_server_chart_metadata(monkeypatch):
    tracker = DeliveryTracker("生成饼图")
    observe_image(monkeypatch, tracker, name="python_execute", metadata=chart_metadata())
    assert tracker.status()["missing"] == ["图表图片"]


def test_inspection_export_has_known_chart_types_without_filename_guess(monkeypatch):
    tracker = DeliveryTracker("生成饼图和柱状图")
    facts = {"checkCount": 6, "counts": {"正常": 4, "异常": 2}}
    observe_image(monkeypatch, tracker, name="export_inspection_report", filename="证据.png",
                  metadata={"source": "saved_inspection_reports", "facts": facts},
                  data={"resultType": "inspection_report_export", "facts": facts})
    assert tracker.status()["missing"] == []
    assert tracker.files["image"]["chart_types"] == ["bar", "pie"]


def test_inspection_analysis_is_not_a_dashboard_image(monkeypatch):
    tracker = DeliveryTracker("制作大屏放入 Word")
    facts = {"checkCount": 6}
    observe_image(monkeypatch, tracker, name="export_inspection_report", filename="大屏分析.png",
                  metadata={"source": "saved_inspection_reports", "facts": facts},
                  data={"resultType": "inspection_report_export", "facts": facts,
                        "dashboard": {"reportId": "dashboard-1"}})
    tracker.files["word"] = inspect_file("报告.docx", word(picture()))
    assert tracker.dashboard is True
    assert not tracker.dashboard_images
    assert any("Word 中嵌入" in item for item in tracker.status()["missing"])


def test_actual_dashboard_renderer_provenance_and_word_reference_complete_delivery(monkeypatch):
    tracker = DeliveryTracker("制作大屏放入 Word")
    observe_image(monkeypatch, tracker, name="render_dashboard_image", filename="快照.png",
                  metadata={"tool": "render_dashboard_image", "delivery_evidence": {
                      "producer": "server_renderer.v1", "kind": "dashboard_image", "dashboard_id": "dashboard-1",
                      "source_query_ids": ["query-1"]}}, data={"dashboard_id": "dashboard-1"})
    tracker.files["word"] = inspect_file("报告.docx", word(picture()))
    assert tracker.status()["missing"] == []


def test_dashboard_receipt_and_stored_provenance_must_match(monkeypatch):
    tracker = DeliveryTracker("制作大屏放入 Word")
    observe_image(monkeypatch, tracker, name="render_dashboard_image", metadata={
        "tool": "render_dashboard_image", "delivery_evidence": {
            "producer": "server_renderer.v1", "kind": "dashboard_image", "dashboard_id": "dashboard-1"}},
                  data={"dashboard_id": "dashboard-2"})
    assert not tracker.dashboard_images


def test_actual_presentation_progress_can_continue_beyond_four_rounds():
    tracker = DeliveryTracker("做 15 页 PPT")
    for written in range(8):
        tracker.progress = {"written_count": written, "body_page_count": 13, "total_page_count": 15}
        assert tracker.continuation()
    assert tracker.continuations == 8


def test_unrelated_files_and_progress_titles_do_not_reset_stall_count():
    tracker = DeliveryTracker("做 15 页 PPT")
    for index in range(2):
        tracker.progress = {"title": f"新标题 {index}"}
        tracker.files[str(index)] = {"kind": "docx", "name": f"无关文档 {index}"}
        assert tracker.continuation()
    tracker.files["new"] = {"kind": "image", "name": "无关图片"}
    tracker.progress = {"title": "再换标题"}
    with pytest.raises(DeliveryIncomplete):
        tracker.continuation()


def test_reducing_then_restoring_written_count_is_not_new_progress():
    tracker = DeliveryTracker("做 15 页 PPT")
    tracker.progress = {"written_count": 3, "body_page_count": 13, "total_page_count": 15}
    assert tracker.continuation()
    tracker.progress["written_count"] = 0
    assert tracker.continuation()
    tracker.progress["written_count"] = 3
    with pytest.raises(DeliveryIncomplete):
        tracker.continuation()
