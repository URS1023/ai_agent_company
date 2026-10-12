"""查询 Excel 使用服务端保存数据，保留范围和宋体，不让单元格字符串变成公式。"""
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from app.modules.openwebui.query_spreadsheet import export_query_spreadsheet


@pytest.mark.parametrize("name,expected", [
    (None, "analysis-query-1.xlsx"),
    ("备品备件库存表.xlsx", "备品备件库存表.xlsx"),
    ("备品备件库存表", "备品备件库存表.xlsx"),
])
def test_source_values_fonts_and_formula_safety_are_preserved(monkeypatch, name, expected):
    from app.modules.openwebui import chart_renderer
    from app.modules.analytics import exports, queries
    from app.modules.artifact import service
    captured = {}
    row = SimpleNamespace(id="query-1", result_json={"truncated": True})
    result = SimpleNamespace(columns=[SimpleNamespace(columnName="设备"), SimpleNamespace(columnName="温度")],
        data=[{"设备": "=UNTRUSTED()", "温度": 23.5}], datasourceId="source-1", datasourceVersion=2,
        totalCount=1, truncated=True, analysis="本次返回的样本", generatedSql="SELECT device, temperature FROM test_table")

    def load(ctx, context, identity):
        assert identity == "query-1" and context["run_id"] == "run-1"
        captured["authorized"] = True
        return row

    def create(db, **kwargs):
        captured.update(kwargs)
        return SimpleNamespace(id="artifact-1")

    def no_query(*args, **kwargs):
        raise AssertionError("Excel 导出不应重新查询")

    monkeypatch.setattr(chart_renderer, "load_authorized_query", load)
    monkeypatch.setattr(exports, "historical_result", lambda *args, **kwargs: result)
    monkeypatch.setattr(exports, "commit_change", lambda *args, **kwargs: None)
    monkeypatch.setattr(queries, "execute_query", no_query)
    monkeypatch.setattr(service, "create_artifact", create)
    monkeypatch.setattr(service, "artifact_reference", lambda artifact: {"id": artifact.id})
    args = {"query_id": "query-1", **({"name": name} if name is not None else {})}
    response = export_query_spreadsheet(SimpleNamespace(db=None, workspace_id="workspace-1", user_id="user-1"),
                                       {"run_id": "run-1"}, args)
    workbook = load_workbook(BytesIO(captured["data"]))
    assert captured["authorized"]
    assert workbook["查询结果"]["A2"].value == "=UNTRUSTED()"
    assert workbook["查询结果"]["A2"].data_type == "s"
    assert workbook["查询结果"]["B2"].value == 23.5
    assert workbook["查询结果"]["A2"].font.name in {"宋体", "SimSun"}
    assert "数据来源" in workbook.sheetnames
    assert response.data["truncated"] is True
    assert captured["metadata"]["delivery_evidence"]["source_query_id"] == "query-1"
    assert captured["name"] == response.data["name"] == expected


@pytest.mark.parametrize("name", ["库存.csv", "库存.xls", "   "])
def test_spreadsheet_name_rejects_false_formats(name):
    from app.errors import ValidationError
    from app.modules.chat_artifacts.query_spreadsheet import spreadsheet_filename
    with pytest.raises(ValidationError):
        spreadsheet_filename(name, "default.xlsx")
