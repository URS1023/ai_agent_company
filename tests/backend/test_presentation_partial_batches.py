"""验证长文稿分批写入时，有效正文不被错误页码和空正文回滚。"""
from __future__ import annotations

import io

import pytest
from pptx import Presentation

from app.common.storage import LocalStorage
from app.errors import ToolExecutionError
from app.modules.artifact.renderers import render_presentation
from app.modules.presentation import draft


def _page(number: int) -> dict:
    """创建不依赖业务数据库或模型服务的合成正文。"""
    return {
        "slide_number": number,
        "title": f"验收主题 {number}",
        "layout": "cards",
        "key_message": f"第 {number} 个主题通过明确范围、证据和行动形成完整的执行闭环。",
        "items": [
            {"title": "明确范围", "description": "列出本次合成验收的输入边界和适用范围，避免把局部结果当作完整结论。"},
            {"title": "核对证据", "description": "保留每项数据的来源与计算方法，并检查所有生成页面使用一致的统计口径。"},
            {"title": "落实行动", "description": "根据已确认的事实安排负责人、执行步骤和复核记录，持续跟踪处理结果。"},
        ],
    }


@pytest.fixture
def staged(monkeypatch, tmp_path):
    """将演示文稿草稿限制在本用例的临时目录中。"""
    storage = LocalStorage(tmp_path)
    monkeypatch.setattr(draft, "get_storage", lambda: storage)
    draft.create_draft(
        "workspace", "run", {"title": "合成验收", "page_count": 15, "slides": [_page(i) for i in range(1, 14)]},
        template_id="executive-ivory", aspect_ratio="16:9",
    )
    return "workspace", "run"


def test_valid_final_body_page_survives_extra_cover_and_closing(staged):
    """复现第13页与越界14/15同批提交；有效页必须保存且最终仍为15页。"""
    draft.write_draft_slide_batches(*staged, [_page(i) for i in range(1, 13)])
    status = draft.write_draft_slide_batches(*staged, [
        _page(13), _page(14), {"slide_number": 15, "layout": "closing", "title": "结束"},
    ])
    assert status["written_now"] == [13]
    assert status["missing_slide_numbers"] == []
    assert [item["slide_number"] for item in status["rejected_slides"]] == [14, 15]
    assert status["next_action"] == "finalize_presentation"
    spec = draft.final_spec(*staged)
    deck = Presentation(io.BytesIO(render_presentation(spec)))
    assert len(deck.slides) == 15
    assert "验收主题 13" in " ".join(shape.text for shape in deck.slides[13].shapes if shape.has_text_frame)


def test_empty_body_is_reported_before_finalize(staged):
    """错误结构不能在写入时宣称已完成，需明确提供可修正字段。"""
    with pytest.raises(ToolExecutionError) as caught:
        draft.write_draft_slide_batches(*staged, [{
            "slide_number": 1, "layout": "content", "title": "验收主题 1",
            "key_message": "只有一句结论", "notes": "这里只是备注", "series": [],
        }])
    assert "items" in caught.value.message
    progress = caught.value.extra["presentation_progress"]
    assert progress["written_count"] == 0
    assert progress["missing_slide_numbers"] == list(range(1, 14))
    assert not draft.load_draft(*staged)["slides"][0]["written"]


def test_bad_page_does_not_discard_other_valid_pages(staged):
    """非法页码、空正文和重复页均只影响各自的输入项。"""
    status = draft.write_draft_slide_batches(*staged, [
        _page(1), {**_page(2), "slide_number": "oops"},
        {"slide_number": 3, "layout": "table", "title": "缺少表格内容"},
        _page(4), _page(4),
    ])
    assert status["written_now"] == [1, 4]
    assert len(status["rejected_slides"]) == 3
    assert status["batch_status"] == "partial"
    assert 1 not in status["missing_slide_numbers"] and 4 not in status["missing_slide_numbers"]
    assert 2 in status["missing_slide_numbers"] and 3 in status["missing_slide_numbers"]


def test_missing_final_page_error_contains_repair_progress(staged):
    """终稿未齐时，错误保留完整进度供原生工具循环继续执行。"""
    draft.write_draft_slide_batches(*staged, [_page(i) for i in range(1, 13)])
    with pytest.raises(ToolExecutionError) as caught:
        draft.final_spec(*staged)
    assert caught.value.extra["presentation_progress"]["missing_slide_numbers"] == [13]


def test_legacy_openwebui_schema_keeps_complete_body_shapes(monkeypatch):
    """固定旧版上游原字节，验证历史兼容合同，不依赖本机未跟踪源码。"""
    import ast
    import copy
    import hashlib
    import json
    from pathlib import Path
    from types import SimpleNamespace

    from app.modules.openwebui.client import OpenWebUIClient
    from app.providers.tool.builtin_provider import builtin_specs
    from app.providers.tool.validator import validate_arguments

    tool = next(item for item in builtin_specs() if item.name == "write_presentation_slides")
    spec = tool.to_openai_function()["function"]
    root = Path(__file__).resolve().parents[2]
    upstream = root / "tests/fixtures/openwebui-0.11.4/tools.py.txt"
    source = upstream.read_bytes()
    manifest = json.loads((root / "services/openwebui/UPSTREAM.json").read_text(encoding="utf-8"))
    assert hashlib.sha256(source).hexdigest() == manifest["files"]["backend/open_webui/utils/tools.py"]
    tree = ast.parse(source.decode("utf-8"))
    functions = [node for node in tree.body if isinstance(node, ast.FunctionDef)
                 and node.name in {"clean_properties", "clean_openai_tool_schema"}]
    assert len(functions) == 2 and {node.name for node in functions} == {"clean_properties", "clean_openai_tool_schema"}
    namespace = {"copy": copy}
    exec(compile(ast.Module(body=functions, type_ignores=[]), str(upstream), "exec"), namespace)
    cleaned = namespace["clean_openai_tool_schema"](spec)
    properties = cleaned["parameters"]["properties"]["slides"]["items"]["properties"]
    assert {"items", "chart", "table", "left", "right"} <= properties.keys()
    assert {"categories", "series"} <= properties["chart"]["properties"].keys()
    assert {"headers", "rows"} <= properties["table"]["properties"].keys()
    assert "description" in properties["left"]["items"]["properties"]
    payload = {"slides": [
        _page(1),
        {"slide_number": 2, "title": "合成饼图", "layout": "chart", "chart": {
            "chart_type": "pie", "categories": ["已完成", "未完成"],
            "series": [{"name": "合成数量", "values": [8, 2]}],
        }},
        {"slide_number": 3, "title": "合成表格", "layout": "table", "table": {
            "headers": ["分类", "合成数量"], "rows": [["已完成", 8], ["未完成", 2]],
        }},
        {"slide_number": 4, "title": "合成对比", "layout": "comparison",
            "left": [{"title": "开始前", "description": "校验输入"}],
            "right": [{"title": "结束后", "description": "核验交付"}]},
    ]}
    assert validate_arguments(payload, cleaned["parameters"], {}) == payload
    native = object.__new__(OpenWebUIClient)
    native.socket = SimpleNamespace(get_sid=lambda: "synthetic-session")
    captured = {}
    monkeypatch.setattr(native, "request", lambda method, path, **kwargs: captured.update(kwargs))
    native.submit(chat_id="synthetic-chat", message_id="synthetic-message",
                  user_message={"content": "合成工具定义验收", "parentId": None},
                  model="synthetic", specs=[cleaned], system="合成验收")
    sent = captured["json"]["tool_servers"][0]["specs"][0]
    assert sent["parameters"] == cleaned["parameters"]
    assert "notes 只是备注" in sent["description"]
