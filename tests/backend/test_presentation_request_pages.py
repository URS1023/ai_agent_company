"""多工件请求中 PPT 文字页数约束的真实草稿与 PPTX 行为验证。"""
from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from pptx import Presentation

from app.common.storage import LocalStorage
from app.modules.artifact.renderers import render_presentation
from app.modules.presentation import draft, service
from app.modules.presentation.request_constraints import infer_page_constraint
from app.providers.tool import builtin_provider


REQUEST = (
    "帮我查看当前仓库的备品备件库存，生成一份名为备品备件库存表的Excel和一份库存数据PPT，再做一份ERP系统介绍PPT。"
    "使用当前选择的数据源，数据文件要带备件名称、仓库名称、计量单位、批次和安全库存，保留空值，给出需要核对的问题。"
    "ERP介绍为通用产品介绍，六到八页即可。"
)


def page(number):
    """合成文字仅用于验收页数与页序，不包含模拟业务数字。"""
    return {"title": f"主题{number}", "layout": "cards", "key_message": "以明确的业务边界、统一流程和可核对的交付说明这一主题。",
            "items": [{"title": "业务范围", "description": "说明这个主题对应的业务目标、适用范围与前置条件，所有内容只作为通用产品介绍，不代表已实施的项目。"},
                      {"title": "具体流程", "description": "用明确步骤说明业务人员如何发起申请、核对信息并完成任务，保持上下游交接清晰并保留必要的执行记录。"},
                      {"title": "验收依据", "description": "围绕实际交付核对资料是否齐全、口径是否一致与流程是否完整，不使用没有来源的数据或未经确认的效果承诺。"}]}


@pytest.fixture
def staging(monkeypatch, tmp_path):
    """工具存储与模板解析均独立于在用数据库和文件。"""
    storage = LocalStorage(tmp_path)
    monkeypatch.setattr(draft, "get_storage", lambda: storage)
    monkeypatch.setattr(service, "resolve_template", lambda *args: {"id": "executive-ivory"})


def create_plan(run_id, title, body_count, *, request=REQUEST, inputs=None, **args):
    """从真实工具分支读取 Run 的当前用户请求；不直接给模型参数补页数。"""
    tool = next(item for item in builtin_provider.builtin_specs() if item.name == "create_presentation_plan")
    run = SimpleNamespace(context_snapshot={"user_request": request, "plugin_inputs": inputs or {}})
    result = builtin_provider.invoke(tool, {"title": title, "slides": [page(i) for i in range(1, body_count + 1)], **args}, {
        "workspace_id": "w", "user_id": "u", "run_id": run_id,
        "db": SimpleNamespace(get=lambda *args: run),
    })
    return result.data


def finish(run_id, status):
    """经过真实草稿写入、完成校验和 PPTX 渲染，再读取物理页数。"""
    draft.write_draft_slide_batches("w", run_id, [{**page(index), "title": item["title"], "slide_number": index}
        for index, item in enumerate(status["outline"], 1)])
    return Presentation(io.BytesIO(render_presentation(draft.final_spec("w", run_id))))


def test_real_multi_artifact_request_keeps_erp_inside_range_and_inventory_unchanged(staging):
    erp = create_plan("erp", "ERP系统介绍", 7)
    inventory = create_plan("stock", "库存数据PPT", 4)
    assert erp["total_page_count"] == 8 and erp["body_page_count"] == 6
    assert erp["page_count_constraint"]["minimum"] == 6
    assert erp["page_count_constraint"]["maximum"] == 8
    assert inventory["total_page_count"] == 6 and inventory["page_count_constraint"] is None
    assert len(finish("erp", erp).slides) == 8
    assert len(finish("stock", inventory).slides) == 6


@pytest.mark.parametrize("phrase, natural, expected", [
    ("六到八页", 9, 8), ("6-8页", 7, 7), ("6～8页", 5, 6),
    ("六至八页", 9, 8), ("六页到八页", 7, 7), ("八页", 9, 8),
    ("二十一到二十三页", 24, 23),
])
def test_chinese_and_arabic_total_page_counts(phrase, natural, expected):
    result = infer_page_constraint(f"ERP系统介绍，{phrase}即可。", "ERP系统介绍", natural_page_count=natural)
    assert result["selected"] == expected


@pytest.mark.parametrize("user_text", [
    "制作ERP系统介绍和库存数据PPT，六到八页。",
    "制作ERP系统介绍。再做库存PPT，六到八页。",
    "ERP介绍六页；ERP介绍八页。",
    "ERP介绍正文六页。",
    "ERP介绍至少六页。",
    "ERP介绍第六页写方案。",
    "ERP介绍六页左右。",
    "生成两份PPT，六到八页。",
    "库存PPT引用ERP的数据，六到八页。",
    "ERP介绍五十一页。",
])
def test_ambiguous_or_unrelated_constraints_are_not_guessed(user_text):
    assert infer_page_constraint(user_text, "ERP系统介绍", natural_page_count=9) is None


@pytest.mark.parametrize("inputs, args, expected", [
    ({"page_count": 5}, {"page_count": 7}, 5),
    ({}, {"page_count": 7}, 7),
    ({"page_count": "auto"}, {"page_count": 7}, 7),
])
def test_explicit_selected_or_tool_page_count_has_priority(staging, inputs, args, expected):
    status = create_plan("explicit", "ERP系统介绍", 7, inputs=inputs, **args)
    assert status["total_page_count"] == expected
    assert status["page_count_constraint"] is None
    assert len(finish("explicit", status).slides) == expected


def test_explicit_no_cover_keeps_all_seven_body_topics(staging):
    status = create_plan("no-cover", "ERP系统介绍", 7, include_cover=False, include_closing=False)
    assert status["total_page_count"] == status["body_page_count"] == 7
    assert [slide["title"] for slide in status["outline"]] == [f"主题{i}" for i in range(1, 8)]
    assert len(finish("no-cover", status).slides) == 7


def test_distinct_named_decks_keep_separate_counts():
    request = "ERP系统介绍六页，库存数据PPT四页。"
    assert infer_page_constraint(request, "ERP系统介绍", natural_page_count=9)["selected"] == 6
    assert infer_page_constraint(request, "库存数据PPT", natural_page_count=9)["selected"] == 4


def test_product_acronym_does_not_bind_another_deck_about_the_same_product():
    assert infer_page_constraint("ERP介绍六到八页。", "ERP实施方案", natural_page_count=9) is None
