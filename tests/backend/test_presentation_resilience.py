"""Regression coverage for resilient staged PPT generation."""
from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.common.storage import LocalStorage
from app.modules.artifact.renderers import render_presentation
from app.modules.presentation import draft as draft_module
from app.modules.presentation.spec_normalizer import normalize_presentation_spec
from app.modules.runtime.agent import (
    PresentationRepairLimitError,
    _friendly_run_error,
    _track_presentation_finalize_failure,
)
from app.providers.tool.builtin_provider import (
    _presentation_quality_issues,
    builtin_specs,
)


def _rich_cards(title: str) -> dict:
    return {
        "title": title,
        "layout": "cards",
        "key_message": f"{title}需要以统一数据、标准流程和持续运营共同支撑。",
        "items": [
            {
                "title": "业务目标",
                "description": "明确范围、责任人与量化指标，确保系统建设直接服务经营目标。",
            },
            {
                "title": "落地机制",
                "description": "统一主数据和跨部门流程，并通过版本、权限和审计保证执行一致。",
            },
            {
                "title": "持续改进",
                "description": "上线后跟踪效率、成本和质量变化，按月复盘并迭代配置。",
            },
        ],
    }


def test_model_chart_aliases_and_top_level_table_are_preserved():
    spec = normalize_presentation_spec({
        "title": "ERP系统介绍",
        "slides": [
            {
                "title": "数字化转型中的ERP",
                "layout": "chart",
                "key_message": "ERP将分散业务数据转化为统一、实时、可分析的经营数据。",
                "chart": {
                    "type": "bar",
                    "categories": ["流程自动化", "数据整合", "智能决策"],
                    "series": [
                        {"label": "传统系统", "data": [30, 25, 15]},
                        {"label": "ERP+数字化", "data": [85, 90, 75]},
                    ],
                    "insight": "ERP在数据整合和流程自动化方面形成明显能力提升。",
                },
            },
            {
                "title": "ROI评估框架",
                "layout": "table",
                "key_message": "ROI需要同时核算效率提升、成本节约和风险降低。",
                "table": {
                    "headers": ["评估维度", "指标", "典型收益", "解释"],
                    "rows": [
                        ["效率", "月结周期", "7天缩短至2天", "实时凭证减少集中处理"],
                        ["成本", "库存资金", "降低25%", "预测和补货协同减少积压"],
                        ["风险", "审计整改", "降低60%", "规则拦截并保留完整轨迹"],
                    ],
                },
            },
            {
                "title": "数据分析模块：决策智能化",
                "layout": "chart",
                "key_message": "实时看板、预测分析和异常预警共同支撑管理决策。",
                "chart_type": "table",
                "categories": ["指标类别", "描述"],
                "series": [
                    {"name": "经营看板", "values": ["营收概览", "业绩排名", "趋势分析"]},
                    {"name": "预测分析", "values": ["销售预测", "库存预测", "资金预测"]},
                    {"name": "异常预警", "values": ["库存异常", "应收逾期", "成本偏差"]},
                ],
                "insight": "数据分析从事后统计转向实时感知和事前预测。",
            },
            {
                "title": "传统管理 vs ERP 数字化管理对比",
                "layout": "comparison",
                "key_message": "ERP将人工经验驱动转变为标准流程与实时数据驱动。",
                "comparison": {
                    "left": {
                        "title": "传统管理",
                        "items": ["重复手工录入", "纸质流程易延误", "数据更新滞后"],
                    },
                    "right": {
                        "title": "ERP 数字化管理",
                        "items": ["业务单据自动生成", "电子流程全程追溯", "经营数据实时更新"],
                    },
                },
            },
        ],
    })

    chart = spec["slides"][0]["items"][0]
    assert chart["type"] == "chart"
    assert chart["chart_type"] == "bar"
    assert chart["series"] == [
        {"name": "传统系统", "values": [30, 25, 15]},
        {"name": "ERP+数字化", "values": [85, 90, 75]},
    ]
    table = spec["slides"][1]["items"][0]
    assert table["type"] == "table"
    assert table["rows"][0] == ["评估维度", "指标", "典型收益", "解释"]
    assert len(table["rows"]) == 4
    series_table = spec["slides"][2]["items"][0]
    assert spec["slides"][2]["layout"] == "table"
    assert series_table["rows"][0] == ["指标类别", "描述"]
    assert series_table["rows"][1] == ["经营看板", "营收概览；业绩排名；趋势分析"]
    comparison = spec["slides"][3]
    assert comparison["layout"] == "comparison"
    assert comparison["left_title"] == "传统管理"
    assert comparison["right_title"] == "ERP 数字化管理"
    assert len(comparison["left"]) == len(comparison["right"]) == 3
    assert _presentation_quality_issues(spec) == []


def test_explicit_rich_layout_wins_over_malformed_auxiliary_chart():
    """Regression for the real ERP payload that previously caused an endless repair loop."""
    spec = normalize_presentation_spec({
        "title": "ERP系统介绍",
        "slides": [
            {
                "title": "国际与国内 ERP 对比",
                "layout": "comparison",
                "key_message": "两类厂商各有侧重，应根据企业规模和全球化要求权衡。",
                "chart": {
                    "type": "bar",
                    "categories": ["国际厂商", "国内厂商"],
                    "series": [
                        {"label": "功能完整性", "data": [90, 85, 80, 75, 70]},
                        {"label": "本地适配性", "data": [70, 80, 90, 88, 85]},
                    ],
                },
                "left": [{
                    "title": "国际厂商",
                    "items": [
                        {"title": "全球能力", "description": "支持跨国家和多会计准则。"},
                        {"title": "行业积累", "description": "大型企业流程覆盖更完整。"},
                        {"title": "实施代价", "description": "周期和总体成本通常较高。"},
                    ],
                }],
                "right": [{
                    "title": "国内厂商",
                    "items": [
                        {"title": "本地规则", "description": "更贴合国内财税和管理习惯。"},
                        {"title": "交付速度", "description": "标准场景可更快完成上线。"},
                        {"title": "适用边界", "description": "复杂全球化能力需要专项评估。"},
                    ],
                }],
            },
            {
                "title": "主流厂商选型",
                "layout": "table",
                "key_message": "选型需联合评估业务覆盖、实施能力与长期总成本。",
                "chart": {
                    "type": "column",
                    "categories": ["SAP", "Oracle", "用友"],
                    "series": [{"label": "综合能力", "data": [92, 90, 86]}],
                },
                "headers": ["厂商", "适用规模", "主要特点"],
                "rows": [
                    ["SAP", "大型集团", "全球化与行业方案完整"],
                    ["Oracle", "中大型企业", "技术平台与财务能力突出"],
                    ["用友", "成长型与大型企业", "本地化适配和服务网络完善"],
                ],
            },
            {
                "title": "ERP 投资回报分析",
                "layout": "chart",
                "key_message": "投入在第二年后逐步转化为效率、库存和风险收益。",
                "categories": ["第1年", "第2年", "第3年", "第4年", "第5年"],
                "series": [
                    {"name": "累计ROI", "data": [-100, -30, 35, 42, 38]},
                    {"name": "年度节约", "data": [15, 25, 35, 40, 38]},
                ],
                "items": [
                    {"point": "初始投入", "text": "第一年集中投入软件、实施和培训。"},
                    {"point": "见效期", "text": "第二年流程优化开始形成节约。"},
                    {"point": "稳定期", "text": "第三年后持续跟踪真实收益。"},
                ],
            },
        ],
    })

    comparison = spec["slides"][0]
    assert comparison["layout"] == "comparison"
    assert len(comparison["left"]) == len(comparison["right"]) == 3
    table = spec["slides"][1]
    assert table["layout"] == "table"
    assert table["items"][0]["rows"][0] == ["厂商", "适用规模", "主要特点"]
    chart = spec["slides"][2]
    assert chart["layout"] == "chart"
    assert chart["items"][0]["series"][0]["values"] == [-100, -30, 35, 42, 38]
    assert _presentation_quality_issues(spec) == []


def test_erp_21_page_draft_infers_page_numbers_splits_batches_and_renders(
    tmp_path,
    monkeypatch,
):
    storage = LocalStorage(str(tmp_path / "storage"))
    monkeypatch.setattr(draft_module, "get_storage", lambda: storage)
    workspace_id = "workspace-erp"
    run_id = "run-erp-21-pages"
    titles = [f"ERP主题{i:02d}" for i in range(1, 20)]

    created = draft_module.create_draft(
        workspace_id,
        run_id,
        {
            "title": "ERP系统介绍",
            "subtitle": "从业务协同到数字化经营",
            "slides": [
                {"title": title, "layout": "cards", "key_message": f"{title}的核心判断"}
                for title in titles
            ],
        },
        template_id="executive-ivory",
        aspect_ratio="16:9",
        selected_page_count=21,
    )
    assert created["body_page_count"] == 19
    assert created["total_page_count"] == 21

    pages = [_rich_cards(title) for title in titles]
    pages[8] = {
        "title": titles[8],
        "layout": "chart",
        "key_message": "ERP使流程自动化、数据整合和实时分析能力同步提升。",
        "chart": {
            "type": "bar",
            "categories": ["流程自动化", "数据整合", "实时分析", "智能决策"],
            "series": [
                {"name": "传统系统", "data": [30, 25, 20, 15]},
                {"name": "ERP+数字化", "data": [85, 90, 80, 75]},
            ],
            "insight": "统一ERP后，数据整合成为实时分析与智能决策的共同基础。",
            "source": "示例实施基线",
        },
    }
    pages[13] = {
        "title": titles[13],
        "layout": "table",
        "key_message": "ERP投资回报必须落实到可追踪、可复核的业务指标。",
        "table": {
            "headers": ["维度", "指标", "目标", "业务解释"],
            "rows": [
                ["效率", "订单周期", "缩短50%", "流程自动触发减少等待"],
                ["成本", "库存资金", "降低25%", "需求预测减少积压"],
                ["质量", "数据错误率", "低于0.1%", "源头录入与规则校验"],
            ],
        },
    }

    first = draft_module.write_draft_slide_batches(
        workspace_id,
        run_id,
        pages[:5],
    )
    assert first["written_now"] == [1, 2, 3, 4, 5]
    assert first["internal_batch_count"] == 2
    assert first["inferred_slide_numbers"] == [1, 2, 3, 4, 5]

    remaining = draft_module.write_draft_slide_batches(
        workspace_id,
        run_id,
        pages[5:],
    )
    assert remaining["missing_slide_numbers"] == []
    assert remaining["internal_batch_count"] == 4
    assert remaining["inferred_slide_numbers"] == list(range(6, 20))

    final = draft_module.final_spec(workspace_id, run_id)
    assert final["page_count"] == 21
    assert _presentation_quality_issues(final) == []
    assert final["slides"][8]["items"][0]["series"][0]["values"] == [30, 25, 20, 15]
    assert final["slides"][13]["items"][0]["rows"][0] == ["维度", "指标", "目标", "业务解释"]

    deck_bytes = render_presentation(final)
    deck = Presentation(io.BytesIO(deck_bytes))
    assert len(deck.slides) == 21
    assert sum(
        1 for slide in deck.slides for shape in slide.shapes
        if getattr(shape, "has_chart", False)
    ) == 1
    assert sum(
        1 for slide in deck.slides for shape in slide.shapes
        if getattr(shape, "has_table", False)
    ) == 1


def test_write_slides_tool_schema_documents_aliases_and_accepts_server_partitioning():
    spec = next(item for item in builtin_specs() if item.name == "write_presentation_slides")
    slides_schema = spec.input_schema["properties"]["slides"]
    item_properties = slides_schema["items"]["properties"]

    assert slides_schema["maxItems"] == 50
    assert "slide_number" not in slides_schema["items"].get("required", [])
    assert {"chart", "table", "categories", "series", "headers", "rows"}.issubset(item_properties)
    series_properties = item_properties["items"]["items"]["properties"]["series"]["items"]["properties"]
    assert {"values", "data"}.issubset(series_properties)


def test_identical_finalize_failures_stop_after_three_attempts_with_specific_message():
    failures: dict[str, int] = {}
    summary = "第 9 页缺少完整的图表分类或数据序列"

    assert _track_presentation_finalize_failure(failures, summary) == 1
    assert _track_presentation_finalize_failure(failures, summary) == 2
    with pytest.raises(PresentationRepairLimitError) as captured:
        _track_presentation_finalize_failure(failures, summary)

    message, code = _friendly_run_error(captured.value, "run-erp")
    assert code == "presentation_repair_limit"
    assert "连续三次" in message
    assert "保留" in message


def test_nested_value_cards_timeline_metrics_and_panels_render_visible_copy():
    """Real providers may wrap label/data in value; no empty cards may survive."""
    spec = normalize_presentation_spec({
        "title": "ERP 系统全面介绍——企业管理数字化转型的核心引擎",
        "subtitle": "面向企业管理者与业务负责人的系统认知指南",
        "template_id": "memphis-studio",
        "slides": [
            {
                "title": "财务模块：业财一体化",
                "layout": "cards",
                "key_message": "统一核算、成本和资金信息，让财务数据实时反映业务。",
                "items": [
                    {"type": "card", "value": {"label": "总账管理", "data": "统一会计科目并自动生成凭证。"}},
                    {"type": "card", "value": {"label": "成本核算", "data": "按订单与产品追踪实际成本。"}},
                    {"type": "card", "value": {"label": "资金管控", "data": "预测现金流并监控预算执行。"}},
                    {"type": "card", "value": {"label": "经营分析", "data": "把财务结果追溯到业务动作。"}},
                ],
            },
            {
                "title": "ERP 实施方法论",
                "layout": "timeline",
                "key_message": "规划、蓝图、建设、上线与优化需要逐阶段验收。",
                "items": [{
                    "type": "timeline",
                    "value": [
                        {"label": "项目规划", "data": "明确目标、范围、组织和里程碑。"},
                        {"label": "蓝图设计", "data": "梳理现状并确认未来业务流程。"},
                        {"label": "系统建设", "data": "完成配置、开发、迁移和集成测试。"},
                        {"label": "切换上线", "data": "培训用户并完成生产环境切换。"},
                        {"label": "持续优化", "data": "按业务指标复盘并迭代配置。"},
                    ],
                }],
            },
            {
                "title": "ERP 带来的业务价值",
                "layout": "metrics",
                "key_message": "示例指标需要用可追踪的口径衡量。",
                "items": [
                    {"type": "metric", "value": {"label": "库存周转", "data": "+20%"}},
                    {"type": "metric", "value": {"label": "月结周期", "data": "7→2 天"}},
                    {"type": "metric", "value": {"label": "订单交付", "data": "+15%"}},
                ],
            },
            {
                "title": "传统管理 vs ERP 管理",
                "layout": "comparison",
                "key_message": "ERP 用统一流程和实时数据替代分散手工协作。",
                "left": [
                    {"type": "card", "value": {"label": "数据", "data": "分散在表格和独立系统。"}},
                    {"type": "card", "value": {"label": "流程", "data": "人工传递且难以追踪。"}},
                ],
                "right": [
                    {"type": "card", "value": {"label": "数据", "data": "统一口径并实时共享。"}},
                    {"type": "card", "value": {"label": "流程", "data": "系统流转并保留审计轨迹。"}},
                ],
            },
        ],
    })

    assert spec["slides"][0]["items"][0] == {
        "type": "card",
        "title": "总账管理",
        "description": "统一会计科目并自动生成凭证。",
    }
    assert spec["slides"][1]["layout"] == "timeline"
    assert len(spec["slides"][1]["items"]) == 5
    assert spec["slides"][1]["items"][0]["title"] == "项目规划"
    assert spec["slides"][2]["items"][0]["value"] == "+20%"
    assert spec["slides"][3]["left_title"] == "传统管理"
    assert spec["slides"][3]["right_title"] == "ERP 管理"
    assert spec["slides"][3]["left"][0]["description"] == "分散在表格和独立系统。"
    assert _presentation_quality_issues(spec) == []

    deck = Presentation(io.BytesIO(render_presentation(spec)))
    all_text = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    for expected in (
        "总账管理", "统一会计科目并自动生成凭证。", "项目规划",
        "明确目标、范围、组织和里程碑。", "库存周转", "+20%",
        "分散在表格和独立系统。", "系统流转并保留审计轨迹。",
    ):
        assert expected in all_text
    assert "要点 1" not in all_text
    assert "{'label'" not in all_text

    cover = deck.slides[0]
    title_shape = next(shape for shape in cover.shapes if getattr(shape, "text", "") == spec["title"])
    subtitle_shape = next(shape for shape in cover.shapes if getattr(shape, "text", "") == spec["subtitle"])
    assert title_shape.top + title_shape.height < subtitle_shape.top


def test_content_rows_use_available_height_without_footer_overlap_or_top_crowding():
    spec = normalize_presentation_spec({
        "title": "ERP 系统全面介绍",
        "subtitle": "面向管理者与业务负责人的 ERP 系统全面指南",
        "template_id": "memphis-studio",
        "slides": [
            {
                "title": "管理者的五大核心痛点",
                "layout": "content",
                "key_message": "信息孤岛、流程断裂和数据滞后会持续放大管理成本。",
                "items": [
                    {"title": f"痛点 {index}", "description": "跨部门信息无法实时协同，需要重复核对并承担错误风险。"}
                    for index in range(1, 9)
                ],
            },
            {
                "title": "管理层行动建议",
                "layout": "content",
                "key_message": "分步实施、组织保障和持续优化缺一不可。",
                "items": [
                    {"title": "第一步：分步实施", "description": "从价值最高的模块入手，验证效果后逐步扩展。"},
                    {"title": "第二步：一把手工程", "description": "建立跨部门项目治理机制并保障关键资源。"},
                    {"title": "第三步：持续优化", "description": "按经营指标复盘，持续调整流程和系统配置。"},
                ],
            },
        ],
    })

    assert spec["slides"][0]["title"] == "管理者的核心痛点"
    deck = Presentation(io.BytesIO(render_presentation(spec)))

    dense_slide = deck.slides[1]
    last_item = next(
        shape for shape in dense_slide.shapes
        if "痛点 8" in getattr(shape, "text", "")
    )
    assert last_item.top + last_item.height <= Inches(6.72)

    action_slide = deck.slides[2]
    action_titles = [
        next(
            shape for shape in action_slide.shapes
            if getattr(shape, "text", "") == title
        )
        for title in ("第一步：分步实施", "第二步：一把手工程", "第三步：持续优化")
    ]
    assert action_titles[1].top - action_titles[0].top >= Inches(1.3)
    assert action_titles[2].top - action_titles[1].top >= Inches(1.3)

    cover = deck.slides[0]
    subtitle_shape = next(
        shape for shape in cover.shapes
        if getattr(shape, "text", "") == spec["subtitle"]
    )
    assert subtitle_shape.height <= Inches(0.68)


def test_two_column_quality_uses_column_titles_instead_of_comparison_defaults():
    spec = normalize_presentation_spec({
        "title": "联合交付方案",
        "slides": [{
            "title": "联合体模式把高概率淘汰风险降为可管理风险",
            "layout": "two_column",
            "column_titles": ["合作方负责", "我方负责"],
            "left": ["一级资质与属地备案", "现场施工与本地协调"],
            "right": ["平台总体架构与集成", "数据治理与质量验收"],
        }],
    })

    assert spec["slides"][0]["column_titles"] == ["合作方负责", "我方负责"]
    assert _presentation_quality_issues(spec) == []
