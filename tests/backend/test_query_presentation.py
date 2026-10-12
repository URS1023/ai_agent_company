"""隔离验证查询到 PPT 的数值绑定、完整分页、来源及授权边界。"""
from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from pptx import Presentation
import pytest
from sqlalchemy import select

from test_openwebui_query_dashboard import dashboard_case

from app.common.storage import get_storage
from app.config import settings
from app.errors import AppError, ValidationError
from app.models import Artifact, ToolCall
from app.modules.analytics.models import QueryHistory
from app.modules.chat_artifacts.query_presentation import (
    export_query_presentation, presentation_filename, query_presentation_spec,
)


def _values(case, rows, *, columns=None, truncated=False):
    """只更新内存平台回执；本夹具禁止业务查询和模型请求。"""
    case.query.result_json = {**case.query.result_json, "data": rows, "totalCount": len(rows),
        "truncated": truncated, "columns": columns or [{"columnName": key} for key in rows[0]]}
    case.ctx.db.commit()


def _export(case, tmp_path, monkeypatch, **args):
    """读取实际生成的工件和 Office 包，不能以回执声明代替文件验证。"""
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    result = export_query_presentation(case.ctx, case.tool, {"query_ids": [case.query.id], "presentation_mode": "snapshot", **args})
    row = case.ctx.db.get(Artifact, result.artifacts[0]["id"])
    path = Path(get_storage().get_path(row.storage_key))
    assert path.is_relative_to(tmp_path.resolve())
    return result, row, Presentation(BytesIO(path.read_bytes()))


def _text(deck):
    """提取页面可见文本，包括原生表格。"""
    return "\n".join(shape.text if shape.has_text_frame else "\n".join(
        cell.text for row in shape.table.rows for cell in row.cells) if shape.has_table else ""
        for slide in deck.slides for shape in slide.shapes)


def test_wide_long_results_preserve_all_cells_on_numbered_pages(dashboard_case, tmp_path, monkeypatch):
    case = dashboard_case
    rows = [{f"字段{column}": f"原值-{row}-{column}" for column in range(12)} for row in range(16)]
    _values(case, rows, truncated=True)
    result, artifact, deck = _export(case, tmp_path, monkeypatch, name="库存完整快照", title="库存记录")
    actual_values = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_table:
                assert len(shape.table.rows) <= 8
                assert len(shape.table.columns) <= 6
                for row in list(shape.table.rows)[1:]:
                    actual_values.extend(cell.text for cell in list(row.cells)[1:])
    assert sorted(actual_values) == sorted(value for row in rows for value in row.values())
    assert result.data["page_count"] == len(deck.slides) == 11
    assert artifact.name == "库存完整快照.pptx"
    assert artifact.run_id == case.run.id and artifact.owner_user_id == case.ctx.user_id
    proof = artifact.metadata_json["delivery_evidence"]
    assert proof["source_query_ids"] == [case.query.id] and proof["truncated"] is True
    assert proof["source_scopes"][0]["displayed_rows"] == 16
    assert proof["source_scopes"][0]["displayed_columns"] == 12
    assert proof["source_scopes"][0]["aggregation"] == "none"
    assert "已截断，仅为部分结果" in _text(deck)
    assert case.query.generated_sql in deck.slides[1].notes_slide.notes_text_frame.text


def test_units_nulls_and_long_values_are_not_rewritten_or_added(dashboard_case, tmp_path, monkeypatch):
    case = dashboard_case
    original = "实际工作记录：" + "这是一段必须完整保留的巡检备注。" * 20
    rows = [{"物料": "机械油", "数量": "2.00", "单位": "桶", "安全库存": None, "说明": original},
            {"物料": "圆钢", "数量": "120.00", "单位": "米", "安全库存": None, "说明": ""}]
    _values(case, rows)
    result, _, deck = _export(case, tmp_path, monkeypatch)
    text = _text(deck)
    assert "2.00" in text and "120.00" in text and "122" not in text
    assert "未提供（空值）" in text and "[注1]" in text
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert original in notes and '"安全库存": null' in notes
    assert result.data["scopes"][0]["notes_cells"] == 1
    assert "不做跨单位合计" in text


def test_no_rows_preserves_empty_fact(dashboard_case, tmp_path, monkeypatch):
    _values(dashboard_case, [], columns=[{"columnName": "物料"}])
    result, _, deck = _export(dashboard_case, tmp_path, monkeypatch)
    assert result.data["page_count"] == len(deck.slides) == 2
    assert result.data["scopes"][0]["returned_rows"] == 0
    assert not any(shape.has_table for slide in deck.slides for shape in slide.shapes)
    assert "本次查询未返回数据" in _text(deck)


def test_multiple_queries_keep_distinct_sources_and_no_synthetic_aggregation(dashboard_case, tmp_path, monkeypatch):
    case = dashboard_case
    second = QueryHistory(id="query-second", workspace_id=case.ctx.workspace_id, owner_user_id=case.ctx.user_id,
        run_id=case.query.run_id, datasource_id=case.source.id, query_text="库存数量", query_mode="natural_language",
        generated_sql="SELECT amount FROM synthetic_inventory", datasource_version=1, status="success",
        created_at=datetime(2026, 9, 24, 10), result_json={"queryId": "query-second", "success": True,
            "truncated": True, "totalCount": 1, "columns": [{"columnName": "数量"}, {"columnName": "单位"}],
            "data": [{"数量": "2.000", "单位": "桶"}]})
    second_call = ToolCall(id="inventory-call", workspace_id=case.ctx.workspace_id, run_id=case.run.id,
        tool_name="query_inventory_data", status="completed", arguments_json={},
        result_json={"status": "success", "data": {"queryId": second.id, "success": True}})
    case.ctx.db.add_all([second, second_call])
    case.ctx.db.commit()
    result, _, deck = _export(case, tmp_path, monkeypatch, query_ids=[case.query.id, second.id])
    assert len(deck.slides) == 5
    scopes = result.data["scopes"]
    assert [scope["returned_rows"] for scope in scopes] == [2, 1]
    assert [scope["truncated"] for scope in scopes] == [False, True]
    assert scopes[0]["queried_at"] != scopes[1]["queried_at"]
    assert "2.000" in _text(deck)


@pytest.mark.parametrize("mutation", ["foreign_owner", "failed_receipt", "unrelated_tool", "new_version", "other_source", "foreign_run"])
def test_unauthorized_or_unrelated_query_cannot_create_artifact(dashboard_case, tmp_path, monkeypatch, mutation):
    case = dashboard_case
    if mutation == "foreign_owner":
        case.query.owner_user_id = "other"
    elif mutation == "failed_receipt":
        case.call.status = "failed"
    elif mutation == "unrelated_tool":
        case.call.tool_name = "invented_query_inventory_data"
    elif mutation == "new_version":
        case.query.datasource_version = 2
    elif mutation == "other_source":
        case.tool["selected_datasource_ids"] = ["other"]
    else:
        case.run.started_by_user_id = "other"
    case.ctx.db.commit()
    with pytest.raises(AppError):
        _export(case, tmp_path, monkeypatch)
    assert case.ctx.db.scalar(select(Artifact)) is None


@pytest.mark.parametrize("changes", [{"totalCount": 100}, {"truncated": None}, {"columns": [] ,"data": []},
    {"columns": [{"columnName": "状态"}, {"columnName": "状态"}]}, {"data": ["wrong"]}])
def test_incomplete_saved_result_rejects_instead_of_inventing_data(dashboard_case, tmp_path, monkeypatch, changes):
    case = dashboard_case
    case.query.result_json = {**case.query.result_json, **changes}
    with pytest.raises(ValidationError):
        _export(case, tmp_path, monkeypatch)
    assert case.ctx.db.scalar(select(Artifact)) is None


def test_page_limit_fails_explicitly_without_dropping_records():
    query = SimpleNamespace(id="query", datasource_id="source", datasource_version=1,
        query_text="大范围快照", query_mode="natural_language", created_at=datetime(2026, 9, 24),
        generated_sql="SELECT value FROM synthetic", result_json={"totalCount": 1200, "truncated": False,
            "columns": [{"columnName": "数值"}], "data": [{"数值": index} for index in range(1200)]})
    with pytest.raises(ValidationError, match="不会省略行列"):
        query_presentation_spec([query], "全部明细")


def test_selected_page_count_is_not_silently_ignored(dashboard_case, tmp_path, monkeypatch):
    case = dashboard_case
    case.run.context_snapshot = {**case.run.context_snapshot, "plugin_inputs": {"page_count": "8"}}
    with pytest.raises(ValidationError, match="本次要求 8 页"):
        _export(case, tmp_path, monkeypatch)


def test_uploaded_template_still_displays_scope_and_all_returned_values(dashboard_case, tmp_path, monkeypatch):
    from app.modules.presentation import service

    template_path = tmp_path / "template.pptx"
    Presentation().save(template_path)
    monkeypatch.setattr(service, "resolve_template", lambda *args: {
        "id": "custom:synthetic", "source": "custom", "template_path": str(template_path)})
    _values(dashboard_case, [{"数值": number} for number in range(10)], truncated=True)
    result, artifact, deck = _export(dashboard_case, tmp_path, monkeypatch)
    actual = [cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
              for row in list(shape.table.rows)[1:] for cell in list(row.cells)[1:]]
    assert actual == [str(number) for number in range(10)]
    assert "已截断，仅为部分结果" in _text(deck)
    assert result.data["page_count"] == len(deck.slides) == 4
    assert artifact.metadata_json["template_source"] == "custom"


def test_renderer_failure_does_not_publish_an_empty_or_partial_artifact(dashboard_case, tmp_path, monkeypatch):
    from app.modules.artifact import renderers

    def fail(spec):
        raise ValueError("合成模板错误")

    monkeypatch.setattr(renderers, "render_presentation", fail)
    with pytest.raises(ValueError, match="合成模板错误"):
        _export(dashboard_case, tmp_path, monkeypatch)
    assert dashboard_case.ctx.db.scalar(select(Artifact)) is None


def test_registered_tool_dispatch_accepts_inventory_receipt(dashboard_case, tmp_path, monkeypatch):
    from app.providers.tool import builtin_provider

    case = dashboard_case
    case.call.tool_name = "query_inventory_data"
    case.ctx.db.commit()
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    specs = {item.name: item for item in builtin_provider.builtin_specs()}
    result = builtin_provider.invoke(specs["export_query_presentation"],
        {"query_ids": [case.query.id], "name": "库存接口报告.pptx"}, case.tool)
    assert result.status == "success", result.to_dict()
    artifact = case.ctx.db.get(Artifact, result.artifacts[0]["id"])
    assert artifact.name == "库存接口报告.pptx"
    assert artifact.metadata_json["delivery_evidence"]["source_query_ids"] == [case.query.id]


def test_default_report_and_scope_are_written_to_real_file(dashboard_case, tmp_path, monkeypatch):
    """默认汇报必须把本次缺失业务写入文件正文，不能只放在回执里。"""
    result, artifact, deck = _export(dashboard_case, tmp_path, monkeypatch, presentation_mode="report",
        title="应付情况汇报", scope_notes=["本次按用户选择只展示应付，应收来源尚未确认。"])
    assert result.data["presentation_mode"] == "report"
    assert artifact.metadata_json["delivery_evidence"]["presentation_mode"] == "report"
    assert "本次按用户选择只展示应付，应收来源尚未确认。" in _text(deck)
    assert "完整返回行列保存在" in _text(deck)
    assert result.data["scopes"][0]["full_data_location"] == "speaker_notes"
    assert artifact.metadata_json["delivery_evidence"]["quality_validation"]["status"] == "passed"


def test_financial_scope_is_derived_from_saved_evidence_without_model_notes(dashboard_case, tmp_path, monkeypatch):
    """模型漏填范围时仍自动区分应付已验证与应收未验证，不能把标题当来源。"""
    case = dashboard_case
    case.query.model_usage_json = {"businessSemantics": {"method": "authorized_metadata_ast", "status": "validated",
        "actualCategories": ["payable", "payment"], "monetaryFields": [{"column": "amount", "categories": ["payable"]}]}}
    case.ctx.db.commit()
    result, artifact, deck = _export(case, tmp_path, monkeypatch, presentation_mode="report", title="应收应付汇报")
    proof = artifact.metadata_json["delivery_evidence"]
    assert proof["financial_scope"]["verified_categories"] == ["payable", "payment"]
    assert proof["financial_scope"]["unverified_title_categories"] == ["receivable"]
    assert "应收未验证" in artifact.name and "应收未验证" in _text(deck)
    assert "不能据此认定数据库不存在该业务" in _text(deck)
    assert result.data["scope_notes"] == proof["automatic_scope_notes"]


def test_unvalidated_model_explanation_cannot_prove_financial_scope(dashboard_case, tmp_path, monkeypatch):
    case = dashboard_case
    case.query.model_usage_json = {"businessSemantics": {"method": "model_explanation", "status": "validated",
        "actualCategories": ["receivable"], "monetaryFields": [{"column": "amount"}]}}
    case.ctx.db.commit()
    _, artifact, deck = _export(case, tmp_path, monkeypatch, presentation_mode="report", title="应收数据报告")
    assert artifact.metadata_json["delivery_evidence"]["financial_scope"]["verified_categories"] == []
    assert "应收未验证" in _text(deck)


@pytest.mark.parametrize("damage", ["table_value", "snapshot_notes"])
def test_matching_page_count_does_not_publish_corrupted_report(dashboard_case, tmp_path, monkeypatch, damage):
    """真实渲染文件被遗漏数据或改写数值时，页数相同也必须阻止交付。"""
    from app.modules.artifact import renderers
    original = renderers.render_presentation

    def corrupt(spec):
        deck = Presentation(BytesIO(original(spec)))
        for slide in deck.slides:
            if damage == "snapshot_notes":
                slide.notes_slide.notes_text_frame.text = "备注已丢失"
            else:
                for shape in slide.shapes:
                    if shape.has_table:
                        shape.table.cell(1, 1).text = "999999"
        output = BytesIO()
        deck.save(output)
        return output.getvalue()

    monkeypatch.setattr(renderers, "render_presentation", corrupt)
    with pytest.raises(ValidationError, match="未发布文件"):
        _export(dashboard_case, tmp_path, monkeypatch, presentation_mode="report")
    assert dashboard_case.ctx.db.scalar(select(Artifact)) is None


@pytest.mark.parametrize("args", [{"presentation_mode": "fake"}, {"scope_notes": "错类型"}, {"scope_notes": [""]}])
def test_invalid_report_scope_rejects_before_artifact(dashboard_case, tmp_path, monkeypatch, args):
    with pytest.raises(ValidationError):
        _export(dashboard_case, tmp_path, monkeypatch, **args)
    assert dashboard_case.ctx.db.scalar(select(Artifact)) is None


@pytest.mark.parametrize("name,expected", [(None, "业务报告.pptx"), ("自定义", "自定义.pptx"), ("自定义.pptx", "自定义.pptx")])
def test_filename_is_fixed_at_generation(name, expected):
    assert presentation_filename(name, "业务报告") == expected


@pytest.mark.parametrize("name", ["自定义.xlsx", "自定义.ppt", "", 17])
def test_invalid_filename_is_not_misrepresented(name):
    with pytest.raises(ValidationError):
        presentation_filename(name, "业务报告")
