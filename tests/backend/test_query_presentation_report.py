"""汇报型查询 PPT 的真实文件回归，全部使用内存中的合成业务快照。"""
from datetime import datetime
from io import BytesIO
import re
from types import SimpleNamespace

from pptx import Presentation

from app.modules.artifact.renderers import render_presentation
from app.modules.chat_artifacts.query_presentation import _source
from app.modules.chat_artifacts.query_presentation_report import report_spec


def _query(query_id, rows, *, sql="SELECT * FROM synthetic", types=None):
    """构造保存结果，不建立任何业务数据库连接。"""
    return SimpleNamespace(id=query_id, datasource_id="synthetic", datasource_version=1,
        query_text=sql, query_mode="sql", generated_sql=sql, created_at=datetime(2026, 9, 29),
        result_json={"truncated": False, "totalCount": len(rows), "data": rows,
                     "columns": [{"columnName": name, "dataType": (types or {}).get(name, "varchar")}
                                 for name in rows[0]]})


def _deck(queries, notes=()):
    spec, scopes = report_spec([_source(query, index) for index, query in enumerate(queries)], "应付情况汇报", list(notes))
    spec["template_id"] = "executive-ivory"
    deck = Presentation(BytesIO(render_presentation(spec)))
    return deck, spec, scopes


def _text(deck):
    return "\n".join(shape.text if shape.has_text_frame else "\n".join(
        cell.text for row in shape.table.rows for cell in row.cells) if shape.has_table else ""
        for slide in deck.slides for shape in slide.shapes)


def test_detail_summary_and_groups_produce_brief_with_real_metrics_and_chart():
    """79 行明细不再变成 31 页台账，汇总值和图表数据必须读取实际 Office 文件。"""
    detail = _query("detail", [{"合同号": f"CN{index}", "供应商名称": f"供应商{index}", "合同金额元": "100.00",
        "已付金额元": "20.00", "未付金额元": "70.00", "本次申请支付金额元": "40.00",
        "部门": "D1", "序号": index, "说明": f"真实明细{index}"} for index in range(79)])
    summary = _query("summary", [{"明细条数": 79, "合同金额合计元": "7900.00", "已付金额合计元": "1580.00",
                                 "未付金额合计元": "5530.00", "本次申请支付金额合计元": "3160.00"}])
    groups = _query("groups", [{"部门": name, "未付金额元": amount} for name, amount in [("D1", "5000.00"), (None, "530.00")]],
                    sql="SELECT department, SUM(amount) FROM synthetic GROUP BY department")
    deck, spec, scopes = _deck([detail, summary, groups], ["本次按用户选择只做应付，应收来源尚未确认。"])
    assert len(deck.slides) == 7 == 1 + len(spec["slides"])
    text = _text(deck)
    assert all(value in text for value in ("7900.00", "1580.00", "5530.00", "3160.00"))
    assert "明细条数：79" in text and not re.search(r"查询结果\s*\d", text)
    metric_slides = [slide for slide in spec["slides"] if slide["layout"] == "metrics"]
    assert len(metric_slides) == 1 and len(metric_slides[0]["items"]) == 4
    assert "明细条数" in next(scope for scope in scopes if scope["query_id"] == "summary")["displayed_fields"]
    assert "本次按用户选择只做应付，应收来源尚未确认。" in text
    charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 1 and list(charts[0].series[0].values) == [5000.0, 530.0]
    assert [category.label for category in charts[0].plots[0].categories] == ["D1", "未提供（空值）"]
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert "真实明细78" in notes and "真实明细78" not in text
    detail_scope = next(scope for scope in scopes if scope["query_id"] == "detail")
    assert detail_scope["displayed_rows"] == 7 and detail_scope["returned_rows"] == 79
    assert detail_scope["aggregation"] == "none"
    assert "预览前 7 / 79 行" in text


def test_group_null_values_and_mixed_units_do_not_turn_into_zero_or_chart():
    """缺值和不同计量单位禁止画成可比较数值，备注保留原值。"""
    for rows in ([{"物料": "油", "数量": None}, {"物料": "钢", "数量": "2"}],
                 [{"物料": "油", "数量": "2", "单位": "桶"}, {"物料": "钢", "数量": "120", "单位": "米"}]):
        query = _query("mixed", rows, sql="SELECT material, SUM(qty) FROM synthetic GROUP BY material")
        deck, _, _ = _deck([query])
        assert not any(shape.has_chart for slide in deck.slides for shape in slide.shapes)
        notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
        assert '"数量": null' in notes if rows[0]["数量"] is None else '"单位": "桶"' in notes


def test_numeric_identifiers_are_not_business_metrics():
    query = _query("ids", [{"客户编码": "000123", "序号": 9, "ID": 21, "未付金额元": "0.00"}],
                   types={"客户编码": "bigint", "序号": "int", "ID": "int", "未付金额元": "decimal"})
    deck, _, scopes = _deck([query])
    assert scopes[0]["displayed_fields"] == ["未付金额元"]
    assert "0.00" in _text(deck) and "000123" not in _text(deck)
    assert "000123" in "\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)


def test_single_inventory_metric_keeps_material_and_unit_visible():
    query = _query("inventory", [{"物料名称": "机械油", "数量": "2.00", "单位": "桶"}])
    deck, _, _ = _deck([query])
    text = _text(deck)
    assert "物料名称：机械油" in text and "单位：桶" in text


def test_five_business_metrics_split_evenly_and_keep_all_values():
    query = _query("five", [{f"指标{index}金额": f"{index}.000000" for index in range(1, 6)}])
    deck, spec, _ = _deck([query])
    metrics = [slide for slide in spec["slides"] if slide["layout"] == "metrics"]
    assert [len(slide["items"]) for slide in metrics] == [3, 2]
    text = _text(deck)
    assert "关键指标" in text and "补充指标" in text and not re.search(r"查询结果\s*\d", text)
    assert all(f"{index}.00" in text for index in range(1, 6))


def test_table_decimal_padding_is_removed_without_rounding_or_changing_codes():
    query = _query("details", [{"合同号": "000123", "合同金额元": "100.000000"},
                               {"合同号": "000124", "合同金额元": "100.123400"}])
    deck, _, _ = _deck([query])
    cells = [cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
             for row in shape.table.rows for cell in row.cells]
    assert "000123" in cells and "000124" in cells
    assert "100.00" in cells and "100.1234" in cells and "100.000000" not in cells
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert '"100.000000"' in notes and '"100.123400"' in notes
