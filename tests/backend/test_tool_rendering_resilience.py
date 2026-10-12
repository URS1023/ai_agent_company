"""Structured tool arguments and document blocks tolerate common model output shapes."""
from __future__ import annotations

import io
import json

import pytest
from docx import Document

from app.errors import SchemaValidationError
from app.modules.artifact.renderers import render_document, render_pdf
from app.providers.tool.validator import validate_arguments


DOCUMENT_SCHEMA = {
    "type": "object",
    "properties": {
        "document_spec": {"type": "object"},
        "name": {"type": "string"},
    },
    "required": ["document_spec"],
}


def document_spec() -> dict:
    return {
        "title": "华北地区销售数据看板报告",
        "sections": [{
            "title": "分析结果",
            "blocks": [
                {
                    "type": "table",
                    "headers": ["指标", "数值"],
                    "rows": [["总销售额", "3,650,000"], ["订单数", "3"]],
                },
                {"type": "bullet", "items": ["加强项目储备", "跟踪回款周期"]},
                {"type": "numbered", "items": ["核对口径", "发布报告"]},
            ],
        }],
    }


def test_schema_boundary_decodes_only_required_json_containers():
    spec = document_spec()
    validated = validate_arguments({
        "document_spec": json.dumps(spec, ensure_ascii=False),
        "name": '{"仍然":"普通文件名文本"}',
    }, DOCUMENT_SCHEMA)
    assert validated["document_spec"] == spec
    assert validated["name"] == '{"仍然":"普通文件名文本"}'

    with pytest.raises(SchemaValidationError) as captured:
        validate_arguments({"document_spec": "{" + ("x" * 1000)}, DOCUMENT_SCHEMA)
    assert "期望 object" in str(captured.value)
    assert len(str(captured.value)) < 300


def test_document_rendering_preserves_headers_and_list_items():
    spec = document_spec()
    data = render_document(spec)
    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "加强项目储备" in paragraphs
    assert "跟踪回款周期" in paragraphs
    assert "核对口径" in paragraphs
    assert "发布报告" in paragraphs
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["指标", "数值"],
        ["总销售额", "3,650,000"],
        ["订单数", "3"],
    ]
    assert render_pdf(spec).startswith(b"%PDF")
