"""以核验过的本机业务数据和真实模型评估当前 DSH 接口、交付及跨轮修改。"""
from __future__ import annotations

import argparse
from collections import Counter
from contextlib import ExitStack
from copy import deepcopy
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
from urllib.parse import urlsplit
import xml.etree.ElementTree as ET
import zipfile

from dotenv import dotenv_values

from real_platform_smoke import Host, ROOT, HERE

STOP = ROOT / ".implementation/local_szh_acceptance/STOP_AFTER_DATABASE_ERROR.json"


def normalized(value):
    """数值按十进制文本比对；保留 NULL 与空字符串的区别。"""
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return format(Decimal(str(value)).normalize(), "f")
    return str(value)


def records(rows, columns):
    """保留重复行的规范化结果，不用集合掩盖行数错误。"""
    return Counter(json.dumps({key: normalized(row.get(key)) for key in columns}, sort_keys=True, ensure_ascii=False)
                   for row in rows)


def seed_platform(path):
    """仅向本轮系统临时 SQLite 注册本地源及已验证元数据，不访问 MySQL。"""
    sys.path.insert(0, str(ROOT / "backend"))
    from app.config import settings
    from app.db import get_engine, new_session
    from app.ids import new_id
    from app.models import DataSource
    from app.modules.analytics.models import SourceMetadata
    from app.modules.secret.service import create_secret
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    engine = get_engine()
    assert engine.url.get_backend_name() == "sqlite"
    assert Path(engine.url.database).resolve().is_relative_to(Path(tempfile.gettempdir()).resolve())
    with new_session() as db:
        identity = new_id()
        secret = create_secret(db, workspace_id=value["workspace_id"], owner_user_id=value["user_id"],
            name="本轮本机只读验收凭据", value=settings.MYSQL_PASSWORD, provider="mysql")
        metadata = value["metadata"]
        row = DataSource(id=identity, workspace_id=value["workspace_id"], owner_user_id=value["user_id"],
            name="本机 szh_new 有界验收", provider="mysql", status="active", visibility="workspace",
            created_at=datetime.now(timezone.utc), updated_at=datetime.now(timezone.utc),
            secret_id=secret.id, connection_config={"version": 1, "host": "127.0.0.1", "port": 3306,
                "databaseName": "szh_new", "username": settings.MYSQL_USER, "timeoutSeconds": 3, "maxRows": 20,
                "allowedTables": [table["tableName"] for table in metadata],
                "allowedColumns": {table["tableName"]: [column["columnName"] for column in table["columns"]]
                                   for table in metadata}})
        db.add(row)
        db.flush()
        db.add(SourceMetadata(workspace_id=value["workspace_id"], owner_user_id=value["user_id"],
            datasource_id=identity, source_version=1, status="ready", tables_json=metadata))
        db.commit()
    engine.dispose()
    print(json.dumps({"source_id": identity}))


def office_text(data, extension):
    """实际解包办公文件，不能只把 HTTP 200 视为成功交付。"""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        paths = ["word/document.xml"] if extension == ".docx" else sorted(
            [path for path in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", path)],
            key=lambda path: int(re.search(r"slide(\d+)", path).group(1)))
        return ["\n".join(node.text or "" for node in ET.fromstring(archive.read(path)).iter()
                           if node.tag.endswith("}t")) for path in paths]


def presentation_checks(data, rows):
    """同时检查实际物理页序及溢出备注，不能用文件数量替代内容核验。"""
    pages = office_text(data, ".pptx")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        notes = "\n".join("\n".join(node.text or "" for node in ET.fromstring(archive.read(path)).iter()
                            if node.tag.endswith("}t")) for path in archive.namelist()
                          if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", path))
    visible = "\n".join(pages)
    all_text = visible + "\n" + notes
    points = [row["point_name"] for row in rows]
    themes = ["概览", "保养要求", "周期", "待确认", "行动"]
    return {"requested_page_topics": len(pages) == 5 and all(topic in page for topic, page in zip(themes, pages)),
            "visible_points": sum(point in visible for point in points),
            "points_including_notes": sum(point in all_text for point in points), "expected_points": len(points),
            "has_notes": bool(notes)}


def package_changes(before_data, after_data, prefix=""):
    """递归核对办公包内容，只把创建/修改时间单独记为元数据差异。"""
    changes, timestamps = [], []
    with zipfile.ZipFile(io.BytesIO(before_data)) as before, zipfile.ZipFile(io.BytesIO(after_data)) as after:
        for path in sorted(set(before.namelist()) | set(after.namelist())):
            label = prefix + path
            if path not in before.namelist() or path not in after.namelist():
                changes.append(label)
                continue
            original, revised = before.read(path), after.read(path)
            if original == revised:
                continue
            if path.endswith(".xlsx") and zipfile.is_zipfile(io.BytesIO(original)) and zipfile.is_zipfile(io.BytesIO(revised)):
                nested_changes, nested_times = package_changes(original, revised, label + "!/")
                changes.extend(nested_changes)
                timestamps.extend(nested_times)
                continue
            if path == "docProps/core.xml":
                normalized = []
                for content in (original, revised):
                    root = ET.fromstring(content)
                    for node in list(root):
                        if node.tag in {"{http://purl.org/dc/terms/}created", "{http://purl.org/dc/terms/}modified"}:
                            root.remove(node)
                    normalized.append(ET.tostring(root))
                if normalized[0] == normalized[1]:
                    timestamps.append(label)
                    continue
            changes.append(label)
    return changes, timestamps


def find_rows(value, columns):
    """查找按约定字段保存的审计行，结果结构不符时返回空。"""
    if isinstance(value, list) and value and all(isinstance(row, dict) and set(columns) <= row.keys() for row in value):
        return value
    for item in value.values() if isinstance(value, dict) else value if isinstance(value, list) else []:
        found = find_rows(item, columns)
        if found:
            return found
    return []


class Acceptance:
    """统一记录有限请求、真实模型耗时和可核验断言。"""
    def __init__(self, host, output, baseline, report):
        self.host, self.output, self.baseline, self.report = host, output, baseline, report
        self.columns = list(baseline["rows"][0])
        self.expected = records(baseline["rows"], self.columns)
        self.started = time.monotonic()

    def api(self, method, path, *, status=200, **kwargs):
        """只访问本轮随机端口；日志不记录凭据和业务原始行。"""
        if STOP.exists():
            raise RuntimeError("数据库出现异常，本轮停止所有后续业务测试")
        started = time.monotonic()
        response = self.host.http.request(method, self.host.base + path, timeout=30, **kwargs)
        entry = {"method": method, "path": path, "status": response.status_code,
                 "milliseconds": round((time.monotonic() - started) * 1000, 2)}
        self.report["api_requests"].append(entry)
        assert response.status_code == status, (entry, response.text[:500])
        return response

    def json(self, method, path, **kwargs):
        return self.api(method, path, **kwargs).json()

    def run(self, conversation, prompt, label, **extra):
        """每轮最多四分钟，数据库异常立即取消；不自动重发用户请求。"""
        started = time.monotonic()
        admitted = self.json("POST", f"/api/chat/conversations/{conversation}/send", status=202,
            headers={"Accept": "application/json"}, json={"content": prompt, "request_id": label,
                "model_profile": "live-acceptance", "mode": "standard", "delivery_mode": "queue", **extra})
        run_id = admitted["run_id"]
        item = {"label": label, "run_id": run_id, "acceptance_seconds": round(time.monotonic() - started, 3),
                "tool_calls": [], "status": "running"}
        self.report["tasks"].append(item)
        print(json.dumps({"phase": label, "status": "started", "run_id": run_id}), flush=True)
        events, first_content, first_activity = [], None, None
        try:
            with self.host.request("GET", f"/api/chat/runs/{run_id}/events", stream=True) as response:
                event = ""
                for line in response.iter_lines(chunk_size=1, decode_unicode=True):
                    elapsed = time.monotonic() - started
                    if STOP.exists() or elapsed > 240 or time.monotonic() - self.started > 450:
                        raise TimeoutError("本轮数据库停止条件或长任务时间预算已到")
                    if line.startswith("event:"):
                        event = line[6:].strip()
                    elif line.startswith("data:"):
                        value = json.loads(line[5:].strip())
                        events.append({"event": event, "data": value, "seconds": round(elapsed, 3)})
                        if event in {"message.delta", "reasoning.delta", "tool.started"} and first_activity is None:
                            first_activity = round(elapsed, 3)
                        if event == "message.delta" and first_content is None:
                            first_content = round(elapsed, 3)
                        if event == "tool.started":
                            item["tool_calls"].append({"name": value.get("tool_name"), "seconds": round(elapsed, 3)})
                            print(json.dumps({"phase": label, "tool": value.get("tool_name"), "seconds": round(elapsed, 3)}), flush=True)
                        if event in {"run.completed", "run.failed", "run.cancelled"}:
                            item.update(status=value.get("status"), task_id=(value.get("task") or {}).get("task_id"),
                                        error=value.get("error"), content=value.get("content"))
                        if event == "done":
                            item["status"] = value["status"]
                            break
        except Exception as error:
            item.update(status="test_stopped", error={"type": type(error).__name__, "message": str(error)[:300]})
            self.host.request("POST", f"/api/chat/runs/{run_id}/cancel")
        finally:
            item.update(total_seconds=round(time.monotonic() - started, 3), first_activity_seconds=first_activity,
                        first_text_seconds=first_content)
            (self.output / (label + "-events.json")).write_text(json.dumps(events, ensure_ascii=False, indent=2), encoding="utf-8")
            print(json.dumps({"phase": label, "status": item["status"], "seconds": item["total_seconds"]}), flush=True)
        return item

    def deliveries(self, conversation, run, *, only_new=False):
        """下载真实工件并逐项检查格式和内容；原始内容仅存本机忽略目录。"""
        detail = self.json("GET", f"/api/chat/conversations/{conversation}")
        items = detail.get("projections", {}).get("deliveries", [])
        findings = []
        for item in items:
            if item.get("status") != "ready" or only_new and item.get("source_run_id") != run["run_id"]:
                continue
            data = self.api("GET", item["download_url"]).content
            extension = Path(item["name"]).suffix.lower()
            target = self.output / "deliveries" / (item["id"] + extension)
            target.parent.mkdir(exist_ok=True)
            target.write_bytes(data)
            finding = {"name": item["name"], "artifact_id": item["id"], "type": item["type"],
                "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), "path": str(target)}
            try:
                if extension == ".json":
                    rows = find_rows(json.loads(data), self.columns)
                    finding.update(rows=len(rows), exact_rows=records(rows, self.columns) == self.expected,
                                   original_order_and_values=rows == self.baseline["rows"])
                elif extension == ".xlsx":
                    from openpyxl import load_workbook
                    book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                    candidates = []
                    for sheet in book:
                        values = list(sheet.iter_rows(values_only=True))
                        for index, headers in enumerate(values):
                            if set(self.columns) <= set(headers):
                                candidates.extend(dict(zip(headers, row)) for row in values[index + 1:] if any(v is not None for v in row))
                                break
                    book.close()
                    finding.update(rows=len(candidates), exact_rows=records(candidates, self.columns) == self.expected)
                elif extension in {".docx", ".pptx"}:
                    texts = office_text(data, extension)
                    combined = "\n".join(texts)
                    finding.update(parseable=True, pages=len(texts) if extension == ".pptx" else None,
                        has_device=self.baseline["rows"][0]["device_name"] in combined, acknowledges_missing_records=any(word in combined for word in ["待确认", "缺少", "未提供", "无法判断"]),
                        slide3_revision=extension == ".pptx" and len(texts) >= 3 and "周期与缺失信息核对" in texts[2])
                    target.with_suffix(".txt").write_text(combined, encoding="utf-8")
                    if extension == ".pptx":
                        finding.update(presentation_checks(data, self.baseline["rows"]))
            except Exception as error:
                finding["validation_error"] = {"type": type(error).__name__, "message": str(error)[:200]}
            findings.append(finding)
        run["deliveries"] = findings
        return findings


def verify_api(acceptance, source_id):
    """先检验真实 HTTP 业务合同及与独立 SQL 基准的逐行一致性。"""
    a, base = acceptance, acceptance.baseline
    listed = a.json("GET", "/api/v1/analytics/datasources/page?page=1&page_size=3")["data"]
    assert any(item["id"] == source_id for item in listed["items"])
    tables = a.json("GET", f"/api/v1/analytics/datasources/{source_id}/tables")["data"]
    assert set(tables) == {item["tableName"] for item in base["metadata"]}
    query = a.json("POST", "/api/v1/analytics/query", json={"datasourceId": source_id, "query": base["sql"],
        "parameters": base["parameters"], "mode": "sql", "maxRows": 10, "needAnalysis": False, "needChart": False})["data"]
    assert records(query["data"], a.columns) == a.expected, "真实分析API结果与独立SQL基准不一致"
    restored = a.json("GET", f"/api/v1/analytics/history/{query['queryId']}")["data"]
    assert records(restored["data"], a.columns) == a.expected
    a.report["assertions"].append({"name": "真实分析API与查询历史逐行比较", "passed": True,
        "rows": len(base["rows"]), "fields_per_row": len(a.columns), "query_id": query["queryId"]})
    graph = {"nodes": [{"id": "start", "kind": "start", "title": "开始", "config": {}},
                       {"id": "value", "kind": "template", "title": "固定发布值", "config": {"template": "正式验收内容"}},
                       {"id": "end", "kind": "end", "title": "结束", "config": {}}],
             "edges": [{"id": "e1", "source": "start", "target": "value", "port": "out"},
                       {"id": "e2", "source": "value", "target": "end", "port": "out"}]}
    flow = a.json("POST", "/api/v1/workflows", status=201, json={"name": "本轮发布契约验收", "domains": ["analysis"], "graph": graph})["data"]
    published = a.json("POST", f"/api/v1/workflows/{flow['id']}/publish", json={"version": flow["version"]})["data"]
    changed = deepcopy(graph)
    changed["nodes"][1]["config"]["template"] = "未发布草稿内容"
    edited = a.json("PUT", f"/api/v1/workflows/{flow['id']}", json={"version": published["version"], "name": flow["name"],
        "domains": flow["domains"], "graph": changed})["data"]
    contract = a.json("GET", f"/api/v1/workflows/{flow['id']}/published-revision")["data"]
    assert contract["revision"] == published["publishedRevision"] and contract["version"] == edited["version"]
    assert contract["graph"]["nodes"][1]["config"]["template"] == "正式验收内容"
    a.report["assertions"].append({"name": "当前HTTP发布契约排除未发布草稿", "passed": True})
    return query


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--baseline", type=Path, help="复用本轮已核验的基准，不重复发出业务查询")
    parser.add_argument("--device-id", help="首次生成基准时显式指定一个已审阅设备业务ID")
    parser.add_argument("--models", action="store_true", help="执行当前配置真实模型的有限样例")
    parser.add_argument("--skip-api", action="store_true", help="复测模型交付时复用已测基准，不重复数据库接口查询")
    parser.add_argument("--model-case", choices=["delivery", "query"], default="delivery")
    args = parser.parse_args()
    if args.seed:
        seed_platform(args.seed)
        return
    if STOP.exists():
        raise SystemExit("已有数据库异常停止记录，需要先处理异常")
    output = args.output or ROOT / ".implementation/live_dsh_acceptance" / datetime.now().strftime("%Y%m%d-%H%M%S")
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    from live_baseline import prepare_baseline
    baseline = json.loads(args.baseline.read_text(encoding="utf-8")) if args.baseline else prepare_baseline(output, device_id=args.device_id)
    assert baseline["status"] == "passed", "基准尚未通过，不能开始接口或模型验收"
    baseline["rows"] = baseline.pop("golden_rows")
    # 接口使用会话级三秒限制；移除基准连接的等价优化器提示，保持相同业务 SQL。
    baseline["sql"] = baseline["sql"].replace("/*+ MAX_EXECUTION_TIME(3000) */", "")
    assert baseline["rows"], "本次固定设备无样本，不自动扩大数据范围"
    values = dotenv_values(ROOT / ".env.local")
    from live_model_proxy import LiveModelProxy
    report = {"started_at": datetime.now(timezone.utc).isoformat(), "scope": "本机szh_new只读；真实DSH HTTP；临时SQLite平台；真实配置模型",
              "baseline": {key: value for key, value in baseline.items() if key not in {"rows", "metadata", "parameters"}},
              "api_requests": [], "tasks": [], "assertions": [], "status": "running"}
    host, proxy = None, None
    started = time.monotonic()
    try:
        with LiveModelProxy.from_opencode(Path.home() / ".config/opencode/opencode.json", model="deepseek-v4.1-flash") as proxy, tempfile.TemporaryDirectory(prefix="test5-live-dsh-") as temporary, ExitStack() as cleanup:
            report["model"] = proxy.model
            directory = Path(temporary)
            host = Host(directory, urlsplit(proxy.base_url).port, env_overrides={
                "DSH_TEST_ISOLATION": "0", "LIVE_DSH_ACCEPTANCE": "1", "PYTHONPATH": str(HERE / "live_python_guard"),
                "LIVE_ALLOWED_TABLES": json.dumps([table["tableName"] for table in baseline["metadata"]]),
                "LIVE_REVIEWED_SMALL_JOIN": "1",
                "LIVE_DB_AUDIT_PATH": str(output / "database.jsonl"), "LIVE_DB_STOP_PATH": str(STOP),
                "MYSQL_HOST": "127.0.0.1", "MYSQL_PORT": "3306", "MYSQL_USER": values.get("MYSQL_USER") or "root",
                "MYSQL_PASSWORD": values["MYSQL_PASSWORD"], "AGNES_BASE_URL": proxy.base_url,
                "AGNES_CHAT_MODEL": proxy.model, "VISION_ENABLED": "0", "EMBEDDING_ENABLED": "0"})
            # 异常时先关闭宿主，再释放临时目录，避免 Windows 文件锁覆盖原始异常。
            cleanup.callback(lambda: host.close() if host else None)
            host.ready()
            print(json.dumps({"phase": "host_ready", "model": proxy.model}), flush=True)
            host.login()
            identity = host.json("GET", "/api/auth/me")
            seed = {"workspace_id": host.http.headers["X-Workspace-Id"], "user_id": identity["user"]["id"], "metadata": baseline["metadata"]}
            seed_file = directory / "seed.json"
            seed_file.write_text(json.dumps(seed, ensure_ascii=False), encoding="utf-8")
            seeded = subprocess.run([sys.executable, "-X", "utf8", str(Path(__file__).resolve()), "--seed", str(seed_file)],
                env=host.env, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=30)
            assert seeded.returncode == 0, seeded.stderr[-1500:]
            source_id = json.loads(seeded.stdout.strip().splitlines()[-1])["source_id"]
            a = Acceptance(host, output, baseline, report)
            query = {"queryId": "independent-baseline-snapshot"} if args.skip_api else verify_api(a, source_id)
            print(json.dumps({"phase": "business_api", "status": "reused" if args.skip_api else "passed"}), flush=True)
            if args.models:
                a.json("POST", "/api/models/profiles", status=201, json={"profile_name": "live-acceptance", "provider": "openai-compatible",
                    "model": proxy.model, "endpoint": proxy.base_url, "api_key": "local-acceptance-proxy", "capabilities": ["chat", "tools"],
                    "runtime_options": {"protocol": "openai-completions", "thinking": "disabled", "context_window": 32768, "max_output_tokens": 4096}})
            if args.models and args.model_case == "query":
                conversation = a.json("POST", "/api/chat/conversations", status=201, json={"title": "真实模型单设备取数验收"})["id"]
                checked = a.run(conversation,
                    f"请查询当前已选择的数据源，设备业务ID为 {baseline['parameters']['device_id']} 的未删除设备名称是什么？"
                    "只取这一条真实记录，回答名称和数据来源。不要生成文件，也不要扫描其他业务。",
                    "model-business-query", datasource_ids=[source_id], skill_ids=["business-analysis", "data-access"])
                used_query = any(tool["name"] == "query_business_data" for tool in checked["tool_calls"])
                events = json.loads((output / "model-business-query-events.json").read_text(encoding="utf-8"))
                results = [((event["data"].get("result") or {}).get("data") or {}) for event in events
                           if event["event"] == "tool.completed" and event["data"].get("tool_name") == "query_business_data"]
                actual = next((result for result in reversed(results) if result.get("generatedSql")), {})
                checked["generated_sql"] = actual.get("generatedSql")
                checked["query_execution_ms"] = actual.get("executionTime")
                checked["sql_model_calls"] = (actual.get("modelUsage") or {}).get("modelCalls")
                checked["sql_repairs"] = (actual.get("modelUsage") or {}).get("sqlRepairs")
                from sqlglot import parse_one, exp
                sql = parse_one(actual["generatedSql"], read="mysql") if actual.get("generatedSql") else None
                predicates = list(sql.find_all(exp.EQ)) if sql is not None else []
                strict_delete = any(isinstance(predicate.this, exp.Column) and predicate.this.name.lower() == "isdeleted"
                                    and isinstance(predicate.expression, exp.Literal) and predicate.expression.this == "0"
                                    for predicate in predicates)
                null_broadening = sql is not None and any(isinstance(node, (exp.Is, exp.Coalesce))
                    and any(column.name.lower() == "isdeleted" for column in node.find_all(exp.Column)) for node in sql.walk())
                actual_name = any(baseline["rows"][0]["device_name"] in row.values() for row in actual.get("data", []))
                report["assertions"].append({"name": "真实模型调用业务查询且名称符合独立基准",
                    "passed": checked["status"] == "completed" and used_query and actual_name
                    and baseline["rows"][0]["device_name"] in (checked.get("content") or "")})
                report["assertions"].append({"name": "本样例未删除条件为严格零且不引入NULL", "passed": strict_delete and not null_broadening})
            if args.models and args.model_case == "delivery":
                snapshot = json.dumps({"source": "本机szh_new本轮独立只读核验", "query_id": query["queryId"], "rows": baseline["rows"]}, ensure_ascii=False)
                uploaded = a.json("POST", "/api/files", status=201, files={"file": ("verified-maintenance.json", snapshot.encode(), "application/json")})
                file_id = uploaded.get("id") or uploaded.get("file_id")
                assert a.api("GET", f"/api/files/{file_id}/download").content == snapshot.encode()
                conversation = a.json("POST", "/api/chat/conversations", status=201, json={"title": "本机真实业务复杂任务验收"})["id"]
                prompt = (f"请根据附件中已经核验的{baseline['rows'][0]['device_name']}保养要求完成一整套交付。本轮只使用附件，不重新查询数据库。"
                    "附件是计划要求，没有实际执行记录；空值保持空，不得把要求写成已完成或已经逾期。"
                    "请先读完整附件，再自主依次完成：1. audit.json，rows逐行保持所有字段名、值和重复行；"
                    "2. 一份Excel，第一行表头严格使用附件rows的原始字段名，全部数据完整；"
                    "3. 一份Word，含设备与计划摘要、每个润滑点的要求表、缺失信息与待确认事项；"
                    "4. 五页PPT，依次为概览、保养要求、周期、待确认、行动建议，保留可继续编辑的草稿。"
                    "四种文件都要实际生成、可下载；在全部完成后给出交付清单。不要把多阶段任务停留在计划，也不要反复确认。")
                complex_run = a.run(conversation, prompt, "complex-delivery", file_ids=[file_id], skill_ids=["document-delivery", "file-work"],
                    required_deliveries=[{"type": value, "count": 1} for value in ["json", "spreadsheet", "document", "presentation"]])
                a.deliveries(conversation, complex_run)
                if complex_run["status"] == "completed" and complex_run.get("task_id") and not STOP.exists():
                    revised = a.run(conversation, "请只把刚才PPT第3页标题修改为“周期与缺失信息核对”，保持五页及其他内容，生成新版PPT。其他已交付文件保留即可。",
                        "cross-turn-revision", task_id=complex_run["task_id"], skill_ids=["document-delivery"],
                        required_deliveries=[{"type": "presentation", "count": 1}])
                    a.deliveries(conversation, revised, only_new=True)
                    report["assertions"].append({"name": "跨轮任务标识保持", "passed": revised.get("task_id") == complex_run["task_id"]})
                    originals = [item for item in complex_run["deliveries"] if item["type"] == "presentation"]
                    revisions = [item for item in revised["deliveries"] if item["type"] == "presentation"]
                    if originals and revisions:
                        changed, timestamps = package_changes(Path(originals[-1]["path"]).read_bytes(), Path(revisions[-1]["path"]).read_bytes())
                        report["assertions"].append({"name": "仅物理第三页内容变化", "changed_entries": changed,
                            "timestamp_only_entries": timestamps,
                            "passed": changed == ["ppt/slides/slide3.xml"] and revisions[-1].get("slide3_revision", False)})
            host.close()
            host = None
            report["models"] = proxy.summary()
            report["model_requests"] = proxy.records
            report["status"] = "measured"
    except Exception as error:
        report.update(status="stopped", error={"type": type(error).__name__, "message": str(error)[:1500]})
        if host:
            report["diagnostics"] = list(host.diagnostics)[-12:]
    finally:
        if host:
            host.close()
        if proxy:
            report["models"] = proxy.summary()
            report["model_requests"] = proxy.records
        audit_path = output / "database.jsonl"
        if audit_path.exists():
            report["database_audit"] = [json.loads(line) for line in audit_path.read_text(encoding="utf-8").splitlines() if line.strip()]
        report["elapsed_seconds"] = round(time.monotonic() - started, 3)
        # 完成测量不代表验收通过；失败和内容断言必须反映在脚本退出码里。
        report["checks_passed"] = (report["status"] == "measured"
            and all(item.get("passed") is True for item in report["assertions"])
            and all(task.get("status") == "completed" for task in report["tasks"])
            and all(not delivery.get("validation_error")
                    and delivery.get("exact_rows", True)
                    and delivery.get("requested_page_topics", True)
                    and delivery.get("points_including_notes", 0) == delivery.get("expected_points", 0)
                    for task in report["tasks"] for delivery in task.get("deliveries", [])))
        (output / "result.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        summary = {"status": report["status"], "checks_passed": report["checks_passed"], "report": str(output / "result.json"), "model": report.get("model"),
            "tasks": [{key: value for key, value in task.items() if key not in {"content", "deliveries"}} for task in report["tasks"]],
            "assertions": report["assertions"], "error": report.get("error"), "elapsed_seconds": report["elapsed_seconds"]}
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if report["checks_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
