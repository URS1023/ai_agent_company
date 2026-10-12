"""真实本机财务与库存的复合任务验收，平台隔离、逐项比对、跨轮继续。"""
from __future__ import annotations

from collections import Counter
import argparse
from contextlib import ExitStack
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
from pathlib import Path
import re
import sqlite3
import subprocess
import sys
import tempfile
import time
from urllib.parse import urlsplit

from dotenv import dotenv_values

from live_baseline import _verify_local_service
from live_database_functional import OUTPUT, STOP, SQL_CASES, write
import live_business_acceptance as business
from live_model_proxy import LiveModelProxy
from real_platform_smoke import Host, ROOT, HERE


def canonical(value):
    """金额和数量按十进制比对，NULL 不与空字符串混同。"""
    if value is None:
        return None
    if not isinstance(value, bool) and re.fullmatch(r"-?\d+(?:\.\d+)?", str(value)):
        return format(Decimal(str(value)).normalize(), "f")
    return str(value)


def rows_equal(left, right, mapping=None):
    """保留重复行，按确定的别名映射比较全部选定字段。"""
    mapping = mapping or {key: key for row in right for key in row}
    def pack(rows, actual):
        return Counter(json.dumps([canonical(row.get(mapping[key] if actual else key)) for key in mapping], ensure_ascii=False)
                       for row in rows)
    return len(left) == len(right) and pack(left, True) == pack(right, False)


def contains_financial_values(visible, baseline):
    """使用十进制数值比较，允许金额保留两位小数，不把显示补零判为数据错误。"""
    tokens = {canonical(token) for token in re.findall(r"(?<![\w.])-?\d+(?:\.\d+)?(?![\w.])", visible.replace(",", ""))}
    return all(canonical(value) in tokens for key, value in baseline.items() if key.endswith("amount"))


def platform_snapshot(directory):
    """只读取当前临时 SQLite 的执行证据，不再连接业务库。"""
    with sqlite3.connect(directory / "platform.sqlite") as db:
        db.row_factory = sqlite3.Row
        output = {}
        for name, sql in {
            "queries": "SELECT id,status,query_mode,generated_sql,result_json,model_usage_json,execution_ms,run_id FROM analytics_query_history",
            "artifacts": "SELECT id,name,type,metadata_json,run_id FROM artifacts",
            "reports": "SELECT id,name,content,lineage,width,height,version,template_id,template_mode,background_color,status FROM report_documents",
        }.items():
            output[name] = [dict(row) for row in db.execute(sql)]
    for name in output:
        for row in output[name]:
            for key in ("result_json", "model_usage_json", "metadata_json", "content", "lineage"):
                if key in row and isinstance(row[key], str):
                    row[key] = json.loads(row[key])
    return output


def query_kind_and_mapping(query):
    """按实际 SQL 字段来源识别查询类别，不以模型回答或数字碰巧相同为依据。"""
    import sqlglot
    from sqlglot import exp
    tree = sqlglot.parse_one(query["generated_sql"], read="mysql")
    sources = {table.name for table in tree.find_all(exp.Table)}
    aliases = {}
    actual_fields = {"c_contractAmount": "contract_amount", "c_pamented": "paid_amount", "c_unpayment": "unpaid_amount",
                     "c_requestPament": "requested_amount", "c_bumen": "department_id", "spInventoryId": "stock_id",
                     "c_productName": "material_name", "c_productModel": "stock_model", "c_availableStock": "quantity",
                     "c_unit": "unit_code", "c_safetyStock": "safety_stock", "c_warehouseName": "warehouse_name",
                     "c_locationId": "location_id", "c_batchNo": "batch"}
    for item in tree.selects:
        fields = list(item.find_all(exp.Column))
        if item.find(exp.Count):
            aliases["record_count"] = item.alias_or_name
        elif len(fields) == 1 and fields[0].name in actual_fields:
            key = actual_fields[fields[0].name]
            if key not in aliases:
                aliases[key] = item.alias_or_name
    if "t_summaryoffundpayments" in sources:
        return ("finance_department" if tree.args.get("group") else "finance_total"), aliases
    if "t_spinventory" in sources:
        return "inventory_sample", aliases
    return None, aliases


def verify_queries(snapshot, baseline, report):
    """把模型执行后的查询与独立 SQL 基准逐项比对。"""
    verified = {}
    for query in snapshot["queries"]:
        if query["status"] != "success" or query["query_mode"] == "profile":
            continue
        kind, mapping = query_kind_and_mapping(query)
        if not kind:
            continue
        fields = set(baseline[kind]["rows"][0])
        complete = fields <= mapping.keys()
        exact = complete and rows_equal(query["result_json"].get("data") or [], baseline[kind]["rows"], {key: mapping[key] for key in fields})
        verified[query["id"]] = {"kind": kind, "exact": exact, "run_id": query["run_id"], "fields_complete": complete}
    report["query_verification"] = verified
    return verified


def audit_deliveries(a, cid, task, directory, baseline, report, label):
    """下载并读取真实文件，检查数量、来源、可见指标和跨轮修改目标。"""
    from openpyxl import load_workbook
    detail = a.json("GET", f"/api/chat/conversations/{cid}")
    snapshot = platform_snapshot(directory)
    write(f"{a.output.name}/{label}-platform.json", snapshot)
    verified = verify_queries(snapshot, baseline, report)
    artifacts = {row["id"]: row for row in snapshot["artifacts"]}
    queries = {row["id"]: row for row in snapshot["queries"]}
    deliveries = []
    for item in detail.get("projections", {}).get("deliveries", []):
        if item.get("status") != "ready" or item.get("source_run_id") != task["run_id"]:
            continue
        checked = {"id": item["id"], "type": item["type"], "name": item["name"]}
        if item["type"] == "dashboard":
            saved = next(row for row in snapshot["reports"] if row["id"] == item["report_id"])
            checked["scope_queries"] = (saved["lineage"] or {}).get("queryIds") or []
            checked["verified_sources"] = all(verified.get(qid, {}).get("exact") for qid in checked["scope_queries"])
            write(f"{a.output.name}/{label}-dashboard.json", saved)
        else:
            data = a.api("GET", item["download_url"]).content
            path = a.output / (item["id"] + Path(item["name"]).suffix)
            path.write_bytes(data)
            checked.update(bytes=len(data), path=str(path))
            metadata = artifacts[item["id"]]["metadata_json"]
            evidence = metadata.get("delivery_evidence") or {}
            if item["type"] == "spreadsheet":
                query = queries.get(evidence.get("source_query_id"))
                checked["verified_source"] = bool(query and verified.get(query["id"], {}).get("exact"))
                book = load_workbook(BytesIO(data), data_only=True, read_only=True)
                expected = (query or {}).get("result_json", {}).get("data") or []
                checked["actual_rows_match"] = False
                if expected:
                    for sheet in book:
                        values = list(sheet.iter_rows(values_only=True))
                        for index, headers in enumerate(values):
                            if set(expected[0]) <= set(headers):
                                rows = [dict(zip(headers, row)) for row in values[index + 1:] if any(value is not None for value in row)]
                                checked["actual_rows_match"] |= rows_equal(rows, expected)
                book.close()
            elif item["type"] == "presentation":
                pages = business.office_text(data, ".pptx")
                checked["pages"] = len(pages)
                checked["page_titles"] = [page.splitlines()[0] if page else "" for page in pages]
                checked["is_erp"] = "ERP" in (item["name"] + (pages[0] if pages else "")).upper()
                visible = "\n".join(pages).replace(",", "")
                checked["actual_financial_values"] = contains_financial_values(visible, baseline["finance_total"]["rows"][0])
                checked["quality_validation"] = metadata.get("quality_validation") or evidence.get("quality_validation")
                checked["third_page_revised"] = len(pages) == 5 and "ERP流程与协同" in pages[2]
        deliveries.append(checked)
    task["deliveries"] = deliveries
    return snapshot


def run(label="live"):
    """最多六次业务 SELECT、四十次模型请求，任一数据库故障停止且不重试。"""
    if STOP.exists():
        raise RuntimeError("已有数据库停止标记，不启动验收")
    target = _verify_local_service(10)
    baseline = json.loads((OUTPUT / "database-baseline.json").read_text(encoding="utf-8"))
    metadata = json.loads((OUTPUT / "authorized-test-metadata.json").read_text(encoding="utf-8"))
    if not re.fullmatch(r"[a-z0-9-]{1,40}", label):
        raise ValueError("验收批次名称只能使用小写字母、数字和短横线")
    output = OUTPUT / label
    if (output / "result.json").exists():
        raise ValueError("批次结果已存在，请保留原始失败证据并指定新批次")
    output.mkdir(exist_ok=True)
    business.STOP = STOP
    values = dotenv_values(ROOT / ".env.local")
    report = {"target": target, "scope": "本机szh_new真实只读；平台临时SQLite；最新DSH/Worker；OpenCode deepseek-v4.1-flash",
              "api_requests": [], "tasks": [], "assertions": [], "status": "running"}
    host = proxy = None
    started = time.perf_counter()
    try:
        with LiveModelProxy.from_opencode(Path.home() / ".config/opencode/opencode.json", model="deepseek-v4.1-flash") as proxy, tempfile.TemporaryDirectory(prefix="test5-complex-db-") as temp, ExitStack() as cleanup:
            directory = Path(temp)
            host = Host(directory, urlsplit(proxy.base_url).port, env_overrides={
                "DSH_TEST_ISOLATION": "0", "LIVE_DSH_ACCEPTANCE": "1", "PYTHONPATH": str(HERE / "live_python_guard"),
                "LIVE_ALLOWED_TABLES": json.dumps([row["tableName"] for row in metadata]),
                "LIVE_REVIEWED_SCAN_LIMITS": json.dumps({"t_summaryoffundpayments": 150}),
                "LIVE_DB_AUDIT_PATH": str(output / "database.jsonl"), "LIVE_DB_STOP_PATH": str(STOP),
                "MYSQL_HOST": "127.0.0.1", "MYSQL_PORT": "3306", "MYSQL_USER": values.get("MYSQL_USER") or "root",
                "MYSQL_PASSWORD": values["MYSQL_PASSWORD"], "AGNES_BASE_URL": proxy.base_url,
                "AGNES_CHAT_MODEL": proxy.model, "VISION_ENABLED": "0", "EMBEDDING_ENABLED": "0"})
            cleanup.callback(host.close)
            host.ready()
            host.login()
            identity = host.json("GET", "/api/auth/me")
            seed = directory / "seed.json"
            seed.write_text(json.dumps({"workspace_id": host.http.headers["X-Workspace-Id"], "user_id": identity["user"]["id"], "metadata": metadata}, ensure_ascii=False), encoding="utf-8")
            child = subprocess.run([sys.executable, "-X", "utf8", str(HERE / "live_business_acceptance.py"), "--seed", str(seed)],
                                   env=host.env, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=30)
            if child.returncode:
                raise RuntimeError(child.stderr[-1200:])
            source = json.loads(child.stdout.strip().splitlines()[-1])["source_id"]
            a = business.Acceptance(host, output, {"rows": baseline["finance_total"]["rows"]}, report)
            for case in ("finance_total", "inventory_sample"):
                result = a.json("POST", "/api/v1/analytics/query", json={"datasourceId": source, "query": SQL_CASES[case], "mode": "sql", "maxRows": 20, "needAnalysis": False, "needChart": False})["data"]
                passed = rows_equal(result.get("data") or [], baseline[case]["rows"])
                report["assertions"].append({"name": case + "_http_exact", "passed": passed, "rows": len(result.get("data") or [])})
                if not passed:
                    raise RuntimeError("实际 HTTP 查询与独立 SQL 基准不符")
            print(json.dumps({"phase": "actual_business_api", "status": "passed", "queries": 2}), flush=True)
            a.json("POST", "/api/models/profiles", status=201, json={"profile_name": "live-acceptance", "provider": "openai-compatible",
                "model": proxy.model, "endpoint": proxy.base_url, "api_key": "local-acceptance-proxy", "capabilities": ["chat", "tools"],
                "runtime_options": {"protocol": "openai-completions", "thinking": "disabled", "context_window": 32768, "max_output_tokens": 4096}})
            cid = a.json("POST", "/api/chat/conversations", status=201, json={"title": "本机财务库存多步长任务验收"})["id"]
            prompt = ("请基于选中的本机数据源完成一个完整任务：查询资金支付汇总明细的全部未删除记录，给出记录数、合同金额、已付金额、未付金额、本次申请支付金额的总计以及按部门汇总；"
                      "再查备品备件库存及关联的备件名称、仓库名称、数量、单位、安全库存、批次和仓位，取未删除记录前10条即可。"
                      "分别生成资金支付Excel和库存Excel，再生成资金支付情况汇报PPT、数字大屏，最后制作另一份5页的ERP系统介绍PPT（五页主题依次为概览、业务模块、流程协同、实施步骤、价值）。"
                      "同一结果用于多个文件时不要重新查询。不同物料单位不合计；缺失阈值不要当作正常。这里要的是资金支付和应付口径，不把它改称应收。"
                      "名称无法关联时保留真实编码并说明。请自主完成全部5项交付，不用逐项等我确认。")
            task = a.run(cid, prompt, "complex-finance-inventory", datasource_ids=[source],
                         skill_ids=["business-analysis", "data-access", "document-delivery", "dashboard-delivery"],
                         required_deliveries=[{"type": "spreadsheet", "count": 2}, {"type": "presentation", "count": 2}, {"type": "dashboard", "count": 1}])
            snapshot = audit_deliveries(a, cid, task, directory, baseline, report, "initial")
            counts = Counter(item["type"] for item in task["deliveries"])
            report["assertions"].append({"name": "exact_requested_deliveries", "passed": counts == {"spreadsheet": 2, "presentation": 2, "dashboard": 1}, "actual": dict(counts)})
            kinds = {row["kind"] for row in report["query_verification"].values() if row["exact"]}
            report["assertions"].append({"name": "all_data_scopes_verified", "passed": kinds == set(SQL_CASES)})
            if task["status"] == "completed" and task.get("task_id") and not STOP.exists():
                before = len(snapshot["queries"])
                revised = a.run(cid, "只修改刚才的ERP系统介绍PPT，把第3页标题改为“ERP流程与协同”，其余四页及五页总数保持不变。重新交付这份PPT，不重新查询数据库，其他文件保留。",
                                "cross-turn-erp-revision", task_id=task["task_id"], datasource_ids=[source], skill_ids=["document-delivery"],
                                required_deliveries=[{"type": "presentation", "count": 1}])
                after = audit_deliveries(a, cid, revised, directory, baseline, report, "revised")
                report["assertions"].append({"name": "revision_no_requery", "passed": len(after["queries"]) == before})
                report["assertions"].append({"name": "revision_correct_target", "passed": any(row.get("third_page_revised") for row in revised["deliveries"])})
                original = next((row for row in task["deliveries"] if row.get("is_erp")), None)
                new = next((row for row in revised["deliveries"] if row.get("is_erp")), None)
                if original and new:
                    changed, stamps = business.package_changes(Path(original["path"]).read_bytes(), Path(new["path"]).read_bytes())
                    report["assertions"].append({"name": "only_slide3_changed", "passed": changed == ["ppt/slides/slide3.xml"], "changed": changed, "timestamp_only": stamps})
            report["status"] = "measured"
    except Exception as error:
        report.update(status="stopped", error={"type": type(error).__name__, "message": str(error)[:1200]})
    finally:
        if proxy:
            report["model"] = proxy.summary()
            report["model_requests"] = proxy.records
        report["elapsed_seconds"] = round(time.perf_counter() - started, 3)
        report["cross_turn_tested"] = any(task["label"] == "cross-turn-erp-revision" for task in report["tasks"])
        report["passed"] = report["status"] == "measured" and bool(report["tasks"]) and all(row["passed"] for row in report["assertions"]) and all(task["status"] == "completed" for task in report["tasks"])
        for task in report["tasks"]:
            for item in task.get("deliveries", []):
                report["passed"] &= all(item.get(key, True) for key in ("actual_rows_match", "verified_source", "verified_sources"))
                if item["type"] == "presentation":
                    report["passed"] &= item["pages"] == 5 if item["is_erp"] else bool(item.get("actual_financial_values") and (item.get("quality_validation") or {}).get("status") == "passed")
        write(f"{label}/result.json", report)
        print(json.dumps({"status": report["status"], "passed": report["passed"], "seconds": report["elapsed_seconds"], "error": report.get("error"), "report": str(output / "result.json")}, ensure_ascii=False), flush=True)
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--label", default="live")
    raise SystemExit(run(parser.parse_args().label))
