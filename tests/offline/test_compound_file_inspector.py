"""仅用内存生成合成附件检查独立验收判定，不调用业务代码或数据库。"""
from io import BytesIO
import importlib.util
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches as PptInches


MODULE = Path(__file__).resolve().parents[1] / 'ui' / 'compound_file_inspector.py'
spec = importlib.util.spec_from_file_location('compound_file_inspector', MODULE)
inspector = importlib.util.module_from_spec(spec)
spec.loader.exec_module(inspector)


def package_change(data, change):
    """只改变指定 XML/成员，保留其余真实包结构。"""
    with ZipFile(BytesIO(data)) as source:
        files = {name: source.read(name) for name in source.namelist()}
    change(files)
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for name, value in files.items():
            target.writestr(name, value)
    return output.getvalue()


def deck(pages=15):
    """生成可读取的普通正文页，用于关系而非排版测试。"""
    presentation = Presentation()
    for index in range(pages):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f'巡检第{index + 1}页'
        slide.placeholders[1].text = '应检8台，已检3台，未检5台；本文件只用于合成验收。'
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def picture(color='blue', size=(320, 180)):
    """测试用非纯色图，不能作为真实业务产物注入浏览器场景。"""
    image = Image.new('RGB', size, 'white')
    ImageDraw.Draw(image).rectangle((10, 10, size[0] // 2, size[1] - 10), fill=color)
    output = BytesIO()
    image.save(output, 'PNG')
    return output.getvalue()


def word(images):
    document = Document()
    document.add_paragraph('应检8台，未检5台，未检占比62.5%。')
    for image in images:
        document.add_picture(BytesIO(image), width=Inches(4))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def sheet(value=3):
    workbook = Workbook()
    workbook.active.title = '查询结果'
    workbook.active.append(['状态', '设备数'])
    workbook.active.append(['已巡检', value])
    workbook.active.append(['未巡检', 5])
    origin = workbook.create_sheet('数据来源')
    origin.append(['查询标识', 'Q'])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_actual_fifteen_pages_are_read_in_relationship_order():
    result = inspector.inspect_presentation(deck(), {'pages': 15, 'titles': ['巡检第13页']})
    assert result['pages'] == len(set(result['targets'])) == 15


@pytest.mark.parametrize('fault', ['orphan', 'duplicate', 'empty'])
def test_slide_files_cannot_fake_complete_delivery(fault):
    def change(files):
        if fault == 'empty':
            node = ET.fromstring(files['ppt/slides/slide13.xml'])
            for text in node.iter(inspector.A + 't'):
                text.text = ''
            files['ppt/slides/slide13.xml'] = ET.tostring(node)
        else:
            node = ET.fromstring(files['ppt/presentation.xml'])
            slides = node.find(inspector.P + 'sldIdLst')
            if fault == 'orphan':
                slides.remove(slides[-1])
            else:
                slides[-1].set(inspector.R + 'id', slides[0].get(inspector.R + 'id'))
            files['ppt/presentation.xml'] = ET.tostring(node)
    with pytest.raises(ValueError, match='页数|重复引用|可见正文'):
        inspector.inspect_presentation(package_change(deck(), change), {'pages': 15})


def test_word_requires_both_real_referenced_images():
    first, second = picture(), picture('red')
    digests = [inspector.sha256(data).hexdigest() for data in [first, second]]
    result = inspector.inspect_word(word([first, second]), {'images': digests, 'text': ['未检5台']})
    assert set(result['images']) == set(digests)


@pytest.mark.parametrize('fault', ['orphan', 'repeated', 'broken'])
def test_word_cannot_substitute_orphan_repeated_or_broken_picture(fault):
    first, second = picture(), picture('red')
    expected = {'images': [inspector.sha256(data).hexdigest() for data in [first, second]]}
    data = word([first, first] if fault == 'repeated' else [first, second])
    def change(files):
        if fault == 'orphan':
            node = ET.fromstring(files['word/document.xml'])
            for drawing in node.findall('.//' + inspector.W + 'drawing'):
                drawing.clear()
            files['word/document.xml'] = ET.tostring(node)
        elif fault == 'broken':
            node = ET.fromstring(files['word/document.xml'])
            next(node.iter(inspector.A + 'blip')).set(inspector.R + 'embed', 'missing')
            files['word/document.xml'] = ET.tostring(node)
    with pytest.raises(ValueError, match='双图|关系缺失'):
        inspector.inspect_word(package_change(data, change), expected)


@pytest.mark.parametrize('value', [3, '3', '=1+2', 4])
def test_spreadsheet_compares_real_cells_and_types(value):
    expected = {'columns': ['状态', '设备数'], 'rows': [['已巡检', 3], ['未巡检', 5]], 'origin': {'查询标识': 'Q'}}
    if type(value) is int and value == 3:
        assert inspector.inspect_spreadsheet(sheet(value), expected)['rows'][1] == ('已巡检', 3)
    else:
        with pytest.raises(ValueError, match='数字|公式'):
            inspector.inspect_spreadsheet(sheet(value), expected)


def test_spreadsheet_orphan_values_do_not_replace_real_result():
    def change(files):
        files['xl/worksheets/orphan.xml'] = files['xl/worksheets/sheet1.xml']
        files['xl/worksheets/sheet1.xml'] = files['xl/worksheets/sheet1.xml'].replace(b'<v>3</v>', b'<v>9</v>')
    with pytest.raises(ValueError, match='数字'):
        inspector.inspect_spreadsheet(package_change(sheet(), change),
            {'columns': ['状态', '设备数'], 'rows': [['已巡检', 3], ['未巡检', 5]]})


@pytest.mark.parametrize('fault', ['tiny', 'blank', 'transparent'])
def test_placeholder_images_are_rejected(fault):
    if fault == 'tiny':
        data = picture(size=(20, 20))
    else:
        image = Image.new('RGBA', (320, 180), (255, 0, 0, 0) if fault == 'transparent' else 'white')
        stream = BytesIO()
        image.save(stream, 'PNG')
        data = stream.getvalue()
    with pytest.raises(ValueError, match='微型|有效可见'):
        inspector.inspect_image(data)


def test_bundle_cannot_read_outside_download_directory(tmp_path):
    outside = tmp_path / 'outside.png'
    outside.write_bytes(picture())
    root = tmp_path / 'downloads'
    root.mkdir()
    with pytest.raises(ValueError, match='本轮下载目录'):
        inspector.inspect_bundle({'download_root': str(root), 'files': {'chart': str(outside)}})


@pytest.mark.parametrize('wrong_pie', [False, True])
def test_complete_download_bundle_checks_real_native_pie_values(tmp_path, wrong_pie):
    """只生成解析器测试包；浏览器夹具永远通过正式工具创建自己的产物。"""
    first, second = picture(), picture('red')
    presentation = Presentation(BytesIO(deck()))
    chart = CategoryChartData()
    chart.categories = ['已巡检', '未巡检']
    chart.add_series('设备数', [4 if wrong_pie else 3, 5])
    presentation.slides[3].shapes.add_chart(XL_CHART_TYPE.PIE, PptInches(1), PptInches(2), PptInches(4), PptInches(3), chart)
    output = BytesIO()
    presentation.save(output)
    data = {'chart': first, 'dashboard': second, 'word': word([first, second]),
            'spreadsheet': sheet(), 'presentation': output.getvalue()}
    files = {}
    for kind, content in data.items():
        location = tmp_path / kind
        location.write_bytes(content)
        files[kind] = str(location)
    pie = {'labels': ['已巡检', '未巡检'], 'values': [3, 5]}
    proof = inspector.sha256(inspector.json.dumps({'labels': pie['labels'], 'values': [3.0, 5.0]},
        ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    request = {'download_root': str(tmp_path), 'files': files, 'expected': {
        'pie': pie, 'rendered_values_sha256': proof,
        'word': {'text': ['未检占比62.5%']},
        'spreadsheet': {'columns': ['状态', '设备数'], 'rows': [['已巡检', 3], ['未巡检', 5]], 'origin': {'查询标识': 'Q'}},
        'presentation': {'pages': 15, 'pie': pie},
    }}
    if wrong_pie:
        with pytest.raises(ValueError, match='饼图'):
            inspector.inspect_bundle(request)
    else:
        result = inspector.inspect_bundle(request)
        assert result['presentation']['pages'] == 15
        assert result['chart']['sha256'] in result['word']['images']
        assert result['dashboard']['sha256'] in result['word']['images']
