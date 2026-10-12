"""离线核验封尾开关、物理页修改与真实 PPTX 溢出备注；禁止数据库和网络。"""
from copy import deepcopy
import io
import socket
from types import SimpleNamespace
import zipfile
from xml.etree import ElementTree

import pytest
from pptx import Presentation
from sqlalchemy.engine import Engine

from app.common.storage import LocalStorage
from app.errors import ToolExecutionError
from app.modules.presentation import draft
from app.modules.artifact.renderers import _render_uploaded_presentation
from app.modules.presentation.engine import render_builtin_presentation, presentation_overflow_report
from app.modules.presentation.service import get_template
from app.modules.presentation.spec_normalizer import normalize_presentation_spec
from app.providers.tool import builtin_provider


@pytest.fixture(autouse=True)
def offline_only(monkeypatch, tmp_path):
    """草稿只写临时目录，连接数据库或外部服务即令测试失败。"""
    def denied(*args, **kwargs):
        raise AssertionError("本测试禁止数据库或网络连接")
    monkeypatch.setattr(Engine, "connect", denied)
    monkeypatch.setattr(socket.socket, "connect", denied)
    monkeypatch.setattr(socket.socket, "connect_ex", denied)
    storage = LocalStorage(tmp_path)
    monkeypatch.setattr(draft, "get_storage", lambda: storage)


def page(title):
    """合成数据只用于确认页序与内容保留，不伪装真实业务结论。"""
    return {"title": title, "layout": "cards", "key_message": f"{title}需要明确来源和复核边界。",
            "items": [
                {"title": "输入范围", "description": "仅使用本次提供的合成材料，明确主题边界和当前已知条件，不把样本外的事实当成既定结论。"},
                {"title": "证据核对", "description": "逐项核对输入材料与输出说明，记录引用来源并将缺失字段明确标为未提供，保证统计口径一致。"},
                {"title": "后续行动", "description": "列出需要进一步确认的责任事项和执行步骤，保留修订记录，并在补充材料后复核相关判断。"},
            ]}


def render(spec):
    return render_builtin_presentation(spec, get_template("executive-ivory"))


def render_route(spec, route, tmp_path):
    """同时覆盖内置、客户原生版式和客户模板无占位符的降级路径。"""
    if route == "builtin":
        return render(spec)
    template = Presentation()
    if route == "uploaded-fallback":
        for layout in template.slide_layouts:
            for shape in list(layout.placeholders):
                shape._element.getparent().remove(shape._element)
    path = tmp_path / "template.pptx"
    template.save(path)
    return _render_uploaded_presentation(spec, str(path))


def xml_pages(data):
    """保留每张物理页的完整 XML，比较修改前后的非目标页。"""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")}


@pytest.mark.parametrize("cover,closing", [(False, False), (True, True), (True, False), (False, True)])
@pytest.mark.parametrize("route", ["builtin", "uploaded", "uploaded-fallback"])
def test_exact_five_pages_respects_independent_cover_flags(cover, closing, route, tmp_path):
    """四种封尾组合均准确生成五页，未启用的页面不占用正文主题。"""
    titles = ["概览", "要求", "周期", "待确认", "行动"][:5 - int(cover) - int(closing)]
    status = draft.create_draft("w", "r", {"title": "精确页数验收", "page_count": 5,
        "include_cover": cover, "include_closing": closing, "slides": [page(title) for title in titles]},
        template_id="executive-ivory", aspect_ratio="16:9")
    assert status["total_page_count"] == 5
    assert status["body_page_count"] == len(titles)
    assert [item["page_number"] for item in status["outline"]] == list(range(1 + int(cover), 1 + int(cover) + len(titles)))
    draft.write_draft_slide_batches("w", "r", [{**page(title), "slide_number": index}
        for index, title in enumerate(titles, start=1)])
    spec = draft.final_spec("w", "r")
    assert spec["include_cover"] is cover and spec["include_closing"] is closing
    data = render_route(spec, route, tmp_path)
    deck = Presentation(io.BytesIO(data))
    assert len(deck.slides) == 5
    for index, title in enumerate(titles, start=int(cover)):
        assert title in " ".join(shape.text for shape in deck.slides[index].shapes if shape.has_text_frame)
    if not cover and not closing:
        assert [item["title"] for item in spec["slides"]] == ["概览", "要求", "周期", "待确认", "行动"]


def test_legacy_defaults_and_resume_keep_cover_and_page_mapping():
    """没有开关的旧调用仍保留封尾，跨轮恢复保留物理页与正文索引映射。"""
    status = draft.create_draft("w", "source", {"page_count": 5, "slides": [page(str(i)) for i in range(3)]},
        template_id="executive-ivory", aspect_ratio="16:9")
    assert status["include_cover"] and status["include_closing"]
    assert status["body_page_count"] == 3
    assert status["pages"][0]["role"] == "cover"
    assert status["pages"][-1]["role"] == "closing"
    resumed = draft.resume_draft("w", "source", "target")
    assert resumed["outline"][1]["slide_number"] == 2
    assert resumed["outline"][1]["page_number"] == 3
    assert resumed["pages"] == status["pages"]
    normalized = normalize_presentation_spec({"slides": [page(str(i)) for i in range(3)]}, page_count=5)
    assert normalized["include_cover"] and normalized["include_closing"]
    assert len(Presentation(io.BytesIO(render(normalized))).slides) == 5


@pytest.mark.parametrize("cover,closing", [(True, True), (False, False)])
def test_edit_physical_third_page_changes_only_that_slide(cover, closing):
    """解包实际 OOXML，确认用户第3页只改该页，其他物理页逐字节不变。"""
    pages = [page(f"原始主题{i}") for i in range(1, 6 - int(cover) - int(closing))]
    draft.create_draft("w", "r", {"title": "逐页修改验收", "page_count": 5, "slides": pages,
        "include_cover": cover, "include_closing": closing}, template_id="executive-ivory", aspect_ratio="16:9")
    draft.write_draft_slide_batches("w", "r", [{**item, "slide_number": i} for i, item in enumerate(pages, 1)])
    before = xml_pages(render(draft.final_spec("w", "r")))
    status = draft.write_draft_slide_batches("w", "r", [{**page("仅修改物理第三页"), "physical_page_number": 3}])
    assert status["written_now"] == [3 - int(cover)]
    after = xml_pages(render(draft.final_spec("w", "r")))
    assert set(before) == set(after)
    assert [name for name in before if before[name] != after[name]] == ["ppt/slides/slide3.xml"]


def test_legacy_page_number_alias_and_conflicting_physical_page():
    """历史 page_number 仍表示正文索引，新物理页字段不能与正文索引相互矛盾。"""
    draft.create_draft("w", "r", {"page_count": 5, "slides": [page(str(i)) for i in range(3)]},
        template_id="executive-ivory", aspect_ratio="16:9")
    status = draft.write_draft_slide_batches("w", "r", [{**page("历史正文第三页"), "page_number": 3}])
    assert status["written_now"] == [3]
    assert draft.load_draft("w", "r")["slides"][2]["title"] == "历史正文第三页"
    with pytest.raises(ToolExecutionError):
        draft.write_draft_slide_batches("w", "r", [{**page("矛盾页码"), "slide_number": 3, "physical_page_number": 3}])
    with pytest.raises(ToolExecutionError):
        draft.write_draft_slide_batches("w", "r", [{**page("误改封面"), "physical_page_number": 1}])


@pytest.mark.parametrize("route", ["builtin", "uploaded", "uploaded-fallback"])
def test_all_ten_table_rows_are_visible_or_in_real_notes_xml(route, tmp_path):
    """七行可见、三行写入真实备注 XML；每条明细和原备注都能从文件读取。"""
    rows = [[f"点位-{i:02}", f"要求-{i:02}", f"周期-{i:02}"] for i in range(1, 11)]
    spec = normalize_presentation_spec({"title": "表格溢出", "page_count": 3,
        "include_cover": False, "include_closing": False, "slides": [
            {"title": "十行明细", "layout": "table", "notes": "已有备注必须保留", "table": {
                "headers": ["点位", "要求", "周期"], "rows": rows}}, page("复核"), page("行动")]})
    data = render_route(spec, route, tmp_path)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        visible = " ".join(ElementTree.fromstring(archive.read("ppt/slides/slide1.xml")).itertext())
        notes_names = [name for name in archive.namelist() if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")]
        assert notes_names
        notes = " ".join(" ".join(ElementTree.fromstring(archive.read(name)).itertext()) for name in notes_names)
    assert "已有备注必须保留" in notes
    assert "其余条目" in visible and "共 3 行" in visible
    for index, row in enumerate(rows):
        for value in row:
            assert value in (visible if index < 7 else notes)
        if index >= 7:
            assert row[0] not in visible
    overflow = presentation_overflow_report(spec)
    assert overflow[0]["page_number"] == 1
    assert overflow[0]["total_data_rows"] == 10
    assert overflow[0]["visible_data_rows"] == 7
    assert overflow[0]["notes_data_rows"] == 3


def test_render_tool_reports_actual_overflow_and_keeps_explicit_flags(monkeypatch):
    """工具入口到真实 PPTX 都遵守封尾开关，并把实际备注溢出情况反馈给模型。"""
    from app.modules.artifact import service
    from app.modules.presentation import service as presentation_service
    captured = {}
    def create(*args, **kwargs):
        captured.update(kwargs)
        return SimpleNamespace(id="artifact", name=kwargs["name"])
    monkeypatch.setattr(service, "create_artifact", create)
    monkeypatch.setattr(presentation_service, "resolve_template", lambda *args: {"id": "executive-ivory"})
    tool = next(item for item in builtin_provider.builtin_specs() if item.name == "render_presentation")
    spec = {"title": "精确交付", "page_count": 3, "include_cover": False, "include_closing": False,
        "slides": [{"title": "数据明细", "layout": "table", "table": {"headers": ["点位", "要求"],
            "rows": [[f"点位{i}", f"要求{i}"] for i in range(10)]}}, page("核对"), page("行动")]}
    result = builtin_provider.invoke(tool, {"name": "audit.pptx", "presentation_spec": spec}, {
        "workspace_id": "w", "user_id": "u", "run_id": "r",
        "db": SimpleNamespace(commit=lambda: None, get=lambda *args: None)})
    assert captured["name"] == "audit.pptx"
    assert len(Presentation(io.BytesIO(captured["data"])).slides) == 3
    assert result.data["presentation_overflow"][0]["visible_data_rows"] == 7
    assert result.data["presentation_overflow"][0]["notes_data_rows"] == 3
    assert captured["metadata"]["presentation_overflow"] == result.data["presentation_overflow"]


@pytest.mark.parametrize("requested,expected", [("audit.json", "audit.json"), ("audit.JSON", "audit.JSON"), ("audit", "audit.json")])
def test_json_artifact_extension_is_not_duplicated(monkeypatch, requested, expected):
    """真实渲染、伪工件存储，不访问数据库即可检查工具最终文件名。"""
    from app.modules.artifact import service
    captured = {}
    def create(*args, **kwargs):
        captured.update(kwargs)
        return SimpleNamespace(id="artifact", name=kwargs["name"])
    monkeypatch.setattr(service, "create_artifact", create)
    tool = next(item for item in builtin_provider.builtin_specs() if item.name == "save_json_artifact")
    result = builtin_provider.invoke(tool, {"name": requested, "data": {"rows": 10}}, {
        "workspace_id": "w", "user_id": "u", "run_id": "r", "db": SimpleNamespace(commit=lambda: None)})
    assert captured["name"] == expected
    assert result.data["name"] == expected
