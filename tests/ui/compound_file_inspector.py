"""独立检查浏览器已下载的文件；不导入应用，不连接数据库、网络或模型。"""
from __future__ import annotations

import base64
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import posixpath
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from PIL import Image, ImageStat
from openpyxl import load_workbook


R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
MAX_BYTES = 64 * 1024 * 1024


def check(condition, message):
    """失败立即中止，不能把损坏附件降级成仅检查文件名。"""
    if not condition:
        raise ValueError(message)


class Package:
    """有界解析 Office 关系；保留图片原始字节，不向磁盘解压成员。"""

    def __init__(self, data, marker):
        check(len(data) <= MAX_BYTES, "文件超出核验预算")
        self.zip = ZipFile(BytesIO(data))
        entries = self.zip.infolist()
        self.names = {item.filename for item in entries}
        check(0 < len(entries) < 10000 and len(entries) == len(self.names), "ZIP 条目为空、重复或过多")
        check(sum(item.file_size for item in entries) <= 128 * 1024 * 1024, "ZIP 累计展开过大")
        for item in entries:
            check(item.file_size <= 16 * 1024 * 1024 and not item.flag_bits & 1, "ZIP 成员过大或加密")
            check(not item.filename.startswith(('/', '\\')) and '..' not in item.filename.split('/'), "ZIP 成员路径越界")
        self.xml('[Content_Types].xml')
        roots = self.relations('')
        check(any(row['target'] == marker and row['type'].endswith('/officeDocument') and not row['external']
                  for row in roots.values()), "缺少真实 Office 包入口")
        self.root = self.xml(marker)
        self.marker = marker

    def xml(self, name):
        """拒绝声明实体的 XML，只处理包内有界内容。"""
        check(name in self.names, '缺少部件：' + name)
        data = self.zip.read(name)
        check(b'<!DOCTYPE' not in data.upper() and b'<!ENTITY' not in data.upper(), 'XML 不允许实体声明')
        return ET.fromstring(data)

    def relations(self, owner):
        """命名空间前缀可变，但关系身份、类型与目标必须真实。"""
        name = posixpath.join(posixpath.dirname(owner), '_rels', posixpath.basename(owner) + '.rels')
        if name not in self.names:
            return {}
        result = {}
        for node in self.xml(name):
            identity, target = node.get('Id'), node.get('Target', '')
            check(identity and identity not in result, '关系标识为空或重复')
            external = node.get('TargetMode') == 'External'
            if not external:
                target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
                check(not target.startswith('../') and '\\' not in target, '关系目标越界')
            result[identity] = {'target': target, 'type': node.get('Type', ''), 'external': external}
        return result

    def target(self, relations, identity, kind):
        """正文使用的外链、断链和错误类型不算交付。"""
        row = relations.get(identity, {})
        check(not row.get('external') and row.get('type', '').endswith('/' + kind)
              and row.get('target') in self.names, '关系缺失、外链或类型不符：' + kind)
        return row['target']


def inspect_image(data):
    """解码并检查可见像素，排除占位微图、透明图和纯色图。"""
    check(len(data) <= MAX_BYTES, '图片超出核验预算')
    with Image.open(BytesIO(data)) as original:
        width, height = original.size
        check(width >= 160 and height >= 100 and width * height <= 40_000_000, '图片是微型占位图或过大')
        check(original.format == 'PNG', '业务图必须是实际 PNG')
        original.verify()
    with Image.open(BytesIO(data)) as original:
        visible = Image.new('RGBA', original.size, 'white')
        visible.alpha_composite(original.convert('RGBA'))
        visible = visible.convert('RGB')
        visible.thumbnail((640, 640))
        check(max(ImageStat.Stat(visible).stddev) > 2, '图片没有有效可见内容')
    return {'sha256': sha256(data).hexdigest(), 'width': width, 'height': height}


def inspect_presentation(data, expected):
    """以真实引用顺序校验页数、每页正文与巡检饼图数值。"""
    package = Package(data, 'ppt/presentation.xml')
    check(package.root.tag == P + 'presentation', '不是演示文稿根节点')
    relations = package.relations(package.marker)
    ids = package.root.findall('./' + P + 'sldIdLst/' + P + 'sldId')
    check(len(ids) == expected.get('pages', 5), '实际引用页数不符合要求')
    targets, texts, charts = [], [], []
    for index, node in enumerate(ids):
        target = package.target(relations, node.get(R + 'id'), 'slide')
        check(target not in targets, '演示顺序重复引用同一页')
        targets.append(target)
        document = package.xml(target)
        check(document.tag == P + 'sld', '幻灯片部件类型错误')
        text = ' '.join(item.text or '' for item in document.iter(A + 't')).strip()
        check(text, f'第{index + 1}页缺少可见正文')
        page_titles = expected.get('page_titles', [])
        if page_titles:
            check(len(page_titles) == len(ids) and page_titles[index] in text, f'第{index + 1}页主题错误')
        if 0 < index < len(ids) - 1:
            check(len(text) >= expected.get('min_body_characters', 1), f'第{index + 1}页正文过少')
        texts.append(text)
        for chart_ref in document.iter(C + 'chart'):
            chart_path = package.target(package.relations(target), chart_ref.get(R + 'id'), 'chart')
            chart = package.xml(chart_path)
            for pie in chart.iter(C + 'pieChart'):
                for series in pie.iter(C + 'ser'):
                    categories = series.find(C + 'cat')
                    values = series.find(C + 'val')
                    charts.append({'labels': [item.text for item in categories.iter(C + 'v')] if categories is not None else [],
                                   'values': [item.text for item in values.iter(C + 'v')] if values is not None else []})
    for title in expected.get('titles', []):
        check(any(title in text for text in texts), '实际幻灯片缺少主题：' + title)
    if expected.get('pie'):
        wanted = expected['pie']
        check(any(row['labels'] == wanted['labels'] and len(row['values']) == len(wanted['values'])
                  and all(Decimal(a) == Decimal(str(b)) for a, b in zip(row['values'], wanted['values']))
                  for row in charts), 'PPT 没有引用正确的巡检饼图数据')
    return {'pages': len(targets), 'targets': targets, 'texts': texts, 'pie_series': charts, 'sha256': sha256(data).hexdigest()}


def inspect_word(data, expected):
    """只认正文可达的指定两张图片，孤立 media 文件不能凑数。"""
    package = Package(data, 'word/document.xml')
    check(package.root.tag == W + 'document', '不是 Word 正文')
    text = ''.join(item.text or '' for item in package.root.iter(W + 't'))
    for term in expected.get('text', []):
        check(term in text, 'Word 缺少正文：' + term)
    images = []
    relations = package.relations(package.marker)
    for node in [*package.root.iter(A + 'blip'), *package.root.iter('{urn:schemas-microsoft-com:vml}imagedata')]:
        check(not node.get(R + 'link'), 'Word 只含外链图片')
        identity = node.get(R + 'embed') or node.get(R + 'id')
        target = package.target(relations, identity, 'image')
        image_data = package.zip.read(target)
        images.append(sha256(image_data).hexdigest())
        if images[-1] in expected.get('images', []):
            inspect_image(image_data)
    wanted = expected.get('images', [])
    check(len(set(wanted)) == len(wanted) and set(wanted).issubset(images), 'Word 缺少指定双图的真实正文引用')
    if expected.get('song_font'):
        for marker in ['word/document.xml', 'word/styles.xml']:
            fonts = list(package.xml(marker).iter(W + 'rFonts'))
            check(fonts, 'Word 缺少宋体字体声明')
            for font in fonts:
                check(all(font.get(W + field) == '宋体' for field in ['ascii', 'hAnsi', 'eastAsia', 'cs']), 'Word 字体不是宋体')
    return {'text': text, 'images': images, 'sha256': sha256(data).hexdigest()}


def inspect_spreadsheet(data, expected):
    """逐格核对真实工作表，禁止孤立表、公式和数字文本冒充结果。"""
    package = Package(data, 'xl/workbook.xml')
    check(package.root.tag == S + 'workbook', '不是 Excel 工作簿')
    relations = package.relations(package.marker)
    targets = []
    for node in package.root.findall('./' + S + 'sheets/' + S + 'sheet'):
        target = package.target(relations, node.get(R + 'id'), 'worksheet')
        check(target not in targets, '工作簿重复引用同一工作表')
        targets.append(target)
        check(package.xml(target).tag == S + 'worksheet', '工作表类型错误')
    workbook = load_workbook(BytesIO(data), data_only=False, read_only=False, keep_links=False)
    check('查询结果' in workbook and '数据来源' in workbook, 'Excel 缺少结果或来源表')
    sheet = workbook['查询结果']
    actual = list(sheet.values)
    wanted = [expected['columns'], *expected['rows']]
    check(len(actual) == len(wanted), 'Excel 结果行数错误')
    for row, wanted_row in zip(sheet.iter_rows(), wanted):
        check(len(row) == len(wanted_row), 'Excel 结果列数错误')
        for cell, value in zip(row, wanted_row):
            check(cell.data_type not in {'f', 'e'}, 'Excel 公式或错误不能替代实际值')
            if type(value) in (int, float):
                check(type(cell.value) in (int, float) and Decimal(str(cell.value)) == Decimal(str(value)), 'Excel 数字类型或数值错误')
            else:
                check(type(cell.value) is type(value) and cell.value == value, 'Excel 文本或空值错误')
    origin = dict(workbook['数据来源'].values)
    for key, value in expected.get('origin', {}).items():
        check(key in origin and str(origin[key]) == str(value), 'Excel 数据来源错误：' + key)
    if expected.get('song_font'):
        for table in workbook.worksheets:
            for row in table:
                for cell in row:
                    if cell.value is not None:
                        check(cell.font.name == '宋体', 'Excel 字体不是宋体')
    return {'rows': actual, 'origin': origin, 'sha256': sha256(data).hexdigest()}


def inspect_bundle(request):
    """所有输入限定在本轮下载目录，只读文件并输出验证事实。"""
    root = Path(request['download_root']).resolve(strict=True)
    files = {}
    for kind in ['chart', 'dashboard', 'word', 'spreadsheet', 'presentation']:
        path = Path(request['files'][kind]).resolve(strict=True)
        check(path.is_relative_to(root) and path.is_file(), '文件不在本轮下载目录')
        check(path.stat().st_size <= MAX_BYTES, '下载文件超出预算')
        files[kind] = path.read_bytes()
    expected = request['expected']
    chart = inspect_image(files['chart'])
    dashboard = inspect_image(files['dashboard'])
    check(chart['sha256'] != dashboard['sha256'], '饼图与大屏不能是同一张图')
    proof = sha256(json.dumps({'labels': expected['pie']['labels'], 'values': [float(value) for value in expected['pie']['values']]},
                             ensure_ascii=False, sort_keys=True, allow_nan=False).encode()).hexdigest()
    check(proof == expected['rendered_values_sha256'], '饼图来源数值摘要不匹配')
    return {'chart': chart, 'dashboard': dashboard,
            'word': inspect_word(files['word'], {**expected['word'], 'images': [chart['sha256'], dashboard['sha256']]}),
            'spreadsheet': inspect_spreadsheet(files['spreadsheet'], expected['spreadsheet']),
            'presentation': inspect_presentation(files['presentation'], expected['presentation'])}


def main():
    """从标准输入接收有界 JSON；只运行独立解析，错误以非零退出返回。"""
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    raw = sys.stdin.buffer.read(96 * 1024 * 1024 + 1)
    check(len(raw) <= 96 * 1024 * 1024, '检查请求过大')
    request = json.loads(raw)
    if request.get('kind') == 'pptx':
        result = inspect_presentation(base64.b64decode(request['base64'], validate=True), request['expected'])
    else:
        result = inspect_bundle(request)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(f'{type(error).__name__}: {error}', file=sys.stderr)
        raise SystemExit(1)
